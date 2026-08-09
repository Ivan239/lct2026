"""Canvas synthesis (plan 9.3): a role the template lacks is synthesized on a
CLONE of a template slide — its background art and chrome survive — instead of
a blank white page. Verified structurally (renders are eyeballed in sweeps)."""

import os

from conftest import SURVEY_31, TEMPLATES_DIR, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from common.synthesis import SYNTHESIZE
from generator.generator import generate

TJ = os.path.join(TEMPLATES_DIR, "custom_f496182bb15f42bb.pptx")

PLAN = [
    ({"type": "title", "title": "Т", "subtitle": "П"}, 4),
    ({"type": "closing", "title": "Спасибо!", "subtitle": "Пока"}, SYNTHESIZE),
]


@requires(TJ)
def test_canvas_keeps_background_art(tmp_path):
    out = str(tmp_path / "canvas.pptx")
    generate(TJ, PLAN, out, synth_canvas={1: 10})
    prs = Presentation(out)
    assert len(prs.slides) == 2
    synth = prs.slides[1]
    # The T-Ж folder background is a full-bleed PICTURE — it must survive.
    assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in synth.shapes)
    texts = " ".join(
        s.text_frame.text for s in synth.shapes if s.has_text_frame
    )
    assert "Спасибо!" in texts
    # The canvas's own content text must NOT leak through.
    assert "что вы есть" not in texts


@requires(TJ)
def test_without_canvas_hint_still_works(tmp_path):
    out = str(tmp_path / "no_canvas.pptx")
    generate(TJ, PLAN, out)  # from-scratch path — package gate must still pass
    assert Presentation(out).slides


@requires(TJ)
def test_image_frame_survives_a_narrow_canvas_and_clears_its_chrome():
    """A canvas is REUSED for a different role, so its own text area is not
    automatically a usable content area. T-Zh study slide 10 is one low line and
    hands back a 1.03in band against the template's 4.04in: the title ate it,
    both candidate bands for the image frame came out NEGATIVE, and the slide
    shipped as a heading alone on a full-page empty card. Every T-Zh template
    has three such canvases, and they are exactly the ones offered for synthesis.

    Falling back to the template-wide band then has to respect the furniture the
    clone still carries — first render after the fallback showed the dashed
    frame crossing the footer rule and boxing in «КОММЕНТАРИЙ»."""
    from pptx.util import Emu

    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.layout_bounds import infer_content_bounds
    from generator.slide_kit import is_chrome_shape
    from generator.synthesizer import IMAGE_PLACEHOLDER_NAME, synthesize_image_caption
    from template_parser.parser import extract_template, extract_theme

    prs = Presentation(TJ)
    theme = apply_observed_style(extract_theme(TJ), observe_deck_style(prs))
    bounds = infer_content_bounds(extract_template(TJ))

    idx = synthesize_image_caption(
        prs, theme, bounds,
        {"title": "Платформа в работе", "image": "экран дашборда"},
        canvas_idx=10,
    )
    slide = prs.slides[idx]
    frames = [s for s in slide.shapes if s.name == IMAGE_PLACEHOLDER_NAME]
    assert frames, "the image slide shipped as a bare heading"

    frame = frames[0]
    footers = [
        s for s in slide.shapes
        if s is not frame and s.top is not None and s.height is not None
        and is_chrome_shape(s, prs.slide_height)
        and s.top > prs.slide_height // 2
    ]
    assert footers, "fixture changed: this canvas is supposed to carry footer chrome"
    lowest_allowed = min(int(s.top) for s in footers)
    assert int(frame.top + frame.height) <= lowest_allowed, (
        f"frame runs into the canvas footer: ends at "
        f"{Emu(frame.top + frame.height).inches:.2f}in, footer starts at "
        f"{Emu(lowest_allowed).inches:.2f}in")


@requires(TJ)
def test_canvas_placeholders_are_blanked_but_managed_chrome_is_kept():
    """A clone keeps the whole footer band, and nothing on this path blanked it:
    «КОММЕНТАРИЙ» shipped verbatim under EVERY synthesized slide of this deck.
    The native fill path has always cleared exactly this class of text.

    The split matters as much as the clearing — the page number and the running
    topic slot are owned by later passes (iter31) and must survive, or this fix
    just re-breaks what that one repaired."""
    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.layout_bounds import infer_content_bounds
    from generator.slide_kit import is_chrome_shape
    from generator.synthesizer import synthesize_image_caption
    from template_parser.parser import extract_template, extract_theme

    prs = Presentation(TJ)
    theme = apply_observed_style(extract_theme(TJ), observe_deck_style(prs))
    bounds = infer_content_bounds(extract_template(TJ))

    # Slide 8 carries all three: the topic slot, a page number and the
    # ownerless «КОММЕНТАРИЙ» — the split this test is about.
    canvas = prs.slides[8]
    before = [
        s.text_frame.text.strip() for s in canvas.shapes
        if s.has_text_frame and is_chrome_shape(s, prs.slide_height)
        and s.text_frame.text.strip()
    ]
    assert any("КОММЕНТАРИЙ" in t for t in before), (
        "fixture changed: this canvas is supposed to carry the placeholder")
    assert any(t.isdigit() for t in before), (
        "fixture changed: this canvas is supposed to carry a page number")

    idx = synthesize_image_caption(
        prs, theme, bounds, {"title": "Платформа", "image": "дашборд"}, canvas_idx=8)
    after = [
        s.text_frame.text.strip() for s in prs.slides[idx].shapes
        if s.has_text_frame and is_chrome_shape(s, prs.slide_height)
        and s.text_frame.text.strip()
    ]
    assert not any("КОММЕНТАРИЙ" in t for t in after), f"placeholder shipped: {after}"
    assert any(t.isdigit() for t in after), f"the page number was wiped too: {after}"


@requires(TJ)
def test_image_frame_does_not_reserve_room_for_a_wrap_that_cannot_happen():
    """The clearance under the title carries a safety margin because the height
    estimate can under-count lines (the renderer wraps 2-4% earlier than the
    metrics predict). That margin has to be EARNED: a title that provably fits
    on one line was still given the two-line floor, leaving 0.47in of dead band
    between the heading and the frame — a visible hole on the render.

    Both directions are checked: the short title gets exactly the gap, the long
    one keeps its reserve and the frame never climbs into it."""
    from pptx.util import Emu, Inches

    from fonts.metrics import FontResolver
    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.layout_bounds import infer_content_bounds
    from generator.synthesizer import IMAGE_PLACEHOLDER_NAME, synthesize_image_caption
    from template_parser.parser import extract_template, extract_theme

    bounds = infer_content_bounds(extract_template(TJ))
    gaps = {}
    for label, title in (
        ("short", "Платформа в работе"),
        ("long", "Как мы перестроили процесс согласования отчётности в компании"),
    ):
        prs = Presentation(TJ)
        theme = apply_observed_style(extract_theme(TJ), observe_deck_style(prs))
        idx = synthesize_image_caption(
            prs, theme, bounds, {"title": title, "image": "дашборд"},
            resolver=FontResolver(TJ, extract_theme(TJ)), canvas_idx=8)
        shapes = prs.slides[idx].shapes
        frame = next(s for s in shapes if s.name == IMAGE_PLACEHOLDER_NAME)
        title_box = next(s for s in shapes if s.name.startswith("TextBox"))
        gaps[label] = int(frame.top) - int(title_box.top + title_box.height)

    assert 0 < gaps["short"] <= int(Inches(0.4)), (
        f"dead band under a title that cannot wrap: {Emu(gaps['short']).inches:.2f}in")
    assert gaps["long"] > 0, "the frame climbed into a title that did wrap"


@requires(SURVEY_31)
def test_synthesized_block_steps_around_the_canvas_decor():
    """The survey-31 canvas keeps a small hand-drawn heart at 4.53-5.46in, and
    the synthesized KPI block landed across it — the render showed the artwork
    sitting on top of «клиентов в месяц».

    Two causes, and neither fix works alone. The canvas's own band is 2.82in on
    a 7.5in slide, while a heading plus three KPI pairs need about 2.9in, so the
    content overflowed the band's bottom — which is exactly where the heart
    lives. Widening the band alone just re-centres the block onto the heart;
    stepping around obstacles alone does nothing, because the 2.82in band offers
    no clear position. Measured in that order, one attempt at a time."""
    from pptx.util import Emu

    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.layout_bounds import infer_content_bounds
    from generator.synthesizer import synthesize_stats_kpi
    from template_parser.parser import extract_template, extract_theme

    prs = Presentation(SURVEY_31)
    theme = apply_observed_style(extract_theme(SURVEY_31), observe_deck_style(prs))
    bounds = infer_content_bounds(extract_template(SURVEY_31))

    idx = synthesize_stats_kpi(
        prs, theme, bounds,
        {"title": "Результаты",
         "stats": [["-40%", "времени на отчёт"], ["+18", "клиентов в месяц"],
                   ["4 дня", "на внедрение"]]},
        canvas_idx=29)
    slide = prs.slides[idx]

    decor = [s for s in slide.shapes if "PICTURE" in str(s.shape_type)]
    assert decor, "fixture changed: this canvas is supposed to carry decor"
    texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert len(texts) >= 4, f"expected a title and KPI pairs, got {len(texts)}"

    for art in decor:
        art_top, art_bottom = int(art.top), int(art.top + art.height)
        for box in texts:
            top, bottom = int(box.top), int(box.top + box.height)
            assert not (top < art_bottom and bottom > art_top), (
                f"{box.text_frame.text.strip()[:24]!r} at "
                f"{Emu(top).inches:.2f}-{Emu(bottom).inches:.2f}in runs across decor at "
                f"{Emu(art_top).inches:.2f}-{Emu(art_bottom).inches:.2f}in")


def test_canvas_picker_skips_the_templates_readme_page():
    """"Fewest text boxes" picks the blandest slide, and the blandest slide is
    sometimes the template's own instructions page. The T-Zh universal deck ends
    with a «Технический слайд» telling whoever uses the template which fonts to
    download; it has 3 boxes and no furniture, the branded candidates have 5 and
    9, so the <=3 cut left it as the only option and ALL THREE synthesized
    slides were cloned from it — bare white pages on a template whose identity
    is black with yellow artwork.

    Its background is (248,248,248) on exactly one slide of the deck, so
    singleton backgrounds are dropped first. The rotation must survive that:
    filtering to the deck's MAJORITY background instead was tried and reverted,
    because it collapsed T-Zh mono from alternating blue/white to white only."""
    import glob
    import json
    import os

    from common.synthesis import SYNTHESIZE
    from design_system.style_profile import build_measured_profile
    from evaluation.loop import _synth_canvas_hints

    def hints_for(tid, model, positions=4):
        cache = f"output/loop/{model}/{tid}/archetypes.json"
        pngs = sorted(glob.glob(f"output/loop/{model}/{tid}/rendered/*.png"))
        if not os.path.exists(cache) or not pngs:
            return None
        archetypes = {int(k): v for k, v in json.load(open(cache)).items()}
        profile = build_measured_profile(pngs, archetypes)
        plan = [({"type": "bullet_list"}, SYNTHESIZE)] * positions
        return _synth_canvas_hints(f"output/templates/{tid}.pptx", plan, profile)

    universal = hints_for("custom_47dfd8952eb47583", "GigaChat-3-Ultra")
    if universal:
        assert 11 not in universal.values(), (
            f"the readme page is being cloned as a canvas: {universal}")

    mono = hints_for("custom_838830368dac3116", "GigaChat-2")
    if mono:
        assert len(set(mono.values())) > 1, (
            f"colour rotation collapsed to a single canvas: {mono}")


@requires(TJ)
def test_synthesized_title_uses_the_templates_own_size():
    """Synthesized slides set their titles at a flat 32pt whatever the deck.
    Measured with the same title picker the rest of the pipeline uses, the real
    templates sit at 42 / 51 / 26 / 40 / 32pt — so on T-Zh mono ours came out a
    third smaller than the designer's, and on the study template a fifth larger.

    Mode, not mean: a designed size repeats across the deck (it equals the
    median on four of the five templates), a one-off display figure does not."""
    import os

    from conftest import TEMPLATES_DIR
    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.generator import _max_font_pt, _template_title_pt
    from generator.layout_bounds import infer_content_bounds
    from generator.synthesizer import synthesize_bullet_list
    from template_parser.parser import extract_template, extract_theme

    mono = os.path.join(TEMPLATES_DIR, "custom_838830368dac3116.pptx")
    if not os.path.exists(mono):
        return

    assert _template_title_pt(Presentation(mono)) == 51
    assert _template_title_pt(Presentation(TJ)) == 26, "the study template is SMALLER than 32"

    prs = Presentation(mono)
    theme = apply_observed_style(extract_theme(mono), observe_deck_style(prs))
    theme["title_pt"] = _template_title_pt(prs)
    # A CONTENT slide. The cover (synthesize_title) keeps its own hero scale on
    # purpose — real covers run to the template's largest size, not its median.
    idx = synthesize_bullet_list(prs, theme, infer_content_bounds(extract_template(mono)),
                                 {"title": "Итоги", "bullets": ["Раз", "Два"]})
    titles = [_max_font_pt(s) for s in prs.slides[idx].shapes
              if s.has_text_frame and "Итоги" in s.text_frame.text]
    assert titles and titles[0] == 51, f"synthesized title shipped at {titles}"


@requires(SURVEY_31)
def test_cover_subtitle_clears_a_title_at_the_templates_own_size():
    """The cover used a flat 40pt while these templates set theirs at
    42 / 68 / 34 / 92 / 92. Taking the template's size is the point — and it
    immediately broke the layout, because the subtitle sat at a hard-coded 60%
    of the slide height. That was clear of a one-line 40pt title and nothing
    else: at 88pt the title wrapped to two lines, ending at 5.94in against a
    subtitle at 4.50in, and the render showed «Итоги внедрения за квартал»
    printed straight through it.

    Same lesson as iter35: reserve from the real box, never from an offset that
    happens to work at one size."""
    from pptx.util import Emu

    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.generator import _template_cover_pt
    from generator.layout_bounds import infer_content_bounds
    from generator.synthesizer import synthesize_title
    from template_parser.parser import extract_template, extract_theme

    prs = Presentation(SURVEY_31)
    assert _template_cover_pt(prs) == 92, "fixture: this deck's cover is 92pt"

    theme = apply_observed_style(extract_theme(SURVEY_31), observe_deck_style(prs))
    theme["cover_pt"] = _template_cover_pt(prs)
    idx = synthesize_title(
        prs, theme, infer_content_bounds(extract_template(SURVEY_31)),
        {"title": "Платформа Поток", "subtitle": "Итоги внедрения за квартал"})

    boxes = [s for s in prs.slides[idx].shapes
             if s.has_text_frame and s.text_frame.text.strip()]
    title = next(s for s in boxes if "Платформа" in s.text_frame.text)
    subtitle = next(s for s in boxes if "Итоги" in s.text_frame.text)
    assert int(subtitle.top) >= int(title.top + title.height), (
        f"subtitle at {Emu(subtitle.top).inches:.2f}in runs into a title ending at "
        f"{Emu(title.top + title.height).inches:.2f}in")


@requires(SURVEY_31)
def test_content_starts_past_a_decor_strip_running_down_the_edge():
    """_centered_top steps around canvas decor VERTICALLY, which cannot help
    when the art is a column. The survey-31 canvas keeps seven small blobs at
    x 0.45-1.16in spanning y 1.96-6.71 — 70% of the band's height once merged —
    and the KPI text started at x 0.45, so the render showed a blob sitting on
    «-40%» and on its caption. There is no clear horizontal band to move into,
    only a narrower one to start from.

    Coverage is a property of the STRIP: each blob covers about 10% of the band
    alone. And the edge test needs a tolerance, not equality — these blobs
    alternate between 0.4507in and 0.4537in, and a strict "starts at or before
    the band's left edge" dropped four of the seven, taking coverage to 31% and
    silently disabling the check.

    The survey blobs no longer reach this rule: iter58 established they are
    bullet MARKERS, not artwork, and drops them with the text they marked — the
    better answer than starting content to their right. Measured then: those four
    canvases (survey-31 24/26, survey-69 62/64) were the rule's only live cases
    in the whole corpus. The rule itself stands for a real side-decor column, so
    it is exercised here on a constructed one rather than deleted: a template
    that pins art down one edge is an obvious thing to meet next, and this keeps
    the calibration (merged span, edge tolerance) from being lost."""
    import base64
    import io

    from pptx.util import Emu, Inches

    from generator.synthesizer import _clip_to_side_decor

    # A 1x1 PNG is enough — the rule reads geometry, never pixels.
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(Inches(13.33)), Emu(Inches(7.5))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # An art column down the left edge: seven blocks whose left edges alternate
    # by 0.003in, exactly as the survey blobs did — the case that showed a strict
    # edge equality silently disables the rule.
    art = [slide.shapes.add_picture(io.BytesIO(png),
                                    Emu(Inches(0.45 + (i % 2) * 0.003)),
                                    Emu(Inches(1.0 + i * 0.8)),
                                    Emu(Inches(0.7)), Emu(Inches(0.6)))
           for i in range(7)]
    band = {"left": Emu(Inches(0.45)), "top": Emu(Inches(1.0)),
            "right": Emu(Inches(12.0)), "bottom": Emu(Inches(6.6))}

    clipped = _clip_to_side_decor(band, slide, prs.slide_width, prs.slide_height)
    strip_right = max(int(s.left + s.width) for s in art)
    assert int(clipped["left"]) >= strip_right, (
        f"content starts at {Emu(clipped['left']).inches:.2f}in, inside a decor strip "
        f"ending at {Emu(strip_right).inches:.2f}in")
    assert int(clipped["right"]) == int(band["right"]), "the clear side must not move"


@requires(TJ)
def test_canvas_band_widens_to_the_template_only_into_clear_space():
    """A canvas's own text extent is a good guide until it is simply narrower
    than the layout. The T-Zh study canvas ends its text at 7.59in on a card
    running to 9.6in, so a synthesized comparison got 3.50in columns and the
    render showed «Автоматический сбор по / расписанию» wrapping while two
    inches of card sat empty. Same argument iter32 makes for a band that is too
    SHORT, applied to width.

    Blind widening would be wrong — this very template has canvases that are
    narrow because a photo fills the rest — so each side moves only into space
    no picture stands in. Both cases live on the same template on purpose."""
    from pptx.util import Emu

    from generator.deck_style import apply_observed_style, observe_deck_style
    from generator.layout_bounds import infer_content_bounds
    from generator.synthesizer import _prepare_blank_slide, _resolve_bounds
    from template_parser.parser import extract_template, extract_theme

    bounds_in = infer_content_bounds(extract_template(TJ))

    prs = Presentation(TJ)
    theme = apply_observed_style(extract_theme(TJ), observe_deck_style(prs))
    template = _resolve_bounds(prs, bounds_in)
    _, _, _, widened = _prepare_blank_slide(prs, theme, bounds_in, canvas_idx=8)
    assert int(widened["right"]) == int(template["right"]), (
        f"clear canvas kept a narrow band: {Emu(widened['right']).inches:.2f}in "
        f"against the template's {Emu(template['right']).inches:.2f}in")

    prs2 = Presentation(TJ)
    theme2 = apply_observed_style(extract_theme(TJ), observe_deck_style(prs2))
    _, _, _, kept = _prepare_blank_slide(prs2, theme2, bounds_in, canvas_idx=5)
    art = [s for s in prs2.slides[5].shapes if "PICTURE" in str(s.shape_type)
           and s.left is not None and s.width]
    assert art, "fixture: canvas 5 is supposed to carry artwork"
    assert int(kept["right"]) < int(template["right"]), (
        f"content was widened into artwork: right={Emu(kept['right']).inches:.2f}in")
