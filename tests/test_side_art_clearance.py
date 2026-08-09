"""Text must not be drawn over a side art panel.

Seen on the deck's FIRST slide: the universal template's cover box runs to
9.44in while its art panel — four pink balloons — starts at 5.91in. The
designer's own «Название / презентации» is short and breaks by hand, so it never
reaches the art; a generated title of ordinary length ran straight through the
balloons and «управленческой» was unreadable.

Only OUR text triggers the narrowing, so a template whose own text stays clear
is untouched, and a full-bleed background (left edge 0, the T-Zh cards and the
survey backdrops) is not a panel at all — every text box on such a slide sits on
it by design.
"""

from conftest import TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu

from generator.generator import _pick_title_shape, generate

ART_LEFT_IN = 5.91  # the balloons on slide 0 of the universal template
LONG = "Платформа управленческой отчётности «Поток»"
SHORT = "Поток"


def _cover(tmp_path, title, name):
    out = str(tmp_path / f"{name}.pptx")
    generate(TJ_UNIVERSAL, [({"type": "title", "title": title, "subtitle": "Итоги квартала"}, 0)], out)
    slide = list(Presentation(out).slides)[0]
    return _pick_title_shape(slide, set())


@requires(TJ_UNIVERSAL)
def test_a_long_title_is_kept_out_of_the_art_panel(tmp_path):
    source = Presentation(TJ_UNIVERSAL)
    art = next(s for s in list(source.slides)[0].shapes if "PICTURE" in str(s.shape_type))
    assert abs(Emu(art.left).inches - ART_LEFT_IN) < 0.05, "fixture: the art moved"

    title = _cover(tmp_path, LONG, "long")
    assert int(title.left + title.width) <= int(art.left), (
        f"the title box runs to {Emu(title.left + title.width).inches:.2f}in, into art "
        f"starting at {Emu(art.left).inches:.2f}in")


@requires(TJ_UNIVERSAL)
def test_a_short_title_keeps_the_designers_box(tmp_path):
    """The narrowing must be a response to our text, not a blanket rule: a title
    that never reaches the panel leaves the template's own box alone."""
    source = Presentation(TJ_UNIVERSAL)
    own = _pick_title_shape(list(source.slides)[0], set())
    title = _cover(tmp_path, SHORT, "short")
    assert int(title.width) == int(own.width), (
        f"box narrowed from {Emu(own.width).inches:.2f}in to "
        f"{Emu(title.width).inches:.2f}in for a title that fits")


@requires(TJ_UNIVERSAL)
def test_a_full_bleed_background_is_not_a_side_panel():
    """A picture starting at the slide's left edge is the background every box on
    the slide sits on. Treating it as an obstacle would leave nowhere to put
    anything — the T-Zh study cards and the survey backdrops are exactly that."""
    from fonts.metrics import FontResolver
    from qa.geometry import _side_art_left
    from template_parser.parser import extract_theme

    prs = Presentation(TJ_UNIVERSAL)
    FontResolver(TJ_UNIVERSAL, extract_theme(TJ_UNIVERSAL))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.left is None or not shape.width:
                continue
            art_left = _side_art_left(slide, shape, prs.slide_width, prs.slide_height)
            assert art_left is None or art_left > int(shape.left), (
                "a panel must start to the right of the box it constrains")
