"""_align_left_edges does not slide a box onto a picture the indent was clearing.

VK Education's cover sets the speaker line at 2.94in, right of a round
speaker photo at 0.72–2.49in; the title starts at 0.68in. _fill_title aligned
the subtitle to the title and printed it under the photo — on every VK
Education deck from iter103 to iter112, five of five.

A cap on the shift was tried long ago and reverted (a body returned to its own
column hung in an empty slide — docs/LESSONS.md), so the rule is about the cause: an
indent that clears a picture is design. Without such a picture the alignment
works as before — the second synthetic case pins that.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from generator.generator import _align_left_edges, generate

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f6e0000000049454e44ae426082")


def _rect(s):
    return (s.left, s.top, s.left + s.width, s.top + s.height)


def _overlap(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _slide(tmp_path, with_picture):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(8), Inches(1))
    body = slide.shapes.add_textbox(Inches(3.0), Inches(5.5), Inches(5), Inches(0.7))
    if with_picture:
        png = tmp_path / "p.png"
        png.write_bytes(PNG_1PX)
        slide.shapes.add_picture(str(png), Inches(0.6), Inches(5.0), Inches(1.8), Inches(1.8))
    return title, body


def test_an_indent_that_clears_a_picture_is_kept(tmp_path):
    title, body = _slide(tmp_path, with_picture=True)
    _align_left_edges(title, body)
    assert body.left == Inches(3.0)


def test_without_a_picture_the_boxes_are_still_aligned(tmp_path):
    title, body = _slide(tmp_path, with_picture=False)
    _align_left_edges(title, body)
    assert body.left == title.left


def _speaker_cover(prs):
    """A slide with exactly two text boxes and a picture beside the lower one."""
    for i, slide in enumerate(prs.slides):
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if len(texts) != 2 or not pics:
            continue
        low = max(texts, key=lambda s: s.top)
        high = min(texts, key=lambda s: s.top)
        if low.left - high.left > Inches(1) and any(
                p.left < low.left and p.top < low.top + low.height and low.top < p.top + p.height
                for p in pics):
            return i
    return None


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_the_cover_subtitle_stays_clear_of_the_speaker_photo(tmp_path):
    idx = _speaker_cover(Presentation(VK_EDU))
    assert idx is not None, "fixture: VK Education's speaker cover not found"
    subtitle = "Объединяем данные. Показываем точки роста вашей команды продаж"
    block = {"type": "title", "title": "Единая платформа для аналитики продаж", "subtitle": subtitle}
    out = str(tmp_path / "deck.pptx")
    generate(VK_EDU, [(block, idx)], out)
    slide = Presentation(out).slides[0]
    box = next(s for s in slide.shapes if s.has_text_frame and subtitle[:20] in s.text_frame.text)
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    under = [p.name for p in pics if _overlap(_rect(box), _rect(p))]
    assert not under, f"subtitle box lies over picture(s) {under}"
