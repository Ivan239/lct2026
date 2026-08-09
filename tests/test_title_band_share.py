"""A heading introduces the content; past a share of the band it replaces it.

Synthesized slides set the title at the size measured from the template, which
is right for the template's OWN titles and wrong for a long one: on the T-Zh
mono deck a 58-character heading came out four lines tall — a 3.51in box on a
5.62in slide, 62% — and squeezed the two comparison columns it introduces into
a strip at the bottom. The templates' own titles run 8-23% of the slide
(median across the corpus).

Found by giving the offline probe a third brief whose TITLES are long: the two
existing ones vary the body text and hold titles short, so anything depending on
a title's length was invisible to it — the last two defects had to be found on
decks assembled by hand.
"""

from conftest import TJ_MONO, requires

from pptx import Presentation
from pptx.util import Emu

from fonts.metrics import FontResolver
from generator.deck_style import apply_observed_style, observe_deck_style
from generator.generator import _template_title_pt
from generator.layout_bounds import infer_content_bounds
from generator.synthesizer import synthesize_two_column_comparison
from template_parser.parser import extract_template, extract_theme

LONG_TITLE = "Как изменился процесс подготовки управленческой отчётности"
SHORT_TITLE = "Как изменился процесс"
COLUMNS = {"left_heading": "Было", "left_points": ["Ручной сбор показателей"],
           "right_heading": "Стало", "right_points": ["Автоматический сбор"]}


def _synthesized(title):
    prs = Presentation(TJ_MONO)
    theme = apply_observed_style(extract_theme(TJ_MONO), observe_deck_style(prs))
    theme["title_pt"] = _template_title_pt(prs)
    idx = synthesize_two_column_comparison(
        prs, theme, infer_content_bounds(extract_template(TJ_MONO)),
        dict(COLUMNS, title=title), resolver=FontResolver(TJ_MONO, extract_theme(TJ_MONO)))
    slide = prs.slides[idx]
    title_box = next(s for s in slide.shapes
                     if s.has_text_frame and s.text_frame.text.strip() == title)
    return title_box, Emu(prs.slide_height).inches


@requires(TJ_MONO)
def test_a_long_heading_does_not_eat_the_slide():
    box, slide_h = _synthesized(LONG_TITLE)
    share = Emu(box.height).inches / slide_h
    assert share <= 0.45, f"the heading takes {share:.0%} of the slide"


@requires(TJ_MONO)
def test_a_short_heading_keeps_the_templates_own_size():
    """The cap may only act on a heading that overruns: a normal one is still
    set at the size measured from the template."""
    prs = Presentation(TJ_MONO)
    template_pt = _template_title_pt(prs)
    box, _ = _synthesized(SHORT_TITLE)
    size = box.text_frame.paragraphs[0].runs[0].font.size.pt
    assert size == template_pt, f"{size}pt against the template's {template_pt}pt"


@requires(TJ_MONO)
def test_the_columns_still_start_below_the_heading():
    box, _ = _synthesized(LONG_TITLE)
    prs = Presentation(TJ_MONO)
    theme = apply_observed_style(extract_theme(TJ_MONO), observe_deck_style(prs))
    theme["title_pt"] = _template_title_pt(prs)
    idx = synthesize_two_column_comparison(
        prs, theme, infer_content_bounds(extract_template(TJ_MONO)),
        dict(COLUMNS, title=LONG_TITLE), resolver=FontResolver(TJ_MONO, extract_theme(TJ_MONO)))
    slide = prs.slides[idx]
    heading = next(s for s in slide.shapes
                   if s.has_text_frame and s.text_frame.text.strip() == LONG_TITLE)
    columns = [s for s in slide.shapes
               if s.has_text_frame and "Было" in s.text_frame.text]
    assert columns, "the comparison columns vanished"
    assert int(columns[0].top) >= int(heading.top + heading.height)
