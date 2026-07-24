"""Cross-clone consistency: two slides filled from the same template slide
must ship with the same body font size, even when only one of them carried
text long enough to force a shrink."""

from conftest import SURVEY_31, requires

from pptx import Presentation

from generator.generator import generate

LONG = [
    "Очень длинный пункт списка, который занимает несколько строк и заставляет подгонку уменьшить шрифт " * 2,
    "Второй столь же длинный пункт списка про автоматизацию корпоративных презентаций и дизайн-систем " * 2,
    "Третий длинный пункт про кластеризацию слайдов на архетипы и переиспользование классификаций " * 2,
    "Четвёртый длинный пункт про экономику токенов и межшаблонный кэш фингерпринтов геометрии " * 2,
]
SHORT = ["Разбор шаблона", "Дизайн-система", "Сборка презентации"]


def _body_size(slide):
    sizes = []
    for shape in slide.shapes:
        for p in shape.text_frame.paragraphs if shape.has_text_frame else []:
            for run in p.runs:
                if run.font.size and run.text.strip():
                    sizes.append(run.font.size.pt)
    return min(sizes) if sizes else None


@requires(SURVEY_31)
def test_clones_share_font_size(tmp_path):
    plan = [
        ({"type": "bullet_list", "title": "Длинные", "bullets": LONG}, 28),
        ({"type": "bullet_list", "title": "Короткие", "bullets": SHORT}, 28),
    ]
    out = str(tmp_path / "harmony.pptx")
    generate(SURVEY_31, plan, out)
    prs = Presentation(out)

    size_long = _body_size(prs.slides[0])
    size_short = _body_size(prs.slides[1])
    assert size_long is not None and size_short is not None
    # The long-text slide had to shrink; the short one must match it instead
    # of towering over its sibling at the template's default size.
    assert size_long < 18.0
    assert size_short == size_long


@requires(SURVEY_31)
def test_icons_follow_harmonized_size(tmp_path):
    """harmonize_clone_font_sizes fires AFTER the fillers positioned the
    bullet-marker icons — the icons must be re-laid-out for the size the deck
    ships with, or they keep the pitch of the pre-harmonize font and drift
    below their lines (the artifact _realign_icons_after_resize exists for)."""
    from pptx.util import Emu, Inches

    from fonts.metrics import FontResolver
    from generator.generator import (
        _find_bullet_icons,
        _metrics_for_run,
        _paragraph_line_height_pt,
        _pick_body_shape,
        _reference_run,
    )
    from template_parser.parser import extract_theme

    plan = [
        ({"type": "bullet_list", "title": "Длинные", "bullets": LONG}, 28),
        ({"type": "bullet_list", "title": "Короткие", "bullets": SHORT}, 28),
    ]
    out = str(tmp_path / "harmony_icons.pptx")
    generate(SURVEY_31, plan, out)
    prs = Presentation(out)

    short = prs.slides[1]
    body = _pick_body_shape(short, set())
    icons = _find_bullet_icons(short, body)
    assert len(icons) == len(SHORT)  # extras deleted, one marker per bullet

    size = _body_size(short)
    resolver = FontResolver(out, extract_theme(out))
    metrics = _metrics_for_run(_reference_run(body), resolver)
    paragraphs = list(body.text_frame.paragraphs)
    # Single-line bullets with zero para spacing: successive icon tops must
    # step by exactly the line pitch of the FINAL (harmonized) font size.
    for i in range(len(icons) - 1):
        expected = Emu(int(Inches(_paragraph_line_height_pt(paragraphs[i], size, metrics=metrics) / 72)))
        step = icons[i + 1].top - icons[i].top
        assert abs(step - expected) <= Emu(int(Inches(0.005))), (
            f"icon {i + 1}->{i + 2} step {Emu(step).inches:.3f}in "
            f"!= pitch {expected.inches:.3f}in of harmonized {size}pt"
        )
