"""Removing someone else's chart must not leave a hole in the slide.

The chart itself has to go — its data goes stale the moment new text is written
around it — but the designer put the summary paragraph UNDER it. On slide 6 of
the survey template the chart runs 2.05-4.37in and the body sits at 4.56, so a
generated slide came out as a heading, a 2.70in void — a third of the slide —
and three lines stranded at the bottom. Verified on the render before and after.

Only shapes wholly below the picture move, and never above a gap under whatever
ends higher up, so nothing slides into another box.
"""

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from generator.generator import generate

CHART_SLIDE = 5
BLOCK = {"type": "bullet_list", "title": "Что мешало собирать отчётность",
         "bullets": ["Ручной сбор показателей", "Разные форматы выгрузок",
                     "Согласование до двух недель"]}


@requires(SURVEY_31)
def test_the_body_moves_up_into_the_space_the_chart_left(tmp_path):
    source = Presentation(SURVEY_31)
    template_slide = list(source.slides)[CHART_SLIDE]
    chart = next(s for s in template_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    body_before = max(s.top for s in template_slide.shapes
                      if s.has_text_frame and s.top is not None)

    out = str(tmp_path / "chart.pptx")
    generate(SURVEY_31, [(BLOCK, CHART_SLIDE)], out)
    slide = list(Presentation(out).slides)[0]

    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE], \
        "the stale chart must still be removed"

    body = next(s for s in slide.shapes
                if s.has_text_frame and "Ручной сбор" in s.text_frame.text)
    assert int(body.top) < int(body_before), (
        f"body stayed at {Emu(body.top).inches:.2f}in, where the chart's summary was")
    # …and it must not have climbed into the box above it.
    above = [s for s in slide.shapes
             if s.top is not None and s.height and s.shape_id != body.shape_id
             and int(s.top) < int(body.top)]
    for other in above:
        assert int(other.top + other.height) <= int(body.top), "the body overlaps a box above it"


@requires(SURVEY_31)
def test_a_slide_without_a_removed_chart_is_untouched(tmp_path):
    """The lift must be a consequence of the removal, not a general reflow."""
    source = Presentation(SURVEY_31)
    # Text boxes only: marker icons are repositioned by the bullet filler by
    # design (iter18), so including them would test that instead.
    tops = {s.shape_id: int(s.top) for s in list(source.slides)[1].shapes
            if s.top is not None and s.has_text_frame}

    out = str(tmp_path / "plain.pptx")
    generate(SURVEY_31, [(BLOCK, 1)], out)
    slide = list(Presentation(out).slides)[0]
    for shape in slide.shapes:
        if shape.shape_id in tops and shape.top is not None:
            assert int(shape.top) == tops[shape.shape_id]
