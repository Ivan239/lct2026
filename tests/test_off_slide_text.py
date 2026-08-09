"""Text running off the slide is scored by SEVERITY, and measured on the text
rather than on its box.

Both halves were wrong, and each one alone hid the other. The survey cover ran
its 90pt title to 8.75in on a 7.5in slide with the last line cut off by the
bottom edge — 9.1 detected it («1 боксов выходят за границы слайда») and scored
it 5/5, because one bad box among fifteen is a 6% rate and a rate rounds to
perfect. The deck reported 98.8 for three iterations while its cover was sliced
in half.

Scoring it by severity alone would then have punished the designers: a box is
routinely taller than its text (survey-69 slide 6 runs a box to 7.83in while
its paragraphs end at 5.11in), so the box-based measure had a false positive on
a pristine template. Measuring the text removes it.
"""

from conftest import SURVEY_31, SURVEY_69, TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from evaluation.deterministic import evaluate


def _text_box(slide, text, size_pt, left_in, top_in, width_in, height_in):
    box = slide.shapes.add_textbox(Emu(int(Inches(left_in))), Emu(int(Inches(top_in))),
                                   Emu(int(Inches(width_in))), Emu(int(Inches(height_in))))
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(size_pt)
    run.font.name = "Arial"
    return box


def _deck_with_cover(tmp_path, template, size_pt, top_in, height_in):
    """A cover set at `size_pt` where the synthesizer used to put it, followed by
    five ordinary slides.

    The filler is the whole point: the rate this criterion used to score by is
    one bad box over EVERY content box in the deck, so a fixture holding only
    the broken cover scores 1/1 = catastrophic on the old code too and proves
    nothing. A real deck dilutes it to about 6%, which rounded to a clean 5."""
    prs = Presentation(template)
    for sld_id in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(sld_id)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    _text_box(cover, "Платформа управленческой отчётности «Поток»", size_pt,
              1.33, top_in, 10.67, height_in)

    for i in range(5):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _text_box(slide, f"Раздел {i + 1}", 32, 1.0, 0.8, 10.0, 0.9)
        _text_box(slide, "Ручной сбор показателей из семи систем", 18, 1.0, 2.2, 10.0, 0.6)
        _text_box(slide, "Разные форматы выгрузок у подразделений", 18, 1.0, 3.0, 10.0, 0.6)

    out = str(tmp_path / f"cover_{size_pt}.pptx")
    prs.save(out)
    return out


@requires(SURVEY_31)
def test_a_title_running_off_the_slide_is_not_scored_perfect(tmp_path):
    broken = _deck_with_cover(tmp_path, SURVEY_31, 90, 2.85, 5.90)
    assert Presentation(broken).slide_height == Emu(int(Inches(7.5)))

    result = evaluate(broken)["9.1"]
    assert result["score"] <= 2, (
        f"a cover cut off by the slide edge scored {result['score']}: {result['detail']}")


@requires(SURVEY_31)
def test_the_same_title_that_fits_scores_clean(tmp_path):
    """The fitted size iter59 produces — the check must not simply dislike big
    covers."""
    fitted = _deck_with_cover(tmp_path, SURVEY_31, 64, 2.85, 3.22)
    assert evaluate(fitted)["9.1"]["score"] == 5


@requires(SURVEY_69)
def test_a_slack_box_on_a_designers_slide_is_not_a_defect():
    """survey-69 slide 6: box to 7.83in on a 7.5in slide, text ending at 5.11in.

    Asserted on the primitive, not on the score: with severity scoring this one
    box would drag a pristine 69-slide template to 2/5, and that is precisely
    what measuring the BOX instead of the text would have done."""
    from fonts.metrics import FontResolver
    from evaluation.deterministic import _text_past_edge
    from template_parser.parser import extract_theme

    prs = Presentation(SURVEY_69)
    resolver = FontResolver(SURVEY_69, extract_theme(SURVEY_69))
    slide = list(prs.slides)[5]
    slack = [s for s in slide.shapes
             if s.has_text_frame and s.text_frame.text.strip() and s.top is not None
             and s.height and int(s.top + s.height) > int(prs.slide_height)]
    assert slack, "fixture changed: this slide is supposed to carry a slack box"
    for box in slack:
        assert _text_past_edge(box, prs.slide_width, prs.slide_height,
                               resolver.metrics_for) == 0

    assert evaluate(SURVEY_69)["9.1"]["score"] == 5


@requires(TJ_UNIVERSAL)
def test_pristine_template_stays_clean():
    assert evaluate(TJ_UNIVERSAL)["9.1"]["score"] == 5
