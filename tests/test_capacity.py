"""get_capacity regressions. Two past bugs live here:
1. Title-first shape picking stole the real body on slides whose caption box
   has no explicit run font size -> capacity 1 instead of the real count.
2. Paragraph count was reported without capping by physical icon markers ->
   plans asked for more items than there are icons to point at them."""

from conftest import SURVEY_31, TJ_MONO, TJ_TEMPLATE, TJ_UNIVERSAL, requires

from pptx import Presentation

from generator.generator import get_capacity


@requires(SURVEY_31)
def test_bullet_capacity_capped_by_icons():
    prs = Presentation(SURVEY_31)
    # slide 28: 8 paragraphs in the body but only 4 icon markers -> 4.
    assert get_capacity(prs.slides[28], "bullet_list") == 4
    # slide 26: 8 paragraphs, 8 icons -> 8.
    assert get_capacity(prs.slides[26], "bullet_list") == 8
    # slide 24: 7 paragraphs, 7 icons -> 7.
    assert get_capacity(prs.slides[24], "bullet_list") == 7


@requires(SURVEY_31)
def test_stats_single_box_capacity_from_icons():
    prs = Presentation(SURVEY_31)
    # Single-text-block stats slides: capacity = number of icon markers.
    assert get_capacity(prs.slides[1], "stats_kpi") == 4
    assert get_capacity(prs.slides[3], "stats_kpi") == 2


@requires(TJ_MONO)
@requires(TJ_TEMPLATE)
def test_item_char_budget_is_learned_from_the_templates_own_sample():
    """get_capacity answers HOW MANY; this answers HOW LONG. The budget comes
    from the designer's own sample text rather than from chars-per-line, because
    the size those slots inherit is not in the run — it resolves through the
    placeholder/master chain, and the theme routinely lies about the deck's real
    look. Where both can be computed they agree: metrics give ~17 chars for the
    T-Zh mono slot at its rendered size, its sample is 15."""
    from generator.generator import get_item_char_budget

    mono = Presentation(TJ_MONO)
    tight = get_item_char_budget(list(mono.slides)[2], "bullet_list")
    assert tight is not None and 12 <= tight <= 24, f"one-line slot budget is {tight}"

    study = Presentation(TJ_TEMPLATE)
    roomy = get_item_char_budget(list(study.slides)[9], "bullet_list")
    assert roomy is not None and roomy > tight * 2, (
        f"a prose box must not get the same budget as a one-line slot: {roomy} vs {tight}")

    assert get_item_char_budget(list(mono.slides)[2], "stats_kpi") is None


@requires(TJ_UNIVERSAL)
def test_grid_row_inside_the_chrome_band_still_counts_as_slots():
    """The chrome test is "short box hugging a slide edge band", and the T-Zh
    universal list is 6 items in a 3x2 grid whose BOTTOM ROW ends 0.32in inside
    that band. Three of the six cells therefore read as furniture, capacity came
    back 3, and half the designed layout could never be used — on a template
    whose whole point is that layout.

    A box in a COMPLETE grid alongside non-chrome cells is content whatever band
    it sits in: page numbers and header strips do not form grids with content
    boxes. Measured over all five real templates, admitting chrome into slot
    detection changes exactly this one slide and creates no all-chrome grid."""
    from generator.generator import _find_slot_boxes, get_capacity
    from generator.slide_kit import is_chrome_shape

    prs = Presentation(TJ_UNIVERSAL)
    slide = list(prs.slides)[2]
    assert get_capacity(slide, "bullet_list") == 6

    slots = _find_slot_boxes(slide, set())
    assert any(is_chrome_shape(s, prs.slide_height) for s in slots), (
        "the bottom row is what this test is about — if nothing here is "
        "chrome-classified the fixture changed and the test proves nothing")
    assert not all(is_chrome_shape(s, prs.slide_height) for s in slots)


@requires(TJ_TEMPLATE)
@requires(TJ_UNIVERSAL)
def test_a_one_big_number_stats_slide_has_capacity_one():
    """The T-Zh study template's stats slide is «20 227 000» set at 82pt under a
    26pt title — a one-big-number design. It reaches get_capacity's
    single-text-block fallback, which answered "no fixed capacity", so three
    stat pairs went in and the render showed three same-size lines reading as
    sentences. The design was gone.

    A body run 1.5x the title is a display figure, and such a slide holds one
    figure. Both sizes must be STATED: the universal template's own single-box
    stats slide leaves its runs unsized (Google Slides export), and guessing
    there is the blindness _pick_title_shape already suffers from — it keeps the
    old answer. Its 8-box KPI board never reaches this fallback at all."""
    from generator.generator import get_capacity

    study = list(Presentation(TJ_TEMPLATE).slides)[8]
    assert get_capacity(study, "stats_kpi") == 1

    universal = list(Presentation(TJ_UNIVERSAL).slides)
    assert get_capacity(universal[9], "stats_kpi") == 4, "the KPI board must be untouched"
    # iter87: this slide holds one PAIR — a 92pt figure and a caption — and
    # says so structurally, from two boxes, not from a guess about sizes. It
    # used to answer None only because the topmost box (the figure itself) was
    # claimed as the title, leaving a single box behind; that claim is what put
    # a heading across the artwork, and it is gone.
    assert get_capacity(universal[8], "stats_kpi") == 1


@requires(TJ_TEMPLATE)
def test_a_one_big_number_slide_sets_number_and_label_apart(tmp_path):
    """iter42 gave this slide capacity 1; the format was still wrong. The single
    pair was written as one run — "-40% — времени на подготовку регулярной
    отчётности" — at one size, so the render showed a two-line headline sentence
    where the template shows «20 227 000» alone at 82pt. A KPI needs the figure
    to read as a figure.

    The pair is fitted to the box UP FRONT, because enforce_text_fits shrinks by
    clamping every run to a single size: a pair that arrives too tall comes back
    flattened, and 82pt over 27pt became 24pt over 24pt."""
    from generator.generator import generate

    out = str(tmp_path / "display.pptx")
    generate(TJ_TEMPLATE,
             [({"type": "stats_kpi", "title": "Результаты",
                "stats": [["-40%", "времени на подготовку регулярной отчётности"]]}, 8)],
             out)

    slide = list(Presentation(out).slides)[0]
    figure = next((s for s in slide.shapes
                   if s.has_text_frame and "-40%" in s.text_frame.text), None)
    assert figure is not None, "the stat vanished"

    # Asserted on the OUTCOME, not on the mechanism: iter70 moved the label out
    # of the figure's box into the empty strip the card leaves beneath it, so
    # "two paragraphs of one shape" is no longer the shape of the answer. What
    # must hold either way is that the figure reads as a figure — displayed
    # large, with a distinctly smaller label somewhere below it, neither of them
    # microtype.
    def _paragraph_pt(shape, needle):
        """The size of the paragraph carrying `needle` — per paragraph, because
        the two may still share one box, where a max over the whole shape would
        report the figure's size for both and compare 24pt against 24pt."""
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs)
            if needle in text:
                sizes = [r.font.size.pt for r in para.runs if r.font.size and r.text.strip()]
                if sizes:
                    return max(sizes)
        return None

    label_shape = next(
        (s for s in slide.shapes
         if s.has_text_frame and "времени на подготовку" in s.text_frame.text), None)
    assert label_shape is not None, "the label vanished"

    figure_pt = _paragraph_pt(figure, "-40%")
    label_pt = _paragraph_pt(label_shape, "времени на подготовку")
    assert figure_pt and label_pt
    assert figure_pt > label_pt * 1.5, (
        f"the figure does not read as a figure: {figure_pt}pt over {label_pt}pt")
    assert label_pt >= 9, f"label shrunk into microtext: {label_pt}pt"
    if label_shape.shape_id != figure.shape_id:
        assert int(label_shape.top) >= int(figure.top), "the label must sit below the figure"
