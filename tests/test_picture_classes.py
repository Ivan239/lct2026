"""The decor-vs-stale-data picture rule (common/pictures.py), verified on the
real corpora it was calibrated on: survey-deck chart screenshots must still be
flagged (their data goes stale next to new text), while the T-Ж template's
full-bleed background art and photos must NOT be — the old size-only rule
excluded 11 of that template's 12 slides and the generator threw the whole
uploaded design away."""

import os

from conftest import SURVEY_31, SURVEY_69, TEMPLATES_DIR, requires

from pptx import Presentation

from common.pictures import has_oversized_picture

TJ = os.path.join(TEMPLATES_DIR, "custom_f496182bb15f42bb.pptx")


@requires(SURVEY_31)
@requires(SURVEY_69)
def test_chart_screenshots_still_flagged():
    for path, chart_slides in ((SURVEY_31, [4, 5, 6, 7, 10, 12, 15, 20]), (SURVEY_69, [9, 20, 28])):
        prs = Presentation(path)
        for idx in chart_slides:
            assert has_oversized_picture(prs.slides[idx], prs.slide_width, prs.slide_height), \
                f"{os.path.basename(path)} slide {idx}: chart slide no longer flagged"


@requires(TJ)
def test_background_art_and_photos_not_flagged():
    prs = Presentation(TJ)
    flagged = [
        i for i in range(len(prs.slides))
        if has_oversized_picture(prs.slides[i], prs.slide_width, prs.slide_height)
    ]
    # Full-bleed folder backgrounds (with text on top) and decorative photos —
    # none of them carry stale data; the template must stay fillable.
    assert flagged == [], f"decor wrongly flagged on slides {flagged}"
