"""Экспорт в .html (ТЗ: три формата). Без LibreOffice: PNG подставляются готовые.

Что обязано держаться: каждый слайд есть и картинкой, и НАСТОЯЩИМ текстом
(включая текст в группах и таблицах — `slide.shapes` их не видит), текст
экранирован, а подмена шрифтов не прячется."""

import base64
import os

from pptx import Presentation
from pptx.util import Inches

from rendering.export import build_html, slide_texts

# 1×1 прозрачный PNG
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII=")


def _deck(tmp_path):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Рост <выручки> & прибыли"
    group = s1.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Текст в группе"
    s2 = prs.slides.add_slide(blank)
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Метрика", "Значение"
    table.cell(1, 0).text, table.cell(1, 1).text = "NPS", "62"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    pngs = []
    for i in (1, 2):
        png = tmp_path / f"deck-{i}.png"
        png.write_bytes(PNG)
        pngs.append(str(png))
    return str(path), pngs


def test_text_is_read_inside_groups_and_tables(tmp_path):
    path, _ = _deck(tmp_path)
    texts = slide_texts(path)
    assert "Текст в группе" in texts[0], texts
    assert "NPS | 62" in texts[1], texts


def test_html_has_every_slide_as_image_and_text(tmp_path):
    path, pngs = _deck(tmp_path)
    out = build_html(path, pngs, str(tmp_path / "deck.html"), title="Дека")
    page = open(out, encoding="utf-8").read()
    assert page.count("data:image/png;base64,") == 2
    assert 'id="slide-1"' in page and 'id="slide-2"' in page
    assert "Текст в группе" in page and "NPS | 62" in page
    assert "Рост &lt;выручки&gt; &amp; прибыли" in page, "текст обязан экранироваться"
    assert "<выручки>" not in page


def test_font_substitution_is_named_not_hidden(tmp_path):
    path, pngs = _deck(tmp_path)
    page = open(build_html(path, pngs, str(tmp_path / "a.html"), substituted=["VK Sans"]),
                encoding="utf-8").read()
    assert "VK Sans" in page and "Arial" in page
    clean = open(build_html(path, pngs, str(tmp_path / "b.html")), encoding="utf-8").read()
    assert "заменены" not in clean


def test_the_fallback_font_is_one_the_host_has(monkeypatch):
    """В Linux-образе Arial нет: подставить в копию отсутствующий шрифт — снова
    отдать выбор LibreOffice, а сообщение «заменены на Arial» было бы неправдой."""
    from rendering import render

    monkeypatch.setattr(render, "_installed_families",
                        lambda: {"liberation sans", "dejavu sans"})
    assert render.fallback_typeface() == "Liberation Sans"
    monkeypatch.setattr(render, "_installed_families", lambda: {"arial", "liberation sans"})
    assert render.fallback_typeface() == "Arial"


def test_the_html_names_the_real_fallback(tmp_path):
    path, pngs = _deck(tmp_path)
    page = open(build_html(path, pngs, str(tmp_path / "c.html"), substituted=["Play"],
                           fallback="Liberation Sans"), encoding="utf-8").read()
    assert "заменены на Liberation Sans" in page


def test_fonts_in_nested_directories_are_seen(tmp_path, monkeypatch):
    """Linux кладёт шрифты по подкаталогам (/usr/share/fonts/truetype/…): плоский
    скан видел там ноль семейств, и в образе сдачи подменялся КАЖДЫЙ шрифт."""
    import glob
    import shutil

    import pytest

    from rendering import render

    candidates = [p for d in ("/System/Library/Fonts/Supplemental", "/Library/Fonts",
                              "/usr/share/fonts")
                  for p in glob.glob(os.path.join(d, "**", "*.ttf"), recursive=True)]
    if not candidates:
        pytest.skip("на хосте нет ни одного .ttf")
    nested = tmp_path / "truetype" / "family"
    nested.mkdir(parents=True)
    shutil.copy(candidates[0], nested / os.path.basename(candidates[0]))

    monkeypatch.setattr(render, "FONT_DIRS", [str(tmp_path)])
    render._installed_families.cache_clear()
    try:
        assert render._installed_families(), "шрифт во вложенном каталоге не найден"
    finally:
        render._installed_families.cache_clear()
