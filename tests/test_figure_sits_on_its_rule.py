"""A display figure stops at the rule the designer drew under it.

VK Tech's board draws a 0.03in hairline at 3.91in and sets «7» at 166pt above
it: a line that tall carries its descent well clear. Our fitted figure sits
lower, and the rule ran through the digits like a strikethrough — three decks
in a row (iter142, iter144, iter147). iter142 already stopped the figure at its
label (4.03in); the rule is higher, so the floor is whichever comes first.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Inches

from generator.generator import generate

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
BOARD = 41
PAIRS = [["+18%", "с наставничеством"], ["-18%", "без наставничества"]]


def _deck(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "stats_kpi", "title": "Цифры", "stats": PAIRS}, BOARD)], out)
    return Presentation(out).slides[0]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_figure_box_stops_at_the_rule(tmp_path):
    slide = _deck(tmp_path)
    figures = [s for s in slide.shapes
               if s.has_text_frame and s.text_frame.text.strip() in {"+18%", "-18%"}]
    rules = [s for s in slide.shapes
             if s.height and int(s.height) <= int(Inches(0.06)) and s.width and s.top is not None
             and not (s.has_text_frame and s.text_frame.text.strip())]
    assert figures and rules, (len(figures), len(rules))
    for figure in figures:
        below = [int(r.top) for r in rules
                 if int(figure.top) < int(r.top)
                 and int(r.left) < int(figure.left + figure.width)
                 and int(r.left + r.width) > int(figure.left)]
        if below:
            assert int(figure.top + figure.height) <= min(below), (
                figure.text_frame.text, int(figure.top + figure.height), min(below))


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_the_label_floor_still_holds(tmp_path):
    """iter142's rule is not lost: where there is no rule, the label is the
    floor."""
    slide = _deck(tmp_path)
    texts = {s.text_frame.text.strip(): s for s in slide.shapes
             if s.has_text_frame and s.text_frame.text.strip()}
    for figure_text, label_text in PAIRS:
        figure, label = texts.get(figure_text), texts.get(label_text)
        assert figure is not None and label is not None, sorted(texts)
        assert figure.top + figure.height <= label.top
