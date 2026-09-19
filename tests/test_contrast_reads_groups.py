"""Contrast and media checks see what sits inside GROUPS.

After iter124 the geometric criteria opened groups; the contrast passes and
the media count still read the top level of slide.shapes. On VK Tech's team
slide (one group of five card columns) the pixel pass measured only the title,
and VK Education's phone mockup — a group of two pictures — did not count as a
picture at all (iter126).

The contrast half needs the slide's render; it uses the template render the
loop caches under output/loop and skips when there is none (tests do not run
LibreOffice).
"""

import glob
import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from evaluation.contrast import evaluate_boxed_contrast
from evaluation.deterministic import deck_media

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
VK_EDUCATION = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
TEAM_SLIDE = 12  # «Слайд-визитка команды»: title + one group of card columns
TEAM_RENDER = sorted(glob.glob(os.path.join(
    os.path.dirname(TEMPLATES_DIR), "loop", "*", "vk_tech", "rendered", "vk_tech-13.png")))
MOCKUP_SLIDE = 31  # phone mockup: a GROUP of frame + screen pictures


def _one_slide(src, keep, out, only_groups=False):
    prs = Presentation(src)
    ids = prs.slides._sldIdLst
    for i, sld_id in reversed(list(enumerate(ids))):
        if i != keep:
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            ids.remove(sld_id)
    if only_groups:  # leave the grouped pictures as the slide's only ones
        for s in list(prs.slides[0].shapes):
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                s._element.getparent().remove(s._element)
    prs.save(out)
    return out


@pytest.mark.skipif(not (os.path.exists(VK_TECH) and TEAM_RENDER),
                    reason="fixture deck or its cached render missing")
def test_pixel_contrast_measures_text_inside_groups(tmp_path):
    deck = _one_slide(VK_TECH, TEAM_SLIDE, str(tmp_path / "team.pptx"))
    res = evaluate_boxed_contrast(deck, TEAM_RENDER[:1])
    # Only the grouped boxes carry a grey («Должность», declared (121,132,146));
    # the title is black, ~18:1. A worst box well under that means the group
    # was measured. (iter126 asserted 1.88 and «low contrast» here — that was
    # the thin-stroke blend, not the grey; see test_contrast_thin_strokes.)
    assert res["per_slide"][0] < 10, res["per_slide"]


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_pictures_inside_groups_count_as_media(tmp_path):
    deck = _one_slide(VK_EDUCATION, MOCKUP_SLIDE, str(tmp_path / "mockup.pptx"), only_groups=True)
    assert deck_media(deck)["substantive_pictures"] == 2
