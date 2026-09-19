"""Figures the designer set at one size stay at one size.

Each figure is fitted to one line of its own box (iter115), and a board's
boxes differ in width: VK Tech's two-figure board sets «7» and «10» both at
166pt in 2.60in and 4.73in boxes, and «27%» / «9%» came out at 82pt and 166pt —
the smaller number shouting over the larger (six VK Tech decks, iter120-132).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import generate

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
BOARD = 41
STATS = [["27%", "Уходят в первые 6 месяцев"], ["9%", "Уходят с наставниками"]]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_figures_of_one_board_share_one_size(tmp_path):
    template_sizes = {r.font.size.pt for s in Presentation(VK_TECH).slides[BOARD].shapes
                      if s.has_text_frame and s.text_frame.text in ("7", "10")
                      for p in s.text_frame.paragraphs for r in p.runs if r.font.size}
    assert len(template_sizes) == 1, "fixture: the designer set both figures at one size"

    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "stats_kpi", "title": "Статистика текучести", "stats": STATS}, BOARD)], out)
    slide = Presentation(out).slides[0]
    sizes = {s.text_frame.text: max(r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs
                                    if r.font.size)
             for s in slide.shapes if s.has_text_frame and s.text_frame.text in ("27%", "9%")}
    assert set(sizes) == {"27%", "9%"}, sizes
    assert len(set(sizes.values())) == 1, sizes
