"""Cards kept as GROUPS are list slots on the native path too.

VK Tech's four-card slide is in the bullet_list family, and each card is a
GROUP (icon + «Заголовок / Подзаголовок / Текст» and more «Текст»). The
native list filler read only the top level of slide.shapes: it saw one text
box — the title's — wrote every bullet into it at 7pt, lost the title and left
28 placeholder strings in the cards (iter120). iter105 fixed the same blindness
on the synthesis path only.

Now a grid of >=3 same-size content groups is a set of slots: one item per
card, the card's other lines (and their marker glyphs) go, the title keeps its
box, and get_capacity reports the card count (the rule that capacity mirrors
the filler).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from evaluation.deterministic import placeholder_hits
from generator.generator import generate, get_capacity
from generator.slide_kit import content_groups

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
BULLETS = ["Аудит данных завершён в июле", "Справочники перенесены полностью",
           "Сделки перенесены без потерь", "Контакты переносятся сейчас"]


def _card_slide(prs):
    return max(range(len(prs.slides)), key=lambda i: len(content_groups(prs.slides[i])))


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_a_list_on_a_card_slide_fills_the_cards(tmp_path):
    prs = Presentation(VK_TECH)
    idx = _card_slide(prs)
    assert len(content_groups(prs.slides[idx])) == 4, "fixture: VK Tech's four-card slide"
    assert get_capacity(prs.slides[idx], "bullet_list") == 4

    out = str(tmp_path / "deck.pptx")
    block = {"type": "bullet_list", "title": "Выполненные этапы миграции", "bullets": BULLETS}
    generate(VK_TECH, [(block, idx)], out)

    slide = Presentation(out).slides[0]
    hits = placeholder_hits(slide)
    assert not hits, f"{len(hits)} placeholder strings left: {sorted(set(hits))}"
    top_level = [s.text_frame.text.strip() for s in slide.shapes
                 if s.has_text_frame and s.text_frame.text.strip()]
    assert "Выполненные этапы миграции" in top_level, f"the title lost its box: {top_level}"
    cards = content_groups(slide)
    card_texts = [[c.text_frame.text for c in g.shapes if getattr(c, "has_text_frame", False)
                   and c.text_frame.text.strip()] for g in cards]
    assert sorted(t for texts in card_texts for t in texts) == sorted(BULLETS), card_texts
    assert all(len(texts) == 1 for texts in card_texts), "one item per card"


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_surplus_cards_are_removed(tmp_path):
    prs = Presentation(VK_TECH)
    idx = _card_slide(prs)
    out = str(tmp_path / "deck.pptx")
    block = {"type": "bullet_list", "title": "Три этапа", "bullets": BULLETS[:3]}
    generate(VK_TECH, [(block, idx)], out)
    slide = Presentation(out).slides[0]
    assert len(content_groups(slide)) == 3
    assert not placeholder_hits(slide)
