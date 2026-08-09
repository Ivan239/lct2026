"""A cloned canvas must not keep the bullet-marker icons of the text it lost.

Seen on a render of the synthesized survey deck: a column of SEVEN little blob
icons marched down the left margin while the three real bullets sat lower with
their own «•» markers. The icons were the canvas's own — decks that draw bullets
as pictures keep them as independent shapes, so stripping the canvas's text
boxes (the clone path's normal first move) leaves them pointing at nothing.

The native fill path has deleted surplus markers since iter18
(_reposition_bullet_icons); the clone path never got the same pass — the same
native/clone divergence CLAUDE.md records for the chrome blanking.

Both directions matter, so both are asserted here: the survey canvas loses its
marker column, and a T-Zh canvas keeps its decor untouched.
"""

from conftest import SURVEY_31, TJ_TEMPLATE, requires

from pptx import Presentation

from generator.layout_bounds import infer_content_bounds
from generator.deck_style import apply_observed_style, observe_deck_style
from generator.synthesizer import synthesize_bullet_list
from template_parser.parser import extract_template, extract_theme

BULLETS = {"title": "Что мешало собирать отчётность",
           "bullets": ["Ручной сбор показателей", "Разные форматы выгрузок",
                       "Согласование до двух недель"]}

# Slide 25 of the survey template: «Top 7», a body of seven lines, each marked
# by a 0.71x0.61in blob picture to its left. Verified on the template's render.
ICON_CANVAS = 24


def _pictures(slide):
    return [s for s in slide.shapes if "PICTURE" in str(s.shape_type)]


@requires(SURVEY_31)
def test_marker_column_does_not_survive_the_canvas_it_marked():
    prs = Presentation(SURVEY_31)
    assert len(_pictures(list(prs.slides)[ICON_CANVAS])) == 7, "fixture changed"

    theme = apply_observed_style(extract_theme(SURVEY_31), observe_deck_style(prs))
    idx = synthesize_bullet_list(prs, theme, infer_content_bounds(extract_template(SURVEY_31)),
                                 BULLETS, canvas_idx=ICON_CANVAS)

    left_over = _pictures(prs.slides[idx])
    assert not left_over, (
        f"{len(left_over)} orphaned bullet markers survived onto the synthesized slide")


@requires(TJ_TEMPLATE)
def test_canvas_decor_is_not_mistaken_for_a_marker_column():
    """The rule must stay blind on decks that mark bullets with real characters:
    every picture on those canvases is artwork, and losing it would trade one
    defect for a worse one."""
    prs = Presentation(TJ_TEMPLATE)
    theme = apply_observed_style(extract_theme(TJ_TEMPLATE), observe_deck_style(prs))
    bounds = infer_content_bounds(extract_template(TJ_TEMPLATE))

    for canvas_idx, before in [(i, len(_pictures(s))) for i, s in enumerate(prs.slides)]:
        if not before:
            continue
        probe = Presentation(TJ_TEMPLATE)
        idx = synthesize_bullet_list(probe, theme, bounds, BULLETS, canvas_idx=canvas_idx)
        assert len(_pictures(probe.slides[idx])) == before, (
            f"canvas {canvas_idx} lost decor: {before} pictures -> "
            f"{len(_pictures(probe.slides[idx]))}")
