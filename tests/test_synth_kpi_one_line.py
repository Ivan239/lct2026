"""A synthesized KPI board sets its figures on one line each, at one size.

The synthesizer drew every figure at a flat 44pt in a column of width /
count. Six figures make 1.5in columns: «≤40 мин.» and «18,4 млн ₽» wrapped,
and the second line printed over the caption under it (iter122, survey-31).
Native boards already keep a figure on one line (iter115); the synthesized
path now does too — one size for the whole board, the largest at which every
figure fits its column.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Emu

from common.synthesis import SYNTHESIZE
from fonts.metrics import FontResolver
from generator.generator import generate
from generator.text_fit import estimate_wrapped_lines, horizontal_margins_in
from template_parser.parser import extract_theme

SURVEY = os.path.join(TEMPLATES_DIR, "custom_30e96c06e2d47ec3.pptx")
STATS = [["4 из 6", "модулей перенесено"], ["86%", "данных клиентов перенесено"],
         ["+25%", "ускорение выпуска"], ["≤40 мин.", "простой при переключении"],
         ["18,4 млн ₽", "план бюджета"], ["1", "высокий риск"]]


@pytest.mark.skipif(not os.path.exists(SURVEY), reason=f"fixture deck missing: {SURVEY}")
def test_every_synthesized_figure_is_one_line_at_one_size(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(SURVEY, [({"type": "stats_kpi", "title": "Статус этапов", "stats": STATS}, SYNTHESIZE)], out)
    slide = Presentation(out).slides[0]
    figures = {n for n, _ in STATS}
    boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text in figures]
    assert len(boxes) == len(STATS), [s.text_frame.text for s in slide.shapes if s.has_text_frame]

    sizes = {r.font.size.pt for b in boxes for p in b.text_frame.paragraphs for r in p.runs}
    assert len(sizes) == 1, f"figures of one board at different sizes: {sorted(sizes)}"

    resolver = FontResolver(out, extract_theme(out))
    for box in boxes:
        run = box.text_frame.paragraphs[0].runs[0]
        metrics = resolver.metrics_for(run.font.name, bold=True)
        lines = estimate_wrapped_lines(box.text_frame.text, Emu(box.width).inches, run.font.size.pt,
                                       metrics=metrics, margins_in=horizontal_margins_in(box))
        assert lines == 1, f"{box.text_frame.text!r} at {run.font.size.pt}pt wraps to {lines} lines"
