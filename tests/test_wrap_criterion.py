"""dop_wrap counts words the renderer must break mid-letter, and dop_text_split
no longer counts a designer's frame as bad content splitting.

Both were built on _overflows — «text taller than its box» — which iter61
measured to be a designer habit rather than a defect: 32 of the 79 content boxes
of the pristine survey-31, 25 marked noAutofit. The verdicts that produced were
the harshest false ones in the harness: dop_text_split gave untouched designer
templates 1/5 and dop_wrap 3/5.

«Некорректный перенос» has a real meaning, and the project has been fighting it
since iter18: a word too long for its box width («Автоматизаци/я» on a 51pt
title, «подключённых клиентов е-/жемесячно» on a KPI caption). Measured on the
corpus that signal fires 0 times across all 217 designer slides, so it is scored
categorically — one such slide is one bad slide.
"""

from conftest import SURVEY_31, SURVEY_69, TJ_MONO, TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from evaluation.deterministic import evaluate

TEMPLATES = (SURVEY_31, SURVEY_69, TJ_UNIVERSAL, TJ_MONO)


def _deck(tmp_path, name, caption_width_in):
    prs = Presentation(TJ_UNIVERSAL)
    for sld_id in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(sld_id)

    def box(slide, text, size_pt, left, top, width, height):
        b = slide.shapes.add_textbox(Emu(int(Inches(left))), Emu(int(Inches(top))),
                                     Emu(int(Inches(width))), Emu(int(Inches(height))))
        b.text_frame.word_wrap = True
        b.text_frame.text = text
        run = b.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(size_pt)
        run.font.name = "Arial"
        return b

    kpi = prs.slides.add_slide(prs.slide_layouts[6])
    box(kpi, "Результаты пилотного внедрения", 32, 0.5, 0.5, 9, 0.8)
    # iter57's defect: a long caption forced into the figure's narrow box at
    # display size, which the renderer split as «е-/жемесячно».
    box(kpi, "подключённых клиентов ежемесячно", 24, 5.0, 2.3, caption_width_in, 0.6)
    for i in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box(slide, f"Раздел {i + 1}", 32, 0.5, 0.5, 9, 0.8)
        box(slide, "Ручной сбор показателей", 18, 0.5, 2.0, 9, 0.5)

    out = str(tmp_path / f"{name}.pptx")
    prs.save(out)
    return out


@requires(TJ_UNIVERSAL)
def test_a_word_that_cannot_fit_its_box_caps_the_score(tmp_path):
    result = evaluate(_deck(tmp_path, "narrow", 2.4))["dop_wrap"]
    assert result["score"] <= 2, f"scored {result['score']}: {result['detail']}"


@requires(TJ_UNIVERSAL)
def test_the_same_caption_in_a_wide_box_is_clean(tmp_path):
    """Only the break is the defect — the identical text at the identical size
    scores clean once its box is wide enough to hold the longest word."""
    assert evaluate(_deck(tmp_path, "wide", 6.0))["dop_wrap"]["score"] == 5


@requires(SURVEY_69)
def test_a_fractional_font_size_is_not_a_break():
    """cap_size_to_longest_word starts from round(size_pt), so a 10.5pt box comes
    back as 10 with nothing shrunk. That artefact alone was the single «designer
    defect» in the corpus — survey-69's footnote, which the render shows on one
    line."""
    from fonts.metrics import FontResolver
    from evaluation.deterministic import _breaks_a_word, _content_shapes, _shape_max_size
    from template_parser.parser import extract_theme

    prs = Presentation(SURVEY_69)
    metrics_for = FontResolver(SURVEY_69, extract_theme(SURVEY_69)).metrics_for
    footnotes = [s for slide in prs.slides for s in _content_shapes(slide)
                 if (_shape_max_size(s) or 0) == 10.5]
    assert footnotes, "fixture changed: the 10.5pt footnote is gone"
    for box in footnotes:
        assert not _breaks_a_word(box, metrics_for)


@requires(SURVEY_31)
def test_pristine_templates_are_clean_on_both_criteria():
    for template in TEMPLATES:
        scores = evaluate(template)
        for cid in ("dop_wrap", "dop_text_split"):
            assert scores[cid]["score"] == 5, f"{template} {cid}: {scores[cid]['detail']}"
