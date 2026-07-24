"""find_sparse_slides: flags near-empty content slides, exempts roles that are
sparse by design, and stays quiet on a normally filled deck."""

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.util import Inches, Pt

from generator.generator import generate
from qa.geometry import find_sparse_slides


def _tiny_text_deck():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2.0), Inches(0.4))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "одинокая строка"
    run.font.size = Pt(14)
    return prs


def test_flags_thin_content_slide():
    prs = _tiny_text_deck()
    issues = find_sparse_slides(prs, {0: "bullet_list"})
    assert len(issues) == 1 and issues[0]["kind"] == "sparse"


def test_design_sparse_roles_exempt():
    prs = _tiny_text_deck()
    for role in ("title", "section_divider", "closing"):
        assert find_sparse_slides(prs, {0: role}) == []


@requires(SURVEY_31)
def test_filled_deck_stays_clean(tmp_path):
    plan = [
        ({"type": "bullet_list", "title": "Проблемы", "bullets": [
            "Ручная адаптация контента", "Перегруженные дизайнеры",
            "Неконсистентный стиль", "Долгое согласование"]}, 28),
        ({"type": "stats_kpi", "title": "Метрики", "stats": [
            ("x10", "ускорение"), ("0 ч", "ручной работы")]}, 1),
    ]
    out = str(tmp_path / "clean.pptx")
    generate(SURVEY_31, plan, out)
    prs = Presentation(out)
    roles = {pos: block["type"] for pos, (block, _) in enumerate(plan)}
    assert find_sparse_slides(prs, roles) == []
