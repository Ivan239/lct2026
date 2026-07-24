"""Sparse-slide mitigation: a body whose text fills well under half its box is
re-anchored to the vertical middle (reads as designed, not as accidentally
empty) — but never when bullet-marker icons are present, because icon rows are
computed from the box top."""

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

from generator.generator import _find_bullet_icons, _pick_body_shape, generate


@requires(SURVEY_31)
def test_underfilled_stats_fallback_centers(tmp_path):
    # Template slide 19: prose text slide (no num/label boxes, no icons), tall
    # body box — two short stat lines fill a small fraction of it.
    plan = [({"type": "stats_kpi", "title": "Результаты",
              "stats": [["x10", "ускорение подготовки"], ["0 ч", "ручного труда"]]}, 19)]
    out = str(tmp_path / "stats.pptx")
    generate(SURVEY_31, plan, out)
    prs = Presentation(out)
    filled = [
        s for s in prs.slides[0].shapes
        if s.has_text_frame and "x10" in s.text_frame.text
    ]
    assert filled and filled[0].text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE

    # KPI emphasis: the number half of each fallback line is its own bold run,
    # the label half stays regular (see generator._embolden_stat_numbers).
    first_para = next(p for p in filled[0].text_frame.paragraphs if p.runs and p.runs[0].text.strip())
    assert first_para.runs[0].text == "x10 — " and first_para.runs[0].font.bold is True
    assert first_para.runs[1].text == "ускорение подготовки" and first_para.runs[1].font.bold is not True


@requires(SURVEY_31)
def test_icon_body_never_centered(tmp_path):
    # Template slide 26: icon-marker bullet list — even badly underfilled it
    # must keep its top anchor, or the icons (laid out from the box top) would
    # point at blank space.
    plan = [({"type": "bullet_list", "title": "Короткий",
              "bullets": ["Один", "Два"]}, 26)]
    out = str(tmp_path / "icons.pptx")
    generate(SURVEY_31, plan, out)
    prs = Presentation(out)
    slide = prs.slides[0]
    body = _pick_body_shape(slide, set())
    assert _find_bullet_icons(slide, body)  # premise: this IS an icon slide
    assert body.text_frame.vertical_anchor in (None, MSO_ANCHOR.TOP)
