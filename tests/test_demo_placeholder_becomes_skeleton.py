"""A demo image sitting in a PICTURE PLACEHOLDER is replaced by the skeleton.

The document itself says «a picture goes here», and what is in the slot is the
template's sample: VK Education's phone mockup showed someone else's feed —
a trip to Iceland — next to our bullets about CRM migration and smart search,
fourteen decks running (iter142-148). image_caption slots have been handled
since iter107; this is the same picture in the same kind of slot on a slide of
any other role.

Measured over the VK templates (iter150): exactly two picture placeholders
still hold an image — this mockup and the closing's 1.18in logo, which the area
floor spares.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from generator.generator import IMAGE_PLACEHOLDER_NAME, generate

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
MOCKUP_LIST = 29   # список с мокапом телефона
CLOSING = 51       # закрытие: логотип 1.18in в таком же слоте
BLOCK = {"type": "bullet_list", "title": "Что уже перенесено",
         "bullets": ["4 из 6 модулей CRM работают в облаке.",
                     "86% данных клиентов перенесено без потерь."]}


def _has_image_placeholder(slide):
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue
        try:
            kind = shape.placeholder_format.type
        except (AttributeError, ValueError):
            continue
        if kind in (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.OBJECT):
            blip = shape._element.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            if blip is not None:
                return True
    return False


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_the_demo_image_is_gone_and_a_skeleton_stands_there(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(VK_EDU, [(BLOCK, MOCKUP_LIST)], out)
    slide = Presentation(out).slides[0]
    assert not _has_image_placeholder(slide), "демо-картинка осталась в слоте"
    frames = [s for s in slide.shapes if s.name == IMAGE_PLACEHOLDER_NAME]
    assert len(frames) == 1, [s.name for s in slide.shapes]
    assert "ИЗОБРАЖЕНИЕ" in frames[0].text_frame.text


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_a_small_slot_picture_is_left_alone(tmp_path):
    """The closing's logo lives in the same kind of slot: below the area floor
    it is chrome, not a demo image."""
    out = str(tmp_path / "deck.pptx")
    generate(VK_EDU, [({"type": "closing", "title": "Спасибо",
                        "subtitle": "Вопросы?"}, CLOSING)], out)
    slide = Presentation(out).slides[0]
    assert _has_image_placeholder(slide), "логотип закрытия не должен был исчезнуть"
    assert not [s for s in slide.shapes if s.name == IMAGE_PLACEHOLDER_NAME]
