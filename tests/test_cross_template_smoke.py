"""One synthetic plan, every real template — the invariants recent iterations
established, checked TOGETHER for the first time.

Each of those fixes was verified on a single template while the network was
down, so nothing so far proves they hold at once, or that a fix made for one
designer's layout does not break another's. No LLM and no rendering: the plan is
hardcoded and every invariant is read straight off the .pptx geometry, so this
stays fast enough to run on every commit."""

import os

from conftest import PRESET_A, SURVEY_31, TJ_MONO, TJ_TEMPLATE, TJ_UNIVERSAL, requires

import pytest
from pptx import Presentation
from pptx.util import Emu

from common.synthesis import SYNTHESIZE
from generator.generator import TOPIC_SLOT_PROMPTS
from generator.slide_kit import is_chrome_shape
from generator.synthesizer import IMAGE_PLACEHOLDER_NAME

PLAN = [
    ({"type": "title", "title": "Платформа «Поток»", "subtitle": "Единая аналитика продаж"}, SYNTHESIZE),
    ({"type": "bullet_list", "title": "Ключевые возможности",
      "bullets": ["Единый дашборд", "Прогноз спроса", "Автоотчёты"]}, SYNTHESIZE),
    ({"type": "stats_kpi", "title": "Результаты",
      "stats": [["+25%", "Рост конверсии"], ["-30 часов", "Экономия времени"]]}, SYNTHESIZE),
    ({"type": "two_column_comparison", "title": "До и после",
      "left_heading": "Разрозненность данных", "left_points": ["Данные в разных системах"],
      "right_heading": "Единая картина", "right_points": ["Всё в одном дашборде"]}, SYNTHESIZE),
    ({"type": "image_caption", "title": "Интерфейс дашборда",
      "image": "Экран ноутбука с графиками продаж"}, SYNTHESIZE),
    ({"type": "closing", "title": "Запустите пилот", "subtitle": "Свяжитесь с менеджером"}, SYNTHESIZE),
]


def _generate(template, tmp_path):
    """Every position gets a CLONED CANVAS on purpose. Without synth_canvas the
    synthesizer draws from scratch, the template's chrome never reaches the deck,
    and the page-number and topic-slot invariants below have nothing to look at —
    the first version of this test passed on all five templates precisely because
    it was checking an empty deck."""
    from generator.generator import generate

    source = Presentation(template)
    canvas_count = len(source.slides._sldIdLst)
    hints = {position: position % canvas_count for position in range(len(PLAN))}
    out = str(tmp_path / (os.path.basename(template).replace(".pptx", "") + "_smoke.pptx"))
    generate(template, PLAN, out, synth_canvas=hints)
    return Presentation(out)


@pytest.mark.parametrize("template", [
    pytest.param(TJ_TEMPLATE, id="tj-teaching", marks=requires(TJ_TEMPLATE)),
    pytest.param(TJ_MONO, id="tj-mono", marks=requires(TJ_MONO)),
    pytest.param(TJ_UNIVERSAL, id="tj-universal", marks=requires(TJ_UNIVERSAL)),
    pytest.param(SURVEY_31, id="survey-31", marks=requires(SURVEY_31)),
    pytest.param(PRESET_A, id="preset-a", marks=requires(PRESET_A)),
])
def test_every_template_holds_the_layout_invariants(template, tmp_path):
    prs = _generate(template, tmp_path)
    width, height = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    assert slides, "generation produced no slides"

    for position, slide in enumerate(slides, start=1):
        for shape in slide.shapes:
            # iter22: a reserved image area that runs off the page renders as a
            # sliver of dashed border at the slide edge.
            if shape.name == IMAGE_PLACEHOLDER_NAME:
                assert shape.top >= 0 and shape.left >= 0, "placeholder starts off-slide"
                assert shape.top + shape.height <= height + Emu(1000), (
                    f"placeholder runs {Emu(shape.top + shape.height - height).inches:.2f}in "
                    f"past the bottom on slide {position}")
                assert shape.left + shape.width <= width + Emu(1000), "placeholder runs off the right"

            if not shape.has_text_frame or not is_chrome_shape(shape, height):
                continue
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            if len(runs) != 1:
                continue
            text = runs[0].text.strip()
            # iter19: a cloned canvas brings the template's own page number.
            # But not every digit box is a page number — the T-Zh universal
            # template footers a brand year, and the first version of the
            # renumbering rewrote "2025" to "0001" on every slide. Page numbers
            # must follow position; anything year-sized must survive verbatim.
            if text.isdigit() and int(text) <= max(len(slides), 99):
                assert int(text) == position, (
                    f"slide {position} shows stale page number {text!r}")

            # iter23: the designer's topic prompt must not ship to the client.
            assert text.lower() not in TOPIC_SLOT_PROMPTS, (
                f"unfilled topic placeholder {text!r} on slide {position}")


@pytest.mark.parametrize("template", [
    pytest.param(TJ_UNIVERSAL, id="tj-universal", marks=requires(TJ_UNIVERSAL)),
])
def test_year_sized_furniture_survives_renumbering(template, tmp_path):
    """The renumbering must not eat furniture that merely looks numeric. The
    T-Zh universal template footers "2025" on ten slides and the first version
    rewrote it to "0001", "0002", ... — invisible to a post-hoc check, because a
    renumbered year is indistinguishable from a real page number. So compare
    against the SOURCE: year-sized digits present in the template must still be
    present in the generated deck."""
    source = Presentation(template)
    src_height = source.slide_height
    years = {
        shape.text_frame.text.strip()
        for slide in source.slides for shape in slide.shapes
        if shape.has_text_frame and is_chrome_shape(shape, src_height)
        and shape.text_frame.text.strip().isdigit()
        and int(shape.text_frame.text.strip()) > 999
    }
    assert years, "premise: this template carries year-sized furniture"

    prs = _generate(template, tmp_path)
    produced = {
        shape.text_frame.text.strip()
        for slide in prs.slides for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    }
    for year in years:
        assert year in produced, f"furniture {year!r} was renumbered away"
