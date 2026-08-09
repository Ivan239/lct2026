"""1.1 counts text printed through other text, not "text taller than its box".

The old measure was _overflows, and it is not a defect measure at all: 32 of the
79 content boxes of the PRISTINE survey-31 overflow, 25 of them marked
noAutofit — the designer keeps the frame smaller than the text on purpose and
the renderer simply draws past it. So 1.1 reported 3/5 on untouched designer
decks while scoring a 4x overflow in a generated one at 5/5.

What a spill can actually do is leave the slide (9.1, iter60) or land on other
text. The threshold for the latter is calibrated, not guessed:

    designers, 217 slides:  worst overlap 0.51 (survey-69 tucks a footnote under
                            a question on purpose); the other four templates 0
    real defect (iter59):   1.00 — the cover subtitle sat entirely inside the
                            title's text band and the two printed through each other
"""

from conftest import SURVEY_31, SURVEY_69, TJ_MONO, TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from evaluation.deterministic import evaluate


def _deck(tmp_path, name, place_second):
    """A cover-like slide with two blocks, plus four ordinary slides so the
    fine-print rate stays realistic."""
    prs = Presentation(SURVEY_31)
    for sld_id in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(sld_id)

    def box(slide, text, size_pt, left, top, width, height):
        b = slide.shapes.add_textbox(Emu(int(Inches(left))), Emu(int(Inches(top))),
                                     Emu(int(Inches(width))), Emu(int(Inches(height))))
        b.text_frame.word_wrap = True
        b.text_frame.text = text
        run = b.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(size_pt)
        run.font.name = "Arial"
        return b

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    box(cover, "Платформа управленческой отчётности «Поток»", 90, 1.33, 2.85, 10.67, 5.9)
    place_second(box, cover)
    for i in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box(slide, f"Раздел {i + 1}", 32, 1.0, 0.8, 10.0, 0.9)
        box(slide, "Ручной сбор показателей из семи систем", 18, 1.0, 2.2, 10.0, 0.6)

    out = str(tmp_path / f"{name}.pptx")
    prs.save(out)
    return out


@requires(SURVEY_31)
def test_a_subtitle_printed_through_the_title_caps_the_score(tmp_path):
    """iter59's defect verbatim: the subtitle clamped to 88% of the slide, which
    is inside a title running to 8.75in."""
    deck = _deck(tmp_path, "collision",
                 lambda box, s: box(s, "Итоги пилотного внедрения за квартал",
                                    20, 2.0, 6.60, 9.33, 0.9))
    result = evaluate(deck)["1.1"]
    assert result["score"] <= 2, f"scored {result['score']}: {result['detail']}"
    assert "поверх другого текста" in result["detail"]


@requires(SURVEY_31)
def test_the_same_subtitle_placed_below_is_clean(tmp_path):
    """Only the collision is the defect — the same two blocks, one under the
    other, must score clean even though the title still overflows its box."""
    deck = _deck(tmp_path, "clean",
                 lambda box, s: box(s, "Итоги пилотного внедрения за квартал",
                                    20, 2.0, 9.20, 9.33, 0.9))
    assert evaluate(deck)["1.1"]["score"] == 5


@requires(SURVEY_69)
def test_a_footnote_tucked_under_a_heading_is_not_a_collision():
    """survey-69's own worst case, 0.51 — deliberate, and it must stay unflagged
    or every real template scores as broken."""
    assert evaluate(SURVEY_69)["1.1"]["score"] == 5


@requires(TJ_UNIVERSAL)
@requires(TJ_MONO)
def test_pristine_templates_report_no_readability_defects():
    for template in (TJ_UNIVERSAL, TJ_MONO, SURVEY_31):
        result = evaluate(template)["1.1"]
        assert result["score"] == 5, f"{template}: {result['detail']}"
