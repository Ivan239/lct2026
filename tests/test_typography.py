"""Plan 10 — baseline typographic hygiene: no mid-word title breaks, code-
enforced text budgets, grid slot detection, font floor."""

import os

from conftest import TEMPLATES_DIR, requires

from pptx import Presentation

from content_parser.two_phase import MAX_BULLET_CHARS, MAX_TITLE_WORDS, _enforce_text_budgets
from generator.generator import _find_slot_boxes, _fit_size_for_shape, _reference_run, get_capacity
from generator.text_fit import cap_size_to_longest_word, soft_hyphenate_long_words, SOFT_HYPHEN

GRID_TPL = os.path.join(TEMPLATES_DIR, "custom_838830368dac3116.pptx")


def test_title_budget_trims_to_first_sentence():
    block = {"title": "Экономьте часы дизайнеров на рутинной вёрстке. "
                      "Получайте консистентные презентации без участия дизайнера.",
             "type": "title"}
    out = _enforce_text_budgets(block, "title")
    assert len(out["title"].split()) <= MAX_TITLE_WORDS
    assert out["title"].startswith("Экономьте")
    assert "Получайте" not in out["title"]


def test_bullet_budget_trims_at_word_boundary():
    long_item = "Очень длинный пункт списка который никак не помещается в отведённый лимит символов никаким образом"
    out = _enforce_text_budgets({"title": "т", "bullets": [long_item], "type": "bullet_list"}, "bullet_list")
    assert len(out["bullets"][0]) <= MAX_BULLET_CHARS
    assert not out["bullets"][0].endswith(" ")


def test_soft_hyphenate_only_long_words():
    text = soft_hyphenate_long_words("Неконсистентность результата")
    assert SOFT_HYPHEN in text.split()[0]
    assert SOFT_HYPHEN not in text.split()[-1]  # 10-letter word untouched


@requires(GRID_TPL)
def test_grid_slots_detected_row_major():
    prs = Presentation(GRID_TPL)
    slots = _find_slot_boxes(prs.slides[2], set())
    assert len(slots) == 6
    assert get_capacity(prs.slides[2], "bullet_list") == 6
    # row-major: first two slots share the top row (equal top bucket)
    assert abs(slots[0].top - slots[1].top) < abs(slots[0].top - slots[2].top)


@requires(GRID_TPL)
def test_title_word_fit_shrinks_below_base(tmp_path):
    prs = Presentation(GRID_TPL)
    slide = prs.slides[1]
    title = next(s for s in slide.shapes if s.has_text_frame and "Заголовок" in s.text_frame.text)
    ref = _reference_run(title)
    base = ref.font.size.pt  # 51pt
    size = _fit_size_for_shape(title, "Автоматизация презентаций", ref)
    assert size.pt < base  # word-fit must shrink from 51
    floor = max(9, round(base * 0.7))
    assert size.pt >= floor  # but never below the floor
