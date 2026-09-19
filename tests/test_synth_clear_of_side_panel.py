"""A synthesized slide stays clear of the canvas's big side panel.

VK Education's title-style layout carries a 6.74in art panel from x 6.60in —
45% of the slide, as a LAYOUT picture. Synthesis stepped around narrow decor
strips only and treated anything bigger as a backdrop, so on that canvas the
title ran under the panel's white shapes («CRM migration sta…») and the image
frame lay across the art (iter127). A big picture standing at the side of the
content band is the other half of the layout: the band stops at its edge.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from common.synthesis import SYNTHESIZE
from generator.generator import generate

VK_EDUCATION = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
CANVAS = 0  # the cover: logo + art panel in its layout, text on the left


def _layout_panel(slide):
    area = slide.part.package.presentation_part.presentation.slide_width * \
        slide.part.package.presentation_part.presentation.slide_height
    pics = [s for s in slide.slide_layout.shapes if "PICTURE" in str(s.shape_type)
            and s.width * s.height / area > 0.3]
    assert pics, "fixture: VK Education's cover layout carries the art panel"
    return pics[0]


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_synthesized_slide_is_drawn_beside_the_panel(tmp_path):
    out = str(tmp_path / "deck.pptx")
    block = {"type": "image_caption", "title": "Статус миграции CRM за третий квартал",
             "image": "график этапов миграции и диаграмма бюджета"}
    generate(VK_EDUCATION, [(block, SYNTHESIZE)], out, synth_canvas={0: CANVAS})
    slide = Presentation(out).slides[0]
    panel_left = _layout_panel(slide).left

    drawn = [s for s in slide.shapes
             if (s.has_text_frame and s.text_frame.text.strip()) or s.name == "img_placeholder"]
    assert any(s.has_text_frame and s.text_frame.text == block["title"] for s in drawn)
    over = [(s.name, s.text_frame.text[:30] if s.has_text_frame else "")
            for s in drawn if s.left + s.width > panel_left]
    assert not over, f"drawn across the art panel (from {panel_left.inches:.2f}in): {over}"
