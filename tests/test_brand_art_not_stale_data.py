"""Artwork the template reuses across slides is brand art, not this slide's data.

The survey deck's contacts slide carries a hand-drawn brand illustration over a
quarter of the slide. is_stale_data_picture called it stale data — a chart or a
screenshot — so the generator would have stripped the brand art out of a deck
built on that slide.

The pixel signature cannot tell them apart: the illustration is two flat colours
and so is a monochrome bar chart, 8 distinct shades against 9 (measured on
slides 31 and 6 of that template, both verified by eye). What separates them is
structural — a chart is made FOR its slide and appears once, while brand art is
placed wherever the designer wants it. Measured over the 138 large pictures of
the corpus that this module would otherwise call stale data: exactly four
repeat, one per survey deck, and each is that same illustration.
"""

import hashlib

from conftest import SURVEY_31, TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from common.pictures import OVERSIZED_PICTURE_AREA_RATIO, is_stale_data_picture

CONTACTS = 30  # brand illustration, repeated from the cover
CHART = 5      # a real monochrome bar chart, unique to its slide


def _pictures(slide):
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _big(slide, prs):
    area = prs.slide_width * prs.slide_height
    return [p for p in _pictures(slide)
            if p.width and (p.width * p.height) / area > OVERSIZED_PICTURE_AREA_RATIO]


@requires(SURVEY_31)
def test_repeated_brand_art_is_kept():
    prs = Presentation(SURVEY_31)
    slide = list(prs.slides)[CONTACTS]
    art = _big(slide, prs)
    assert art, "fixture: the contacts slide is supposed to carry a large picture"

    digest = hashlib.sha1(art[0].image.blob).hexdigest()
    elsewhere = sum(1 for s in prs.slides for p in _pictures(s)
                    if hashlib.sha1(p.image.blob).hexdigest() == digest)
    assert elsewhere > 1, f"fixture: this image should repeat, found {elsewhere}"

    assert not is_stale_data_picture(art[0], slide), "brand art would be deleted"


@requires(SURVEY_31)
def test_a_unique_chart_is_still_stale_data():
    """The rule must not blanket-rescue: a chart of the same flatness, used once,
    stays deletable."""
    prs = Presentation(SURVEY_31)
    slide = list(prs.slides)[CHART]
    charts = _big(slide, prs)
    assert charts, "fixture: this slide is supposed to carry a large chart"
    assert is_stale_data_picture(charts[0], slide)


@requires(TJ_UNIVERSAL)
def test_template_without_charts_is_unaffected():
    prs = Presentation(TJ_UNIVERSAL)
    for slide in prs.slides:
        for picture in _big(slide, prs):
            assert not is_stale_data_picture(picture, slide)
