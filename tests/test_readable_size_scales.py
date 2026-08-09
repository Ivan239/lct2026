"""«Мелкий шрифт» is relative to the canvas, not an absolute point size.

A point is absolute on paper, but a deck is projected to fill a screen: 10pt on
a 10in-wide slide renders exactly as large as 13.3pt on a 13.33in one. All three
T-Zh templates are authored at 10x5.62in, and a flat 11pt threshold reported
NINE boxes of the pristine mono template as fine print — every one of them the
deck's own 10pt body copy, at a size that reads perfectly on screen.

Same class as the three false alarms CLAUDE.md already records: the harness
judging a designer's layout by a ruler that does not apply to it.
"""

from conftest import SURVEY_69, TJ_MONO, requires

from pptx import Presentation
from pptx.util import Emu

from evaluation.deterministic import MIN_READABLE_PT, _content_shapes, _shape_max_size, evaluate

# REFERENCE_SLIDE_WIDTH_EMU is imported inside the tests that need it: at module
# level its absence would abort the whole file with an ImportError, and «the
# constant is missing» is a far weaker statement than «pristine mono was called
# fine print».


@requires(TJ_MONO)
def test_ten_point_body_on_a_ten_inch_canvas_is_not_fine_print():
    prs = Presentation(TJ_MONO)
    assert Emu(prs.slide_width).inches < 11, "fixture: mono is the 10in-wide canvas"

    small = [s for slide in prs.slides for s in _content_shapes(slide)
             if 9.5 <= (_shape_max_size(s) or 99) < MIN_READABLE_PT]
    assert len(small) >= 5, (
        f"fixture changed: expected the deck's 10pt body copy, found {len(small)} boxes")

    assert evaluate(TJ_MONO)["1.1"]["score"] == 5, evaluate(TJ_MONO)["1.1"]["detail"]


@requires(SURVEY_69)
def test_a_footnote_on_a_full_width_canvas_is_still_fine_print():
    """The threshold must not simply go away: the same 10.5pt on a 13.33in slide
    IS fine print, and survey-69's «*new question in survey» stays counted."""
    from evaluation.deterministic import REFERENCE_SLIDE_WIDTH_EMU

    prs = Presentation(SURVEY_69)
    assert int(prs.slide_width) == REFERENCE_SLIDE_WIDTH_EMU
    tiny = [s for slide in prs.slides for s in _content_shapes(slide)
            if (_shape_max_size(s) or 99) < MIN_READABLE_PT]
    assert len(tiny) == 1, f"fixture changed: {len(tiny)} boxes under the threshold"


@requires(SURVEY_69)
def test_the_reference_width_is_exact():
    """13.333in is SHORT of the real 12192000 EMU, and the 1.000025 ratio it
    produced pushed the threshold to 11.0003pt — enough to reclassify 29 boxes
    set at exactly 11pt as fine print. Caught by re-measuring after the change,
    not by any test, so here it is."""
    from evaluation.deterministic import REFERENCE_SLIDE_WIDTH_EMU

    prs = Presentation(SURVEY_69)
    scaled = MIN_READABLE_PT * int(prs.slide_width) / REFERENCE_SLIDE_WIDTH_EMU
    assert scaled == MIN_READABLE_PT
    assert not [s for slide in prs.slides for s in _content_shapes(slide)
                if (_shape_max_size(s) or 99) == 11.0 and 11.0 < scaled]
