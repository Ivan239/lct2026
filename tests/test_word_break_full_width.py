"""A word breaks mid-letter only when it does not fit — the harness must not
flag a word that fits inside the generator's 5% safety margin.

That margin (SINGLE_LINE_SAFETY) is for SETTING a multi-word line, which the
renderer wraps a few percent early. A lone word has nowhere to wrap: it breaks
only past the full width. With the margin, the harness flagged the pristine VK
Education legend «DAU»/«MAU» (99% and 98% of the width) and six KPI figures
on loop decks — all rendered whole (iter132).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from evaluation.deterministic import evaluate

VK_EDUCATION = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
LEGEND_SLIDE = 47  # «Пример оформления графика»: legend labels DAU / MAU in 0.55in boxes


def _one_slide(src, keep, out):
    prs = Presentation(src)
    ids = prs.slides._sldIdLst
    for i, sld_id in reversed(list(enumerate(ids))):
        if i != keep:
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            ids.remove(sld_id)
    return prs


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_a_word_that_fits_is_not_a_break(tmp_path):
    out = str(tmp_path / "legend.pptx")
    _one_slide(VK_EDUCATION, LEGEND_SLIDE, out).save(out)
    assert evaluate(out)["dop_wrap"]["score"] == 5, evaluate(out)["dop_wrap"]["detail"]


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_a_word_that_does_not_fit_still_is(tmp_path):
    out = str(tmp_path / "narrow.pptx")
    prs = _one_slide(VK_EDUCATION, LEGEND_SLIDE, out)
    box = prs.slides[0].shapes.add_textbox(Inches(3), Inches(3), Inches(0.8), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Необходимые"
    run.font.size, run.font.name = Pt(24), "Arial"
    prs.save(out)
    assert evaluate(out)["dop_wrap"]["score"] < 5
