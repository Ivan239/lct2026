"""The harness measures text inside GROUPS, at its place on the slide.

The deterministic criteria read the top level of slide.shapes only. VK Tech
keeps whole layouts as groups — its team slide is one group holding five
columns of cards — so a card pushed past the slide edge was invisible to 9.1
and the slide scored as clean (iter124).

Opening the group is not enough: python-pptx reports a member's position in
the group's own coordinate space. On that slide the raw positions run to
11.4in down a 5.6in slide; mapped through the group transform they land
where the render draws the text.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from evaluation.deterministic import evaluate

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
TEAM_SLIDE = 12  # «Слайд-визитка команды»: title + one group of five card columns


def _one_slide(src, keep, out, shift_group_in=0.0):
    prs = Presentation(src)
    ids = prs.slides._sldIdLst
    for i, sld_id in reversed(list(enumerate(ids))):
        if i != keep:
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            ids.remove(sld_id)
    groups = [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP]
    assert len(groups) == 1, "fixture: the team slide is one group"
    groups[0].left = groups[0].left + int(Inches(shift_group_in))
    prs.save(out)
    return out


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_grouped_text_past_the_edge_is_caught(tmp_path):
    pristine = evaluate(_one_slide(VK_TECH, TEAM_SLIDE, str(tmp_path / "a.pptx")))
    assert pristine["9.1"]["score"] == 5, pristine["9.1"]["detail"]

    shifted = evaluate(_one_slide(VK_TECH, TEAM_SLIDE, str(tmp_path / "b.pptx"), shift_group_in=1.0))
    assert shifted["9.1"]["score"] < 5, shifted["9.1"]["detail"]
