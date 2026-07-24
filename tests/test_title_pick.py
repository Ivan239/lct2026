"""_pick_title_shape on placeholder-less slides: font size alone must not win —
on the preset stats slide the 48pt KPI numbers out-size the 30pt title, and
mistaking a number box for the title shifts every num/label pair one box over
(the shipped deck showed the slide title inside KPI position #1)."""

from conftest import PRESET_A, requires

from pptx import Presentation

from generator.generator import _pick_title_shape, generate

STATS = {
    "type": "stats_kpi",
    "title": "Результаты пилота",
    "stats": [("x10", "ускорение подготовки"), ("3", "шаблона компаний"), ("0 ч", "ручной вёрстки")],
}


def _shape_texts(slide):
    return {
        sh.shape_id: "".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
        for sh in slide.shapes if sh.has_text_frame
    }


@requires(PRESET_A)
def test_title_band_beats_bigger_font():
    prs = Presentation(PRESET_A)
    picked = _pick_title_shape(prs.slides[3], set())
    texts = _shape_texts(prs.slides[3])
    assert texts[picked.shape_id] == "Метрики за квартал"


@requires(PRESET_A)
def test_stats_fill_lands_in_right_boxes(tmp_path):
    out = str(tmp_path / "stats.pptx")
    generate(PRESET_A, [(STATS, 3)], out)
    slide = Presentation(out).slides[0]
    by_top = sorted(
        (sh for sh in slide.shapes if sh.has_text_frame),
        key=lambda s: (s.top or 0, s.left or 0),
    )
    texts = ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs) for sh in by_top]
    assert texts[0] == "Результаты пилота"  # topmost box is the title
    assert texts[1:4] == ["x10", "3", "0 ч"]  # number row, left to right
    assert texts[4:7] == ["ускорение подготовки", "шаблона компаний", "ручной вёрстки"]
