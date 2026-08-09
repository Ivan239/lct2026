"""Where a synthesized block sits under its title — the TEMPLATE decides.

History, because this file has held both answers. Centring came first: a
comparison glued to its title with the lower half of the slide empty (deck
f78f14b1, slide 6) read as an accidentally half-empty slide, so short content
was moved to the middle of the band left under the title.

Measured later (iter95-97): the designers do the opposite. The median gap
between a title and the content under it is 0.00 / 0.06 / 0.20 / 0.25 / 1.12in
on the five real templates — four of them keep content right under the heading,
and only T-Zh mono holds it low. Our synthesized slides sat 14-21% of the slide
height below their titles against the templates' own 0-4%, and on the render the
heading looked detached from the list it introduces.

So neither rule is flat: the drop is CAPPED at the gap the template itself
keeps. survey-31 (0.00in) hugs, mono (1.12in) still drops, and content tall
enough to fill the band stays where it was in either case.

The renders of both variants were compared side by side before this was
changed — and only after iter96 fixed the bold-title measurement, because the
first attempt at the cap landed a bullet on top of a two-line title that the
model thought was one line.
"""

from conftest import SURVEY_31, TJ_MONO, requires

from pptx import Presentation
from pptx.util import Emu, Inches

from common.synthesis import SYNTHESIZE
from generator.generator import generate

TWO_COL = {
    "type": "two_column_comparison",
    "title": "Автоматизация против ручной работы",
    "left_heading": "Традиционный подход",
    "left_points": ["Долгая ручная вёрстка", "Неоднородность дизайна", "Зависимость от дизайнеров"],
    "right_heading": "Подход СлайдоГена",
    "right_points": ["Автоматическая генерация", "Сохранение фирменного стиля", "Независимость от человека"],
}

STATS = {
    "type": "stats_kpi",
    "title": "Результаты пилота",
    "stats": [("x10", "ускорение подготовки"), ("0 ч", "ручной вёрстки")],
}

SHORT_BULLETS = {
    "type": "bullet_list",
    "title": "Решение",
    "bullets": ["Разбор шаблона", "Извлечение дизайн-системы", "Сборка презентации"],
}

LONG_BULLETS = {
    "type": "bullet_list",
    "title": "Очень плотный слайд",
    "bullets": [
        f"Пункт номер {i}: развёрнутое описание с достаточно длинным текстом, "
        "чтобы строка переносилась и занимала заметную высоту на слайде"
        for i in range(1, 13)
    ],
}

PLAN = [
    (TWO_COL, SYNTHESIZE),
    (STATS, SYNTHESIZE),
    (SHORT_BULLETS, SYNTHESIZE),
    (LONG_BULLETS, SYNTHESIZE),
]


def _boxes(slide):
    """Shapes in creation order: synthesizers add the title box first."""
    return [s for s in slide.shapes if s.has_text_frame]


@requires(SURVEY_31)
def test_content_hugs_the_title_when_the_template_does(tmp_path):
    """survey-31 keeps its own content right under its titles (median gap 0.00in),
    so ours does too — instead of floating in the middle with the heading alone
    at the top."""
    out = str(tmp_path / "synth.pptx")
    generate(SURVEY_31, PLAN, out)
    prs = Presentation(out)

    title, col_a, col_b = _boxes(prs.slides[0])[:3]
    assert col_a.top == col_b.top
    glued_top = title.top + title.height + Inches(0.3)
    assert col_a.top <= glued_top + Inches(0.25), (
        f"columns sit {Emu(int(col_a.top - glued_top)).inches:.2f}in below the title "
        "on a template that keeps its content against it")

    stat_boxes = _boxes(prs.slides[1])[1:]
    num_a, num_b = stat_boxes[0], stat_boxes[2]
    assert num_a.top == num_b.top
    stats_title = _boxes(prs.slides[1])[0]
    stats_glued = stats_title.top + stats_title.height + Inches(0.35)
    assert num_a.top <= stats_glued + Inches(0.25)

    b_title, b_body = _boxes(prs.slides[2])[:2]
    assert b_body.top <= b_title.top + b_title.height + Inches(0.25) + Inches(0.25)


@requires(TJ_MONO)
def test_content_stays_low_when_the_template_holds_it_low(tmp_path):
    """The other side, and the reason the cap is the template's number and not a
    flat "hug the title": mono's own median gap is 1.12in, and its synthesized
    slides keep that airy layout."""
    out = str(tmp_path / "synth.pptx")
    generate(TJ_MONO, PLAN, out)
    prs = Presentation(out)
    title, col_a, _ = _boxes(prs.slides[0])[:3]
    glued_top = title.top + title.height + Inches(0.3)
    assert col_a.top > glued_top + Inches(0.5), (
        "mono's content should still drop: its own layout keeps content low")


@requires(SURVEY_31)
def test_tall_content_stays_where_the_layout_put_it(tmp_path):
    """Estimated height exceeds the available area — there is nothing to move."""
    out = str(tmp_path / "synth.pptx")
    generate(SURVEY_31, PLAN, out)
    prs = Presentation(out)
    l_title, l_body = _boxes(prs.slides[3])[:2]
    assert abs(int(l_body.top) - int(l_title.top + l_title.height + Inches(0.25))) <= Inches(0.02)
