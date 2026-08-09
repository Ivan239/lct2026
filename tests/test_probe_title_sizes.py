"""Title sizes across a deck, against the template's own.

Criterion 9.3 («одинаковые размеры заголовков») is scored by eye alone, and the
eye is right to be bothered: the universal deck runs 24 / 29 / 42 / 42 on its
content slides, so one heading reads at half the size of another two slides
later.

The template's size for the same slide is printed beside ours, because the two
cases look identical in our file and are not: on universal slide 2 the 24pt is
the designer's own heading on their numbered-grid canvas, while universal slide
3 (29pt against the template's 42pt) and mono slide 2 (31pt against 51pt) are
sizes WE shrank to make our longer text fit. Only the second kind is ours to
fix, and a bare list of sizes cannot tell them apart.

Covers are excluded: a cover is deliberately louder than a content slide — that
is what cover_pt exists for — and the first run of this reported survey-31's
72pt cover as a defect.
"""

import importlib.util
import os

from conftest import ROOT

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

spec = importlib.util.spec_from_file_location(
    "offline_probe", os.path.join(ROOT, "scripts", "offline_probe.py"))
probe = importlib.util.module_from_spec(spec)
spec.loader.exec_module(probe)


def _deck(path, sizes):
    """One slide per size: a title box at the top plus a body box under it."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(5.62)))
    for size in sizes:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for top, text, pt in ((0.4, "Заголовок слайда", size), (2.0, "Контент", 14)):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(6), Inches(0.9))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(pt)
    prs.save(path)
    return path


def _plan(types):
    return [({"type": t}, i) for i, t in enumerate(types)]


def test_a_shrunk_heading_is_reported_with_the_templates_own_size(tmp_path):
    deck = _deck(str(tmp_path / "deck.pptx"), [42, 29, 42, 42])
    source = _deck(str(tmp_path / "src.pptx"), [42, 42, 42, 42])
    plan = _plan(["bullet_list"] * 4)
    known, mode, stray = probe.title_sizes(deck, source, plan)
    assert mode == 42
    assert stray == [(2, 29.0, 42.0)]


def test_an_even_deck_reports_nothing(tmp_path):
    deck = _deck(str(tmp_path / "deck.pptx"), [42, 40, 42])
    source = _deck(str(tmp_path / "src.pptx"), [42, 42, 42])
    assert probe.title_sizes(deck, source, _plan(["bullet_list"] * 3))[2] == []


def test_the_cover_is_not_measured_against_the_content_mode(tmp_path):
    """A 72pt cover over 40pt content slides is the design, not a defect."""
    deck = _deck(str(tmp_path / "deck.pptx"), [72, 40, 40, 40])
    source = _deck(str(tmp_path / "src.pptx"), [72, 40, 40, 40])
    plan = _plan(["title", "bullet_list", "stats_kpi", "closing"])
    assert probe.title_sizes(deck, source, plan)[2] == []
