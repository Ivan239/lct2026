"""The image skeleton keeps clear of the heading box.

The designer may start the artwork INSIDE the heading's box: VK Education 31
runs its screenshots from 3.8in under a 0.74-5.04in heading box, which is fine
while the heading is «Слайды со скриншотами» — short, one line, ending well
before the artwork. Ours is a sentence that wraps to two lines and reaches the
box edge, and the dashed frame printed straight across it (iter145). The frame
now starts past the heading box, the same line the designer's own second
picture keeps (5.14in against a box ending at 5.04in).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import IMAGE_PLACEHOLDER_NAME, generate

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
SCREENSHOTS = 30
BLOCK = {"type": "image_caption", "title": "Процесс миграции данных CRM в облако",
         "image": "Диаграмма этапов миграции с процентами завершения"}


def _slide(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(VK_EDU, [(BLOCK, SCREENSHOTS)], out)
    return Presentation(out).slides[0]


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_frame_starts_past_the_heading_box(tmp_path):
    slide = _slide(tmp_path)
    frame = next(s for s in slide.shapes if s.name == IMAGE_PLACEHOLDER_NAME)
    title = next(s for s in slide.shapes
                 if s.has_text_frame and s.text_frame.text.startswith("Процесс"))
    assert frame.left >= title.left + title.width, (frame.left, title.left + title.width)


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_the_frame_stays_large(tmp_path):
    """Clipped, not squeezed: a frame narrower than 40% of the artwork it
    replaces says less than the artwork did."""
    slide = _slide(tmp_path)
    frame = next(s for s in slide.shapes if s.name == IMAGE_PLACEHOLDER_NAME)
    template = Presentation(VK_EDU).slides[SCREENSHOTS]
    pictures = [s for s in template.shapes if s.shape_type == 13]
    designed = max(int(p.left + p.width) for p in pictures) - min(int(p.left) for p in pictures)
    assert frame.width >= 0.4 * designed, (frame.width, designed)
