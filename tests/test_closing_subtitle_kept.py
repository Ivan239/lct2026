"""A subtitle the brief asked for must not vanish because the slide has one box.

The survey template's closing is a single text box and a decorative heart, so
_pick_body_shape finds nothing to write the subtitle into and it was dropped
without a trace: «Две недели на подключение и обучение команды» never reached
any deck built on that slide, on any brief.

It goes into a box of its own under the title, not a second paragraph of the
title's box: the harness reads a title as its box's whole text, and appending
there made a 99-character heading that criterion 2.1 duly reported as too long.
Same shape iter70 gave the display stat's caption.
"""

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.util import Emu

from generator.generator import generate

CLOSING_SLIDE = 29
BLOCK = {"type": "closing",
         "title": "Запустим пилот в вашем подразделении",
         "subtitle": "Две недели на подключение и обучение команды"}


def _closing(tmp_path, name="closing"):
    out = str(tmp_path / f"{name}.pptx")
    generate(SURVEY_31, [(BLOCK, CLOSING_SLIDE)], out)
    return list(Presentation(out).slides)[0]


@requires(SURVEY_31)
def test_the_subtitle_reaches_the_slide(tmp_path):
    source = Presentation(SURVEY_31)
    boxes = [s for s in list(source.slides)[CLOSING_SLIDE].shapes
             if s.has_text_frame and s.text_frame.text.strip()]
    assert len(boxes) == 1, f"fixture: this closing is supposed to have one text box, got {len(boxes)}"

    slide = _closing(tmp_path)
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert any(BLOCK["subtitle"] in t for t in texts), texts


@requires(SURVEY_31)
def test_it_is_a_separate_shape_smaller_than_the_title(tmp_path):
    slide = _closing(tmp_path, "separate")
    title = next(s for s in slide.shapes
                 if s.has_text_frame and BLOCK["title"] in s.text_frame.text)
    subtitle = next(s for s in slide.shapes
                    if s.has_text_frame and BLOCK["subtitle"] in s.text_frame.text)
    assert subtitle.shape_id != title.shape_id, "the subtitle must not join the title's box"
    assert BLOCK["subtitle"] not in title.text_frame.text

    def _pt(shape):
        return max(r.font.size.pt for p in shape.text_frame.paragraphs
                   for r in p.runs if r.font.size)

    assert _pt(subtitle) < _pt(title)
    assert int(subtitle.top) >= int(title.top)
    assert int(subtitle.top + subtitle.height) <= int(title.top + title.height), \
        "the subtitle must stay inside the area the designer gave the closing"


@requires(SURVEY_31)
def test_a_closing_without_a_subtitle_adds_nothing(tmp_path):
    out = str(tmp_path / "bare.pptx")
    generate(SURVEY_31, [({"type": "closing", "title": BLOCK["title"]}, CLOSING_SLIDE)], out)
    slide = list(Presentation(out).slides)[0]
    filled = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert len(filled) == 1, [s.text_frame.text for s in filled]
