"""A check that does NOT share the generator's model.

Everything else in the harness counts lines with the same fontTools model the
generator uses, so when that model was wrong — bold titles measured with the
regular face, 7-8% narrow on Cyrillic (iter96) — every check agreed with the
mistake: the deck scored 99.5 with a bullet printed through the title's second
line.

Pixels do not share the model. They also do not come labelled, and two earlier
shapes of this measure died on that: counting ink INSIDE the box found nothing
(the unpredicted line is outside the box by definition), and looking BELOW the
box fired on every healthy deck (the next box's text is there, and pixels cannot
say whose line it is). What pixels can say is whether the ink reaches the box's
own bottom edge — a box sized for its text leaves the last line clear of it.

The renders here are DRAWN, not rendered: LibreOffice has no place in the fast
suite, and what is under test is the pixel logic, not the renderer. Silence on
the real decks is verified by running the probe, which the suite does not do.
"""

import importlib.util
import os

from conftest import ROOT

from PIL import Image, ImageDraw

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

spec = importlib.util.spec_from_file_location(
    "offline_probe", os.path.join(ROOT, "scripts", "offline_probe.py"))
probe = importlib.util.module_from_spec(spec)
spec.loader.exec_module(probe)

SLIDE_W, SLIDE_H = 10.0, 5.62
PX_PER_IN = 150

TITLE_TOP, TITLE_LEFT, TITLE_W = 0.5, 0.5, 6.0


def _deck(path, title_height_in):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(SLIDE_W))), Emu(int(Inches(SLIDE_H)))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for top, height, text in ((TITLE_TOP, title_height_in, "Заголовок слайда"),
                              (3.0, 1.0, "Контент под заголовком")):
        box = slide.shapes.add_textbox(Inches(TITLE_LEFT), Inches(top),
                                       Inches(TITLE_W), Inches(height))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(28)
    prs.save(path)
    return path


def _render(path, ink_rows_in):
    """A white slide with black bars at the given (top, bottom) inch spans."""
    img = Image.new("RGB", (int(SLIDE_W * PX_PER_IN), int(SLIDE_H * PX_PER_IN)), "white")
    draw = ImageDraw.Draw(img)
    for top_in, bottom_in in ink_rows_in:
        draw.rectangle([int(TITLE_LEFT * PX_PER_IN), int(top_in * PX_PER_IN),
                        int((TITLE_LEFT + TITLE_W * 0.9) * PX_PER_IN),
                        int(bottom_in * PX_PER_IN)], fill="black")
    img.save(path)
    return path


PLAN = [({}, probe.SYNTHESIZE)]


def test_a_second_line_running_into_the_box_edge_is_reported(tmp_path):
    """The defect: the box was sized for one line and the renderer drew two, so
    the second one starts inside the box and runs off its bottom."""
    deck = _deck(str(tmp_path / "deck.pptx"), title_height_in=0.6)
    render = _render(str(tmp_path / "r1.png"),
                     [(0.55, 0.85), (0.95, 1.15)])  # box ends at 1.10in
    assert probe.title_ink_cut_by_its_box(deck, [render], PLAN) == [(1, 0.0)]


def test_a_box_that_holds_its_text_is_silent(tmp_path):
    """Same two lines, in a box tall enough for them."""
    deck = _deck(str(tmp_path / "deck.pptx"), title_height_in=1.2)
    render = _render(str(tmp_path / "r2.png"), [(0.55, 0.85), (0.95, 1.25)])
    assert probe.title_ink_cut_by_its_box(deck, [render], PLAN) == []


def test_a_native_slide_is_not_judged_by_our_measure(tmp_path):
    """On a native slide the title box is the designer's, and a designer's text
    exceeds its frame by intent — the first run of this flagged two untouched
    T-Zh study slides for exactly that."""
    deck = _deck(str(tmp_path / "deck.pptx"), title_height_in=0.6)
    render = _render(str(tmp_path / "r3.png"), [(0.55, 0.85), (0.95, 1.15)])
    assert probe.title_ink_cut_by_its_box(deck, [render], [({}, 0)]) == []
