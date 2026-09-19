"""A slot grid is judged by its REAL steps, not by steps between bucket floors.

T-Zh mono has a four-column slide: a title box at the top (51pt, «Заголовок /
слайда» in two paragraphs), a year in each column (51pt) and a 10pt caption
under it. The captions sit at 0.31 / 2.67 / 5.14 / 7.49in — steps 2.36 / 2.47 /
2.35in, 0.12 apart, inside the 0.15in jitter allowance. _uniform_axis floored
every coordinate to 0.1in first (0.3 / 2.6 / 5.1 / 7.4, steps 2.3 / 2.5 / 2.3,
0.2 apart) and called it "not a grid". With no slots, the list filler took the
box with the most paragraphs — the two-paragraph TITLE — for the body: the
bullets printed at 51pt in the title box, the title went into a column (the
first 12-slide probe deck, iter118).

Measured over the whole corpus, the fix changes the slots of exactly this one
slide.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Inches

from generator.generator import _uniform_axis, generate, get_capacity

MONO = os.path.join(TEMPLATES_DIR, "custom_838830368dac3116.pptx")


def test_real_steps_inside_the_allowance_make_a_grid():
    lefts = [Inches(0.31), Inches(2.67), Inches(5.14), Inches(7.49)]
    assert _uniform_axis(lefts) is not None
    # and a genuinely ragged row is still refused
    assert _uniform_axis([Inches(0.3), Inches(2.3), Inches(5.3), Inches(7.4)]) is None


def _column_slide(prs):
    """Title box on top and four same-size caption boxes in one row."""
    for i, slide in enumerate(prs.slides):
        boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        rows = {}
        for b in boxes:
            rows.setdefault((round(b.top / 91440), round(b.width / 91440), round(b.height / 91440)), []).append(b)
        if any(len(r) == 4 for r in rows.values()) and sum(len(r) == 4 for r in rows.values()) == 2:
            return i
    return None


@pytest.mark.skipif(not os.path.exists(MONO), reason=f"fixture deck missing: {MONO}")
def test_a_list_on_the_column_slide_fills_the_captions_not_the_title(tmp_path):
    prs = Presentation(MONO)
    idx = _column_slide(prs)
    assert idx is not None, "fixture: mono's four-column slide not found"
    assert get_capacity(prs.slides[idx], "bullet_list") == 4

    bullets = ["Права доступа настраиваются по ролям", "История версий отчёта хранится целиком",
               "Ручной сбор показателей из семи систем"]
    block = {"type": "bullet_list", "title": "Что мешало собирать отчётность", "bullets": bullets}
    out = str(tmp_path / "deck.pptx")
    generate(MONO, [(block, idx)], out)

    slide = Presentation(out).slides[0]
    texts = {s.text_frame.text.strip(): s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()}
    title = texts.get("Что мешало собирать отчётность")
    assert title is not None, f"title not in a box of its own: {list(texts)}"
    for bullet in bullets:
        box = texts.get(bullet)
        assert box is not None, f"{bullet!r} is not in a slot of its own: {list(texts)}"
        assert box.top > title.top, f"{bullet!r} sits above the title"
