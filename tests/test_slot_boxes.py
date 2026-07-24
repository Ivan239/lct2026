"""Multi-box list slots (T-Ж class, plan item 8а): lists built as N separate
one-line text boxes with tiny marker auto-shapes. The filler must put one item
per box and physically remove surplus boxes together with their markers —
"dots without text" was a real user report."""

import os

from conftest import TEMPLATES_DIR, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from generator.generator import _fill_bullet_list, _find_slot_boxes, get_capacity

TJ = os.path.join(TEMPLATES_DIR, "custom_f496182bb15f42bb.pptx")


@requires(TJ)
def test_capacity_counts_slot_boxes():
    prs = Presentation(TJ)
    assert get_capacity(prs.slides[2], "bullet_list") == 5


@requires(TJ)
def test_slot_fill_and_surplus_removal():
    prs = Presentation(TJ)
    slide = prs.slides[2]
    assert len(_find_slot_boxes(slide, set())) == 5

    data = {"title": "Почему это важно", "bullets": ["Первый", "Второй", "Третий"]}
    _fill_bullet_list(slide, data, set())

    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert "Первый" in " ".join(texts) and "Третий" in " ".join(texts)
    assert "Длинное название пункта" not in " ".join(texts)  # template text gone

    dots = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(dots) == 3  # one marker per remaining item, two removed
    # Exactly 3 slot boxes remain (surplus physically deleted, not blanked).
    assert len(_find_slot_boxes(slide, set())) == 3
