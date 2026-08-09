"""A running header that had to be cut must LOOK cut.

The deck's topic is printed on every slide of the T-Zh templates. Trimming it at
a word boundary alone ships a phrase that looks whole and is not: «Единая
платформа управленческой отчётности и аналитики для розничной сети» became
«ЕДИНАЯ ПЛАТФОРМА УПРАВЛЕНЧЕСКОЙ» — an adjective governing a noun that is no
longer there — and «Итоги пилотного внедрения в трёх подразделениях за квартал»
became «…В ТРЁХ». Seen on the render of a long-title cover.

An ellipsis is what tells the reader the topic is shortened. The trailing
function word goes too, the same rule bullets have had since iter56 — which is
why it now lives in common/phrases.py: generator cannot import two_phase
(generator -> two_phase -> template_spec.builder -> generator).
"""

from conftest import TJ_TEMPLATE, requires

from pptx import Presentation

from generator.generator import MAX_RUNNING_HEADER_CHARS, generate

# The new names are imported inside the tests that need them: at module level
# their absence aborts the file with an ImportError, and "the helper is missing"
# says far less than "the header on a real deck reads as a chopped phrase".
ELLIPSIS = "…"

LONG = "Единая платформа управленческой отчётности и аналитики для розничной сети"


def test_a_title_that_fits_is_untouched():
    from generator.generator import _shorten_running_topic

    short = "Платформа управленческой отчётности"
    assert len(short) <= MAX_RUNNING_HEADER_CHARS
    assert _shorten_running_topic(short) == short


def test_a_cut_title_is_marked_as_cut():
    from generator.generator import _shorten_running_topic

    out = _shorten_running_topic(LONG)
    assert out.endswith(ELLIPSIS), out
    assert len(out) <= MAX_RUNNING_HEADER_CHARS
    assert out.startswith("Единая платформа")


def test_the_cut_does_not_end_on_a_preposition():
    from generator.generator import _shorten_running_topic

    out = _shorten_running_topic("Итоги внедрения платформы отчётности для розничной сети")
    body = out[:-len(ELLIPSIS)]
    assert body.split()[-1].lower() not in {"для", "в", "на", "и", "с"}, out


def test_one_very_long_word_is_left_to_the_width_check():
    from generator.generator import _shorten_running_topic

    """No boundary to cut at: the caller's _fits_box_width decides, and shipping
    a word chopped mid-letter is worse than shipping it whole (iter18/25)."""

    word = "Клиентоориентированностьвышевсегоиэтоправдаправдаправда"
    assert _shorten_running_topic(word) == word


@requires(TJ_TEMPLATE)
def test_the_header_of_a_generated_deck_shows_the_ellipsis(tmp_path):
    out = str(tmp_path / "cover.pptx")
    generate(TJ_TEMPLATE, [({"type": "title", "title": LONG, "subtitle": "Итоги"}, 0)], out)
    texts = [s.text_frame.text for s in list(Presentation(out).slides)[0].shapes
             if s.has_text_frame]
    header = [t for t in texts if t.strip().isupper() and t.strip()]
    assert header, f"no running header found in {texts}"
    assert any(ELLIPSIS in t for t in header), header
