"""Bold text has to be measured with the bold face.

Synthesized titles are drawn bold (_style_paragraph(..., bold=True)) and were
measured with the regular face. On Cyrillic that is not a rounding error:
measured on Arial at 40pt, «Что мешало собирать отчётность вовремя» is 796pt
regular and 858pt bold, against a title box budget of 853pt. The model promised
one line, the renderer drew two, and the title box was sized for one — so
everything positioned after the title (a body block, an image frame) was placed
against a title that ends lower than we think.

Seen on the render: survey-31's synthesized bullet list printed its first bullet
straight through the title's second line, the moment a layout change removed the
slack that had been hiding it.

BOLD_WIDTH_FACTOR (1.05, in cap_size_to_longest_word) was the earlier answer to
the same problem, and it is both flat and low — the real faces measure 7-8%
apart on this text. It stays where it is: it still guards the no-metrics path.
"""

import os

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.util import Emu

from common.synthesis import SYNTHESIZE
from fonts.metrics import FontResolver
from generator.generator import generate
from generator.text_fit import SINGLE_LINE_SAFETY, line_height_pt
from template_parser.parser import extract_theme

TITLE = "Что мешало собирать отчётность вовремя"

BULLETS = {
    "type": "bullet_list",
    "title": TITLE,
    "bullets": ["Ручной сбор показателей из семи независимых систем",
                "Разные форматы выгрузок у каждого подразделения",
                "Согласование занимало до двух недель"],
}


@requires(SURVEY_31)
def test_the_bold_face_is_wider_than_the_regular_one(tmp_path):
    resolver = FontResolver(SURVEY_31, extract_theme(SURVEY_31))
    regular = resolver.metrics_for("Aeroport")
    bold = resolver.metrics_for("Aeroport", bold=True)
    assert bold.text_width_pt(TITLE, 40) > regular.text_width_pt(TITLE, 40) * 1.05, (
        "the bold donor is not being used: "
        f"{bold.text_width_pt(TITLE, 40):.0f}pt vs {regular.text_width_pt(TITLE, 40):.0f}pt")


@requires(SURVEY_31)
def test_a_two_line_bold_title_gets_a_two_line_box(tmp_path):
    """The box is what everything below is positioned against, so a title the
    renderer draws on two lines must not be boxed for one.

    The bold face is loaded here DIRECTLY, not through metrics_for(bold=True):
    a test that measures with the API it is testing fails on HEAD with a
    TypeError, and "the argument is missing" says far less than "the box is one
    line where the text needs two"."""
    import pytest
    from fontTools.ttLib import TTFont

    from fonts.metrics import FontMetrics

    # Spelled out here rather than imported from the module under change, for
    # the same reason the measurement is direct: an ImportError is not a defect
    # report.
    bold_faces = ("/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                  "/Library/Fonts/Arial Bold.ttf",
                  "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    bold = None
    for path in bold_faces:
        if os.path.exists(path):
            bold = FontMetrics(TTFont(path, lazy=True))
            break
    if bold is None:
        pytest.skip("no bold donor face on this machine")

    out = str(tmp_path / "synth.pptx")
    generate(SURVEY_31, [(BULLETS, SYNTHESIZE)], out)
    title = [s for s in list(Presentation(out).slides)[0].shapes
             if s.has_text_frame and TITLE in s.text_frame.text][0]
    size_pt = title.text_frame.paragraphs[0].runs[0].font.size.pt

    # The budget the LAYOUT works to: the renderer wraps a few percent early, so
    # a single-line guarantee is only earned below SINGLE_LINE_SAFETY.
    budget = Emu(title.width).inches * 72 * SINGLE_LINE_SAFETY
    drawn = bold.text_width_pt(TITLE, size_pt)
    assert drawn > budget, (
        f"fixture no longer wraps: {drawn:.0f}pt against {budget:.0f}pt — pick a longer title")
    two_lines_in = 2 * line_height_pt(size_pt, metrics=bold) / 72
    assert Emu(title.height).inches >= two_lines_in * 0.9, (
        f"title box is {Emu(title.height).inches:.2f}in where two lines need "
        f"{two_lines_in:.2f}in")
