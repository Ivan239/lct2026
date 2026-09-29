"""A synthesized cover must end above the slide's bottom edge, and its subtitle
must not be printed through it.

Seen on a render of the survey deck: the cover title took the template's own
90pt, wrapped to three lines, and its box ran to 8.75in on a 7.5in slide — the
last line («отчётности «Поток»») was cut off by the bottom edge. The subtitle's
safety clamp then made it worse: pinned at 88% of the slide height, it landed
INSIDE the overflowing title and the two printed through each other.

The width cap (cap_size_to_longest_word, iter18/25) only prevents mid-word
breaks; nothing capped the block's HEIGHT. That is the same lesson docs/LESSONS.md
records after the study deck's cover — change a size and recompute the layout —
applied to the axis that fix did not cover.
"""

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.util import Emu

from fonts.metrics import FontResolver
from generator.deck_style import apply_observed_style, observe_deck_style
from generator.generator import _template_cover_pt
from generator.layout_bounds import infer_content_bounds
from generator.synthesizer import synthesize_title
from template_parser.parser import extract_template, extract_theme

COVER = {"title": "Платформа управленческой отчётности «Поток»",
         "subtitle": "Итоги пилотного внедрения в трёх подразделениях за квартал"}


@requires(SURVEY_31)
def test_cover_title_stays_on_the_slide_and_clear_of_its_subtitle():
    prs = Presentation(SURVEY_31)
    theme = apply_observed_style(extract_theme(SURVEY_31), observe_deck_style(prs))
    # generate() measures the cover size off the template and puts it in the
    # theme; the defect only exists at that measured size (90pt here), so a
    # theme without it tests the 40pt default and proves nothing.
    theme["cover_pt"] = _template_cover_pt(prs)
    idx = synthesize_title(prs, theme, infer_content_bounds(extract_template(SURVEY_31)), COVER,
                           resolver=FontResolver(SURVEY_31, extract_theme(SURVEY_31)))
    slide = prs.slides[idx]

    # With real metrics, as generate() runs it: without a resolver the height
    # estimate falls back to an average-width guess that predicts fewer lines,
    # and the overflow this test exists for never appears.
    boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    title = next(s for s in boxes if s.text_frame.text.strip() == COVER["title"])
    subtitle = next(s for s in boxes if s.text_frame.text.strip() == COVER["subtitle"])

    assert int(title.top + title.height) <= int(prs.slide_height), (
        f"the title box ends at {Emu(title.top + title.height).inches:.2f}in on a "
        f"{Emu(prs.slide_height).inches:.2f}in slide")
    assert int(subtitle.top) >= int(title.top + title.height), (
        f"the subtitle starts at {Emu(subtitle.top).inches:.2f}in, inside a title "
        f"running to {Emu(title.top + title.height).inches:.2f}in")
    assert int(subtitle.top + subtitle.height) <= int(prs.slide_height)


@requires(SURVEY_31)
def test_a_short_cover_keeps_the_templates_own_size():
    """The fit only ever shrinks. A title that fits must still be set at the
    size measured from the template — the whole point of measuring it."""
    prs = Presentation(SURVEY_31)
    theme = apply_observed_style(extract_theme(SURVEY_31), observe_deck_style(prs))
    theme["cover_pt"] = _template_cover_pt(prs)
    idx = synthesize_title(prs, theme, infer_content_bounds(extract_template(SURVEY_31)),
                           {"title": "Поток", "subtitle": "Итоги квартала"},
                           resolver=FontResolver(SURVEY_31, extract_theme(SURVEY_31)))
    title = next(s for s in prs.slides[idx].shapes
                 if s.has_text_frame and s.text_frame.text.strip() == "Поток")
    size = title.text_frame.paragraphs[0].runs[0].font.size.pt
    assert size >= (theme.get("cover_pt") or 40) - 1, (
        f"a one-word cover was shrunk to {size}pt")
