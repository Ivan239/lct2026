"""Soft hyphens inserted for a bigger size are removed once the size is final.

_set_run_text adds them when the longest word cannot fit at the size it is about
to use — a visible hyphen beats a word chopped mid-letter (iter18/25). But
enforce_text_fits then shrinks the text further, and at the smaller size the
word fits with room to spare. The renderer prefers breaking at a soft hyphen to
wrapping, so a stale one keeps hyphenating a word that no longer needs it: the
T-Zh mono heading rendered as «Что мешало собирать упра-вленческую отчётность
вовремя», where «управленческую» measures 232.6pt against a 293.9pt budget at
its final 28pt.
"""

from conftest import TJ_MONO, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from fonts.metrics import FontResolver
from generator.text_fit import SOFT_HYPHEN, soft_hyphenate_long_words
from qa.geometry import drop_needless_soft_hyphens
from template_parser.parser import extract_theme

WORD = "управленческую"


def _deck(tmp_path, width_in, size_pt):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(5.62)))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(int(Inches(0.2))), Emu(int(Inches(0.5))),
                                   Emu(int(Inches(width_in))), Emu(int(Inches(2.0))))
    box.text_frame.word_wrap = True
    box.text_frame.text = soft_hyphenate_long_words(f"Что мешало собирать {WORD} отчётность")
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(size_pt)
    run.font.name = "Arial"
    path = str(tmp_path / f"deck_{width_in}_{size_pt}.pptx")
    prs.save(path)
    return path


@requires(TJ_MONO)
def test_a_hyphen_the_text_outgrew_is_removed(tmp_path):
    path = _deck(tmp_path, width_in=4.08, size_pt=28)
    prs = Presentation(path)
    assert SOFT_HYPHEN in list(prs.slides)[0].shapes[0].text_frame.text, "fixture: no hyphen to drop"

    cleaned = drop_needless_soft_hyphens(prs, FontResolver(path, extract_theme(path)))
    assert cleaned, "nothing was cleaned"
    text = list(prs.slides)[0].shapes[0].text_frame.text
    assert SOFT_HYPHEN not in text
    assert WORD in text, text


@requires(TJ_MONO)
def test_a_hyphen_the_text_still_needs_is_kept(tmp_path):
    """The escape hatch must survive where it is still doing its job: the same
    word in a box too narrow for it at any readable size."""
    path = _deck(tmp_path, width_in=1.2, size_pt=28)
    prs = Presentation(path)
    drop_needless_soft_hyphens(prs, FontResolver(path, extract_theme(path)))
    assert SOFT_HYPHEN in list(prs.slides)[0].shapes[0].text_frame.text
