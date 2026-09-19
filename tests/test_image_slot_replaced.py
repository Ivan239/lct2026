"""A native image_caption slide does not ship the template's sample picture.

The image slot of a template slide holds the DESIGNER's sample, never ours: on
VK Education a laptop mockup of the VK Store community page (a girl in a
hoodie, «место встречи, Невский 28») stood under our heading «Единая картина
продаж», a phone mockup of someone's travel post next to it, and a VK
Education post screenshot on another slide (iter103, iter106). The synthesized
image_caption path has drawn a dashed «ИЗОБРАЖЕНИЕ» skeleton all along; the
native path now swaps the sample for the same skeleton.

The phone is a GROUP of two pictures (frame + screen), not a PICTURE — a
filter by shape type left it standing next to the skeleton, with the skeleton's
caption printed across it (iter107, first attempt).

Production path on the real template: generate() with image_caption on the
template slides that carry a large picture next to a title. On HEAD every
sample picture is still there and no skeleton exists.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from generator.generator import generate

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
SLOT_SHARE = 0.10


def _leaves(group):
    for child in group.shapes:
        if child.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _leaves(child)
        else:
            yield child


def _sample_pictures(slide, area):
    """Pictures and picture-only groups of image-slot size."""
    found = []
    for shape in slide.shapes:
        if not (shape.width and shape.height) or shape.width * shape.height / area < SLOT_SHARE:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            found.append(shape.name)
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            leaves = list(_leaves(shape))
            if leaves and all(leaf.shape_type == MSO_SHAPE_TYPE.PICTURE for leaf in leaves):
                found.append(shape.name)
    return found


def _mockup_slides(prs):
    """Template slides whose only content is a heading and big picture(s)."""
    area = prs.slide_width * prs.slide_height
    picked = []
    for i, slide in enumerate(prs.slides):
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if texts and len(texts) <= 2 and _sample_pictures(slide, area) \
                and any(s.shape_type == MSO_SHAPE_TYPE.PICTURE
                        and s.width * s.height / area > 0.25 for s in slide.shapes):
            picked.append(i)
    return picked[:3]


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_native_image_caption_replaces_the_sample_picture_with_a_skeleton(tmp_path):
    prs = Presentation(VK_EDU)
    area = prs.slide_width * prs.slide_height
    slides = _mockup_slides(prs)
    assert len(slides) == 3, f"fixture: expected three mockup slides, got {slides}"
    assert any(shape.shape_type == MSO_SHAPE_TYPE.GROUP
               for i in slides for shape in prs.slides[i].shapes
               if shape.name in _sample_pictures(prs.slides[i], area)), \
        "fixture: one of the mockups is expected to be a group of pictures (the phone)"

    block = {"type": "image_caption", "title": "Единая картина продаж",
             "image": "Экран дашборда с графиком выручки по неделям"}
    out = str(tmp_path / "deck.pptx")
    generate(VK_EDU, [(block, i) for i in slides], out)

    deck = Presentation(out)
    width, height = deck.slide_width, deck.slide_height
    for number, slide in enumerate(deck.slides, 1):
        left = _sample_pictures(slide, width * height)
        assert not left, f"slide {number}: template sample picture(s) left: {left}"
        frames = [s for s in slide.shapes if s.name == "img_placeholder"]
        assert len(frames) == 1, f"slide {number}: {len(frames)} image skeletons"
        frame = frames[0]
        assert frame.left >= 0 and frame.top >= 0 \
            and frame.left + frame.width <= width and frame.top + frame.height <= height, \
            f"slide {number}: the skeleton is cut by the slide edge"
        assert "ИЗОБРАЖЕНИЕ" in frame.text_frame.text
