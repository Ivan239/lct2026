"""A canvas's content GROUPS go with its content text boxes.

VK Tech keeps each card of its four-card slide as a GROUP (icon + «Заголовок /
Подзаголовок / Текст» and four more «Текст»). _prepare_blank_slide removed the
canvas's own text boxes before placing ours, but text_shapes() reads only the
top level of slide.shapes, so all four cards stayed — 28 placeholder strings
under our title and image frame, on both synthesized slides of the first VK
Tech deck (iter102) and again on iter105's.

Production path on the real template: generate() with synth_canvas pointing at
that slide, for the two roles that landed on it. On HEAD each slide carries 28
placeholder strings; fixed, none.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from common.synthesis import SYNTHESIZE
from evaluation.deterministic import placeholder_hits
from generator.generator import generate
from generator.slide_kit import content_groups

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")


def _card_canvas(prs):
    """The template slide with the most content groups (the four-card slide)."""
    return max(range(len(prs.slides)), key=lambda i: len(content_groups(prs.slides[i])))


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_synthesized_slides_carry_no_leftover_card_placeholders(tmp_path):
    canvas = _card_canvas(Presentation(VK_TECH))
    assert len(content_groups(Presentation(VK_TECH).slides[canvas])) >= 4, \
        "fixture: the four-card canvas is expected to keep its cards as groups"
    plan = [({"type": "image_caption", "title": "Экран единого дашборда",
              "image": "Панель управления с графиками"}, SYNTHESIZE),
            ({"type": "two_column_comparison", "title": "До и после внедрения",
              "left_heading": "Было", "left_points": ["Данные в разных системах"],
              "right_heading": "Стало", "right_points": ["Единый дашборд"]}, SYNTHESIZE)]
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, plan, out, synth_canvas={0: canvas, 1: canvas})
    for number, slide in enumerate(Presentation(out).slides, 1):
        hits = placeholder_hits(slide)
        assert not hits, f"slide {number}: {len(hits)} placeholder strings left: {sorted(set(hits))}"
