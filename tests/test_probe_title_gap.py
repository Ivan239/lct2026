"""Our gap under the title, next to the template's own.

An open question, printed rather than decided. Synthesized slides centre their
block in the band that is left, and on short content that opens a hole: on
survey-31 ours ran 24-30% of the slide height where the template's own slides
run 0-8%. Capping the drop at the template's gap was tried and reverted — it
contradicts test_synth_center, which records the opposite decision, made on a
real deck whose comparison sat glued to the title with the lower half empty.

Both observations are true at once: the designer's content is glued to the title
AND dense, ours is sparse. Neither rule settles it, so the two numbers go into
the probe's line and the next iteration can argue with data instead of taste.
"""

import importlib.util
import os

from conftest import ROOT

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

spec = importlib.util.spec_from_file_location(
    "offline_probe", os.path.join(ROOT, "scripts", "offline_probe.py"))
probe = importlib.util.module_from_spec(spec)
spec.loader.exec_module(probe)

SLIDE_H = 5.62


def _deck(path, gap_in, chrome=False):
    """Title at 0.3-1.0in, one content box `gap_in` below it."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(SLIDE_H)))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for top, height, text, size in ((0.3, 0.7, "Заголовок", 28),
                                    (1.0 + gap_in, 0.6, "Контент слайда", 18)):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(6), Inches(height))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
    if chrome:
        box = slide.shapes.add_textbox(Inches(9.4), Inches(5.3), Inches(0.4), Inches(0.16))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "02"
        run.font.size = Pt(8)
    prs.save(path)
    return path


def test_reports_both_gaps_as_percentages_of_the_slide(tmp_path):
    ours = _deck(str(tmp_path / "deck.pptx"), gap_in=1.5)
    theirs = _deck(str(tmp_path / "src.pptx"), gap_in=0.0)
    assert probe.title_gap_vs_template(ours, theirs) == (round(1.5 / SLIDE_H * 100), 0)


def test_chrome_does_not_pass_for_content(tmp_path):
    """A page number is furniture: counted as the second box it would report a
    gap to the FOOTER, which is nobody's layout decision."""
    ours = _deck(str(tmp_path / "deck.pptx"), gap_in=1.5, chrome=True)
    theirs = _deck(str(tmp_path / "src.pptx"), gap_in=0.0, chrome=True)
    assert probe.title_gap_vs_template(ours, theirs) == (round(1.5 / SLIDE_H * 100), 0)


def test_silent_when_no_slide_has_two_content_boxes(tmp_path):
    """A cover-only deck says nothing about the gap, and a measure with nothing
    to measure must say so rather than print a zero."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(SLIDE_H)))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.7))
    box.text_frame.paragraphs[0].add_run().text = "Только заголовок"
    cover = str(tmp_path / "cover.pptx")
    prs.save(cover)
    assert probe.title_gap_vs_template(cover, _deck(str(tmp_path / "src.pptx"), 0.0)) is None
