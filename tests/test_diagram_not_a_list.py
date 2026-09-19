"""A flowchart the designer drew is not a list slide.

VK Education's «Оформление схем» — a guide slide with an example flowchart:
11 boxes, 15 connector lines, 9 of them arrowed — was classified as a list.
The one node with text, a 1.26in circle, took the list; the title slot took
the first bullet and ran off the slide edge; the other nodes stayed empty
(iter130). Like a table or a chart, a diagram is not offered until diagrams
can be filled (backlog item 5).

The rule must not take legitimate layouts that happen to have lines: a Gantt
chart's gridlines (6, no arrows), T-Zh mono's column separators (4), a
«01 → 02 → 03 → 04» timeline (3 arrows).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from template_spec.builder import build_spec

VK_EDUCATION = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
MONO = os.path.join(TEMPLATES_DIR, "custom_838830368dac3116.pptx")
WORKSPACE = os.path.join(TEMPLATES_DIR, "vk_workspace.pptx")
FLOWCHARTS = {12, 13}
GANTT = 42


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_flowchart_slides_are_not_offered_as_lists():
    n = len(Presentation(VK_EDUCATION).slides)
    spec = build_spec(VK_EDUCATION, {i: "bullet_list" for i in range(n)})
    offered = {s["idx"] for f in spec["families"] for s in f["slides"]}
    assert offered, "the rule must not empty the family"
    assert not offered & FLOWCHARTS, sorted(offered & FLOWCHARTS)


@pytest.mark.skipif(not all(os.path.exists(p) for p in (VK_EDUCATION, MONO, WORKSPACE)),
                    reason="fixture decks missing")
def test_lines_alone_do_not_make_a_diagram():
    # Imported here: at module level a missing helper would fail the whole file
    # with ImportError on HEAD instead of the behaviour test above failing.
    from common.pictures import has_connector_diagram

    assert not has_connector_diagram(Presentation(VK_EDUCATION).slides[GANTT])
    assert not has_connector_diagram(Presentation(MONO).slides[9])
    assert not has_connector_diagram(Presentation(WORKSPACE).slides[26])
