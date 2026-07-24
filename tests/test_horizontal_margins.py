"""Width budgets must honor the shape's real horizontal insets and the
paragraph-level bullet indent (a:pPr marL/indent) — a flat 0.2" guess ate most
of the single-line safety margin on real icon-bullet slides (lIns=0 + rIns=0.1"
+ marL=0.25"), letting an 18pt line through that the renderer then wrapped,
which knocked the whole icon column off its rows (deck f78f14b1, slide 2)."""

from conftest import SURVEY_31, requires

from pptx import Presentation

from fonts.metrics import FontResolver
from generator import generator as G
from generator.text_fit import (
    fit_font_size_single_line,
    frame_horizontal_margins_in,
    horizontal_margins_in,
    paragraph_indent_in,
)
from template_parser.parser import extract_theme

# The exact line that wrapped in production: 418.6pt at 18pt in Arial-donor
# metrics, against a 6.15"-wide body with lIns=0, rIns≈0.1", marL=0.25".
WRAPPED_LINE = "Неконсистентность оформления между авторами"
BULLETS = [
    "Долгая ручная подготовка презентаций вручную",
    "Нехватка ресурсов дизайнеров на рутине",
    WRAPPED_LINE,
    "Потеря уникального фирменного стиля",
]
ICON_SLIDE_IDX = 28  # 4 icon-marker bullets


@requires(SURVEY_31)
def test_margins_read_real_insets_and_indent():
    prs = Presentation(SURVEY_31)
    body = next(s for s in prs.slides[ICON_SLIDE_IDX].shapes if "439" in s.name)
    # lIns=0 + rIns≈0.1" — well below the two-side default of 0.2".
    assert 0.05 < frame_horizontal_margins_in(body) < 0.15
    # Every content paragraph reserves marL≈0.25" for the icon column.
    p = body.text_frame.paragraphs[0]
    assert 0.2 < paragraph_indent_in(p) < 0.3
    # Combined: the budget shrinks by ~0.35", not the flat 0.2".
    assert 0.3 < horizontal_margins_in(body) < 0.4


@requires(SURVEY_31)
def test_single_line_fit_rejects_size_that_renderer_wraps():
    prs = Presentation(SURVEY_31)
    body = next(s for s in prs.slides[ICON_SLIDE_IDX].shapes if "439" in s.name)
    resolver = FontResolver(SURVEY_31, theme=extract_theme(SURVEY_31))
    metrics = resolver.metrics_for("+mn-lt")
    assert metrics is not None
    fitted = fit_font_size_single_line(
        [WRAPPED_LINE], body.width, 18,
        metrics=metrics, margins_in=horizontal_margins_in(body),
    )
    # 18pt measures 418.6pt against a real usable width of 417.6pt — any
    # budget that accepts 18 here is lying about the margins.
    assert fitted is not None and fitted.pt < 18


@requires(SURVEY_31)
def test_icon_slide_fill_lands_single_line_sizes():
    prs = Presentation(SURVEY_31)
    slide = prs.slides[ICON_SLIDE_IDX]
    resolver = FontResolver(SURVEY_31, theme=extract_theme(SURVEY_31))
    G._fill_bullet_list(
        slide, {"title": "Проблемы", "bullets": BULLETS}, set(), resolver=resolver
    )
    body = next(s for s in slide.shapes if "439" in s.name)
    sizes = {
        r.font.size.pt
        for p in body.text_frame.paragraphs
        for r in p.runs
        if r.font.size and r.text.strip()
    }
    assert sizes and max(sizes) < 18
