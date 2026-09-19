"""A template slide built around a native TABLE is not offered for filling.

Nothing fills tables yet (backlog item 5), so a slide matched for its text
boxes ships its table as the designer left it: VK Education's slide 37 went out
as our bullets next to «Заголовок столбца, млн», «Акцент», «Строка» and demo
figures 15/10/14/4 (iter109). The matcher picked it second for bullet_list,
and two more table slides sat in the title family.

The exclusion lives in build_spec, which runs on every generation — a rule at
parse time would never reach templates whose classification is cached. The
role map here is synthetic (every slide offered as every role's candidate), so
the test does not depend on a cached parse.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from template_spec.builder import build_spec

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")


def _table_slides(prs):
    return {i for i, slide in enumerate(prs.slides)
            if any(getattr(s, "has_table", False) and s.has_table for s in slide.shapes)}


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
@pytest.mark.parametrize("role", ["bullet_list", "title"])
def test_no_family_offers_a_slide_with_a_table(role):
    prs = Presentation(VK_EDU)
    tables = _table_slides(prs)
    assert tables, "fixture: VK Education is expected to carry table slides"

    spec = build_spec(VK_EDU, {i: role for i in range(len(prs.slides))})
    offered = {s["idx"] for fam in spec["families"] for s in fam["slides"]}
    assert offered, "the family must not be emptied by the rule"
    assert not offered & tables, f"table slides offered as {role}: {sorted(offered & tables)}"
