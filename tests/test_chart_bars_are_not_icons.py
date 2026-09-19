"""The bars of a chart built from pictures are not bullet icons.

VK Tech's chart slide draws its bars as 22 separate pictures, 0.76in wide and
1.75–3.29in tall, with the values and captions as white text boxes on top of
them. _find_bullet_icons took every picture narrower than 1in to the left of
the body for a list marker, moved the «icons» line by line and deleted the
«surplus»: 19 of 22 bars went, and the white labels were left white on white —
eight invisible numbers on iter108's slide 7, two thirds of the slide blank.

Production path on the real template: generate() with a stats block on the
template slide carrying the most tall narrow pictures. On HEAD three bars are
left; fixed, all of them.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from generator.generator import generate

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")


def _bars(slide):
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            and s.width < Inches(1) and s.height >= Inches(1)]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_a_stats_fill_keeps_every_bar_of_a_picture_chart(tmp_path):
    prs = Presentation(VK_TECH)
    idx = max(range(len(prs.slides)), key=lambda i: len(_bars(prs.slides[i])))
    template_bars = len(_bars(prs.slides[idx]))
    assert template_bars >= 10, "fixture: expected VK Tech's chart slide built of bar pictures"

    pairs = [[f"+{n}%", f"Показатель номер {n}"] for n in range(11, 23)]
    block = {"type": "stats_kpi", "title": "Рост показателей после запуска", "stats": pairs}
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [(block, idx)], out)

    left = len(_bars(Presentation(out).slides[0]))
    assert left == template_bars, f"{template_bars - left} of {template_bars} bars removed as «icons»"
