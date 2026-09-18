"""Leftover template placeholder text is a categorical defect, groups included.

The first deck generated on the VK Tech template shipped two slides with four
untouched «Заголовок / Подзаголовок / Текст» cards under our own text, and the
harness scored the deck 94.9: every criterion reads top-level slide.shapes, and
the template keeps each card as a GROUP. Appendix 1 of the VK Tech brief lists
«остался текст-заглушка» under integrity; this is that check.
"""

from conftest import ROOT  # noqa: F401

from pptx import Presentation
from pptx.util import Inches, Pt

from evaluation import deterministic


def _deck(path, groups, loose=()):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.62)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.8))
    box.text_frame.text = "Наш настоящий заголовок"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    for i, text in enumerate(loose):
        b = slide.shapes.add_textbox(Inches(0.5), Inches(1.5 + i * 0.6), Inches(6), Inches(0.5))
        b.text_frame.text = text
    for g, texts in enumerate(groups):
        grp = slide.shapes.add_group_shape()
        for i, text in enumerate(texts):
            b = grp.shapes.add_textbox(Inches(1 + g * 3), Inches(2 + i * 0.5), Inches(2.5), Inches(0.4))
            b.text_frame.text = text
    prs.save(path)
    return path


def test_placeholders_inside_groups_are_found(tmp_path):
    deck = _deck(str(tmp_path / "d.pptx"),
                 groups=[["Заголовок", "Подзаголовок", "Текст"]])
    scores = deterministic.evaluate(deck)
    assert scores["dop_no_placeholders"]["score"] == 1
    assert "слайд 1" in scores["dop_no_placeholders"]["detail"]


def test_appendix_patterns_are_found_at_top_level(tmp_path):
    deck = _deck(str(tmp_path / "d.pptx"), groups=[],
                 loose=["Lorem ipsum dolor sit amet", "Срок: XXX дней"])
    assert deterministic.evaluate(deck)["dop_no_placeholders"]["score"] == 1


def test_real_text_that_merely_starts_with_a_label_is_not_flagged(tmp_path):
    """«Текст договора» is a heading, not a slot label: the match is on the
    whole normalised string, not a prefix."""
    deck = _deck(str(tmp_path / "d.pptx"), groups=[["Текст договора", "Пункт 4.2 расторгнут"]],
                 loose=["Заголовок отчёта за квартал"])
    assert deterministic.evaluate(deck)["dop_no_placeholders"]["score"] == 5
