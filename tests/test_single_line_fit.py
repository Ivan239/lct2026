"""Single-line guarantee for icon-marker lists: the renderer wraps a few
percent earlier than fontTools predicts, so instead of predicting the wrap we
shrink until no item CAN wrap (with margin) — and fall back to the wrapped
layout when that's unreachable."""

from conftest import SURVEY_31, requires

from pptx import Presentation

from fonts.metrics import FontResolver
from generator.generator import _fill_bullet_list, _pick_body_shape
from generator.text_fit import SINGLE_LINE_SAFETY, _usable_width_in, fit_font_size_single_line
from template_parser.parser import extract_theme

# The bullet that actually wrapped in a rendered preview while the estimator
# said "fits": 418.6pt measured vs 428.1pt budget at 18pt — a 2% margin the
# renderer ate. The single-line fitter must therefore pick < 18pt for it.
BORDERLINE = "Неконсистентность оформления между авторами"


@requires(SURVEY_31)
def test_borderline_bullet_forced_below_base():
    resolver = FontResolver(SURVEY_31, extract_theme(SURVEY_31))
    metrics = resolver.metrics_for("+mn-lt")
    prs = Presentation(SURVEY_31)
    body = _pick_body_shape(prs.slides[28], set())

    size = fit_font_size_single_line([BORDERLINE], body.width, 18, metrics=metrics)
    assert size is not None and size.pt < 18
    budget = _usable_width_in(body.width / 914400) * 72 * SINGLE_LINE_SAFETY
    assert metrics.text_width_pt(BORDERLINE, size.pt) <= budget


@requires(SURVEY_31)
def test_absurdly_long_item_falls_back():
    resolver = FontResolver(SURVEY_31, extract_theme(SURVEY_31))
    metrics = resolver.metrics_for("+mn-lt")
    prs = Presentation(SURVEY_31)
    body = _pick_body_shape(prs.slides[28], set())
    long_text = "очень длинный пункт " * 20
    assert fit_font_size_single_line([long_text], body.width, 18, metrics=metrics) is None


@requires(SURVEY_31)
def test_filler_applies_single_line_size():
    prs = Presentation(SURVEY_31)
    resolver = FontResolver(SURVEY_31, extract_theme(SURVEY_31))
    data = {"title": "Проблемы", "bullets": [
        BORDERLINE, "Ручная адаптация", "Перегруженные дизайнеры", "Долгое согласование"]}
    _fill_bullet_list(prs.slides[28], data, set(), resolver=resolver)

    body = _pick_body_shape(prs.slides[28], set())
    sizes = {
        run.font.size.pt
        for p in body.text_frame.paragraphs for run in p.runs
        if run.font.size and run.text.strip()
    }
    assert sizes and max(sizes) < 18  # shrunk so the borderline line can't wrap
