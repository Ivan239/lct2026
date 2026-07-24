"""Regression for a real crash report: "Не удалось собрать презентацию:
not enough values to unpack (expected 2, got 1)". Root cause — the legacy
parse_brief/resize_block path (content_parser/parser.py) validated the
LENGTH of "stats" but never the shape of each pair, so a malformed single-
element item (["x10"] instead of ["x10", "ускорение"]) sailed through and
crashed generator._fill_stats_kpi's `for num, label in stats` unpack.

Fixed at both ends: _sanitize_pairs strips malformed pairs at the source
(parser.py), and _fill_stats_kpi filters defensively too, so no future
content source can reintroduce the same crash."""

import json

from conftest import SURVEY_31, requires

from pptx import Presentation

from content_parser.parser import RESIZE_FIELD_BY_TYPE, _sanitize_pairs, parse_brief, resize_block
from generator.generator import _fill_stats_kpi, generate, get_capacity


class FakeClient:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = 0

    def chat(self, messages, model=None, **kwargs):
        self.calls += 1
        return {"choices": [{"message": {"content": self.responses.pop(0)}}]}


def test_sanitize_pairs_drops_malformed_stat():
    block = {"type": "stats_kpi", "title": "т", "stats": [["x10", "ускорение"], ["3"], ["0", "часов"]]}
    cleaned = _sanitize_pairs(block)
    assert cleaned["stats"] == [["x10", "ускорение"], ["0", "часов"]]


def test_parse_brief_sanitizes_malformed_stats():
    raw = json.dumps([
        {"type": "title", "title": "Продукт", "subtitle": "Подзаголовок"},
        {"type": "stats_kpi", "title": "Метрики", "stats": [["x10", "ускорение"], ["3"]]},
    ])
    blocks = parse_brief(FakeClient([raw]), "бриф", models=["GigaChat"])
    stats_block = next(b for b in blocks if b["type"] == "stats_kpi")
    assert stats_block["stats"] == [["x10", "ускорение"]]


def test_resize_block_sanitizes_result():
    # Original has 1 stat, target wants 2 -> resize_block must actually call
    # the model (its own short-circuit only skips when counts already match).
    original = {"type": "stats_kpi", "title": "т", "stats": [["x10", "у"]], "order": 0}
    bad_resize = json.dumps({"title": "т", "stats": [["x10", "у"], ["3"], ["0", "ч"]]})
    resized = resize_block(FakeClient([bad_resize]), original, target_count=2, models=["GigaChat"])
    # 3 raw items, 1 malformed -> sanitized length 2 matches target -> accepted
    assert resized["stats"] == [["x10", "у"], ["0", "ч"]]


def test_fill_stats_kpi_never_crashes_on_malformed_pair():
    """Direct unit check on the consumer, independent of which upstream path
    produced the bad data."""
    from pptx import Presentation as _Presentation

    prs = _Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = {"title": "Метрики", "stats": [["x10", "ускорение"], ["3"], ["0", "часов"]]}
    _fill_stats_kpi(slide, data, set())  # must not raise


@requires(SURVEY_31)
def test_generate_survives_malformed_stats_through_full_pipeline(tmp_path):
    """End-to-end: a plan carrying the exact malformed shape from the crash
    report must not blow up generate()."""
    plan = [
        ({"type": "stats_kpi", "title": "Метрики", "stats": [["x10", "ускорение"], ["3"]]}, 1),
    ]
    out = str(tmp_path / "malformed.pptx")
    generate(SURVEY_31, plan, out)  # must not raise
    assert Presentation(out).slides
