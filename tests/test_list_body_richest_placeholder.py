"""A list goes into the list, not into the label above it.

VK Tech's contents slide: title, a one-line label «Содержание» (BODY idx 1)
over a four-item list (BODY idx 2), a photo box on the left. The body picker
took idx 1 — the label: capacity 1, the list blanked, the single bullet cut to
«Средне», and the label then aligned onto the title's left edge, printed
through the title (iter123).

Two rules, both general: another body placeholder holding MORE content slots
is the body; and a box beside the title is not moved onto it.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import generate, get_capacity

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
CONTENTS_SLIDE = 8
BULLETS = ["Ручной сбор показателей", "Разные форматы выгрузок",
           "Согласование до двух недель", "Сверка вручную перед отчётом"]


def _rect(s):
    return s.left, s.top, s.left + s.width, s.top + s.height


def _overlap(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_contents_slide_list_is_filled_and_title_left_alone(tmp_path):
    slide = Presentation(VK_TECH).slides[CONTENTS_SLIDE]
    assert slide.slide_layout.name.endswith("Содержание"), "fixture: VK Tech contents slide"
    assert get_capacity(slide, "bullet_list") == 4

    out = str(tmp_path / "deck.pptx")
    block = {"type": "bullet_list", "title": "Проблема поиска в мессенджере", "bullets": BULLETS}
    generate(VK_TECH, [(block, CONTENTS_SLIDE)], out)
    slide = Presentation(out).slides[0]

    holders = [s for s in slide.shapes if s.has_text_frame
               and all(b in s.text_frame.text for b in BULLETS)]
    assert holders, [s.text_frame.text for s in slide.shapes if s.has_text_frame]

    title = next(s for s in slide.shapes if s.has_text_frame
                 and s.text_frame.text == block["title"])
    on_title = [s.text_frame.text for s in slide.shapes
                if s.has_text_frame and s.text_frame.text.strip()
                and s.shape_id != title.shape_id and _overlap(_rect(s), _rect(title))]
    assert not on_title, f"printed through the title: {on_title}"
