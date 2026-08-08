"""Slot-based paragraph filling (deck custom_78dc..., slide 2 — the reference
defect): hand-built templates hold list content in a few CONTENT paragraphs
separated by spacer paragraphs and trailing junk (including '\\x0b'-only runs
that render as extra blank lines). Sequential filling wrote items into the
spacers (ragged marL, collapsed rhythm), and clearing-but-keeping the tail
left enough phantom lines to overfill the box — the renderer then silently
autofit-shrank the text, dragging bullet icons off their rows."""

import pytest

from conftest import SURVEY_69, TJ_MONO, TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from fonts.metrics import FontResolver
from generator import generator as G
from template_parser.parser import extract_theme

STATS_SLIDE_IDX = 2  # single-box stats fallback: 4 content slots, spacers, 3 icons
STATS = [("15 мин", "на презентацию"), ("2 мес", "разработки"), ("1 GPU", "на прод")]


def _fill(prs):
    slide = prs.slides[STATS_SLIDE_IDX]
    resolver = FontResolver(SURVEY_69, theme=extract_theme(SURVEY_69))
    G._fill_stats_kpi(
        slide, {"title": "Экономика", "stats": STATS}, set(), resolver=resolver
    )
    return slide


@requires(SURVEY_69)
def test_no_phantom_paragraphs_after_fill():
    prs = Presentation(SURVEY_69)
    slide = _fill(prs)
    body = next(s for s in slide.shapes if "354" in s.name)
    paragraphs = list(body.text_frame.paragraphs)
    texts = ["".join(r.text for r in p.runs) for p in paragraphs]
    # The last paragraph must be our last item — no trailing spacers/junk left
    # to overfill the box and trigger renderer-side autofit shrinking.
    assert texts[-1].strip()
    assert sum(1 for t in texts if t.strip()) == len(STATS)
    # The '\x0b'-only leftover paragraph (renders as 3 blank lines) is gone.
    assert not any("\x0b" in t for t in texts)


@requires(SURVEY_69)
def test_items_share_left_edge_and_rhythm():
    prs = Presentation(SURVEY_69)
    slide = _fill(prs)
    body = next(s for s in slide.shapes if "354" in s.name)
    paragraphs = list(body.text_frame.paragraphs)
    content = [
        i for i, p in enumerate(paragraphs)
        if "".join(r.text for r in p.runs).strip()
    ]
    # One left edge: the template's mixed marL leftovers (0.25" on the first
    # prose paragraph, none on later ones) must be unified.
    marls = {
        paragraphs[i]._p.pPr.get("marL") if paragraphs[i]._p.pPr is not None else None
        for i in content
    }
    assert len(marls) == 1
    # Uniform rhythm: the same number of spacer paragraphs between every pair
    # of consecutive items (the template's own gaps were 2-then-1).
    gaps = {b - a for a, b in zip(content, content[1:])}
    assert len(gaps) == 1


@requires(SURVEY_69)
def test_icons_track_spacer_rhythm():
    prs = Presentation(SURVEY_69)
    slide = _fill(prs)
    body = next(s for s in slide.shapes if "354" in s.name)
    icons = [
        s for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width < Inches(1.0)
        and s.left <= body.left
    ]
    assert len(icons) == len(STATS)
    tops = sorted(i.top for i in icons)
    pitches = [b - a for a, b in zip(tops, tops[1:])]
    # With one spacer line between items every icon pitch is equal, and it must
    # exceed a single content line (24pt ≈ 0.46" at natural spacing) because it
    # spans content + spacer.
    assert max(pitches) - min(pitches) < Inches(0.05)
    assert pitches[0] > Inches(0.5)


@pytest.mark.parametrize("template", [TJ_UNIVERSAL, TJ_MONO], ids=["universal", "mono"])
def test_slot_numbering_badges_survive_and_renumber(template, tmp_path):
    """The template numbers its list items («01 Название пункта … 06»), but
    _clear_unclaimed_text blanks every text shape nobody claimed, so generated
    decks shipped bare items with the numbering silently erased — brand
    furniture, gone.

    Renumbering matters as much as keeping them: a grid is not always numbered
    in reading order. The T-Zh universal slide is a 3x2 grid numbered DOWN the
    columns, so its top row reads «01 03 05» — right for six items, and a
    visible bug for the three a bullet deck actually ships."""
    block = {"type": "bullet_list", "title": "Итоги квартала",
             "bullets": ["Забота о клиенте", "Единый стандарт", "Быстрый отклик"]}
    out = str(tmp_path / "badges.pptx")
    G.generate(template, [(block, 2)], out)

    slide = list(Presentation(out).slides)[0]
    badges = sorted(
        shape.text_frame.text.strip() for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip().isdigit()
    )
    assert badges == ["01", "02", "03"], f"numbering shipped as {badges}"


@requires(TJ_UNIVERSAL)
def test_full_grid_numbers_every_item_in_reading_order(tmp_path):
    """Six items in the 3x2 grid must read 01..06 down the page. The bottom
    row's badges live in the chrome band, where _renumber_static_slide_numbers
    rewrites any all-digit box to the PAGE number — it turned «04»/«05»/«06»
    into three «01»s until badges we manage started carrying a marker name."""
    block = {"type": "bullet_list", "title": "Итоги квартала",
             "bullets": ["Забота о клиенте", "Единый стандарт", "Быстрый отклик",
                         "Прозрачный отчёт", "Гибкий тариф", "Поддержка 24/7"]}
    out = str(tmp_path / "grid6.pptx")
    G.generate(TJ_UNIVERSAL, [(block, 2)], out)

    slide = list(Presentation(out).slides)[0]
    badges = sorted(
        shape.text_frame.text.strip() for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip().isdigit()
    )
    assert badges == ["01", "02", "03", "04", "05", "06"], f"numbering: {badges}"
