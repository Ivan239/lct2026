"""A display figure is never the slide's title, however large it is set.

The position discount alone stopped being enough once iter70 gave the figure the
size the template sets it at: 72pt against an 18pt heading on the T-Zh study
deck, and half of 72 still beats 18. So «-40%» was read as the title, the real
heading was counted as body text and reported as a stray widow, and criterion
2.1 never saw that slide's title at all — it scored 2 long titles where there
were 3.

The predicate is deliberately NOT "short and contains a digit": «Top 7» is a
heading on a real template and would have been caught by that. Half the
non-space characters must be figure characters — 25% for «Top 7» and «4 дня»,
100% for «-40%» and «20 227 000».
"""

from conftest import SURVEY_31, TJ_TEMPLATE, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from generator.generator import _pick_title_shape

# _is_display_figure is imported inside the tests that need it: at module level
# its absence aborts the file, and "the helper is missing" says far less than
# "«-40%» was picked as the slide's title".


def _slide_with(tmp_shapes):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(5.62)))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for text, size_pt, top_in in tmp_shapes:
        box = slide.shapes.add_textbox(Emu(int(Inches(0.4))), Emu(int(Inches(top_in))),
                                       Emu(int(Inches(7.2))), Emu(int(Inches(0.8))))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size_pt)
    return slide


def test_a_figure_loses_to_the_heading_even_at_four_times_the_size():
    slide = _slide_with([("Результаты пилотного внедрения", 18, 1.22),
                         ("-40%", 72, 3.21)])
    title = _pick_title_shape(slide, set())
    assert title.text_frame.text == "Результаты пилотного внедрения", title.text_frame.text


def test_a_heading_that_happens_to_contain_a_digit_is_still_a_heading():
    from generator.generator import _is_display_figure

    for text in ("Top 7", "4 дня", "Итоги 2025 года"):
        assert not _is_display_figure(_slide_with([(text, 40, 1.0)]).shapes[0]), text


def test_real_figures_are_recognised():
    from generator.generator import _is_display_figure

    for text in ("-40%", "+18", "x3", "20 227 000"):
        assert _is_display_figure(_slide_with([(text, 40, 1.0)]).shapes[0]), text


@requires(SURVEY_31)
def test_the_surveys_own_heading_is_untouched():
    """«Top 7» heads a real slide of a real template — the case the naive
    predicate would have broken."""
    slide = list(Presentation(SURVEY_31).slides)[24]
    title = _pick_title_shape(slide, set())
    assert title is not None and title.text_frame.text.strip() == "Top 7"


@requires(TJ_TEMPLATE)
def test_a_slide_whose_only_text_is_a_figure_still_gets_a_title():
    """The discount lowers the score, it does not disqualify: with nothing else
    to pick, the figure is still the most title-like shape on the slide."""
    slide = _slide_with([("20 227 000", 92, 3.31)])
    assert _pick_title_shape(slide, set()) is not None
