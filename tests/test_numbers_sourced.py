"""«Все цифры и факты со слайда есть в исходных материалах?» (Appendix 1 of the
VK Tech brief) as a deterministic check, dop_numbers_sourced.

Until iter108 only the eye asked it. iter105 found «+5 позиций» and «средний
чек +12%» by reading the slides; iter108's VK Tech deck carried twenty numbers
against a brief with none — eight of them white on white under a chart whose
bars had been removed, i.e. invisible on the render and still in the file the
client edits — and the harness said 89.4.

The reference is the content package's numbers, or the brief's when there is
no package. What these tests pin: numbers are found where the reader finds
them (inside groups and table cells, which the top-level criteria never read),
compared by value (a sign or a unit is rephrasing, not invention), and the
template's furniture (a page number, «01» of a step) is not taken for a claim.
"""

import os

from conftest import ROOT

from pptx import Presentation
from pptx.util import Inches, Pt

from content_package import extract_numbers, load_package
from evaluation import rubric
from evaluation.deterministic import numbers_sourced_score, unsourced_numbers


def _box(shapes, left, top, width, height, text):
    box = shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    return box


def _deck(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide.shapes, 0.5, 0.5, 8, 1, "Рост конверсии 25 % за квартал")
    _box(slide.shapes, 0.5, 1.6, 8, 1, "Экономия +30 часов, шаг 01 из плана")
    group = slide.shapes.add_group_shape()
    _box(group.shapes, 0.5, 2.8, 4, 1, "Стоимость лида −40%")
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(4.0), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Метрика", "Значение"
    table.cell(1, 0).text, table.cell(1, 1).text = "Точность прогноза", "95%"
    # Chrome: a small box at the bottom edge — the template's year stamp.
    _box(slide.shapes, 12.0, 7.2, 1.0, 0.25, "2025")
    prs.save(path)
    return path


def test_numbers_in_groups_and_tables_are_found_and_chrome_is_not(tmp_path):
    deck = _deck(str(tmp_path / "deck.pptx"))
    source = extract_numbers("Рост конверсии на +25%, экономия 30 часов в месяц.")
    found = unsourced_numbers(deck, source)
    assert found == [(1, "-40%"), (1, "95%")], found


def test_every_number_from_the_source_scores_five(tmp_path):
    deck = _deck(str(tmp_path / "deck.pptx"))
    source = extract_numbers("25%, 30 часов, стоимость лида −40%, точность 95%")
    result = numbers_sourced_score(deck, source)
    assert result["score"] == 5, result


def test_invented_numbers_are_categorical_and_named(tmp_path):
    deck = _deck(str(tmp_path / "deck.pptx"))
    one = numbers_sourced_score(deck, extract_numbers("25% 30 часов 95%"))
    assert one["score"] == 3 and "слайд 1: -40%" in one["detail"], one
    none = numbers_sourced_score(deck, [])
    assert none["score"] == 1, none      # four claims; «01» and «2025» are not claims
    assert "слайд 1: 25%, +30часов, -40%, 95%" in none["detail"], none


def test_weeks_are_a_unit_so_two_weeks_is_a_claim():
    """«2 недели» used to come out as a bare «2» — an enumerator, not checked."""
    assert extract_numbers("срок внедрения 2 недели") == ["2недели"]


def test_the_check_counts_in_the_content_bucket():
    assert "dop_numbers_sourced" in rubric.BUCKETS["content"]["criteria"]
    assert rubric.CRITERIA["dop_numbers_sourced"][1] == "det"


def test_a_package_fact_on_a_slide_is_sourced(tmp_path):
    """The package's numbers are the reference: a fact copied onto a slide
    passes, one that is not in the package does not."""
    pkg = load_package(os.path.join(ROOT, "samples", "content_packages", "feature_smart_search"))
    fact = next(f for f in pkg["facts"] if extract_numbers(f))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide.shapes, 0.5, 0.5, 8, 1, fact)
    _box(slide.shapes, 0.5, 2.0, 8, 1, "Выручка выросла на 777%")
    path = str(tmp_path / "pkg_deck.pptx")
    prs.save(path)
    assert unsourced_numbers(path, pkg["numbers"]) == [(1, "777%")]
