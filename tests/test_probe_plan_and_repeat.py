"""The probe's plan builder is public, and it can repeat roles.

Every ad-hoc script in this project that assembled a plan by hand got it wrong
the same way: iter42 ignored the plan's `count` and read a capacity fix as a
no-op, iter56 skipped the character budgets, and iter85 skipped them again and
read the renderer's own autofit shrink as a product defect — three times the
same class of mistake, each costing an iteration. build_plan is the one path,
so the correct way is now also the easy way.

repeat_outline covers the axis no content preset reaches: a role used twice
sends the matcher back to the same template slide, i.e. the CLONE path, where
this project's nastiest bugs lived (orphaned sldId entries and a notesSlide
claimed by several slides, both of which made PowerPoint demand repair).
"""

import importlib.util
import os

import pytest

from conftest import ROOT, TJ_TEMPLATE, requires

spec = importlib.util.spec_from_file_location(
    "offline_probe", os.path.join(ROOT, "scripts", "offline_probe.py"))
probe = importlib.util.module_from_spec(spec)
spec.loader.exec_module(probe)


def test_repeat_multiplies_only_content_roles():
    out = probe.repeat_outline(probe.OUTLINE, 3)
    roles = [item["role"] for item in out]
    assert roles.count("title") == 1, "the cover must not repeat"
    assert roles.count("closing") == 1, "the closing must not repeat"
    assert roles.count("bullet_list") == 3
    assert roles.count("stats_kpi") == 3


def test_repeat_of_one_changes_nothing():
    assert probe.repeat_outline(probe.OUTLINE, 1) == probe.OUTLINE


@requires(TJ_TEMPLATE)
def test_build_plan_applies_count_and_budget():
    """The two things hand-rolled plans keep forgetting, asserted together."""
    from design_system.style_profile import build_measured_profile
    from template_spec.builder import build_spec

    archetypes, pngs = probe._cached_parse("custom_f496182bb15f42bb")
    if archetypes is None:
        pytest.skip("no cached parse for the study template")
    profile = build_measured_profile(pngs, archetypes) if pngs else None
    spec_ = build_spec(TJ_TEMPLATE, archetypes, style_profile=profile)

    plan = probe.build_plan(spec_, probe.OUTLINE, probe.LONG)
    bullets = next(b for b, _ in plan if b.get("type") == "bullet_list")

    original = probe.LONG["bullet_list"]["bullets"]
    assert len(bullets["bullets"]) <= len(original), "count was ignored"
    assert any(len(item) < len(orig) for item, orig in zip(bullets["bullets"], original)), \
        "no bullet was trimmed: the character budget was not applied"


@requires(TJ_TEMPLATE)
def test_a_repeated_role_reuses_a_template_slide():
    """The point of the axis: the same slide index appears more than once, which
    is what makes the generator clone it."""
    from design_system.style_profile import build_measured_profile
    from template_spec.builder import build_spec

    archetypes, pngs = probe._cached_parse("custom_f496182bb15f42bb")
    if archetypes is None:
        pytest.skip("no cached parse for the study template")
    profile = build_measured_profile(pngs, archetypes) if pngs else None
    spec_ = build_spec(TJ_TEMPLATE, archetypes, style_profile=profile)

    plan = probe.build_plan(spec_, probe.repeat_outline(probe.OUTLINE, 3), probe.LONG)
    indices = [idx for _, idx in plan]
    assert len(indices) > len(set(indices)), f"no slide is reused: {indices}"


def test_repeated_blocks_carry_different_items():
    """A repeat that only changes the heading is not a repeat: the per-slot
    character budget trims a suffix away, and the slides end up with identical
    bullets under different titles — a fixture that stops tripping the duplicate
    check without actually varying. The items are replaced outright.
    """
    first = probe.vary_block(probe.LONG["bullet_list"], 0)
    second = probe.vary_block(probe.LONG["bullet_list"], 1)

    assert first["bullets"] != second["bullets"]
    assert first["title"] != second["title"]
    # …and the difference must survive a trim to a short slot budget.
    from content_parser.two_phase import _enforce_text_budgets

    trimmed = [_enforce_text_budgets(dict(b), "bullet_list", 28)["bullets"]
               for b in (first, second)]
    assert trimmed[0] != trimmed[1], trimmed


def test_variation_leaves_the_first_block_alone():
    assert probe.vary_block(probe.LONG["stats_kpi"], 0) == dict(probe.LONG["stats_kpi"])


def test_the_wrap_overflow_count_ignores_the_designers_own_boxes(tmp_path):
    """A template's own text exceeds its frame 27-43 times per deck by intent
    (iter61). Counting those would report the designer's habits as our defect —
    the first version of this number printed 23 for a six-slide deck, most of it
    not ours."""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    from fonts.metrics import FontResolver
    from generator.text_fit import SINGLE_LINE_SAFETY, _usable_width_in
    from template_parser.parser import extract_theme

    # The box has to sit ON the wrap boundary, since that is what the counter
    # measures: a box whose text is simply taller than its frame is the
    # designer's own habit and is not counted at all (see
    # test_probe_wrap_overflow.py). The first fixture here was a 28pt line in a
    # 3in box — over its frame at every width, so it stopped being counted when
    # the measure was narrowed, and the test failed for a reason that had
    # nothing to do with what it guards.
    text = "Собственный текст дизайнера, и это его право"
    size_pt = 18
    source = str(tmp_path / "src.pptx")

    def _build(width_in):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(5.62)))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(int(Inches(0.3))), Emu(int(Inches(0.3))),
                                       Emu(int(Inches(width_in))), Emu(int(Inches(0.4))))
        box.text_frame.word_wrap = True
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size_pt)
        box.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        prs.save(source)

    _build(3.0)
    metrics = FontResolver(source, extract_theme(source)).metrics_for("Arial")
    insets = 3.0 - _usable_width_in(3.0)
    drawn_in = metrics.text_width_pt(text, size_pt) / 72
    _build(drawn_in / ((1 + SINGLE_LINE_SAFETY) / 2) + insets)

    counted_all = probe.boxes_over_at_render_wrap(source)
    assert counted_all >= 1, "fixture: the box is supposed to overflow"

    counted_ours = probe.boxes_over_at_render_wrap(source, source, [({}, 0)])
    assert counted_ours == 0, "the designer's own text must not be counted as ours"
