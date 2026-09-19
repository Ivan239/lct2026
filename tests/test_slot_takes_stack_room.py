"""A list slot takes the room of the stack under it.

VK Tech's card slide sets «Заголовок» (14pt, one line, 0.16in) over «Текст»
in each card. The list goes into the heading boxes and the «Текст» boxes are
cleared, yet the item budget was one line — 17 characters — and items were cut
to «Требуются 2», «Планируется» (iter138). The slot now takes the stack's
room: the budget counts its lines, the box grows into it, its 64% leading —
harmless on one line — becomes 100% so wrapped lines do not overprint, and a
widow-guard glue that cannot fit the line is released (iter139).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Inches, Pt

from fonts.metrics import FontResolver
from generator.generator import generate
from template_parser.parser import extract_theme
from template_spec.builder import build_spec

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
CARDS = 25
NBSP = "\u00a0"  # a literal, not an import: on HEAD the test must fail on behaviour
ITEMS = ["Требуются 2 дополнительных сервера", "Планируется раскатка на всех пользователей",
         "Поддержка команды поиска", "Комитет утверждает бюджет"]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_card_heading_budget_counts_the_stack():
    spec = build_spec(VK_TECH, {CARDS: "bullet_list"})
    (entry,) = [s for f in spec["families"] for s in f["slides"] if s["idx"] == CARDS]
    assert entry["item_chars"] >= len("Требуются 2 дополнительных сервера"), entry


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_filled_slot_grows_with_normal_leading(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "bullet_list", "title": "Ресурсы", "bullets": ITEMS}, CARDS)], out)
    slide = Presentation(out).slides[0]
    filled = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.replace(NBSP, " ") in {i.replace(NBSP, " ") for i in ITEMS}]
    assert len(filled) == 4, [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    for s in filled:
        assert s.height > Inches(0.3), (s.text_frame.text, s.height)
        for p in s.text_frame.paragraphs:
            assert p.line_spacing is None or p.line_spacing >= 1.0, (s.text_frame.text, p.line_spacing)


def test_glue_wider_than_its_line_is_released(tmp_path):
    from pptx import Presentation as New

    from qa.geometry import unglue_overwide_pairs
    prs = New()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1.2), Inches(1))
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Требуются 2 дополнительных сервера"
    run.font.size, run.font.name = Pt(14), "Arial"
    path = str(tmp_path / "g.pptx")
    prs.save(path)
    prs = Presentation(path)
    resolver = FontResolver(path, extract_theme(path))
    assert unglue_overwide_pairs(prs, resolver)
    assert NBSP not in prs.slides[0].shapes[0].text_frame.text
