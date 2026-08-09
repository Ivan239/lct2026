"""A stats slide whose topmost box is a display FIGURE has no title slot.

_stats_title_shape takes the topmost content box, which is right where the
designer put a heading there. On the universal template's «20 227 000» slide
there is no heading at all — the layout is a 92pt figure over a caption — and
the topmost box IS the figure. Writing our heading into it produced the worst
render this project has seen in weeks: «Результаты пилотного внедрения» drawn
straight across the green sculpture, «пилотного» unreadable, on a deck scoring
91.9.

Measured across the corpus, exactly one slide picks a figure this way, so the
rule is as narrow as the defect. With the box left alone the filler puts the
number in it at the template's own 92pt and the label in the caption below —
the designer's layout.
"""

from conftest import SURVEY_31, TJ_TEMPLATE, TJ_UNIVERSAL, requires

from pptx import Presentation

from generator.generator import _content_text_shapes, _stats_title_shape, get_capacity

FIGURE_SLIDE = 8   # «20 227 000», no heading
BOARD_SLIDE = 9    # the 8-box KPI board, headed


@requires(TJ_UNIVERSAL)
def test_a_figure_is_not_taken_as_the_stats_title():
    slide = list(Presentation(TJ_UNIVERSAL).slides)[FIGURE_SLIDE]
    assert _stats_title_shape(slide, set()) is None


@requires(TJ_UNIVERSAL)
def test_the_slide_then_offers_one_pair():
    """The consequence that matters: the figure's box is available again, so the
    number goes where the template sets it at 92pt."""
    slide = list(Presentation(TJ_UNIVERSAL).slides)[FIGURE_SLIDE]
    assert get_capacity(slide, "stats_kpi") == 1


@requires(TJ_UNIVERSAL)
def test_a_headed_stats_slide_keeps_its_title():
    """The rule must not swallow real headings: the KPI board is unaffected."""
    slide = list(Presentation(TJ_UNIVERSAL).slides)[BOARD_SLIDE]
    title = _stats_title_shape(slide, set())
    assert title is not None
    assert not title.text_frame.text.strip().startswith("20 227")


@requires(TJ_TEMPLATE)
@requires(SURVEY_31)
def test_no_other_template_loses_its_stats_title():
    for template in (TJ_TEMPLATE, SURVEY_31):
        prs = Presentation(template)
        # CONTENT boxes, not any text: a full-bleed photo slide carries only
        # chrome (its page number), has no title to lose, and counting it made
        # this test fail on the fix it was written to protect.
        lost = [i for i, s in enumerate(prs.slides)
                if _content_text_shapes(s, set()) and _stats_title_shape(s, set()) is None]
        assert not lost, f"{template}: slides {lost} lost their stats title"
