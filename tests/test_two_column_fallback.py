"""_fill_two_column_comparison on a hand-built slide (plain textboxes, no
placeholders) — the geometry fallback must fill both columns instead of
silently shipping them empty."""

import pytest

from pptx import Presentation
from pptx.util import Inches, Pt

from generator.generator import _fill_two_column_comparison

DATA = {
    "title": "Сравнение подходов",
    "left_heading": "СлайдоГен",
    "left_points": ["Учится на вашем шаблоне", "Фирменный стиль точно"],
    "right_heading": "Конкуренты",
    "right_points": ["Стандартные темы", "Без брендинга клиента"],
}


@pytest.fixture
def handmade_slide():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank, no placeholders

    def box(left, top, w, h, text, size):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
        return shape

    box(0.5, 0.4, 9.0, 0.8, "Старый заголовок", 32)
    box(0.5, 1.5, 4.2, 0.5, "Колонка А", 20)
    box(0.5, 2.2, 4.2, 3.0, "старый пункт", 14)
    box(5.2, 1.5, 4.2, 0.5, "Колонка Б", 20)
    box(5.2, 2.2, 4.2, 3.0, "старый пункт", 14)
    return prs, slide


def _texts(slide):
    return {
        shape.shape_id: " | ".join(
            "".join(r.text for r in p.runs) for p in shape.text_frame.paragraphs
        )
        for shape in slide.shapes if shape.has_text_frame
    }


def test_geometry_fallback_fills_both_columns(handmade_slide):
    prs, slide = handmade_slide
    _fill_two_column_comparison(slide, DATA, set())
    all_text = " ".join(_texts(slide).values())
    assert "СлайдоГен" in all_text
    assert "Конкуренты" in all_text
    assert "Учится на вашем шаблоне" in all_text
    assert "Стандартные темы" in all_text
    # Old template text must be gone from the boxes we claimed.
    assert "Колонка А" not in all_text
    assert "старый пункт" not in all_text
