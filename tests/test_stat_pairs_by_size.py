"""A KPI board with its figures in a row and captions in a row under them.

VK Tech's two-figure slide sets «7» and «10» in display boxes (153pt and
116pt) and a 10pt «Описание» under each. _fill_stats_kpi paired boxes as
neighbours in reading order — the two display boxes first — so «4 из 6» and
its caption «модулей CRM» both went into display boxes and wrapped over each
other at 116-153pt, and «86%» landed in a 10pt caption (iter114, slide 4).

Two rules, both pinned here on the production path:
- figures and captions are told apart by size, and each figure takes the
  nearest caption;
- a figure is one line: «4 из 6» at the template's 153pt is three lines in a
  2.60in box, so its size comes down to what one line holds.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from fonts.metrics import FontResolver
from pptx.util import Emu

from generator.generator import _max_font_pt, _metrics_for_run, _reference_run, generate
from generator.text_fit import estimate_wrapped_lines, horizontal_margins_in
from template_parser.parser import extract_theme

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
STATS = [["4 из 6", "модулей CRM"], ["86%", "клиентских данных"]]


def _figure_row_slide(prs):
    """A slide with exactly two boxes ≥1.5x the size of two others, the big
    ones side by side above the small ones."""
    for i, slide in enumerate(prs.slides):
        boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()
                 and s.shape_type is not None and s.shape_type.name == "TEXT_BOX"]
        sized = [(b, _max_font_pt(b) or 0) for b in boxes]
        if len(sized) != 4 or not all(pt for _, pt in sized):
            continue
        sized.sort(key=lambda bp: -bp[1])
        big, small = sized[:2], sized[2:]
        if (min(pt for _, pt in big) >= 1.5 * max(pt for _, pt in small)
                and abs(big[0][0].top - big[1][0].top) < 91440
                and all(s.top > b.top for s, _ in small for b, _ in big)):
            return i
    return None


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_figures_go_to_the_display_boxes_on_one_line(tmp_path):
    idx = _figure_row_slide(Presentation(VK_TECH))
    assert idx is not None, "fixture: VK Tech's two-figure board not found"
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "stats_kpi", "title": "Статистика", "stats": STATS}, idx)], out)

    slide = Presentation(out).slides[0]
    text = {s.text_frame.text: s for s in slide.shapes if s.has_text_frame}
    for number, label in STATS:
        figure, caption = text.get(number), text.get(label)
        assert figure is not None and caption is not None, f"{number!r}/{label!r} not in their own boxes"
        assert _max_font_pt(figure) > 1.5 * _max_font_pt(caption), \
            f"{number!r} is set at {_max_font_pt(figure)}pt, its caption at {_max_font_pt(caption)}pt"
        assert caption.top > figure.top, f"caption of {number!r} is not under it"

    # Same metrics the filler uses (inherited bold included); the render of
    # this slide shows «4 из 6» on one line at the size this check accepts.
    resolver = FontResolver(out, extract_theme(out))
    figure = text["4 из 6"]
    size = _max_font_pt(figure)
    assert size < 153, "the figure kept the template's 153pt"
    lines = estimate_wrapped_lines("4 из 6", Emu(figure.width).inches, size,
                                   metrics=_metrics_for_run(_reference_run(figure), resolver),
                                   margins_in=horizontal_margins_in(figure))
    assert lines == 1, f"«4 из 6» at {size}pt wraps to {lines} lines"
