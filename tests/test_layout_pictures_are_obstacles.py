"""A layout's pictures are part of the canvas.

slide.shapes is the slide's own list. A logo or art panel the LAYOUT draws
renders underneath just the same, and every obstacle check in the synthesizer
read only the slide — so on VK Education's cover-style canvas (layout
«1_Титульный слайд»: the «education» wordmark at 0.72/0.76in, 3.03x0.54in, plus
a 6.74in art panel, both layout pictures; the slide itself has none) the
synthesized image_caption title was drawn straight over the logo (iter103,
seen on the render).

The check runs the production path — generate() with synth_canvas — on the real
template, and asserts geometry, not helper names: on HEAD the title lands at
T=0.76in, on top of the logo; fixed, it starts below the logo's bottom edge.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Emu

from common.synthesis import SYNTHESIZE
from generator.generator import generate

VK_EDU = os.path.join(TEMPLATES_DIR, "vk_education.pptx")


def _cover_canvas(prs):
    """First template slide on the cover layout with no text of its own."""
    return next(i for i, s in enumerate(prs.slides)
                if s.slide_layout.name.startswith("1_")
                and not any(sh.has_text_frame and sh.text_frame.text.strip()
                            for sh in s.shapes))


@pytest.mark.skipif(not os.path.exists(VK_EDU), reason=f"fixture deck missing: {VK_EDU}")
def test_synthesized_title_clears_the_layouts_logo(tmp_path):
    prs = Presentation(VK_EDU)
    canvas = _cover_canvas(prs)
    layout = list(prs.slides)[canvas].slide_layout
    logos = [s for s in layout.shapes if "PICTURE" in str(s.shape_type)
             and s.height < prs.slide_height * 0.12]
    assert logos, "fixture: the cover layout is expected to carry a small logo picture"
    logo_bottom = max(int(s.top + s.height) for s in logos)

    out = str(tmp_path / "deck.pptx")
    block = {"type": "image_caption", "title": "Единая аналитика продаж",
             "image": "экран дашборда"}
    generate(VK_EDU, [(block, SYNTHESIZE)], out, synth_canvas={0: canvas})
    slide = Presentation(out).slides[0]
    title = next(s for s in slide.shapes
                 if s.has_text_frame and "аналитика" in s.text_frame.text)
    assert int(title.top) >= logo_bottom, (
        f"title at {Emu(int(title.top)).inches:.2f}in sits on the layout logo "
        f"ending at {Emu(logo_bottom).inches:.2f}in")
