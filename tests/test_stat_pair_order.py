"""A KPI pair must land figure-first no matter which box reading order hands
over first.

Found by eye on a render of the T-Zh universal template: the stats slide read
«времени на подготовку регулярной отчётности» ABOVE «-40%», and the caption
«подключённых клиентов ежемесячно», squeezed into the figure's 24pt box, broke
mid-word as «е-/жемесячно». One cause for both: the filler wrote the number into
boxes[i*2] and the label into boxes[i*2+1], which assumes every board lists
figure then caption. This board is the other way round — its 11pt caption sits
BELOW a 24pt figure, and _content_text_shapes hands the caption over first.

The pairing is decided by font size now, so the board's own order is irrelevant.
"""

from conftest import TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu

from generator.generator import (
    _content_text_shapes,
    _fill_stats_kpi,
    _max_font_pt,
    _stats_title_shape,
)

STATS = [["-40%", "времени на подготовку регулярной отчётности"],
         ["+18", "подключённых клиентов ежемесячно"]]


@requires(TJ_UNIVERSAL)
def test_figure_goes_into_the_big_box_on_a_caption_first_board():
    prs = Presentation(TJ_UNIVERSAL)
    slide = list(prs.slides)[9]

    claimed = set()
    title = _stats_title_shape(slide, claimed)
    if title is not None:
        claimed.add(title.shape_id)
    boxes = _content_text_shapes(slide, claimed)
    sizes = [_max_font_pt(b) for b in boxes]
    tops = [b.top for b in boxes]
    # Guard the premise: this board really does hand the small caption over
    # first. Should the template change, the test below would prove nothing.
    assert sizes[0] < sizes[1] and tops[0] > tops[1], (
        f"expected a caption-first board, got sizes={sizes[:2]}")

    _fill_stats_kpi(slide, {"title": "Результаты", "stats": STATS}, set())

    for i, (num, label) in enumerate(STATS):
        # Big/small decided from the sizes measured BEFORE the fill, so this
        # asserts on the rendered outcome rather than on the helper under test:
        # on HEAD it fails with the caption's text found in the 24pt box.
        pair = sorted(((sizes[i * 2], boxes[i * 2]), (sizes[i * 2 + 1], boxes[i * 2 + 1])),
                      key=lambda kv: -kv[0])
        big, small = pair[0][1], pair[1][1]
        assert big.text_frame.text.strip() == num
        assert small.text_frame.text.strip() == label
        assert Emu(big.top).inches < Emu(small.top).inches, (
            "the figure must sit above its caption, as the designer drew it")


def test_ties_on_size_are_broken_by_position():
    """Same-size boxes still have an answer: the upper one is the figure."""
    from generator.generator import _order_stat_pair

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    lower = slide.shapes.add_textbox(Emu(0), Emu(914400 * 3), Emu(914400), Emu(914400))
    upper = slide.shapes.add_textbox(Emu(0), Emu(914400), Emu(914400), Emu(914400))
    assert _order_stat_pair(lower, upper) == (upper, lower)
    assert _order_stat_pair(upper, lower) == (upper, lower)
