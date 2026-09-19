"""A filled figure box ends where its own label begins.

Designers stack the two with an overlap and rely on the descent of a huge
display figure to keep the digits clear: VK Tech's figure box runs to 4.21in
over a label that starts at 4.03in, and «7» at 166pt sits well above it. Our
figure is fitted down to what the box's line holds, bottom-anchored text
follows the box down, and «94 секунды» printed through «среднее время поиска»
(iter141). Measured: 15 of the 55 figure/label pairs in the VK templates and
5 of 18 elsewhere are stacked like this (iter142).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import generate

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
BOARD = 41
PAIRS = [["94 секунды", "среднее время поиска"],
         ["38%", "поисков заканчивались без результата"]]


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_figure_box_stops_at_its_label(tmp_path):
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "stats_kpi", "title": "Цифры", "stats": PAIRS}, BOARD)], out)
    slide = Presentation(out).slides[0]
    texts = {s.text_frame.text.strip(): s for s in slide.shapes
             if s.has_text_frame and s.text_frame.text.strip()}
    for figure_text, label_text in PAIRS:
        figure = texts.get(figure_text)
        label = next((s for t, s in texts.items() if t.replace(" ", " ") == label_text), None)
        assert figure is not None and label is not None, sorted(texts)
        assert figure.top + figure.height <= label.top, (
            figure_text, figure.top + figure.height, label.top)


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_the_trim_never_starves_the_figure(tmp_path):
    """The box may not be cut below the line it has to print: a box shorter
    than its own text is how the renderer's autofit gets invited back in."""
    out = str(tmp_path / "deck.pptx")
    generate(VK_TECH, [({"type": "stats_kpi", "title": "Цифры", "stats": PAIRS}, BOARD)], out)
    slide = Presentation(out).slides[0]
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        size = max((r.font.size.pt for p in shape.text_frame.paragraphs
                    for r in p.runs if r.font.size), default=0)
        if size:
            assert shape.height >= size / 72 * 914400, (shape.text_frame.text, size, shape.height)
