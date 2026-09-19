"""A left-aligned text box is kept inside the slide.

VK Education's KPI slide sets its right column in boxes that end 0.45-1.46in
past the slide's right edge; the designer's short «Объяснение этого
показателя» ends inside, a generated caption ran on and was cut by the edge
(«…из-за точност», iter136). The box now ends at the slide's right margin and
the caption wraps there.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import generate

VK_EDUCATION = os.path.join(TEMPLATES_DIR, "vk_education.pptx")
KPI = 17
STATS = [["-63%", "Снижение среднего времени поиска одного сообщения"],
         ["94 с", "Среднее время поиска одного сообщения до пилота"]]


@pytest.mark.skipif(not os.path.exists(VK_EDUCATION), reason=f"fixture deck missing: {VK_EDUCATION}")
def test_filled_boxes_end_inside_the_slide(tmp_path):
    template = Presentation(VK_EDUCATION)
    assert any(s.has_text_frame and s.left + s.width > template.slide_width
               for s in template.slides[KPI].shapes), "fixture: boxes overhang the edge"

    out = str(tmp_path / "deck.pptx")
    generate(VK_EDUCATION, [({"type": "stats_kpi", "title": "Поиск до внедрения", "stats": STATS}, KPI)], out)
    prs = Presentation(out)
    over = [(s.text_frame.text[:30], round((s.left + s.width - prs.slide_width) / 914400, 2))
            for s in prs.slides[0].shapes
            if s.has_text_frame and s.text_frame.text.strip() and s.left + s.width > prs.slide_width]
    assert not over, over
