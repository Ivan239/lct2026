"""A «KPI board» of more than six pairs is a table, not a KPI board.

VK Tech offered stats slides of capacity 10-40: a table drawn with text boxes
(«Текст / 100 / 100», a header row, «Итого:»), calendars at 4-8pt and a chart
built of pictures. The plan asked the model for as many pairs as the slide
held, and the figure/caption pairing laid them across the table's cells —
scrambled on seven decks in a row (iter105-iter120). The VK Tech brief fails a
slide with more than six items; the corpus's real boards hold 1-4 pairs.

The role map here is synthetic (every slide a stats_kpi candidate), so the
test does not depend on a cached classification.
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation

from generator.generator import get_capacity
from template_spec.builder import build_spec

# The brief's density limit, stated here rather than imported: a test that
# imports the new constant fails on HEAD with ImportError, not on behaviour.
MAX_KPI_PAIRS = 6

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_no_stats_slide_is_offered_beyond_six_pairs():
    prs = Presentation(VK_TECH)
    big = [i for i, s in enumerate(prs.slides) if (get_capacity(s, "stats_kpi") or 0) > MAX_KPI_PAIRS]
    assert big, "fixture: VK Tech is expected to carry table-like boards"

    spec = build_spec(VK_TECH, {i: "stats_kpi" for i in range(len(prs.slides))})
    offered = [s for fam in spec["families"] if fam["role"] == "stats_kpi" for s in fam["slides"]]
    assert offered, "the rule must not empty the family"
    assert all((s["capacity"] or 0) <= MAX_KPI_PAIRS for s in offered), \
        [(s["idx"], s["capacity"]) for s in offered if (s["capacity"] or 0) > MAX_KPI_PAIRS]
