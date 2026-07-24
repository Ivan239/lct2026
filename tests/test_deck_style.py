"""Synthesized slides must follow the deck's OBSERVED style, not the theme:
hand-built decks (and our synthetic presets) style runs directly and leave
theme1.xml at Office defaults — synthesizing from that theme produced a white
Calibri divider inside a navy Georgia deck (preset A render sweep)."""

from conftest import PRESET_A, requires

from pptx import Presentation
from pptx.oxml.ns import qn

from common.synthesis import SYNTHESIZE
from generator.deck_style import apply_observed_style, observe_deck_style
from generator.generator import generate

NAVY, ORANGE = "#182A4A", "#F08C28"


@requires(PRESET_A)
def test_observed_style_reads_real_deck_look():
    obs = observe_deck_style(Presentation(PRESET_A))
    assert obs["bg"] == NAVY
    assert obs["accent"] == ORANGE
    assert obs["major_font"] == "Georgia"
    assert obs["minor_font"] == "Georgia"


def test_contrast_veto_replaces_unreadable_text():
    theme = {"fonts": {}, "palette": {}}
    merged = apply_observed_style(theme, {"bg": "#182A4A", "text": "#10203A"})
    assert merged["palette"]["dk1"] == "#FFFFFF"  # near-bg text replaced


@requires(PRESET_A)
def test_synthesized_divider_matches_deck_style(tmp_path):
    out = str(tmp_path / "divider.pptx")
    generate(PRESET_A, [({"type": "section_divider", "title": "На пути к решению"}, SYNTHESIZE)], out)
    slide = Presentation(out).slides[0]

    bg = slide._element.cSld.find(qn("p:bg"))
    srgb = bg.find(".//" + qn("a:srgbClr"))
    assert f"#{srgb.get('val')}" == NAVY

    title_runs = [r for sh in slide.shapes if sh.has_text_frame
                  for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
    run = title_runs[0]
    assert run.font.name == "Georgia"
    assert f"#{run.font.color.rgb}" == ORANGE
