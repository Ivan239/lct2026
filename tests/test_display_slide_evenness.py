"""A slide that is one big figure and its caption is sparse BY DESIGN.

The evenness criteria exempt the cover, the divider and the closing for that
reason. The universal template's «20 227 000» layout belongs in the same
company — it has no heading at all — but it carries the stats_kpi role, and
exempting that whole role would be wrong: the same template's other stats slide
is an eight-box KPI board, as dense as any content slide.

So the test is on the SLIDE: at most two content boxes, one of them a display
figure. Measured across every probe deck and every template, that selects the
two big-number slides of the repeat deck and the template's own slide 8, and
nothing else — a stats slide that does carry a heading has three content boxes
and stays in.

Became visible only after iter88 made the probe pass slide_roles at all; before
that the cover and closing were counted as content too and drowned it.
"""

from conftest import TJ_UNIVERSAL, requires

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from evaluation.deterministic import _content_shapes, evaluate
from generator.generator import _is_display_figure


def _deck(tmp_path, figure_text, extra_box):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(10))), Emu(int(Inches(5.62)))

    def slide_with(pairs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for text, size, top in pairs:
            box = slide.shapes.add_textbox(Emu(int(Inches(0.4))), Emu(int(Inches(top))),
                                           Emu(int(Inches(7.0))), Emu(int(Inches(0.9))))
            box.text_frame.text = text
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
            box.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        return slide

    # Dense content slides on purpose: the point is that ONE sparse figure slide
    # alongside them swings the spread past the threshold. With mild fixtures the
    # deck is even either way and the test proves nothing (it passed on HEAD).
    body = ("Ручной сбор показателей из семи независимых систем, разные форматы "
            "выгрузок у каждого подразделения и согласование до двух недель — "
            "всё это делало отчётность медленной и ненадёжной")
    for _ in range(3):
        slide_with([("Заголовок раздела", 32, 0.6), (body, 18, 2.0)])
    figure = [(figure_text, 72, 1.5), ("времени на подготовку отчётности", 18, 3.0)]
    if extra_box:
        figure.insert(0, ("Результаты внедрения", 32, 0.6))
    slide_with(figure)

    path = str(tmp_path / f"deck_{extra_box}.pptx")
    prs.save(path)
    return path


@requires(TJ_UNIVERSAL)
def test_the_templates_own_big_number_slide_is_recognised():
    slide = list(Presentation(TJ_UNIVERSAL).slides)[8]
    boxes = _content_shapes(slide)
    assert len(boxes) <= 2 and any(_is_display_figure(b) for b in boxes)


def test_a_big_number_slide_does_not_make_a_deck_uneven(tmp_path):
    lean = evaluate(_deck(tmp_path, "-40%", extra_box=False))
    assert lean["dop_distribution"]["score"] == 5, lean["dop_distribution"]["detail"]


def test_a_headed_stats_slide_still_counts(tmp_path):
    """The exemption must not swallow an ordinary content slide: the same
    figure WITH a heading has three boxes and is measured like any other."""
    headed = _deck(tmp_path, "-40%", extra_box=True)
    slide = list(Presentation(headed).slides)[3]
    assert len(_content_shapes(slide)) == 3


def test_a_short_text_slide_is_not_mistaken_for_a_figure(tmp_path):
    """Two boxes are not enough on their own — one of them must be a figure."""
    path = _deck(tmp_path, "Итоги квартала", extra_box=False)
    slide = list(Presentation(path).slides)[3]
    assert not any(_is_display_figure(b) for b in _content_shapes(slide))
