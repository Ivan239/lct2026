"""Text printed across a table's cells is a collision.

A table is a graphicFrame, so _text_collisions never saw its cells: the
WorkSpace deck whose synthesized slide and closing were drawn straight over a
table canvas scored 92.4 with a clean geometry line (iter143). Measured over
the corpus (iter144): the rule adds zero hits on all 13 untouched templates
and fires on exactly the two decks that printed our text over a template
table.
"""

from pptx import Presentation
from pptx.util import Inches, Pt

from evaluation.deterministic import evaluate


def _deck(tmp_path, text_top_in):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(3)).table
    for row in table.rows:
        for cell in row.cells:
            cell.text = "Текст"
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size, run.font.name = Pt(12), "Arial"
    box = slide.shapes.add_textbox(Inches(1.2), Inches(text_top_in), Inches(4), Inches(0.6))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Решите вопрос с интеграцией до ноября"
    run.font.size, run.font.name = Pt(28), "Arial"
    path = str(tmp_path / f"deck{text_top_in}.pptx")
    prs.save(path)
    return path


def _readability(path):
    """Criterion 1.1 — the one a text collision caps (deterministic, no render)."""
    return evaluate(path)["1.1"]


def test_text_over_a_table_is_seen(tmp_path):
    score = _readability(_deck(tmp_path, 1.5))
    assert score["score"] <= 2, score
    assert "поверх другого текста" in score["detail"], score


def test_text_beside_a_table_is_not(tmp_path):
    """The table ends at 4in; a box below it is ordinary stacking, and the
    criterion must stay clean — the same line the text-vs-text rule holds."""
    score = _readability(_deck(tmp_path, 4.6))
    assert "поверх другого текста" not in score["detail"], score


def test_cell_rects_follow_the_grid(tmp_path):
    """Cell geometry is the frame's origin plus accumulated widths/heights —
    a whole-table rectangle would swallow boxes the designer put in a gap."""
    from evaluation.deterministic import _table_cell_rects

    prs = Presentation(_deck(tmp_path, 4.6))
    rects = _table_cell_rects(prs.slides[0].shapes)
    assert len(rects) == 6, rects
    lefts = sorted({r[0] for r in rects})
    tops = sorted({r[1] for r in rects})
    assert len(lefts) == 2 and len(tops) == 3, (lefts, tops)
