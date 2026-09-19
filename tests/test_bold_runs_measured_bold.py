"""A bold run is measured in the bold face — everywhere, not just in titles.

Cyrillic bold runs ~7% wider than the regular advances a family name resolves
to, and generator._metrics_for_run — the helper behind the figure fit, the slot
budgets and the body text — asked for the family only. VK Tech's board fitted
«94 секунды» at 31pt, 101.2% of that box's line: the renderer wrapped it, and
the template's 16% display leading printed «94» over «секунды» (iter141).
"""

import os

import pytest
from conftest import TEMPLATES_DIR

from pptx import Presentation
from pptx.util import Emu

from fonts.metrics import FontResolver
from generator.generator import generate
from generator.text_fit import horizontal_margins_in
from template_parser.parser import extract_theme

VK_TECH = os.path.join(TEMPLATES_DIR, "vk_tech.pptx")
BOARD = 41  # two display figures, 166pt, line spacing 16%
FIGURE = "94 секунды"


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_bold_metrics_are_wider_than_the_family_default():
    resolver = FontResolver(VK_TECH, extract_theme(VK_TECH))
    plain = resolver.metrics_for("Play").text_width_pt(FIGURE, 31)
    bold = resolver.metrics_for("Play", bold=True).text_width_pt(FIGURE, 31)
    assert bold > plain, (plain, bold)


@pytest.mark.skipif(not os.path.exists(VK_TECH), reason=f"fixture deck missing: {VK_TECH}")
def test_a_filled_figure_fits_its_line(tmp_path):
    """The one-line rule is only kept if the size it picks really holds: at
    31pt this figure measures past its box and the renderer wraps it."""
    out = str(tmp_path / "deck.pptx")
    block = {"type": "stats_kpi", "title": "Цифры до внедрения",
             "stats": [[FIGURE, "среднее время поиска"],
                       ["38%", "поисков заканчивались без результата"]]}
    generate(VK_TECH, [(block, BOARD)], out)
    prs = Presentation(out)
    resolver = FontResolver(out, extract_theme(out))
    figures = [s for s in prs.slides[0].shapes
               if s.has_text_frame and s.text_frame.text.strip().startswith("94")]
    assert figures, [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    for shape in figures:
        run = shape.text_frame.paragraphs[0].runs[0]
        usable = (Emu(shape.width).inches - horizontal_margins_in(shape)) * 72
        metrics = resolver.metrics_for(run.font.name, bold=bool(run.font.bold))
        assert metrics.text_width_pt(run.text, run.font.size.pt) <= usable, (
            run.text, run.font.size.pt, usable)
