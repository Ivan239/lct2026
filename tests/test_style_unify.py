"""Filled bullet lists must ship with ONE run style per body: hand-designed
templates color accent phrases inside bullets (theme colors, not RGB), and
writing each new line into its paragraph's first run used to leave whole lines
randomly accent-colored depending on which run the original line started with."""

from conftest import SURVEY_31, requires

from pptx import Presentation

from generator.generator import _pick_body_shape, _run_style_key, generate

BULLETS = [
    "Автоматическая адаптация контента под фирменный стиль",
    "Снижение затрат времени на подготовку презентаций",
    "Устранение перегрузки дизайнеров рутинной работой",
    "Единообразие дизайна корпоративных презентаций",
    "Быстрая интеграция с существующими шаблонами",
    "Минимизация ошибок оформления и повышение качества",
    "Расширение поддержки новых корпоративных клиентов",
]


@requires(SURVEY_31)
def test_filled_bullets_share_one_style(tmp_path):
    # Template slide 24 is the accent-phrase case: its paragraphs mix a
    # TEXT_1-colored body style with theme-accent runs, and several lines
    # START with the accent run.
    plan = [({"type": "bullet_list", "title": "Стили", "bullets": BULLETS}, 24)]
    out = str(tmp_path / "styles.pptx")
    generate(SURVEY_31, plan, out)

    prs = Presentation(out)
    body = _pick_body_shape(prs.slides[0], set())
    styles = {
        _run_style_key(run)
        for p in body.text_frame.paragraphs
        for run in p.runs
        if run.text.strip()
    }
    assert len(styles) == 1, f"expected one unified run style, got {styles}"
