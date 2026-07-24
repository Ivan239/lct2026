"""End-to-end generate() invariants on a real uploaded deck — no LLM, no
renderer, pure python-pptx: slide cloning for repeated roles, orphan-icon
removal, and the geometric QA pass leaving a healthy result alone."""

import os

from conftest import SURVEY_31, requires

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from fonts.metrics import FontResolver
from generator.generator import generate
from qa.geometry import enforce_text_fits
from template_parser.parser import extract_theme

PLAN = [
    ({"type": "title", "title": "СлайдоГен", "subtitle": "Автогенерация презентаций"}, 23),
    ({"type": "bullet_list", "title": "Проблемы", "bullets": [
        "Ручная адаптация контента", "Перегруженные дизайнеры",
        "Неконсистентный стиль", "Долгое согласование"]}, 28),
    ({"type": "bullet_list", "title": "Решение", "bullets": [
        "Разбор шаблона", "Извлечение дизайн-системы", "Сборка презентации"]}, 28),
    ({"type": "stats_kpi", "title": "Метрики", "stats": [
        ("x10", "ускорение"), ("0 ч", "ручной работы")]}, 1),
]


def _small_pictures(slide):
    limit = Inches(1.0)
    return [
        s for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width and s.width < limit
    ]


@requires(SURVEY_31)
def test_generate_clones_orphans_and_qa(tmp_path):
    out = str(tmp_path / "generated.pptx")
    generate(SURVEY_31, PLAN, out)
    assert os.path.exists(out)

    prs = Presentation(out)
    assert len(prs.slides) == len(PLAN)

    # Slide 28 was assigned twice: the second occurrence must be a clone that
    # kept the template's icon markers (that's the whole point of cloning).
    first_bullets, second_bullets = prs.slides[1], prs.slides[2]
    assert len(_small_pictures(first_bullets)) == 4  # 4 bullets -> 4 icons kept
    assert len(_small_pictures(second_bullets)) == 3  # 3 bullets -> orphan removed

    # Stats slide: 2 stats on a 4-icon slide -> 2 icons survive.
    assert len(_small_pictures(prs.slides[3])) >= 2

    # The shipped deck must already satisfy the geometric QA pass.
    resolver = FontResolver(out, extract_theme(out))
    assert enforce_text_fits(prs, resolver) == []
