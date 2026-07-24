import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output", "templates")
os.makedirs(OUT_DIR, exist_ok=True)

TITLE_LAYOUT, CONTENT_LAYOUT, COMPARISON_LAYOUT, BLANK_LAYOUT = 0, 1, 4, 6


def style_run(run, font_name, size, color, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def set_background(slide, bg_color):
    # Every slide, not just the cover: the light text color is chosen for this
    # background, and a slide left on the default white master renders that
    # text near-invisible (the original version set it on the title slide only
    # — preset A shipped white content slides with (220,220,225) body text).
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(*bg_color)


def build(name, bg_color, accent_color, font_name, text_color):
    prs = Presentation()

    # --- title slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
    slide.shapes.title.text = "Название продукта"
    style_run(slide.shapes.title.text_frame.paragraphs[0].runs[0], font_name, 44, accent_color, bold=True)
    subtitle = slide.placeholders[1]
    subtitle.text = "Питч для инвесторов"
    style_run(subtitle.text_frame.paragraphs[0].runs[0], font_name, 20, text_color)
    set_background(slide, bg_color)

    # --- bullets slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[CONTENT_LAYOUT])
    set_background(slide, bg_color)
    slide.shapes.title.text = "Ключевые преимущества"
    style_run(slide.shapes.title.text_frame.paragraphs[0].runs[0], font_name, 32, accent_color, bold=True)
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Автоматизация подготовки презентаций"
    for extra in ["Сохранение фирменного стиля", "Сокращение времени с часов до минут"]:
        p = tf.add_paragraph()
        p.text = extra
    for p in tf.paragraphs:
        if p.runs:
            style_run(p.runs[0], font_name, 20, text_color)

    # --- two-column comparison slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[COMPARISON_LAYOUT])
    set_background(slide, bg_color)
    slide.shapes.title.text = "Продукт vs Конкуренты"
    style_run(slide.shapes.title.text_frame.paragraphs[0].runs[0], font_name, 32, accent_color, bold=True)

    left_heading, left_body = slide.placeholders[1], slide.placeholders[2]
    right_heading, right_body = slide.placeholders[3], slide.placeholders[4]
    left_heading.text_frame.text = "Наш подход"
    right_heading.text_frame.text = "Обычные AI-тулы"
    for ph in (left_heading, right_heading):
        style_run(ph.text_frame.paragraphs[0].runs[0], font_name, 20, accent_color, bold=True)

    left_points = ["Учится на вашем шаблоне", "Точная фирменная стилистика", "RU LLM-стек"]
    right_points = ["Свои generic темы", "Без учёта бренда", "Зарубежные модели"]
    for body, points in ((left_body, left_points), (right_body, right_points)):
        tf = body.text_frame
        tf.text = points[0]
        for extra in points[1:]:
            p = tf.add_paragraph()
            p.text = extra
        for p in tf.paragraphs:
            if p.runs:
                style_run(p.runs[0], font_name, 16, text_color)

    # --- stats/KPI slide (blank layout, manual textboxes) ---
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    set_background(slide, bg_color)
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
    title_box.text_frame.text = "Метрики за квартал"
    style_run(title_box.text_frame.paragraphs[0].runs[0], font_name, 30, accent_color, bold=True)

    stats = [("x10", "Ускорение подготовки"), ("3", "Шаблона в пилоте"), ("0", "Часов ручной вёрстки")]
    x = 0.5
    for num, label in stats:
        num_box = slide.shapes.add_textbox(Inches(x), Inches(2.2), Inches(2.8), Inches(1.2))
        num_box.text_frame.text = num
        style_run(num_box.text_frame.paragraphs[0].runs[0], font_name, 48, accent_color, bold=True)
        label_box = slide.shapes.add_textbox(Inches(x), Inches(3.3), Inches(2.8), Inches(0.8))
        label_box.text_frame.text = label
        style_run(label_box.text_frame.paragraphs[0].runs[0], font_name, 16, text_color)
        x += 3.1

    out_path = os.path.join(OUT_DIR, f"{name}.pptx")
    prs.save(out_path)
    print("Saved", out_path)


if __name__ == "__main__":
    build(
        "template_a_corporate",
        bg_color=(24, 42, 74),
        accent_color=(240, 140, 40),
        font_name="Georgia",
        text_color=(220, 220, 225),
    )
    build(
        "template_b_startup",
        bg_color=(250, 250, 250),
        accent_color=(120, 60, 220),
        font_name="Verdana",
        text_color=(40, 40, 40),
    )
