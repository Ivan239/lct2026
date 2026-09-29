"""Экспорт колоды в .pdf и .html (ТЗ: три формата — .pptx, .html, .pdf).

.pptx остаётся главным: в нём нативные, редактируемые объекты. PDF и HTML —
форматы для ПРОСМОТРА, и оба строятся из того же рендера LibreOffice, что и
превью: так пользователь видит в них ровно то, что видел в превью.

HTML — один самодостаточный файл: картинки слайдов вшиты в base64 (его можно
переслать и открыть без сервера), а текст каждого слайда лежит рядом НАСТОЯЩИМ
текстом — его можно выделить, найти поиском и прочитать скринридером. Одна
картинка без текста была бы скриншотом, а не документом.

Если на хосте нет шрифтов шаблона, рендер подменяет их (см.
`render._prepare_render_copy`); подмену НЕ прятать — список заменённых
семейств возвращается вызывающему и пишется в сам HTML.
"""

import base64
import html
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from rendering.render import convert_to_pdf, render_pptx_to_pngs, substituted_typefaces


def _shape_lines(shape):
    """Текст фигуры построчно — с заходом в группы и таблицы: `slide.shapes`
    видит только верхний уровень, а карточки шаблонов часто сгруппированы."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from _shape_lines(child)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                yield " | ".join(cells)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(r.text for r in paragraph.runs).replace("\x0b", " ").strip()
            if text:
                yield text


def slide_texts(pptx_path):
    """Список слайдов, у каждого — строки его текста в порядке фигур."""
    prs = Presentation(pptx_path)
    return [[line for shape in slide.shapes for line in _shape_lines(shape)]
            for slide in prs.slides]


_PAGE = """<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{ color-scheme: light dark; --bg: #f4f4f6; --card: #fff; --ink: #1d1d22; --muted: #6b6b76; }}
  @media (prefers-color-scheme: dark) {{ :root {{ --bg: #16161a; --card: #222228; --ink: #ececf1; --muted: #9a9aa6; }} }}
  body {{ margin: 0; background: var(--bg); color: var(--ink); font: 16px/1.5 system-ui, sans-serif; }}
  main {{ max-width: 1100px; margin: 0 auto; padding: 24px 16px 64px; }}
  h1 {{ font-size: 22px; margin: 0 0 4px; }}
  .note {{ color: var(--muted); font-size: 14px; margin: 0 0 24px; }}
  section {{ background: var(--card); border-radius: 10px; margin: 0 0 24px; overflow: hidden;
             box-shadow: 0 1px 3px rgba(0,0,0,.12); scroll-margin-top: 16px; }}
  section img {{ display: block; width: 100%; height: auto; }}
  .text {{ padding: 12px 16px 16px; }}
  .num {{ color: var(--muted); font-size: 13px; }}
  .text p {{ margin: 4px 0; }}
</style>
</head>
<body>
<main>
<h1>{title}</h1>
<p class="note">{count} слайдов. Под каждым слайдом — его текст. Стрелки ← → листают слайды.{fonts}</p>
{sections}
</main>
<script>
  const s = [...document.querySelectorAll('section')];
  const cur = () => s.findIndex(e => e.getBoundingClientRect().bottom > 1);
  addEventListener('keydown', e => {{
    const d = e.key === 'ArrowRight' || e.key === 'PageDown' ? 1 : e.key === 'ArrowLeft' || e.key === 'PageUp' ? -1 : 0;
    if (!d) return;
    e.preventDefault();
    const i = Math.min(s.length - 1, Math.max(0, cur() + d));
    s[i].scrollIntoView({{behavior: 'smooth'}});
  }});
</script>
</body>
</html>
"""


def build_html(pptx_path, png_paths, out_html, title=None, substituted=()):
    """Собирает самодостаточный HTML из уже отрендеренных PNG (без LibreOffice)."""
    texts = slide_texts(pptx_path)
    title = title or os.path.splitext(os.path.basename(pptx_path))[0]
    sections = []
    for i, png in enumerate(png_paths):
        with open(png, "rb") as f:
            data = base64.b64encode(f.read()).decode("ascii")
        lines = texts[i] if i < len(texts) else []
        alt = html.escape(" — ".join(lines[:2]) or f"Слайд {i + 1}", quote=True)
        body = "".join(f"<p>{html.escape(line)}</p>" for line in lines)
        sections.append(
            f'<section id="slide-{i + 1}"><img src="data:image/png;base64,{data}" alt="{alt}">'
            f'<div class="text"><div class="num">Слайд {i + 1}</div>{body}</div></section>')
    fonts = (f" Шрифтов шаблона нет на сервере рендера, в картинках они заменены на "
             f"Arial: {html.escape(', '.join(substituted))}; в .pptx — оригинальные."
             if substituted else "")
    page = _PAGE.format(title=html.escape(title), count=len(png_paths),
                        fonts=fonts, sections="\n".join(sections))
    with open(out_html, "w", encoding="utf-8") as f:
        f.write(page)
    return out_html


def export_all(pptx_path, out_dir, png_paths=None, title=None):
    """.pdf и .html рядом с .pptx. `png_paths` — уже готовые превью (тогда
    LibreOffice не запускается повторно: PDF остаётся от их рендера)."""
    name = os.path.splitext(os.path.basename(pptx_path))[0]
    pdf_path = os.path.join(out_dir, f"{name}.pdf")
    if png_paths is None:
        png_paths = render_pptx_to_pngs(pptx_path, out_dir)
    if not os.path.exists(pdf_path):
        convert_to_pdf(pptx_path, out_dir)
    substituted = substituted_typefaces(pptx_path)
    html_path = build_html(pptx_path, png_paths, os.path.join(out_dir, f"{name}.html"),
                           title=title, substituted=substituted)
    return {"pdf": pdf_path, "html": html_path, "substituted_fonts": substituted}
