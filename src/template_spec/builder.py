"""Builds a slot spec — a compact, machine-usable description of what the template
can actually hold — from data we already have for free (cached archetype map +
slide geometry). This is the interface between "what the template offers" and
"what content we generate": content generation happens AFTER a concrete slide
(with a concrete capacity) is chosen, not before, so text is born the right size
instead of being resized after the fact.

Family = a group of slides sharing a role. v1 groups purely by classified
archetype; the structure deliberately leaves room to split families further by
visual similarity (CLIP clustering, see docs/IMPROVEMENT_PLAN.md item 3) without
changing consumers."""

from pptx import Presentation

from generator.generator import get_capacity, get_item_char_budget
from generator.slide_kit import content_text_shapes

# Roles that make sense as fill targets. "other" is everything we explicitly
# refuse to reuse (e.g. chart-anchored slides) — never part of a family.
FILLABLE_ROLES = {
    "title", "section_divider", "closing", "bullet_list", "stats_kpi",
    "two_column_comparison", "quote", "image_caption",
}

# Roles where "capacity" (how many list items fit) is a meaningful constraint.
COUNTED_ROLES = {"bullet_list", "stats_kpi"}


def build_spec(template_path, archetype_map, style_profile=None):
    """archetype_map: {slide_index(int): archetype(str)} -> spec dict:
    {"families": [{"id", "role", "slides": [{"idx", "capacity", "bg"}]}],
     "rotation": {"order": [...], "bookends": {...}} | None}

    capacity is an int for counted roles (None = the slide holds a single
    flexible text block, any count fits), and None for uncounted roles.
    style_profile (design_system.style_profile.build_measured_profile output),
    when given, annotates each family member with its measured background
    color and passes the discovered rotation through — the matcher uses both
    to keep generated decks on the template's own color rhythm."""
    prs = Presentation(template_path)
    backgrounds = (style_profile or {}).get("backgrounds", {})

    by_role = {}
    for idx, role in archetype_map.items():
        if role not in FILLABLE_ROLES:
            continue
        # A slide with no text shape at all cannot carry ANY role's content —
        # there is nowhere to put a title, a bullet or a caption. Real templates
        # have these: the T-Zh mono deck's "image_caption" member is a full-bleed
        # decorative photo with zero text boxes, and offering it crashed the run
        # (the filler had nothing to fill). Deliberately not the same thing as
        # common.pictures' stale-data exclusion: that photo IS legitimate decor,
        # it just isn't fillable, so it drops out of the family and the role gets
        # synthesized instead.
        if not content_text_shapes(prs.slides[idx]):
            continue
        entry = {"idx": idx, "capacity": None, "bg": backgrounds.get(idx)}
        if role in COUNTED_ROLES:
            entry["capacity"] = get_capacity(prs.slides[idx], role)
            # How LONG one item may be here, not just how many fit — a 1.93in
            # one-line slot and a full-width prose box are both "a list".
            entry["item_chars"] = get_item_char_budget(prs.slides[idx], role)
        by_role.setdefault(role, []).append(entry)

    families = []
    for n, (role, slides) in enumerate(sorted(by_role.items())):
        slides.sort(key=lambda s: s["idx"])
        families.append({"id": f"f{n}", "role": role, "slides": slides})

    rotation = None
    if style_profile and style_profile["rotation"]["order"]:
        rotation = {
            "order": style_profile["rotation"]["order"],
            "bookends": style_profile.get("bookends", {}),
        }
    return {"families": families, "rotation": rotation}


def describe_for_prompt(spec, offered_roles):
    """One line per family — compact enough that even a 30-slide template adds
    only a handful of lines to the outline prompt. Only offered_roles (the set
    the outline validator will actually accept) are listed: advertising a
    family the validator then rejects (e.g. image_caption) just baits the
    model into wasting a retry."""
    lines = []
    native_roles = set()
    for fam in spec["families"]:
        if fam["role"] not in offered_roles:
            continue
        native_roles.add(fam["role"])
        caps = sorted({s["capacity"] for s in fam["slides"] if s["capacity"]})
        cap_note = f", ёмкость пунктов: {', '.join(map(str, caps))}" if caps else ""
        lines.append(f"- {fam['role']}: {len(fam['slides'])} слайд(ов){cap_note}")
    synth_only = sorted(set(offered_roles) - native_roles)
    if synth_only:
        lines.append(f"- можно создать с нуля: {', '.join(synth_only)}")
    return "\n".join(lines)
