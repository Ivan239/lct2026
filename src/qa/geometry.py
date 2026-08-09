"""Post-generation geometric QA — the LLM-free safety net of
docs/IMPROVEMENT_PLAN.md item 2.

Every filler already sizes its own text to its own box up front, but each does
it locally, one shape at a time, from its own assumptions. This pass re-checks
the finished deck as a whole with the same real font metrics and shrinks
whatever still overflows — catching the cases no individual filler could see:
a run whose size came from template inheritance rather than a filler, a
synthesized slide's estimate drifting from a late edit, or simply a future
filler bug. Belt and suspenders, purely local, no LLM.
"""

from pptx.util import Emu, Inches, Pt

from generator.text_fit import (
    SOFT_HYPHEN,
    cap_size_to_longest_word,
    estimate_wrapped_lines,
    estimate_block_height_in,
    fit_font_size,
    horizontal_margins_in,
    vertical_insets_emu,
)

# The estimator and LibreOffice/PowerPoint never agree to the point — a few
# percent of slack keeps this pass from "fixing" text that actually fits.
OVERFLOW_TOLERANCE = 1.05

# Below this a shrink stops being a fix and becomes microtext.
MIN_READABLE_PT = 9


def _shape_texts(shape):
    return [
        "".join(run.text for run in p.runs)
        for p in shape.text_frame.paragraphs
    ]


def _shape_line_spacing(shape):
    """First explicit paragraph line_spacing — the overflow estimate must use
    the same spacing rule the renderer will (see text_fit.paragraph_pitch_pt),
    or a 1.5x-spaced body passes the check by bare font height and ships
    overflowing."""
    for p in shape.text_frame.paragraphs:
        if p.line_spacing is not None:
            return p.line_spacing
    return None


def harmonize_clone_font_sizes(prs, clone_groups):
    """clone_groups: iterable of lists of slide indices that were filled from
    the same pristine template slide (the original plus its clones — see
    generator._clone_slide). Clones preserve shape ids, so the same logical
    box exists on every slide of a group; after individual fitting, sibling
    slides can end up with different body sizes (one got longer text), which
    reads as sloppy the moment the slides sit next to each other. Unifying to
    the group's minimum is safe by construction: the smallest size already
    proved it fits the identical geometry.

    Returns a list of (slide_idx, shape_id, new_size_pt) for every shape whose
    text it actually shrank — callers must re-derive anything positioned from
    the old font size (bullet-marker icons), or those keep the stale layout."""
    changes = []
    for group in clone_groups:
        if len(group) < 2:
            continue
        min_size_by_shape = {}
        for slide_idx in group:
            for shape in prs.slides[slide_idx].shapes:
                if not shape.has_text_frame:
                    continue
                sizes = [
                    run.font.size.pt
                    for p in shape.text_frame.paragraphs
                    for run in p.runs
                    if run.font.size and run.text.strip()
                ]
                if not sizes:
                    continue
                current = min_size_by_shape.get(shape.shape_id)
                size = max(sizes)
                min_size_by_shape[shape.shape_id] = size if current is None else min(current, size)

        for slide_idx in group:
            for shape in prs.slides[slide_idx].shapes:
                target = min_size_by_shape.get(shape.shape_id)
                if target is None or not shape.has_text_frame:
                    continue
                shrank = False
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size and run.font.size.pt > target:
                            run.font.size = Pt(target)
                            shrank = True
                if shrank:
                    changes.append((slide_idx, shape.shape_id, target))
    return changes


SPARSE_COVERAGE_FRACTION = 0.15
# Roles that are sparse BY DESIGN — a divider is one phrase on a big field of
# brand color; warning about it would train users to ignore warnings.
SPARSE_EXEMPT_ROLES = {"title", "section_divider", "closing"}

# NOTE deliberately absent here: box-rectangle overlap detection. Calibrated
# on the corpus and rejected — designers routinely draw text boxes with
# generous rectangles that overlap while the visible text doesn't (51 hits on
# the ORIGINAL 69-slide survey deck, 3 on each pristine preset), so rectangle
# intersection is noise, not signal. The overlap bug class this would have
# watched for was fixed at the source (body-before-title claiming in fillers).


def _is_content_shape(shape):
    from pptx.enum.shapes import PP_PLACEHOLDER

    if not shape.has_text_frame:
        return False
    chrome = (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE)
    if shape.is_placeholder and shape.placeholder_format.type in chrome:
        return False
    return bool(shape.text_frame.text.strip())


def find_sparse_slides(prs, slide_roles):
    """slide_roles: {final_slide_position(0-based): role_str}. Returns
    [{slide: position, kind: "sparse", details}] for content slides whose text
    and pictures together cover almost none of the slide. Detection only — a
    near-empty slide has no safe automatic fix, so the user judges."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    issues = []
    slide_area = prs.slide_width * prs.slide_height
    if not slide_area:
        return issues
    for position, role in slide_roles.items():
        if role in SPARSE_EXEMPT_ROLES:
            continue
        slide = prs.slides[position]
        covered = sum(
            (s.width or 0) * (s.height or 0)
            for s in slide.shapes
            if _is_content_shape(s)
        )
        covered += sum(
            (s.width or 0) * (s.height or 0)
            for s in slide.shapes
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        if covered / slide_area < SPARSE_COVERAGE_FRACTION:
            issues.append({
                "slide": position,
                "kind": "sparse",
                "details": f"контент занимает {covered / slide_area:.0%} площади слайда",
            })
    return issues


# A picture is a side ART PANEL — text must stay out of it — when it starts to
# the RIGHT of a text box's left edge and covers a real share of the slide. The
# "starts to the right" part is what separates it from a full-bleed background
# (left edge 0), which every text box on the slide legitimately sits on: the
# T-Zh folder cards and the survey backdrops are exactly that, and treating them
# as obstacles would leave nowhere to put anything.
SIDE_ART_MIN_AREA_FRACTION = 0.06
SIDE_ART_GAP_IN = 0.2
# Below this share of the box's own width the clear zone is a sliver: narrowing
# to it would trade a legible overlap for an illegible column.
SIDE_ART_MIN_CLEAR_FRACTION = 0.4


def _side_art_left(slide, shape, slide_w, slide_h):
    left = None
    for art in slide.shapes:
        if "PICTURE" not in str(art.shape_type):
            continue
        if art.left is None or art.top is None or not art.width or not art.height:
            continue
        if (art.width * art.height) < SIDE_ART_MIN_AREA_FRACTION * slide_w * slide_h:
            continue
        if int(art.top) >= int(shape.top + shape.height) or int(art.top + art.height) <= int(shape.top):
            continue
        if not int(shape.left) < int(art.left) < int(shape.left + shape.width):
            continue
        left = int(art.left) if left is None else min(left, int(art.left))
    return left


def keep_text_clear_of_side_art(prs, resolver, slide_indices=None):
    """Narrow a text box whose text would otherwise be drawn over side artwork.

    The universal template's cover box runs to 9.44in while its art panel starts
    at 5.91in: the designer's own «Название / презентации» is short and breaks by
    hand, so it never reaches the art, and a generated title of ordinary length
    ran straight through the balloons with «управленческой» unreadable — on the
    deck's first slide.

    Only acts when OUR text actually reaches the panel, so a template whose own
    text stays clear is untouched. Runs before enforce_text_fits, which then
    re-wraps against the narrowed box. Returns the shapes it narrowed."""
    changed = []
    indices = range(len(prs.slides._sldIdLst)) if slide_indices is None else slide_indices
    for slide_idx in indices:
        slide = prs.slides[slide_idx]
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.width or not shape.height:
                continue
            if shape.left is None or shape.top is None:
                continue
            texts = [t for t in _shape_texts(shape) if t.strip()]
            if not texts:
                continue
            art_left = _side_art_left(slide, shape, prs.slide_width, prs.slide_height)
            if art_left is None:
                continue
            clear = art_left - int(Inches(SIDE_ART_GAP_IN)) - int(shape.left)
            if clear < SIDE_ART_MIN_CLEAR_FRACTION * int(shape.width):
                continue
            size_pt = max((r.font.size.pt for p in shape.text_frame.paragraphs
                           for r in p.runs if r.font.size and r.text.strip()), default=None)
            font_name = next((r.font.name for p in shape.text_frame.paragraphs
                              for r in p.runs if r.font.name), None)
            metrics = resolver.metrics_for(font_name) if font_name else None
            if not size_pt or metrics is None:
                continue
            margins = horizontal_margins_in(shape)
            clear_in = Emu(clear).inches
            reaches = any(
                estimate_wrapped_lines(t, clear_in, size_pt, metrics=metrics,
                                       margins_in=margins) > 1
                for t in texts)
            if not reaches:
                continue
            shape.width = Emu(clear)
            # Re-cap to the NEW width: narrowing a box leaves the size that was
            # fitted to the old one, and the corpus test caught exactly that —
            # «Высокопроизводительная» at 36pt no longer fits 5.3in, so the
            # renderer would break it mid-letter. A smaller title beats a
            # chopped word (iter18/25); enforce_text_fits only checks height.
            capped = cap_size_to_longest_word(
                texts, shape.width, size_pt, metrics=metrics, margins_in=margins,
                bold=any(r.font.bold for p in shape.text_frame.paragraphs
                         for r in p.runs if r.font.bold))
            if float(capped) < size_pt:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size:
                            run.font.size = Pt(float(capped))
            changed.append((slide_idx, shape.shape_id))
    return changed


def drop_needless_soft_hyphens(prs, resolver, slide_indices=None):
    """Remove soft hyphens from words that fit after the final size is known.

    _set_run_text inserts them when the longest word cannot fit at the size it
    is about to use — a visible hyphen beats a word chopped mid-letter
    (iter18/25). But enforce_text_fits then shrinks the text further, and at the
    smaller size the word fits with room to spare: the T-Zh mono heading came
    out as «Что мешало собирать упра-вленческую отчётность вовремя», where
    «управленческую» measures 232.6pt against a 293.9pt budget at its final
    28pt. The renderer breaks at a soft hyphen in preference to wrapping, so a
    stale one keeps hyphenating a word that no longer needs it.

    Returns the shapes it cleaned."""
    cleaned = []
    indices = range(len(prs.slides._sldIdLst)) if slide_indices is None else slide_indices
    for slide_idx in indices:
        slide = prs.slides[slide_idx]
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.width:
                continue
            if SOFT_HYPHEN not in shape.text_frame.text:
                continue
            margins = horizontal_margins_in(shape)
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if SOFT_HYPHEN not in run.text or not run.font.size:
                        continue
                    metrics = resolver.metrics_for(run.font.name) if run.font.name else None
                    if metrics is None:
                        continue
                    plain = run.text.replace(SOFT_HYPHEN, "")
                    fits = cap_size_to_longest_word(
                        plain, shape.width, run.font.size.pt, min_size_pt=1,
                        metrics=metrics, margins_in=margins,
                        bold=bool(run.font.bold))
                    if float(fits) >= run.font.size.pt:
                        run.text = plain
                        cleaned.append((slide_idx, shape.shape_id))
    return cleaned


def enforce_text_fits(prs, resolver, slide_indices=None):
    """Shrinks any text still estimated to overflow its box. Returns a list of
    (slide_idx, shape_id, old_size_pt, new_size_pt) describing what changed —
    empty on a healthy deck.

    Only acts when real glyph metrics are resolvable for the shape's font:
    the char-count fallback estimator overestimates tall/narrow boxes enough
    to "shrink" perfectly fine template text, and a QA pass that damages
    healthy slides is worse than no pass at all."""
    fixes = []
    indices = range(len(prs.slides._sldIdLst)) if slide_indices is None else slide_indices
    for slide_idx in indices:
        slide = prs.slides[slide_idx]
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.width or not shape.height:
                continue
            texts = _shape_texts(shape)
            if not any(t.strip() for t in texts):
                continue
            # Interior blank paragraphs (designer spacers between items) render
            # as a full line each and must count toward the height estimate —
            # dropping them hid exactly the phantom-line overflow this pass
            # exists to catch. Trailing blanks are bottom padding; trim those.
            while texts and not texts[-1].strip():
                texts.pop()
            runs = [run for p in shape.text_frame.paragraphs for run in p.runs]
            sizes = [run.font.size.pt for run in runs if run.font.size]
            if not sizes:
                continue
            size_pt = max(sizes)

            font_name = next((run.font.name for run in runs if run.font.name), None)
            metrics = resolver.metrics_for(font_name) if resolver and font_name else None
            if metrics is None:
                continue

            width_in = Emu(shape.width).inches
            top_inset, bottom_inset = vertical_insets_emu(shape)
            usable_height_emu = max(0, shape.height - top_inset - bottom_inset)
            height_in = Emu(usable_height_emu).inches
            line_spacing = _shape_line_spacing(shape)
            margins_in = horizontal_margins_in(shape)
            estimated = estimate_block_height_in(
                texts, width_in, size_pt, metrics=metrics, line_spacing=line_spacing,
                margins_in=margins_in,
            )
            if estimated <= height_in * OVERFLOW_TOLERANCE:
                continue

            # Readability floor, the same absolute one the filler keeps
            # (generator._fit_size_for_shape): past it the problem is that the
            # CONTENT is too long, and microtext is worse than the slight
            # overflow this pass exists to remove. Without it this pass silently
            # undid that rule — a 49-character bullet in a one-line 2.86x0.24in
            # slot was fitted at 13pt and then taken to 8pt here, unreadable on
            # the render.
            #
            # ABSOLUTE, not a fraction of the current size: a relative floor was
            # tried and it broke idempotence, because each pass would re-floor
            # against its own output (60 -> 42 -> 29 on the survey fixture) and
            # a genuinely oversized block would stop mid-way instead of fitting.
            new_size = fit_font_size(
                texts, shape.width, usable_height_emu, size_pt,
                min_size_pt=MIN_READABLE_PT,
                metrics=metrics, line_spacing=line_spacing, margins_in=margins_in,
            )
            if new_size.pt >= size_pt:
                continue
            for run in runs:
                if run.font.size and run.font.size.pt > new_size.pt:
                    run.font.size = new_size
            fixes.append((slide_idx, shape.shape_id, size_pt, new_size.pt))
    return fixes
