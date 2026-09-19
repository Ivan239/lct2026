import statistics
from collections import Counter

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Length, Pt

from common.pictures import (OVERSIZED_PICTURE_AREA_RATIO, _appears_on_other_slides,
                             _covers_slide_text, is_stale_data_picture)
from common.phrases import drop_dangling_function_words
from common.synthesis import SYNTHESIZE
from fonts.metrics import FontResolver
from generator.deck_style import apply_observed_style, observe_deck_style
from generator.layout_bounds import infer_content_bounds
from generator.synthesizer import IMAGE_PLACEHOLDER_NAME, SYNTHESIZERS
from generator.text_fit import (
    AVG_CHAR_WIDTH_FACTOR,
    SINGLE_LINE_SAFETY,
    _usable_width_in,
    estimate_wrapped_lines,
    cap_size_to_longest_word,
    soft_hyphenate_long_words,
    estimate_block_height_in,
    estimate_wrapped_lines,
    fit_font_size,
    fit_font_size_single_line,
    frame_horizontal_margins_in,
    horizontal_margins_in,
    line_height_pt,
    paragraph_indent_in,
    paragraph_pitch_pt,
    vertical_insets_emu,
)
from qa.geometry import (MIN_READABLE_PT, drop_needless_soft_hyphens, enforce_text_fits,
                         harmonize_clone_font_sizes, keep_text_clear_of_side_art,
                         keep_text_inside_slide, unglue_overwide_pairs)
from qa.package_check import assert_valid_package
from template_parser.parser import extract_template, extract_theme

# Shared with synthesizer via slide_kit (plan 9.3); private aliases keep the
# many existing call sites and tests stable.
from generator.slide_kit import (
    content_groups,
    content_text_shapes as _content_text_shapes,
    slide_height as _slide_height,
    slide_width as _slide_width,
    is_boring_placeholder as _is_boring_placeholder,
    is_chrome_shape,
    text_shapes as _text_shapes,
)


def _max_font_pt(shape):
    sizes = [
        run.font.size.pt
        for p in shape.text_frame.paragraphs
        for run in p.runs
        if run.font.size
    ]
    return max(sizes) if sizes else 0


def _placeholder(slide, idx, claimed_ids):
    """Real hand-designed templates often use plain textboxes instead of proper
    PowerPoint placeholders — slide.placeholders[idx] raises KeyError whenever
    there's no such placeholder at all. Also refuses a placeholder that's already
    claimed by another field of the same block (e.g. title and "body idx=1" can
    be the very same shape on some layouts — without this check the second write
    silently clobbers the first)."""
    try:
        ph = slide.placeholders[idx]
    except KeyError:
        return None
    return None if ph.shape_id in claimed_ids else ph


_FIGURE_CHARS = set("0123456789+-%×x/.,")
# A display figure is never the title, however large it is set. The position
# discount below is not enough on its own: iter70 gave the figure the size the
# template sets it at (72pt against an 18pt heading on the T-Zh study deck), and
# half of 72 still beats 18 — so the harness read «-40%» as the slide's title
# and reported the real heading as a stray widow. Measured across the corpus,
# exactly three slides pick a figure today, including the universal template's
# own «20 227 000».
_FIGURE_SCORE_DISCOUNT = 0.2


def _is_display_figure(shape):
    """True when the shape's text is mostly digits and arithmetic marks.

    Deliberately not "short and contains a digit": «Top 7» is a heading on a
    real template and would have been caught by that. Half the non-space
    characters must be figure characters, which «Top 7» (25%) and «4 дня» (25%)
    are not, and «-40%», «+18», «x3», «20 227 000» are."""
    text = shape.text_frame.text.strip() if shape.has_text_frame else ""
    core = "".join(text.split())
    if not core:
        return False
    return sum(1 for c in core if c in _FIGURE_CHARS) / len(core) >= 0.5


def _pick_title_shape(slide, claimed_ids):
    """Prefer a real title placeholder; otherwise the largest text shape in the
    TOP BAND of the slide's content. Font size alone is not enough: on a
    placeholder-less stats slide the KPI numbers (48pt) out-size the actual
    title (30pt), and picking a number box as "the title" shifts every
    number/label pair one box over (template_a_corporate slide 3, found by
    preset render sweep). Shapes below the top quarter of the content span
    keep competing — a cover slide's mid-slide hero title still wins over
    small captions — but at a discount, so a same-band real title beats them."""
    title = slide.shapes.title
    if title is not None and title.shape_id not in claimed_ids:
        return title
    candidates = _content_text_shapes(slide, claimed_ids)
    if not candidates:
        return None
    tops = [s.top or 0 for s in candidates]
    bottoms = [(s.top or 0) + (s.height or 0) for s in candidates]
    band_limit = min(tops) + (max(bottoms) - min(tops)) * 0.25
    def score(shape):
        size = _max_font_pt(shape) or 0
        if (shape.top or 0) > band_limit:
            size *= 0.5
        if _is_display_figure(shape):
            size *= _FIGURE_SCORE_DISCOUNT
        return size, -(shape.top or 0)

    return max(candidates, key=score)


def _body_placeholders(slide, claimed_ids):
    from pptx.enum.shapes import PP_PLACEHOLDER
    return [ph for ph in slide.placeholders
            if ph.shape_id not in claimed_ids and ph.has_text_frame
            and ph.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)]


def _pick_body_shape(slide, claimed_ids):
    """Prefer the native content placeholder; otherwise the remaining text shape
    with the most paragraphs is the best guess for "the body text".

    idx 1 is the first body placeholder, not always the body: VK Tech's contents
    slide has a one-line label «Содержание» at idx 1 over a four-item list at
    idx 2. The label was taken for the body — capacity 1, the list blanked, the
    single bullet cut to «Средне» (iter123). Another body placeholder holding
    MORE content slots is the body; on ties idx 1 keeps it. Measured over the
    corpus this moves exactly that slide and the two-column slide of the
    startup/corporate fixtures (column header 1 slot, its items 3)."""
    native = _placeholder(slide, 1, claimed_ids)
    if native is not None:
        richer = [ph for ph in _body_placeholders(slide, claimed_ids)
                  if len(_content_slot_indices(ph)) > len(_content_slot_indices(native))]
        if richer:
            return max(richer, key=lambda ph: (len(_content_slot_indices(ph)),
                                               -ph.placeholder_format.idx))
        return native
    candidates = _content_text_shapes(slide, claimed_ids)
    if not candidates:
        return None
    return max(candidates, key=lambda s: len(s.text_frame.paragraphs))


def _reference_run(shape):
    """A run elsewhere in the same shape whose formatting we can copy onto a
    brand-new run — without this, a run created for a previously-empty paragraph
    (add_run()) has no explicit font at all and renders with a different default
    than its siblings, i.e. the font visibly "jumps" between bullets."""
    for p in shape.text_frame.paragraphs:
        if p.runs:
            return p.runs[0]
    return None


def _run_color_key(run):
    """Comparable identity of a run's explicit color: theme colors by scheme
    slot (their .rgb raises), absolute colors by RGB, inherited → None."""
    color = run.font.color
    try:
        if color is None or color.type is None:
            return None
        if color.type == MSO_COLOR_TYPE.SCHEME:
            return ("theme", color.theme_color, color.brightness)
        return ("rgb", color.rgb)
    except (AttributeError, TypeError):
        return None


def _run_style_key(run):
    return (run.font.name, run.font.size.pt if run.font.size else None, run.font.bold, _run_color_key(run))


def _dominant_reference_run(shape):
    """The run whose style covers the most characters in the shape — i.e. the
    body style, not an accent. Hand-designed decks routinely color a key phrase
    inside a bullet (black text, purple accent words); since we write each new
    line into its paragraph's FIRST run, a line whose original happened to
    START with the accent run would come out entirely accent-colored while its
    siblings stay plain — whole lines randomly purple, reading as a glitch.
    Unifying every written line to the shape's dominant style keeps uniform
    templates unchanged and turns mixed ones consistent."""
    weights = {}
    representative = {}
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            if not run.text.strip():
                continue
            key = _run_style_key(run)
            weights[key] = weights.get(key, 0) + len(run.text.strip())
            representative.setdefault(key, run)
    if not weights:
        return None
    best = max(weights, key=lambda k: weights[k])
    return representative[best]


def _clear_run_fill(run):
    """Removes any explicit fill from the run so it falls back to the
    inherited (theme/placeholder) color — assigning some guessed RGB instead
    would bake in whatever the current theme resolves to."""
    from pptx.oxml.ns import qn

    rPr = run._r.get_or_add_rPr()
    for tag in ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:noFill"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)


def _copy_run_format(source, target):
    """Makes target render like source. Where source styles a property
    EXPLICITLY, copy it; where source inherits (color type None), strip the
    target's own explicit value too — merely "not copying" would leave e.g. an
    accent color sitting on the target while its siblings inherit plain text
    color, which is exactly the mixed-style artifact this exists to prevent."""
    if source is None:
        return
    target.font.name = source.font.name
    if source.font.size:
        target.font.size = source.font.size
    target.font.bold = source.font.bold
    try:
        color = source.font.color
        if color is not None and color.type == MSO_COLOR_TYPE.SCHEME:
            target.font.color.theme_color = color.theme_color
            if color.brightness:
                target.font.color.brightness = color.brightness
        elif color is not None and color.type is not None:
            target.font.color.rgb = color.rgb
        else:
            _clear_run_fill(target)
    except (AttributeError, TypeError):
        pass


def _set_paragraph_text(paragraph, text, reference_run=None, font_size=None, unify_format=False):
    """A single visual line is often split across several runs (common in
    Google Slides exports, even with uniform formatting) — put the new text in
    the first run and blank every other run, or its original content leaks
    through as a trailing tail next to our new text. Empty spacer paragraphs
    (no runs at all) get a new run, formatted to match its siblings, rather
    than being silently skipped. font_size, when given, is set as an absolute
    value on every paragraph so sibling bullets end up the same size — scaling
    each paragraph's own original size individually would leave them at
    different absolute sizes whenever the original template already had uneven
    sizing between paragraphs."""
    if not paragraph.runs:
        run = paragraph.add_run()
        _copy_run_format(reference_run, run)
        run.text = text
    else:
        run = paragraph.runs[0]
        run.text = text
        for extra in paragraph.runs[1:]:
            extra.text = ""
        if unify_format:
            # Sibling lines must share one style — see _dominant_reference_run.
            _copy_run_format(reference_run, run)
    if font_size is not None:
        run.font.size = font_size
    elif run.font.size:
        # The source file's own PowerPoint autofit sometimes caches odd values
        # like 17.69pt (a cached shrink percentage of whatever the original text
        # needed) — normalize to a clean whole point even when we didn't shrink
        # anything ourselves, or the file just looks sloppy for no reason of ours.
        run.font.size = Pt(round(run.font.size.pt))
    return True


def _clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ""


def _unify_paragraph_indents(paragraphs, used_indices):
    """Filled list items must share one left edge: prose-era templates carry
    different per-paragraph marL/indent leftovers (0.25" on one item, none on
    the next), which renders as a visibly ragged column once short list items
    replace the original text. Majority value wins; ties go to the first item."""
    from collections import Counter

    pairs = [
        (
            paragraphs[j]._p.pPr.get("marL") if paragraphs[j]._p.pPr is not None else None,
            paragraphs[j]._p.pPr.get("indent") if paragraphs[j]._p.pPr is not None else None,
        )
        for j in used_indices
    ]
    if len(set(pairs)) <= 1:
        return
    target_marL, target_indent = Counter(pairs).most_common(1)[0][0]
    for j in used_indices:
        pPr = paragraphs[j]._p.get_or_add_pPr()
        for attr, value in (("marL", target_marL), ("indent", target_indent)):
            if value is None:
                pPr.attrib.pop(attr, None)
            else:
                pPr.set(attr, value)


def _content_slot_indices(shape):
    """Indices of paragraphs that carry actual template text — the designer's
    content slots. Hand-built decks space their lists with EMPTY paragraphs
    between items (see the survey fixtures: content, blank, content, blank);
    writing sequentially into paragraphs[0..N] stuffs our items into those
    spacers, collapsing the airy rhythm the designer built into a dense clump
    AND mixing per-paragraph spacing (a 1.5-spaced original next to a
    single-spaced spacer renders as visibly uneven line pitch). Filling only
    the original content slots keeps the template's own rhythm. Falls back to
    every paragraph when the box is entirely empty (nothing to learn from)."""
    slots = [
        i for i, p in enumerate(shape.text_frame.paragraphs)
        if "".join(run.text for run in p.runs).strip()
    ]
    return slots or list(range(len(shape.text_frame.paragraphs)))


def _extend_content_slots(shape, n_extra):
    """Appends n_extra new content slots after the last one, replicating the
    template's own inter-item rhythm (any spacer paragraphs between the last
    two slots are cloned along with the slot itself). Without this, a text
    list longer than the template's slot count silently DROPS the tail items —
    real content loss, not a cosmetic defect. New paragraphs are deep copies
    of the last content slot, so they inherit its spacing and run formatting;
    inserted right after it, BEFORE any trailing blank paragraphs, so the new
    items stay visually contiguous with the list."""
    import copy as _copy

    paragraphs = shape.text_frame.paragraphs
    slots = _content_slot_indices(shape)
    if not slots:
        return
    template_el = paragraphs[slots[-1]]._p
    gap_els = []
    if len(slots) >= 2:
        gap_els = [paragraphs[j]._p for j in range(slots[-2] + 1, slots[-1])]
    cursor = template_el
    for _ in range(n_extra):
        for gap in gap_els:
            el = _copy.deepcopy(gap)
            cursor.addnext(el)
            cursor = el
        el = _copy.deepcopy(template_el)
        cursor.addnext(el)
        cursor = el


def _enable_autofit(shape):
    """Belt-and-suspenders alongside the direct font shrink above — flags the
    shape so PowerPoint keeps it that way if the user edits the text later."""
    try:
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


def _metrics_for_run(reference_run, resolver):
    """Metrics of the face the run is actually SET in, bold included.

    Cyrillic bold runs ~7% wider than the regular advances a family name
    resolves to (CLAUDE.md), and every caller here — the one-line figure fit,
    the slot budgets, the title — measured its bold text with regular metrics.
    VK Tech's board printed «94» over «секунды»: the figure fitted «94 секунды»
    at 31pt, which is 101.2% of that box's line, so the renderer wrapped it and
    the template's 16% display leading laid the two lines on top of each other
    (iter141)."""
    if resolver is None or reference_run is None:
        return None
    return resolver.metrics_for(reference_run.font.name, bold=bool(reference_run.font.bold))


# Used only when the shape's own text has no explicit run-level size at all —
# common for short single-line captions in hand-built/Google-Slides-exported
# decks, which often lean on inherited placeholder/theme defaults instead.
# fit_font_size only ever shrinks from this baseline, never grows past it, so
# it just needs to be a plausible upper bound, not an exact guess.
_FALLBACK_BASE_SIZE_PT = 18


def _fit_size_for_shape(shape, lines, reference_run, resolver=None):
    """Sized to the shape's actual, known box dimensions — not a ratio against
    however long the old text happened to be. The old text may itself already
    have been an awkward fit (real slides in hand-designed decks sometimes
    already overflow their own box before we ever touch them), so scaling
    relative to it just perpetuates whatever was already wrong.

    When the reference run carries no explicit size (it's relying on inherited
    placeholder/master defaults, which vary by renderer and can be far larger
    than what actually fits), leaving the new run unsized too would let the
    text overflow its box with hard, mid-word line breaks — falling back to a
    fixed baseline and still fitting it properly is strictly safer, and as a
    side effect gives every run an explicit size, which also happens to dodge
    a LibreOffice glyph-overlap rendering bug that only shows up on unsized
    runs in a substituted font."""
    base_size_pt = reference_run.font.size.pt if reference_run and reference_run.font.size else _FALLBACK_BASE_SIZE_PT
    metrics = _metrics_for_run(reference_run, resolver)
    top_inset, bottom_inset = vertical_insets_emu(shape)
    usable_height = max(0, (shape.height or 0) - top_inset - bottom_inset)
    # Font floor (plan 10г): shrinking below ~70% of the template's own size
    # produces the "microtext" artifact a human spots instantly — past that
    # point the CONTENT is too long (title/bullet budgets, plan 10б), and a
    # 6pt fix is worse than a slight overflow the QA pass can still judge.
    floor_pt = max(9, round(base_size_pt * 0.7))
    margins = horizontal_margins_in(shape)
    size = fit_font_size(
        lines, shape.width, usable_height, base_size_pt,
        min_size_pt=floor_pt,
        metrics=metrics, line_spacing=_shape_line_spacing(shape),
        margins_in=margins,
    )
    # Word-fit (plan 10а): never accept a size whose longest word can't fit
    # the box width — the renderer would break it mid-letter.
    capped = cap_size_to_longest_word(
        lines, shape.width, size.pt, min_size_pt=floor_pt,
        metrics=metrics, margins_in=margins,
        bold=bool(reference_run.font.bold) if reference_run else False,
    )
    return Pt(capped)


def _set_run_text(shape, text, claimed_ids, resolver=None):
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    reference = _reference_run(shape)
    font_size = _fit_size_for_shape(shape, text, reference, resolver=resolver)
    # Escape hatch for words the box can't hold even at the font floor
    # (plan 10а, «Неконсистентность»): the probe (min_size 1, i.e. "what size
    # would this word actually need?") coming back BELOW the size we're about
    # to use means the renderer will chop the word mid-letter — soft-hyphenate
    # so it breaks with a visible hyphen at a vowel instead.
    if font_size is not None and text:
        metrics = _metrics_for_run(reference, resolver)
        word_fit_size = cap_size_to_longest_word(
            text, shape.width, font_size.pt, min_size_pt=1,
            metrics=metrics, margins_in=horizontal_margins_in(shape),
            bold=bool(reference.font.bold) if reference else False,
        )
        if word_fit_size < font_size.pt:
            text = soft_hyphenate_long_words(text)
    if _set_paragraph_text(tf.paragraphs[0], text, reference_run=reference, font_size=font_size):
        claimed_ids.add(shape.shape_id)
        _enable_autofit(shape)
    # Claiming this shape means owning all of it — any further paragraph is
    # leftover from the original template. Delete rather than clear: cleared
    # paragraphs still render as blank lines that can overfill the box and
    # trigger silent renderer-side autofit shrinking.
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


# Content filling less of its box than this reads as an accidentally
# half-empty slide when left hanging at the top edge.
_CENTER_FILL_RATIO = 0.45


def _center_if_underfilled(shape, texts, font_size, reference, resolver=None):
    """Vertically centers a body whose text occupies well under half its box —
    the "2 short lines at the top of a 5-inch box" look is the single most
    common sparse-slide complaint, and a centered block reads as intentional
    where a top-clinging one reads as broken. Never called for icon-marker
    bodies: icon rows are computed from the box TOP, and PowerPoint offers no
    way to know where a middle-anchored first line will land across renderers."""
    if shape is None or not shape.width or not shape.height or font_size is None:
        return
    tf = shape.text_frame
    if tf.vertical_anchor not in (None, MSO_ANCHOR.TOP):
        return  # the template designer already chose an anchor — keep it
    top_inset, bottom_inset = vertical_insets_emu(shape)
    usable_in = Emu(max(0, shape.height - top_inset - bottom_inset)).inches
    if usable_in <= 0:
        return
    metrics = _metrics_for_run(reference, resolver)
    est = estimate_block_height_in(
        texts, Emu(shape.width).inches, font_size.pt,
        metrics=metrics, line_spacing=_shape_line_spacing(shape),
        margins_in=horizontal_margins_in(shape),
    )
    if est / usable_in < _CENTER_FILL_RATIO:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def _set_paragraph_texts(shape, texts, claimed_ids, resolver=None, allow_center=False,
                         prefer_single_line=False):
    """Returns the font size actually used (or None), so callers that also need
    to reposition something else relative to the text (e.g. bullet-marker icon
    graphics — see _reposition_bullet_icons) don't have to recompute it.

    prefer_single_line: for icon-marker lists — shrink (never below the
    single-line fitter's own floor) until every item is guaranteed one line,
    so icon rows land on exact pitch multiples instead of depending on a wrap
    prediction that's only accurate to a few percent. Falls back to the normal
    wrapped fit when a guarantee isn't achievable (abnormally long item)."""
    if shape is None:
        return None
    reference = _dominant_reference_run(shape) or _reference_run(shape)
    # One size for the whole list — sizing each line individually against its own
    # original paragraph would leave sibling bullets at different font sizes,
    # which itself reads as a "jumping" font even when every line technically fits.
    font_size = _fit_size_for_shape(shape, texts, reference, resolver=resolver)
    if prefer_single_line and font_size is not None:
        single = fit_font_size_single_line(
            texts, shape.width, font_size.pt,
            metrics=_metrics_for_run(reference, resolver),
            margins_in=horizontal_margins_in(shape),
        )
        if single is not None:
            font_size = single  # <= the height-fitted size by construction

    if not texts:
        for p in shape.text_frame.paragraphs:
            _clear_paragraph(p)
        return None

    # Fill the template's own content slots (paragraphs that carried text),
    # not paragraphs[0..N] blindly: hand-built decks space their lists with
    # empty spacer paragraphs, and sequential filling stuffs items into those
    # spacers — collapsing the designer's rhythm AND mixing per-paragraph
    # styles (mismatched marL renders as a ragged left edge).
    if len(texts) > len(_content_slot_indices(shape)):
        _extend_content_slots(shape, len(texts) - len(_content_slot_indices(shape)))
    paragraphs = list(shape.text_frame.paragraphs)
    slots = _content_slot_indices(shape)
    used = slots[: len(texts)]

    any_set = False
    for text, slot_idx in zip(texts, used):
        if _set_paragraph_text(paragraphs[slot_idx], text, reference_run=reference, font_size=font_size, unify_format=True):
            any_set = True
    _unify_paragraph_indents(paragraphs, used)
    # Spacer paragraphs keep the designer's rhythm but must not leak old text —
    # and their count is normalized to the SMALLEST gap the template uses
    # between items: prose-era templates carry uneven leftovers (2 blanks here,
    # 1 there) that read as random jumps once short list items replace the
    # original multi-line prose. Everything AFTER the last used slot is deleted
    # outright — cleared paragraphs still render as blank lines (a '\x0b'-only
    # "empty" run is even worth several), and enough phantom lines overfill the
    # box so the renderer silently autofit-shrinks the text, dragging every
    # position computed from our font size (bullet icons) off its row.
    gaps = [slots[k + 1] - slots[k] - 1 for k in range(len(slots) - 1)]
    keep_gap = min(gaps) if gaps else 0
    doomed = list(paragraphs[used[-1] + 1:])
    for a, b in zip(used, used[1:]):
        between = paragraphs[a + 1: b]
        doomed.extend(between[keep_gap:])
        for p in between[:keep_gap]:
            _clear_paragraph(p)
    for p in paragraphs[: used[0]]:
        _clear_paragraph(p)
    for p in doomed:
        p._p.getparent().remove(p._p)
    if any_set:
        claimed_ids.add(shape.shape_id)
        _enable_autofit(shape)
        if allow_center:
            _center_if_underfilled(shape, texts, font_size, reference, resolver=resolver)
    return font_size.pt if font_size is not None else (reference.font.size.pt if reference and reference.font.size else None)


def _is_managed_chrome(shape, slide_height):
    """Chrome text that a LATER pass owns and will rewrite, so blanking it here
    destroys the thing that pass exists to fix.

    Measured on the real templates: every generated deck shipped with its chrome
    band completely empty — no page numbers, no brand year, no running topic —
    because this blanking runs first. Both _renumber_static_slide_numbers
    (iter19) and _fill_running_topic (iter23) were therefore dead code on the
    native fill path: they looked for their shapes and found empty boxes.

    The predicate is not a new guess about what a number means; it is exactly
    those two passes' own tests, so nothing is spared that nobody rewrites.
    A placeholder with no owner («КОММЕНТАРИЙ») keeps being blanked — shipping
    the literal word would be worse than shipping nothing."""
    if not is_chrome_shape(shape, slide_height):
        return False
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    if len(runs) != 1:
        return False
    text = runs[0].text.strip()
    return text.isdigit() or text.lower() in TOPIC_SLOT_PROMPTS


def _clear_unclaimed_text(slide, claimed_ids, slide_height=None):
    """Any text shape we didn't deliberately fill keeps whatever was in the
    original template slide — for hand-designed decks that's often unrelated
    leftover content (e.g. old survey questions). Blank it out rather than
    let it leak into the generated deck looking like garbled AI output.

    Exception: chrome a later pass manages (_is_managed_chrome).

    History worth keeping, because the first attempt at this was WRONG: iter29
    spared all-digit chrome and resurrected three orphan «01»s on the T-Zh
    universal grid. Its bottom row sits inside the chrome band, _find_slot_boxes
    could not see those slots, so their badges were spared as page numbers and
    renumbered. Blanking had been masking that defect. It is safe now only
    because iter30 made such rows detectable — so their badges are claimed and
    named — and that is verified by measurement, not assumed."""
    for shape in _text_shapes(slide):
        if shape.shape_id in claimed_ids:
            continue
        if slide_height is not None and _is_managed_chrome(shape, slide_height):
            continue
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.text = ""


# A small decorative icon (a bullet glyph, a logo accent) is harmless to leave in
# place even if it no longer perfectly matches the new text next to it. A big
# STALE-DATA picture — a chart, a screenshot of survey results — is a leftover
# piece of the OLD content's actual substance: showing real numbers from a
# different context next to our new content is actively misleading. Background
# art and photos are scenery and MUST survive the fill (deleting a full-bleed
# brand background guts the template — see common/pictures.py module doc for
# the measured decision rule). (Ideally a chart-anchored slide is never matched
# in the first place — has_oversized_picture excludes it at match time — this
# removal is just the safety net for whatever slips through.)
def _remove_orphan_marker_columns(slide):
    """After _clear_unclaimed_text, list-marker shapes whose slot text is gone
    keep floating next to blank space (dots-with-no-text — the exact artifact
    from the first T-Ж generation, resurfacing on any slide whose slot texts
    were cleared rather than filled). Removes ALIGNED COLUMNS (>=2 markers
    sharing a left edge) that have no non-empty text to their right in their
    row band; a lone decorative accent shape never forms a column and is left
    alone."""
    small = Emu(int(Inches(0.5)))
    markers = [
        s for s in slide.shapes
        if s.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE)
        and s.left is not None and s.top is not None
        and s.width and s.height and s.width <= small and s.height <= small
    ]
    columns = {}
    for m in markers:
        columns.setdefault(int(m.left // Emu(int(Inches(0.1)))), []).append(m)

    texts = [s for s in _text_shapes(slide) if s.left is not None and s.top is not None]
    for column in columns.values():
        if len(column) < 2:
            continue
        for marker in column:
            center = marker.top + marker.height // 2
            has_text = any(
                t.left > marker.left
                and t.top <= center <= t.top + (t.height or 0)
                for t in texts
            )
            if not has_text:
                marker._element.getparent().remove(marker._element)


_HOLE_MIN_GAP_EMU = Emu(int(Inches(0.3)))


def _close_hole_left_by(slide, top, bottom):
    """Pull the text that sat BELOW a removed picture up into the space it left.

    Removing someone else's chart is right (its data goes stale next to new
    text), but it leaves a hole the size of the chart, and the designer put the
    summary paragraph UNDER it. Measured on the survey template's slide 6: the
    chart ran 2.05-4.37in and the body stayed at 4.56, so a generated slide had
    a 2.70in void — a third of the slide — between the heading and three lines
    stranded at the bottom.

    Only shapes wholly below the picture move, and only far enough to leave a
    gap under whatever ends above it: nothing may slide into another box."""
    ceiling = top
    for shape in slide.shapes:
        if shape.top is None or not shape.height:
            continue
        shape_bottom = int(shape.top + shape.height)
        if shape_bottom <= top:
            ceiling = max(ceiling, shape_bottom + int(_HOLE_MIN_GAP_EMU))
    lift = int(bottom - ceiling)
    if lift <= 0:
        return
    for shape in slide.shapes:
        if shape.top is None or not shape.height or int(shape.top) < bottom:
            continue
        shape.top = Emu(max(int(ceiling), int(shape.top) - lift))


def _remove_oversized_pictures(slide, slide_width, slide_height):
    slide_area = (slide_width or 0) * (slide_height or 0)
    if not slide_area:
        return
    for shape in list(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        area = (shape.width or 0) * (shape.height or 0)
        if area / slide_area > OVERSIZED_PICTURE_AREA_RATIO and is_stale_data_picture(shape, slide):
            top, bottom = int(shape.top or 0), int((shape.top or 0) + (shape.height or 0))
            shape._element.getparent().remove(shape._element)
            _close_hole_left_by(slide, top, bottom)


def _align_left_edges(anchor_shape, *other_shapes, tolerance_emu=45720):
    """Native templates sometimes have the title and body textboxes at slightly
    different left positions (off by a fraction of an inch) — invisible in the
    original content, but jarring once we point at both with our own text.
    tolerance_emu defaults to ~0.05in, below which a gap isn't worth touching.

    An indent that clears a PICTURE is design, not a slip: VK Education's cover
    sets the speaker line at 2.94in, right of a round speaker photo at
    0.72–2.49in, while the title starts at 0.68in. Aligned, the subtitle went
    under the photo on every VK Education deck (iter103–iter112, five times).
    A cap on the shift is not the answer — one was tried and reverted: a body
    returned to its own column hung in an empty slide (CLAUDE.md). The rule is
    about the cause: the box stays where it is when moving it would put it over
    a picture it does not touch now.

    Nor onto the anchor itself: a box in the column BESIDE the title is not a
    slip. VK Tech's contents slide keeps the list at 5.07in right of a title
    ending at 4.79in; aligned, the label above the list printed through the
    title («Средне» over «Проблема…», iter123)."""
    if anchor_shape is None or anchor_shape.left is None:
        return
    for shape in other_shapes:
        if shape is not None and shape.left is not None and abs(shape.left - anchor_shape.left) > tolerance_emu:
            if _move_lands_on_picture(shape, anchor_shape.left):
                continue
            if _move_lands_on(shape, anchor_shape.left, anchor_shape):
                continue
            shape.left = anchor_shape.left


def _overlaps(a, b):
    """(left, top, right, bottom) rectangles share an area."""
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _rect(shape):
    return (int(shape.left), int(shape.top or 0),
            int(shape.left + (shape.width or 0)), int((shape.top or 0) + (shape.height or 0)))


def _move_lands_on(shape, new_left, other):
    """Would moving `shape` to `new_left` put it over `other`, which it does
    not overlap where it stands?"""
    if other.left is None or not other.width:
        return False
    here = _rect(shape)
    moved = (int(new_left), here[1], int(new_left) + here[2] - here[0], here[3])
    return _overlaps(moved, _rect(other)) and not _overlaps(here, _rect(other))


def _move_lands_on_picture(shape, new_left):
    """Would moving `shape` to `new_left` put it over a picture of its slide
    that it does not overlap where it stands?"""
    try:
        slide = shape.part.slide
    except AttributeError:
        return False
    top, bottom = int(shape.top or 0), int((shape.top or 0) + (shape.height or 0))
    here = (int(shape.left), top, int(shape.left + (shape.width or 0)), bottom)
    moved = (int(new_left), top, int(new_left + (shape.width or 0)), bottom)
    for pic in slide.shapes:
        if pic.shape_type != MSO_SHAPE_TYPE.PICTURE or pic.left is None or not pic.width:
            continue
        rect = (int(pic.left), int(pic.top), int(pic.left + pic.width), int(pic.top + pic.height))
        if _overlaps(moved, rect) and not _overlaps(here, rect):
            return True
    return False


_SUBTITLE_SIZE_RATIO = 0.45


def _fill_title(slide, data, claimed_ids, resolver=None):
    title_shape = _pick_title_shape(slide, claimed_ids)
    _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)
    subtitle_shape = _pick_body_shape(slide, claimed_ids)
    if subtitle_shape is None:
        _append_subtitle_paragraph(title_shape, data.get("subtitle", ""), claimed_ids,
                                   resolver=resolver, slide=slide)
        return
    _align_left_edges(title_shape, subtitle_shape)
    _set_run_text(subtitle_shape, data.get("subtitle", ""), claimed_ids, resolver=resolver)


def _append_subtitle_paragraph(shape, subtitle, claimed_ids, resolver=None, slide=None):
    """Give the subtitle its own box under the title when the slide has no second one.

    The survey template's closing is a single text box and a decorative heart,
    so _pick_body_shape returns nothing and the subtitle the brief asked for was
    dropped without a trace — «Две недели на подключение и обучение команды»
    never reached any deck built on that slide. The box is 10.79x2.82in and the
    title takes 1.2in of it, so there is room to say it.

    A separate SHAPE, not a second paragraph of the title's box: the harness
    reads a title as its box's whole text, so appending there made
    «Запустим пилот… + Две недели…» a 99-character heading and criterion 2.1
    reported a long title that does not exist. Same shape iter70 gave the
    display stat's caption.

    Sized up front rather than left to enforce_text_fits, which shrinks by
    clamping every run to ONE size and would erase the hierarchy (iter55)."""
    if shape is None or not subtitle or not shape.height or slide is None:
        return
    reference = _reference_run(shape)
    title_pt = _max_font_pt(shape)
    if not title_pt:
        return
    subtitle_pt = max(MIN_READABLE_PT, round(title_pt * _SUBTITLE_SIZE_RATIO))
    metrics = _metrics_for_run(reference, resolver)
    width_in = Emu(shape.width).inches
    texts = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
    top_inset, bottom_inset = vertical_insets_emu(shape)
    usable_in = Emu(max(0, shape.height - top_inset - bottom_inset)).inches
    needed = (estimate_block_height_in(texts, width_in, title_pt, metrics=metrics)
              + estimate_block_height_in([subtitle], width_in, subtitle_pt, metrics=metrics))
    if needed > usable_in:
        return  # no room: better to drop it than to squash the closing line

    title_in = estimate_block_height_in(texts, width_in, title_pt, metrics=metrics)
    gap = int(Inches(0.25))
    top = int(shape.top) + int(Inches(title_in)) + gap
    height = int(Inches(estimate_block_height_in([subtitle], width_in, subtitle_pt,
                                                 metrics=metrics) + 0.1))
    if top + height > int(shape.top + shape.height):
        return

    box = slide.shapes.add_textbox(Emu(int(shape.left)), Emu(top),
                                   Emu(int(shape.width)), Emu(height))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = shape.text_frame.margin_left
    box.text_frame.margin_right = shape.text_frame.margin_right
    box.text_frame.margin_top = 0
    box.text_frame.text = subtitle
    run = box.text_frame.paragraphs[0].runs[0]
    if reference is not None:
        _copy_run_format(reference, run)
    run.font.size = Pt(subtitle_pt)
    run.font.bold = False
    # Claim it, or _clear_unclaimed_text blanks it later — the trap CLAUDE.md
    # records for the page numbers and the running topic.
    claimed_ids.add(box.shape_id)


def _find_bullet_icons(slide, body_shape):
    """Small pictures sitting to the left of a bullet-list body, roughly within
    its vertical span, that serve as bullet markers in decks that use icon
    graphics instead of a real PowerPoint bullet character. These are separate,
    independently-positioned shapes — they don't reflow when the text does, so
    filling different text into the body without moving them the same amount
    leaves icons pointing at the wrong line (or no line at all).

    Small in BOTH dimensions. With the width checked alone, the bars of a
    chart built from pictures passed for icons: VK Tech's chart slide has 22
    bars 0.76in wide and 1.75–3.29in tall, the filler moved them line by line
    and deleted the «surplus» — 19 of 22 — leaving the chart's white value
    labels white on white (iter108). Measured over the corpus: every other
    picture under 1in wide is at most 0.95in tall; the bars are the only
    narrow pictures taller than 1in."""
    if body_shape is None or body_shape.left is None:
        return []
    icon_max_size = Emu(int(Inches(1.0)))
    candidates = [
        s for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        and s.left is not None and s.top is not None
        and s.width and s.width < icon_max_size
        and s.height and s.height < icon_max_size
        and s.left <= body_shape.left
    ]
    candidates.sort(key=lambda s: s.top)
    return candidates




def _shape_line_spacing(shape):
    """The explicit line_spacing the shape's text actually carries (first
    paragraph that sets one) — the text-fitting estimate must model the same
    spacing rule the renderer will apply, or dense slides that "fit" by bare
    font height overflow at 1.5x spacing and get silently autofit-shrunk by
    the renderer, dragging every dependent position (bullet icons) with it."""
    if shape is None or not shape.has_text_frame:
        return None
    for p in shape.text_frame.paragraphs:
        if p.line_spacing is not None:
            return p.line_spacing
    return None


def _paragraph_line_height_pt(paragraph, font_size_pt, metrics=None):
    """The vertical rhythm the renderer will actually use for this paragraph
    (hand-made decks routinely set 1.5x spacing — the single biggest reason
    icon positions computed from bare font height drift further down every
    row). Percent spacing follows the renderers' own 1.2-em rule — see
    paragraph_pitch_pt."""
    return paragraph_pitch_pt(
        font_size_pt, metrics=metrics, line_spacing=paragraph.line_spacing
    )


def _spacing_pt(value, line_h_pt):
    if value is None:
        return 0.0
    # Length first — it subclasses int, so the numeric (multiple-of-line)
    # check would misread absolute spcPts values as huge multiples.
    if isinstance(value, Length):
        return value.pt
    if isinstance(value, (int, float)):
        return line_h_pt * value
    return value.pt


def _reposition_bullet_icons(icons, body_shape, texts, font_size_pt, metrics=None):
    """Moves each icon to line up with where its corresponding bullet actually
    starts now, estimated the same way the text sizing itself is (see
    text_fit.py) — better than leaving icons at their original fixed spots,
    which were only ever correct for the original text's line count and size.

    Handles a count mismatch instead of demanding exactly one icon per line:
    fewer texts than icons is the everyday case (the plan asked for 3 bullets
    on a 5-icon slide), and every icon beyond the last text is deleted — an
    orphaned marker pointing at blank space is the single most visible "AI
    slop" artifact a generated slide can have. With more texts than icons the
    icons that do exist still get aligned to their own lines."""
    if not icons or not texts or not font_size_pt or body_shape.width is None:
        return
    width_in = Emu(body_shape.width).inches
    frame_margins_in = frame_horizontal_margins_in(body_shape)
    top_inset, _ = vertical_insets_emu(body_shape)
    top = body_shape.top + top_inset
    # Walk EVERY paragraph in document order, not texts[i] -> paragraphs[i]:
    # slot-based filling (see _set_paragraph_texts) leaves the template's
    # spacer paragraphs between our items, and each spacer advances the
    # renderer's layout by its own blank line — an icon walk that skips them
    # drifts a full line per spacer.
    text_idx = 0
    for paragraph in body_shape.text_frame.paragraphs:
        is_content = any(r.text.strip() for r in paragraph.runs)
        # A cleared spacer's runs keep their own (template) size, and that —
        # not our fitted content size — is the height of its blank line.
        own_size = next(
            (r.font.size.pt for r in paragraph.runs if r.font.size), None
        )
        para_size_pt = font_size_pt if is_content else (own_size or font_size_pt)
        line_h_pt = _paragraph_line_height_pt(paragraph, para_size_pt, metrics=metrics)
        line_h_emu = int(Inches(line_h_pt / 72))
        before_emu = int(Inches(_spacing_pt(paragraph.space_before, line_h_pt) / 72))
        after_emu = int(Inches(_spacing_pt(paragraph.space_after, line_h_pt) / 72))

        top += before_emu
        if is_content and text_idx < len(texts):
            if text_idx < len(icons):
                icons[text_idx].top = Emu(int(top + max(0, line_h_emu - icons[text_idx].height) / 2))
            margins_in = frame_margins_in + paragraph_indent_in(paragraph)
            lines = estimate_wrapped_lines(
                texts[text_idx], width_in, font_size_pt, metrics=metrics, margins_in=margins_in
            )
            text_idx += 1
        else:
            lines = 1
        top += line_h_emu * lines + after_emu

    for icon in icons[len(texts):]:
        icon._element.getparent().remove(icon._element)


# Illustration-driven templates (T-Ж class) build lists as N SEPARATE one-line
# text boxes with tiny marker shapes beside them — not one multi-paragraph
# body. Filling such a slide through the body path writes everything into one
# box and _clear_unclaimed_text then blanks the real item boxes, leaving
# orphaned markers ("dots without text", a real user report).
_SLOT_ALIGN_TOLERANCE_EMU = Emu(int(Inches(0.1)))
_SLOT_PITCH_JITTER_EMU = Emu(int(Inches(0.15)))
_MARKER_MAX_SIZE_EMU = Emu(int(Inches(0.5)))
_MIN_SLOT_BOXES = 3


def _uniform_axis(values):
    """Distinct sorted axis coordinates (tolerance-bucketed), or None when the
    gaps between neighbors aren't near-constant.

    The steps are measured on the REAL coordinates (the smallest in each
    bucket), not on the bucket floors: flooring to 0.1in adds up to 0.1in of
    error per coordinate, i.e. up to 0.2in to a step difference — more than the
    whole jitter allowance. T-Zh mono's four-column slide sits at 0.31 / 2.67 /
    5.14 / 7.49in: real steps 2.36 / 2.47 / 2.35 (0.12 apart, a grid), floored
    steps 2.3 / 2.5 / 2.3 (0.2 apart, "not a grid") — so its captions were never
    seen as list slots and the bullets went into the title box (iter118 probe,
    12-slide deck)."""
    buckets = {}
    for v in values:
        key = int(v // _SLOT_ALIGN_TOLERANCE_EMU)
        buckets[key] = min(int(v), buckets.get(key, int(v)))
    coords = [key * _SLOT_ALIGN_TOLERANCE_EMU for key in sorted(buckets)]
    real = [buckets[key] for key in sorted(buckets)]
    if len(real) > 2:
        steps = [real[i + 1] - real[i] for i in range(len(real) - 1)]
        if max(steps) - min(steps) > _SLOT_PITCH_JITTER_EMU:
            return None
    return coords


# Real list grids measure 18-31% of the slide width; a numbering badge row is 3%.
_MIN_SLOT_WIDTH_FRACTION = 0.10


def _find_slot_boxes(slide, claimed_ids):
    """A GRID (R rows × C columns, either may be 1) of >=3 same-size text
    boxes with near-constant steps along each occupied axis — the multi-box
    list pattern in both its shapes: a vertical column of items (T-Ж) or a
    numbered 2×3 grid (plan 10в: the vertical-only detector filled the left
    column and blanked the right one into "invisible" leftovers). Returns
    boxes in reading order (row-major), or []."""
    # A grid of same-size boxes is only a LIST pattern if its cells can actually
    # hold list text. The T-Zh universal template numbers its items with a row of
    # 0.26in badges ("01".."06") — a perfect grid by every geometric test, and
    # filling it stuffed each bullet into a 19pt-wide column that rendered as an
    # 8pt stack of three-letter fragments. Measured across the corpus, real list
    # grids are 18-31% of the slide width and the badge row is 3%, so the cut is
    # nowhere near either side.
    # Chrome is NOT filtered out of the candidates here, unlike everywhere else.
    # The chrome test is "short box hugging a slide edge band", and the T-Zh
    # universal list is 6 items in a 3x2 grid whose BOTTOM ROW ends 0.32in inside
    # that band — so three of the six cells read as furniture, get_capacity
    # reported 3, and half the designed layout could never be used. A box in a
    # COMPLETE grid alongside non-chrome cells is content whatever band it sits
    # in: page numbers and header strips do not form grids with content boxes.
    # The all-chrome guard below keeps a genuine repeated footer strip out.
    # Measured over all five real templates: this changes exactly one slide
    # (universal #2, 3 -> 6) and creates no all-chrome grid anywhere.
    min_slot_width = _slide_width(slide) * _MIN_SLOT_WIDTH_FRACTION
    boxes = [
        s for s in _text_shapes(slide)
        if s.shape_id not in claimed_ids
        and s.left is not None and s.top is not None and s.width and s.height
        and s.width >= min_slot_width
    ]
    by_shape = {}
    for b in boxes:
        key = (
            int(b.width // _SLOT_ALIGN_TOLERANCE_EMU),
            int(b.height // _SLOT_ALIGN_TOLERANCE_EMU),
        )
        by_shape.setdefault(key, []).append(b)

    best = []
    for group in by_shape.values():
        if len(group) < _MIN_SLOT_BOXES:
            continue
        xs = _uniform_axis([b.left for b in group])
        ys = _uniform_axis([b.top for b in group])
        if xs is None or ys is None:
            continue
        # Every box must occupy its own grid cell, and the grid must be full —
        # a ragged scatter that happens to share a size is not a slot pattern.
        cells = {}
        for b in group:
            cell = (int(b.top // _SLOT_ALIGN_TOLERANCE_EMU), int(b.left // _SLOT_ALIGN_TOLERANCE_EMU))
            if cell in cells:
                break
            cells[cell] = b
        else:
            if len(group) != len(xs) * len(ys) or len(group) <= len(best):
                continue
            # A grid made ENTIRELY of chrome is furniture, not a list — a footer
            # strip repeated across the slide would otherwise become fill
            # targets now that chrome is allowed in as a candidate.
            height = _slide_height(slide)
            if all(is_chrome_shape(b, height) for b in group):
                continue
            best = sorted(group, key=lambda b: (
                int(b.top // _SLOT_ALIGN_TOLERANCE_EMU),
                int(b.left // _SLOT_ALIGN_TOLERANCE_EMU),
            ))
    return best


def _slot_markers(slide, slot_box):
    """Marker shapes belonging to this specific list item: small dot
    auto-shapes / icon pictures AND narrow numbering textboxes («01», «02» in
    grid templates, plan 10в) sitting in the same row band, immediately to the
    slot's left. Deleted together with a surplus slot — a leftover number next
    to blank space is the same orphan artifact as a dot without text."""
    markers = []
    for s in slide.shapes:
        if s.left is None or s.top is None or not s.width or not s.height:
            continue
        if s.width > _MARKER_MAX_SIZE_EMU or s.height > _MARKER_MAX_SIZE_EMU:
            continue
        is_glyph = s.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE)
        is_number_box = (
            s.has_text_frame and len(s.text_frame.text.strip()) <= 3
            and s.shape_id != slot_box.shape_id
        )
        if not (is_glyph or is_number_box):
            continue
        if s.left >= slot_box.left:
            continue
        # only the marker directly beside the slot, not one from a neighboring
        # grid column: it must sit closer than one slot-width away
        if slot_box.left - s.left > slot_box.width:
            continue
        center = s.top + s.height // 2
        if slot_box.top - _SLOT_ALIGN_TOLERANCE_EMU <= center <= slot_box.top + slot_box.height + _SLOT_ALIGN_TOLERANCE_EMU:
            markers.append(s)
    return markers


def _harmonize_slot_sizes(prs, slide_indices):
    """Give every slot in a row the same font size, AFTER the final fit pass.

    Each slot is fitted on its own and enforce_text_fits then shrinks each box
    independently, so a short item keeps the template size while a long
    neighbour drops — measured on real templates, one row shipped 8/9/13pt and
    another 9/11/15pt inside identical boxes, which reads as sloppy the moment
    the items sit side by side. Same argument as harmonize_clone_font_sizes:
    slot boxes are equal-size by construction (that is how they are detected),
    so the smallest size has already proved it fits every one of them. Only ever
    shrinks. Returns (slide_idx, shape_id, new_size_pt) so markers positioned
    from the old size get re-laid-out, as CLAUDE.md requires of any post-QA
    resize."""
    changes = []
    for slide_idx in slide_indices:
        slide = prs.slides[slide_idx]
        # Re-detecting the grid here would miss most rows: _fill_list_slots
        # physically removes the surplus boxes, so what is left is no longer the
        # full R×C grid the detector insists on (that is why the T-Zh mono row
        # stayed 9/11/15pt on the first attempt). Group by box SIZE instead —
        # equal-size boxes on one slide are the row, however many survived.
        by_size = {}
        for shape in _content_text_shapes(slide, set()):
            if not shape.width or not shape.height or not shape.text_frame.text.strip():
                continue
            key = (int(shape.width // _SLOT_ALIGN_TOLERANCE_EMU),
                   int(shape.height // _SLOT_ALIGN_TOLERANCE_EMU))
            by_size.setdefault(key, []).append(shape)
        # Every same-size group, not just the biggest one. Taking the largest
        # was wrong as soon as slot numbering badges survived the fill: a
        # three-item row then has THREE item boxes and THREE badge boxes, and
        # the badges won the tie, leaving the items ragged again (caught by
        # test_slots_in_one_row_ship_a_single_font_size). A slide can also hold
        # two rows of different-size slots, and both deserve harmonising.
        for group in by_size.values():
            if len(group) < 2:
                continue
            sizes = [
                run.font.size.pt
                for slot in group
                for para in slot.text_frame.paragraphs for run in para.runs
                if run.font.size and run.text.strip()
            ]
            if len(sizes) < 2 or max(sizes) == min(sizes):
                continue
            target = min(sizes)
            for slot in group:
                shrank = False
                for para in slot.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size and run.font.size.pt > target:
                            run.font.size = Pt(target)
                            shrank = True
                if shrank:
                    changes.append((slide_idx, slot.shape_id, target))
    return changes


# Marks a list-item numbering badge as OURS. _renumber_static_slide_numbers
# rewrites any chrome all-digit box to the page number, and once the universal
# grid's bottom row became fillable its badges «02»/«04»/«06» qualified — three
# of them were rewritten to «01». Claiming only protects a shape from
# _clear_unclaimed_text; this survives to the end of the pipeline.
SLOT_BADGE_NAME = "SlotBadge"


def _renumber_slot_badge(marker, position):
    """Renumber a list item's numbering badge to its position in the SHIPPED
    list, keeping the template's zero padding («01», not «1»).

    Needed because a grid is not always numbered in reading order: the T-Zh
    universal slide is 6 items in a 3x2 grid numbered down the columns, so its
    top row reads «01 03 05». That is correct for six items and looks like a
    bug for three — which is what a three-bullet deck ships. Same narrowness as
    _renumber_static_slide_numbers: a single run, digits only, so a dot glyph or
    a lettered marker is left exactly as the designer drew it."""
    if not marker.has_text_frame:
        return
    runs = [r for p in marker.text_frame.paragraphs for r in p.runs]
    if len(runs) != 1:
        return
    text = runs[0].text.strip()
    if not text.isdigit():
        return
    runs[0].text = str(position).zfill(len(text))
    marker.name = SLOT_BADGE_NAME


# A box directly under a slot continues its stack when the gap between them is
# under this: VK Tech 26 sets «Текст» 0.10in under «Заголовок».
_SLOT_STACK_MAX_GAP_EMU = int(Inches(0.3))


def _slot_room(slide, slot, slots):
    """Height from the slot's top to the bottom of the stack of text boxes
    right under it in its own column — boxes that are not slots themselves,
    so filling the list blanks them and their room is the slot's.

    VK Tech's card slide sets «Заголовок» (14pt, one line, 0.16in) over
    «Текст» (0.22in) in each card. The list goes into the heading boxes, the
    «Текст» boxes are cleared, and the item budget came out as one 1.77in line
    — 17 characters: «Требуются 2», «Планируется» (iter138) while half the card
    stood empty. Only same-width boxes in the same column count; badges, the
    next row of slots and anything across the slide do not."""
    slot_ids = {s.shape_id for s in slots}
    left, width = int(slot.left), int(slot.width)
    bottom = int(slot.top + slot.height)
    below = sorted(
        (s for s in _text_shapes(slide)
         if s.shape_id not in slot_ids and s.shape_id != slot.shape_id
         and s.left is not None and s.top is not None and s.width and s.height
         and abs(int(s.left) - left) <= _SLOT_ALIGN_TOLERANCE_EMU
         and abs(int(s.width) - width) <= _SLOT_ALIGN_TOLERANCE_EMU
         and int(s.top) >= bottom - _SLOT_ALIGN_TOLERANCE_EMU),
        key=lambda s: int(s.top))
    for s in below:
        if int(s.top) - bottom > _SLOT_STACK_MAX_GAP_EMU:
            break
        bottom = max(bottom, int(s.top + s.height))
    return bottom - int(slot.top)


def _fill_list_slots(slide, slots, texts, claimed_ids, resolver=None):
    """One text per slot box; surplus slot boxes are physically removed along
    with their markers (a blanked box would still hold layout space, and its
    orphaned dot is the exact artifact this path exists to prevent)."""
    position = 1
    # Each slot takes the room of the stack under it (see _slot_room) BEFORE
    # its text is fitted: the item budget counted those lines, and a fit to the
    # one-line box would shrink them back into it.
    rooms = [_slot_room(slide, slot, slots) for slot in slots[:len(texts)]]
    for slot, room in zip(slots, rooms):
        if room > int(slot.height):
            slot.height = Emu(room)
            # A leading under 100% is only harmless on one line: VK Tech sets
            # its card headings at 64%, and a wrapped item printed its lines
            # over each other (iter139).
            for paragraph in slot.text_frame.paragraphs:
                spacing = paragraph.line_spacing
                if isinstance(spacing, float) and not isinstance(spacing, Length) and spacing < 1.0:
                    paragraph.line_spacing = 1.0
    for slot, text in zip(slots, texts):
        _set_run_text(slot, text, claimed_ids, resolver=resolver)
        # A numbering badge beside a FILLED slot is the template's furniture,
        # not stale content — but _clear_unclaimed_text blanks every text shape
        # nobody claimed, so the T-Zh grid shipped bare items where the template
        # reads «01 Название пункта … 06». Claiming them is exact: the same
        # _slot_markers that DELETES a badge next to a removed slot keeps the
        # one next to a filled slot, with no new heuristic about what a number
        # on a slide means.
        for marker in _slot_markers(slide, slot):
            claimed_ids.add(marker.shape_id)
            _renumber_slot_badge(marker, position)
        position += 1
    for surplus in slots[len(texts):]:
        for marker in _slot_markers(slide, surplus):
            marker._element.getparent().remove(marker._element)
        surplus._element.getparent().remove(surplus._element)


def _group_texts(group):
    """Text shapes inside a group, groups opened, top to bottom in the group's
    own space."""
    found = []

    def walk(g):
        for child in g.shapes:
            if child.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(child)
            elif getattr(child, "has_text_frame", False) and child.text_frame.text.strip():
                found.append(child)

    walk(group)
    return sorted(found, key=lambda s: (int(s.top or 0), int(s.left or 0)))


def _find_group_slots(slide, claimed_ids):
    """A GRID of >=3 same-size content GROUPS — cards — as list slots, in
    reading order, or [].

    VK Tech keeps each card of its four-card slide as a group (icon +
    «Заголовок / Подзаголовок / Текст» and more «Текст»). The native list
    filler reads only the top level of slide.shapes, so it saw one text box —
    the title's — wrote every bullet into it at 7pt, lost the title and left
    28 placeholder strings in the cards (iter120). Same grid test as
    _find_slot_boxes, on the groups' own boxes."""
    groups = [g for g in content_groups(slide) if g.shape_id not in claimed_ids
              and g.width and g.height]
    by_size = {}
    for g in groups:
        key = (int(g.width // _SLOT_ALIGN_TOLERANCE_EMU), int(g.height // _SLOT_ALIGN_TOLERANCE_EMU))
        by_size.setdefault(key, []).append(g)
    best = []
    for same in by_size.values():
        if len(same) < _MIN_SLOT_BOXES or len(same) <= len(best):
            continue
        xs = _uniform_axis([g.left for g in same])
        ys = _uniform_axis([g.top for g in same])
        if xs is None or ys is None or len(same) != len(xs) * len(ys):
            continue
        best = sorted(same, key=lambda g: (int(g.top // _SLOT_ALIGN_TOLERANCE_EMU),
                                           int(g.left // _SLOT_ALIGN_TOLERANCE_EMU)))
    return best


def _fill_group_slots(slide, groups, texts, claimed_ids, resolver=None):
    """One item per card: into the card's first text box; the card's other
    text boxes go (a blanked box keeps its bullet glyphs and holds space), the
    card's icon and backing stay. Surplus cards are removed whole."""
    filled = []
    for group, text in zip(groups, texts):
        inner = _group_texts(group)
        if not inner:
            continue
        _set_run_text(inner[0], text, claimed_ids, resolver=resolver)
        filled.append(inner[0])
        for other in inner[1:]:
            # The markers beside a removed line (the cards' grey squares are
            # shapes of their own) go with it, as _slot_markers does for slots.
            for marker in _group_markers(group, other):
                marker._element.getparent().remove(marker._element)
            other._element.getparent().remove(other._element)
    for surplus in groups[len(texts):]:
        surplus._element.getparent().remove(surplus._element)
    # One size across the cards: each box is fitted to its own text, and a
    # longer item came out a size smaller than its neighbours.
    sizes = [_max_font_pt(box) for box in filled if _max_font_pt(box)]
    if len(sizes) > 1 and min(sizes) < max(sizes):
        for box in filled:
            for paragraph in box.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(min(sizes))


def _group_markers(group, box):
    """Small non-text shapes of the group in `box`'s row band, left of it —
    the list-marker glyphs of a card line (group space throughout)."""
    center_band = (int(box.top or 0) - _SLOT_ALIGN_TOLERANCE_EMU,
                   int((box.top or 0) + (box.height or 0)) + _SLOT_ALIGN_TOLERANCE_EMU)
    found = []
    for child in group.shapes:
        if child.shape_id == box.shape_id or child.left is None or not child.width:
            continue
        if getattr(child, "has_text_frame", False) and child.text_frame.text.strip():
            continue
        if child.width > _MARKER_MAX_SIZE_EMU or child.height > _MARKER_MAX_SIZE_EMU:
            continue
        center = int(child.top) + int(child.height) // 2
        if int(child.left) < int(box.left or 0) and center_band[0] <= center <= center_band[1]:
            found.append(child)
    return found


def _fill_bullet_list(slide, data, claimed_ids, resolver=None):
    bullets = data.get("bullets", [])
    slots = _find_slot_boxes(slide, claimed_ids)
    if slots:
        _fill_list_slots(slide, slots, bullets, claimed_ids, resolver=resolver)
        title_shape = _pick_title_shape(slide, claimed_ids)
        _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)
        return
    cards = _find_group_slots(slide, claimed_ids)
    if cards:
        _fill_group_slots(slide, cards, bullets, claimed_ids, resolver=resolver)
        title_shape = _pick_title_shape(slide, claimed_ids)
        _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)
        return

    # Body must be picked (and claimed) before title: _pick_body_shape's native
    # placeholder idx=1 check is a deterministic signal straight from the
    # template's own XML, while _pick_title_shape falls back to "biggest
    # explicit run font size" whenever there's no real title placeholder — and
    # that heuristic is blind to a shape whose size is inherited rather than
    # set on the run (common for short secondary captions in hand-built decks).
    # Picking title first let it misread the real multi-paragraph body as the
    # title on such slides, stranding every bullet in a caption-sized box.
    body_shape = _pick_body_shape(slide, claimed_ids)
    bullets = data.get("bullets", [])
    # Icon detection must precede any left-alignment: in icon-bullet designs the
    # body is intentionally indented to the right of the icon column, so snapping
    # it to the title's left edge would shove the text onto the icons AND break
    # the "icon is left of the body" detection itself.
    icons = _find_bullet_icons(slide, body_shape)
    font_size_pt = _set_paragraph_texts(
        body_shape, bullets, claimed_ids, resolver=resolver, allow_center=not icons,
        prefer_single_line=bool(icons),
    )

    title_shape = _pick_title_shape(slide, claimed_ids)
    _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)
    if not icons:
        _align_left_edges(title_shape, body_shape)

    if body_shape is not None and icons:
        metrics = _metrics_for_run(_reference_run(body_shape), resolver)
        _reposition_bullet_icons(icons, body_shape, bullets, font_size_pt, metrics=metrics)


def _embolden_stat_numbers(shape, stats):
    """In the single-text-block stats fallback each KPI renders as a plain
    "num — label" line, visually indistinguishable from body prose. Splitting
    the number into its own bold run (same size — growing it would change the
    line pitch under every position we've already computed) gives the KPI the
    emphasis a dedicated number box would have had."""
    if shape is None:
        return
    import copy as _copy

    paragraphs = [p for p in shape.text_frame.paragraphs if p.runs and p.runs[0].text.strip()]
    for p, (num, label) in zip(paragraphs, stats):
        first = p.runs[0]
        if first.text != f"{num} — {label}":
            continue  # e.g. resized/merged content — don't guess at splitting
        # Clone the run element and insert the clone RIGHT AFTER the number run
        # (add_run() would append past the blanked leftover runs at the tail),
        # so the label inherits identical formatting minus the bold.
        label_r = _copy.deepcopy(first._r)
        first._r.addnext(label_r)
        first.text = f"{num} — "
        first.font.bold = True
        p.runs[1].text = label


def _stats_title_shape(slide, claimed_ids):
    """For stats slides the title is the TOPMOST content box, not the biggest
    font: KPI digits are the largest text on the slide by design, so the
    generic font-size heuristic reliably crowns a number box (seen live on the
    T-Ж «20 227 000» slide — "Метрики" landed inside the giant number)."""
    candidates = [s for s in _content_text_shapes(slide, claimed_ids) if s.top is not None]
    if not candidates:
        return _pick_title_shape(slide, claimed_ids)
    topmost = min(candidates, key=lambda s: s.top)
    if _is_display_figure(topmost):
        return None
    return topmost


# A display figure's caption reads as a caption at about a third of it; below
# the readability floor the number is too big for its box anyway and the fitters
# will pull both down together.
_DISPLAY_LABEL_RATIO = 0.33


_DISPLAY_LABEL_MIN_STRIP_IN = 0.22


def _free_strip_below(slide, shape):
    """Height of the empty band under `shape`, in EMU, or 0.

    The band ends at the nearest shape that overlaps this one horizontally —
    on the T-Zh study stat slide that is the card's rule at 5.00in, leaving
    0.41in under a box ending at 4.59in."""
    if shape.top is None or not shape.height:
        return 0
    bottom = int(shape.top + shape.height)
    left, right = int(shape.left), int(shape.left + shape.width)
    limit = int(_slide_height(slide))
    for other in slide.shapes:
        if other.shape_id == shape.shape_id:
            continue
        if other.top is None or other.left is None or not other.width:
            continue
        if int(other.left) >= right or int(other.left + (other.width or 0)) <= left:
            continue
        top = int(other.top)
        if bottom <= top < limit:
            limit = top
    return max(0, limit - bottom)


def _fill_display_stat(shape, pair, claimed_ids, resolver=None, slide=None):
    """Number big, label under it — the treatment a one-big-number slide has."""
    number, label = str(pair[0]), str(pair[1])
    reference = _dominant_reference_run(shape) or _reference_run(shape)
    base = _max_font_pt(shape) or (reference.font.size.pt if reference and reference.font.size else 40)
    metrics = _metrics_for_run(reference, resolver)
    number_pt = cap_size_to_longest_word(
        number, shape.width, base, metrics=metrics,
        margins_in=horizontal_margins_in(shape), bold=True)

    # Fit the PAIR up front. enforce_text_fits shrinks by clamping every run to
    # one size, so a pair that arrives too tall comes back flattened — measured:
    # 82pt over 27pt became 24pt over 24pt, and the hierarchy this whole branch
    # exists for was gone.
    top_inset, bottom_inset = vertical_insets_emu(shape)
    usable_in = Emu(max(0, shape.height - top_inset - bottom_inset)).inches
    width_in = Emu(shape.width).inches
    margins = horizontal_margins_in(shape)
    spacing = _shape_line_spacing(shape)

    # The designer's box holds the figure ALONE: on the T-Zh study stat slide it
    # is 1.38in tall and one line of its own 82pt is 1.37in. Adding the label as
    # a second paragraph therefore forced the fit down to 24pt over 18pt — a
    # bold word in a sentence, not a KPI. If the template left an empty strip
    # under that box (0.41in there, down to the card's rule), the label goes in
    # the strip and the figure keeps the size the template set.
    strip = _free_strip_below(slide, shape) if slide is not None else 0
    number_in = estimate_block_height_in([number], width_in, number_pt, metrics=metrics,
                                         line_spacing=spacing, margins_in=margins)
    # The label is FITTED to the strip, not held at a ratio of the figure: a
    # third of 82pt is 27pt, which wraps to two lines and needs 0.9in against
    # the 0.41in the card leaves. Shrink until it fits; if even the readable
    # floor will not, the old in-box pair stands.
    strip_in = Emu(strip).inches
    label_pt = max(MIN_READABLE_PT, round(number_pt * _DISPLAY_LABEL_RATIO))
    label_in = estimate_block_height_in([label], width_in, label_pt, metrics=metrics,
                                        line_spacing=spacing, margins_in=margins)
    while label_pt > MIN_READABLE_PT and label_in > strip_in:
        label_pt -= 1
        label_in = estimate_block_height_in([label], width_in, label_pt, metrics=metrics,
                                            line_spacing=spacing, margins_in=margins)
    # The gate is ONE LINE, not a height comparison. Our own model says the
    # designer's 82pt line needs 1.64in in its 1.38in box (spacing 1.2 counts
    # 1.2 x size x percent, CLAUDE.md), i.e. the template "overflows" itself —
    # measuring the figure against the box would refuse the very layout the
    # template ships. What actually matters is that the figure stays a single
    # line, exactly as the designer set it.
    number_lines = estimate_wrapped_lines(number, width_in, number_pt,
                                          metrics=metrics, margins_in=margins)
    if (number_lines <= 1
            and Emu(strip).inches >= max(_DISPLAY_LABEL_MIN_STRIP_IN, label_in)):
        _set_run_text(shape, number, claimed_ids, resolver=resolver)
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(number_pt)
                run.font.bold = True
        caption = slide.shapes.add_textbox(
            Emu(int(shape.left)), Emu(int(shape.top + shape.height)),
            Emu(int(shape.width)), Emu(int(Inches(label_in))))
        caption.text_frame.word_wrap = True
        # Same insets as the figure's box, or the caption starts 0.1in to its
        # right (the default inset of a fresh textbox) and the two no longer
        # share a left edge — visible on the render.
        caption.text_frame.margin_left = shape.text_frame.margin_left
        caption.text_frame.margin_right = shape.text_frame.margin_right
        caption.text_frame.margin_top = 0
        caption.text_frame.text = label
        run = caption.text_frame.paragraphs[0].runs[0]
        reference_run = _reference_run(shape)
        if reference_run is not None:
            _copy_run_format(reference_run, run)
        run.font.size = Pt(label_pt)
        run.font.bold = False
        claimed_ids.add(caption.shape_id)
        return

    while number_pt > MIN_READABLE_PT:
        label_pt = max(MIN_READABLE_PT, round(number_pt * _DISPLAY_LABEL_RATIO))
        est = (estimate_block_height_in([number], width_in, number_pt, metrics=metrics,
                                        line_spacing=spacing, margins_in=margins)
               + estimate_block_height_in([label], width_in, label_pt, metrics=metrics,
                                          line_spacing=spacing, margins_in=margins))
        if est <= usable_in:
            break
        number_pt -= 2
    label_pt = max(MIN_READABLE_PT, round(number_pt * _DISPLAY_LABEL_RATIO))

    _set_paragraph_texts(shape, [number, label], claimed_ids, resolver=resolver)
    paragraphs = [p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
    for para, size, bold in zip(paragraphs, (number_pt, label_pt), (True, False)):
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.bold = bold


def _order_stat_pair(first, second):
    """(figure_box, label_box) for one KPI pair — decided by SIZE, not by order.

    Writing the number into the first box and the label into the second assumes
    every board lists figure-then-caption. The T-Zh universal board does the
    opposite: its caption sits at 11pt BELOW a 24pt figure, and reading order
    hands the caption over first. The render showed the consequence twice on one
    slide — «времени на подготовку регулярной отчётности» printed above «-40%»,
    and a caption forced into the figure's 24pt box broke mid-word as
    «подключённых клиентов е-/жемесячно».

    A figure is set larger than its caption in every board in the corpus, so the
    bigger box wins; on a tie the higher one does, which is where a figure sits
    when both are the same size."""
    first_pt = _max_font_pt(first) or 0
    second_pt = _max_font_pt(second) or 0
    if first_pt != second_pt:
        return (first, second) if first_pt > second_pt else (second, first)
    if (first.top or 0) != (second.top or 0):
        return (first, second) if (first.top or 0) < (second.top or 0) else (second, first)
    return first, second


# Figures and captions are told apart by size when the largest box is at least
# this many times the smallest: display figures on the VK Tech board are
# 116-153pt over 10pt captions, the T-Zh universal board 24pt over 11pt.
_FIGURE_SIZE_RATIO = 1.5


def _center(shape):
    return (int(shape.left or 0) + int(shape.width or 0) / 2,
            int(shape.top or 0) + int(shape.height or 0) / 2)


def _stat_pairs(boxes):
    """(figure_box, label_box) pairs for a KPI board, in reading order.

    Pairing neighbours in reading order — boxes[0] with boxes[1] — assumes each
    figure is followed by its own caption. A board that sets the figures in a
    ROW and the captions in a row under them breaks that: on VK Tech the two
    display boxes (153pt, 116pt) come first, so «4 из 6» and its caption
    «модулей CRM» both went into display boxes and wrapped over each other at
    116-153pt, while the next figure, «86%», went into a 10pt caption box
    (iter114, slide 4).

    When the boxes split cleanly into big figures and small captions, one
    caption per figure, each figure takes the nearest free caption. Otherwise
    the old neighbour pairing stands (ordered inside the pair by size)."""
    sizes = [(_max_font_pt(b) or 0) for b in boxes]
    known = [s for s in sizes if s]
    if len(known) == len(boxes) and known and max(known) >= _FIGURE_SIZE_RATIO * min(known):
        cut = (max(known) * min(known)) ** 0.5
        figures = [b for b, s in zip(boxes, sizes) if s >= cut]
        captions = [b for b, s in zip(boxes, sizes) if s < cut]
        if len(figures) == len(captions):
            pairs, free = [], list(captions)
            for figure in figures:
                fx, fy = _center(figure)
                caption = min(free, key=lambda c: abs(_center(c)[0] - fx) + abs(_center(c)[1] - fy))
                free.remove(caption)
                pairs.append((figure, caption))
            return pairs
    return [_order_stat_pair(boxes[i], boxes[i + 1]) for i in range(0, len(boxes) - 1, 2)]


def _keep_figure_on_one_line(box, text, resolver=None):
    """A KPI figure is ONE line — the rule _fill_display_stat already follows
    for a lone figure, applied to every figure box of a board. The template's
    own figure is «7» at 153pt in a 2.60in box; ours was «4 из 6», which at
    that size wraps to three lines that print over each other. The size comes
    down to the largest one line holds; enforce_text_fits does not do it,
    because the designer's box is shorter than its own line by design and QA
    leaves such boxes alone."""
    size = _max_font_pt(box)
    if not size or not text:
        return
    metrics = _metrics_for_run(_reference_run(box), resolver)
    # A Length (EMU), not points — compared as such it never looked smaller.
    fitted = fit_font_size_single_line([text], box.width, size, min_size_pt=MIN_READABLE_PT,
                                       metrics=metrics, margins_in=horizontal_margins_in(box))
    if fitted is not None and fitted.pt < size:
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = fitted


def _unify_figure_sizes(designed):
    """Figures the designer set at ONE size stay at one size.

    _keep_figure_on_one_line fits each figure to its own box, and a board's
    boxes differ in width: VK Tech's two-figure board sets «7» and «10» both at
    166pt in 2.60in and 4.73in boxes, and ours came out «27%» at 82pt beside
    «9%» at 166pt — the smaller number shouting (iter120-132, six decks).
    Every figure of a same-size group takes the group's smallest fitted size;
    figures the designer sized differently keep their hierarchy.
    `designed`: [(figure_box, size_pt_before_filling)]."""
    groups = {}
    for box, size in designed:
        if size:
            groups.setdefault(round(size), []).append(box)
    for boxes in groups.values():
        if len(boxes) < 2:
            continue
        smallest = min(_max_font_pt(b) or 0 for b in boxes)
        if not smallest:
            continue
        for box in boxes:
            for paragraph in box.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        run.font.size = Pt(smallest)


def _fill_stats_kpi(slide, data, claimed_ids, resolver=None):
    title_shape = _stats_title_shape(slide, claimed_ids)
    _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)

    boxes = _content_text_shapes(slide, claimed_ids)
    # Defensive, not redundant: two_phase.py's _validate_block rejects a
    # malformed [num, label] pair before it ever reaches here, but the legacy
    # parse_brief/resize_block path (content_parser/parser.py) historically
    # didn't — a single stray ["x10"] item crashed the whole generate() call
    # with "not enough values to unpack" deep inside this loop. A block this
    # far downstream should never be able to take the pipeline down over one
    # bad item; dropping it degrades to "one fewer stat," which the capacity
    # logic already handles gracefully.
    stats = [s for s in data.get("stats", []) if isinstance(s, (list, tuple)) and len(s) == 2]
    max_pairs = len(boxes) // 2

    if max_pairs >= 1:
        designed = []
        for (num, label), (figure_box, label_box) in zip(stats[:max_pairs], _stat_pairs(boxes)):
            designed.append((figure_box, _max_font_pt(figure_box)))
            _set_run_text(figure_box, num, claimed_ids, resolver=resolver)
            _keep_figure_on_one_line(figure_box, num, resolver=resolver)
            _set_run_text(label_box, label, claimed_ids, resolver=resolver)
        _unify_figure_sizes(designed)
    elif boxes:
        # Real decks often hold this kind of content as one text block next to
        # small bullet/icon graphics, not as separate number+label boxes — with
        # fewer than 2 boxes there's nowhere to split num/label into pairs, but
        # writing nothing here would leave those icons pointing at blank space,
        # and _clear_unclaimed_text would then blank out whatever was there before.
        # Same icon-marker pattern as bullet_list (see _fill_bullet_list) — reuse
        # the same detection/repositioning so the icons actually track the new
        # line count instead of staying planted at the original text's spacing.
        icons = _find_bullet_icons(slide, boxes[0])
        if (len(stats) == 1 and not icons
                and _is_display_sized(boxes[0], title_shape)):
            # A display-sized single box is a ONE BIG NUMBER slide (the same
            # slide get_capacity gives capacity 1, iter42). Writing
            # "«-40%» — времени на подготовку регулярной отчётности" as one run
            # at one size loses exactly that: the render showed a two-line
            # headline sentence where the template shows «20 227 000» alone at
            # 82pt. Number at the box's own display size, label under it at a
            # third of it — the fitters shrink from there if it does not fit.
            _fill_display_stat(boxes[0], stats[0], claimed_ids, resolver=resolver,
                               slide=slide)
            return
        lines = [f"{num} — {label}" for num, label in stats]
        font_size_pt = _set_paragraph_texts(
            boxes[0], lines, claimed_ids, resolver=resolver, allow_center=not icons,
            prefer_single_line=bool(icons),
        )
        _embolden_stat_numbers(boxes[0], stats)
        if icons:
            metrics = _metrics_for_run(_reference_run(boxes[0]), resolver)
            _reposition_bullet_icons(icons, boxes[0], lines, font_size_pt, metrics=metrics)


def _two_column_shapes_by_geometry(slide, claimed_ids):
    """Fallback for hand-built comparison slides with plain textboxes instead
    of placeholders: split the remaining text shapes at the slide's horizontal
    midpoint into two columns, and inside each column the topmost shape is the
    column heading, the tallest of the rest is the points body. Returns
    (left_heading, left_points, right_heading, right_points) — None where the
    geometry doesn't offer a matching shape."""
    candidates = [
        s for s in _content_text_shapes(slide, claimed_ids)
        if s.left is not None and s.top is not None
    ]
    if len(candidates) < 4:
        return None, None, None, None
    mid = min(s.left for s in candidates) + (
        max(s.left + (s.width or 0) for s in candidates) - min(s.left for s in candidates)
    ) // 2
    columns = {"left": [], "right": []}
    for s in candidates:
        center = s.left + (s.width or 0) // 2
        columns["left" if center < mid else "right"].append(s)
    if not columns["left"] or not columns["right"]:
        return None, None, None, None

    result = []
    for side in ("left", "right"):
        shapes = sorted(columns[side], key=lambda s: s.top)
        heading = shapes[0]
        rest = shapes[1:]
        points = max(rest, key=lambda s: (s.height or 0)) if rest else None
        result.extend([heading, points])
    return tuple(result)


def _fill_two_column_comparison(slide, data, claimed_ids, resolver=None):
    title_shape = _pick_title_shape(slide, claimed_ids)
    _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)

    left_head = _placeholder(slide, 1, claimed_ids)
    left_points = _placeholder(slide, 2, claimed_ids)
    right_head = _placeholder(slide, 3, claimed_ids)
    right_points = _placeholder(slide, 4, claimed_ids)
    if not all((left_head, left_points, right_head, right_points)):
        # No native placeholder set — a hand-built deck. Fall back to geometry
        # so the columns don't silently ship empty.
        left_head, left_points, right_head, right_points = _two_column_shapes_by_geometry(slide, claimed_ids)

    _set_run_text(left_head, data.get("left_heading", ""), claimed_ids, resolver=resolver)
    _set_paragraph_texts(left_points, data.get("left_points", []), claimed_ids, resolver=resolver)
    _set_run_text(right_head, data.get("right_heading", ""), claimed_ids, resolver=resolver)
    _set_paragraph_texts(right_points, data.get("right_points", []), claimed_ids, resolver=resolver)


# A body run this much larger than the slide's title is a display figure, not
# body copy. Measured: the T-Zh study stats slide is 82pt against a 26pt title
# (3.2x), its universal sibling's KPI board runs 11/24pt against a 42pt title.
_DISPLAY_SIZE_RATIO = 1.5


def _is_display_sized(shape, title_shape):
    """True when the box carries a display-size figure. Both sizes must be
    STATED: Google-Slides exports leave runs unsized, and guessing there is what
    _pick_title_shape already suffers from — an unsized box simply keeps the
    old "no fixed capacity" answer."""
    body = _max_font_pt(shape)
    title = _max_font_pt(title_shape) if title_shape is not None else None
    if not body or not title:
        return False
    return body >= title * _DISPLAY_SIZE_RATIO


def get_capacity(slide, archetype):
    """How many content items (bullets / stat pairs) the native slide structure
    can actually hold for this archetype — read-only, doesn't write anything.
    Returns None when there's no fixed capacity to match against (e.g. an
    icon-free single-text-block fallback, which can hold any number of lines,
    or an archetype this doesn't apply to). Must mirror the corresponding
    FILLERS entry's own shape-picking order exactly, or the capacity it
    reports and the shape that actually gets filled can disagree."""
    if archetype == "bullet_list":
        claimed_ids = set()
        # Same order as _fill_bullet_list: slot boxes first (multi-box lists),
        # then the native placeholder, then the title heuristic.
        slots = _find_slot_boxes(slide, claimed_ids)
        if slots:
            return len(slots)
        cards = _find_group_slots(slide, claimed_ids)
        if cards:
            return len(cards)
        body_shape = _pick_body_shape(slide, claimed_ids)
        if body_shape is None:
            return None
        # Content slots, not raw paragraph count: spacer paragraphs between
        # items are rhythm, not capacity (mirrors _set_paragraph_texts).
        capacity = len(_content_slot_indices(body_shape))
        icons = _find_bullet_icons(slide, body_shape)
        # Icon markers are fixed graphics, not text — they can't stretch to
        # match a paragraph count higher than how many of them physically
        # exist, so they're the real ceiling whenever they're present.
        return min(capacity, len(icons)) if icons else capacity

    if archetype == "stats_kpi":
        claimed_ids = set()
        # Mirrors _fill_stats_kpi exactly (topmost content box is the title,
        # chrome excluded) — capacity and filler must see the same boxes.
        title_shape = _stats_title_shape(slide, claimed_ids)
        if title_shape is not None:
            claimed_ids.add(title_shape.shape_id)
        boxes = _content_text_shapes(slide, claimed_ids)
        max_pairs = len(boxes) // 2
        if max_pairs >= 1:
            return max_pairs
        if not boxes:
            return None
        # Single-text-block fallback (see _fill_stats_kpi): still capped by
        # icon markers when present, same reasoning as bullet_list above.
        icons = _find_bullet_icons(slide, boxes[0])
        if icons:
            return len(icons)
        # One box set at DISPLAY size is a one-big-number slide, and its capacity
        # is exactly one. The T-Zh study template's stats slide is «20 227 000»
        # at 82pt under a 26pt title; reporting "no fixed capacity" let three
        # pairs in, and the render showed three same-size lines reading as
        # sentences — the design gone. The T-Zh universal KPI board (8 boxes at
        # 11/24pt) is unaffected: it never reaches this fallback.
        return 1 if _is_display_sized(boxes[0], title_shape) else None

    return None


# How much longer than the designer's own sample a generated item may run.
# The sample is an EXAMPLE, not a hard maximum, so a little slack is right; but
# the measured failure (a 24-char word in a slot whose sample is 15) has to land
# outside it, which puts the ceiling well under 1.6x.
ITEM_CHARS_SLACK = 1.2
# Never ask for less than a few real words, whatever a terse sample says.
MIN_ITEM_CHARS = 12


# Average character width is taken from real Russian body text, spaces
# included — they are part of every line.
_LINE_SAMPLE = "ручной сбор показателей из разных систем"


def _one_line_chars(shape, resolver=None):
    """Characters one line of `shape` holds at the size its own text is set
    (the fallback base size when the size is inherited)."""
    run = _reference_run(shape)
    size_pt = run.font.size.pt if run is not None and run.font.size else _FALLBACK_BASE_SIZE_PT
    width_pt = (Emu(shape.width).inches - horizontal_margins_in(shape)) * 72 * SINGLE_LINE_SAFETY
    metrics = _metrics_for_run(run, resolver)
    if metrics is not None:
        per_char = metrics.text_width_pt(_LINE_SAMPLE, size_pt) / len(_LINE_SAMPLE)
    else:
        per_char = size_pt * AVG_CHAR_WIDTH_FACTOR
    return max(0, int(width_pt / per_char)) if per_char else 0


def get_item_char_budget(slide, archetype, resolver=None):
    """How LONG a single list item may be on this slide — the companion to
    get_capacity, which only answers HOW MANY.

    Measured defect this exists for: the T-Zh mono list slot is 1.93x0.25in with
    the designer's own sample text "Название пункта" (15 chars). Content is
    generated against a FLAT 72-char budget, so a 24-char single word lands in
    it, cannot be wrapped (the box is one line tall) and cannot be broken, and
    the fitter shrinks it to its 9pt floor — a third of the ~14pt the template
    itself renders that label at. _fit_size_for_shape already says as much at
    that floor: past this point the CONTENT is too long.

    The budget comes from the template's OWN sample text, not from geometry:
    the size those slots inherit is not in the run (it resolves through the
    placeholder/master chain, and the theme routinely lies about the deck's real
    look), so measuring chars-per-line means guessing the very number that is
    unreliable. The designer already answered the question by writing an example
    of the right length into the box. The two agree where both can be computed:
    metrics give ~17 chars for that slot at its rendered size, the sample is 15.

    Returns None when there is no sample to learn from — the caller keeps its
    flat default, so a template that ships empty slots is no worse off.

    A sample SHORTER than one line of its slot says nothing about length: VK
    Tech writes «Пункт» and «Текст» into slots that hold 20-28 characters a
    line, the budget came out at 12, and every item was cut to one word —
    «Снижение», «Обязательная», «Выход» (iter125, iter126). There the line is
    the limit. Where the designer's sample fills the line or more (T-Zh mono
    «Название пункта», universal «Длинное название пункта»), it still decides."""
    if archetype != "bullet_list":
        return None
    claimed_ids = set()
    shapes = _find_slot_boxes(slide, claimed_ids)  # same order as _fill_bullet_list
    if shapes:
        samples = [s.text_frame.text.strip() for s in shapes if s.has_text_frame]
        line_box = shapes[0]
        lines = _room_lines(line_box, _slot_room(slide, line_box, shapes))
    else:
        body_shape = _pick_body_shape(slide, claimed_ids)
        if body_shape is None:
            return None
        samples = [
            "".join(run.text for run in body_shape.text_frame.paragraphs[i].runs).strip()
            for i in _content_slot_indices(body_shape)
        ]
        line_box = body_shape
        lines = 1
    lengths = [len(t) for t in samples if t]
    if not lengths:
        return None
    sampled = max(MIN_ITEM_CHARS, round(max(lengths) * ITEM_CHARS_SLACK))
    return max(sampled, lines * _one_line_chars(line_box, resolver))


def _room_lines(box, room_emu):
    """Whole lines of the box's own text that fit a room of room_emu height
    (at least one: the box itself is the floor)."""
    run = _reference_run(box)
    size_pt = run.font.size.pt if run is not None and run.font.size else _FALLBACK_BASE_SIZE_PT
    top_inset, bottom_inset = vertical_insets_emu(box)
    usable_pt = Emu(max(0, int(room_emu) - top_inset - bottom_inset)).pt
    spacing = _shape_line_spacing(box)
    if isinstance(spacing, float) and not isinstance(spacing, Length) and spacing < 1.0:
        spacing = 1.0  # as _fill_list_slots sets it once the slot wraps
    pitch = paragraph_pitch_pt(size_pt, line_spacing=spacing)
    return max(1, int(usable_pt // pitch)) if pitch else 1


# A template picture smaller than this share of the slide is an icon or a
# badge, not part of the slide's image slot. Measured on the VK templates: the
# image slots of VK Education's image_caption slides are 17% (the phone mockup)
# to 60% of the slide; the largest non-slot pictures — VK Tech's card
# illustrations — are 5.8–9.7%, and a threshold of 5% would have taken them.
IMAGE_SLOT_MIN_AREA_FRACTION = 0.10


def _is_picture_group(shape):
    """A GROUP made of pictures only, with no text anywhere inside — a mockup
    assembled from parts (VK Education's phone: a frame picture and a screen
    picture). python-pptx reports it as GROUP, so a PICTURE filter misses it."""
    if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
        return False
    leaves = []

    def walk(group):
        for child in group.shapes:
            if child.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(child)
            else:
                leaves.append(child)

    walk(shape)
    return bool(leaves) and all(leaf.shape_type == MSO_SHAPE_TYPE.PICTURE for leaf in leaves)


def _image_slot_pictures(slide):
    """The pictures that make up the image slot of an image_caption slide.

    The template's image slot never holds OUR image: it holds the designer's
    sample. On VK Education that was a laptop mockup of the VK Store community
    page (a girl in a hoodie, «место встречи, Невский 28») under our heading
    «Единая картина продаж», and a VK Education post screenshot under ours
    elsewhere (iter103, iter106).

    Every picture of a meaningful size that is not a background under the
    slide's text belongs to the slot. Repetition across slides does NOT spare
    it here, unlike in is_stale_data_picture: the VK Store screen is reused by
    the template's laptop and phone mockups alike, and brand art in an image
    slot is not the author's image either. A phone next to a laptop is two
    mockups and one slot, so all of them go — the phone being a group of
    pictures, not a picture."""
    area = (_slide_width(slide) or 0) * (_slide_height(slide) or 0)
    if not area:
        return []
    return [shape for shape in slide.shapes
            if (shape.shape_type == MSO_SHAPE_TYPE.PICTURE or _is_picture_group(shape))
            and shape.width and shape.height
            and shape.width * shape.height / area >= IMAGE_SLOT_MIN_AREA_FRACTION
            and not _covers_slide_text(shape, slide)]


# Skeleton frame and label colour on a native slide: the slide's own text colour
# is usually inherited from its placeholder and cannot be copied (nor resolved —
# CLAUDE.md), and an autoshape's default text is white. Mid-grey reads on the
# light and the dark templates alike.
_SKELETON_GREY = RGBColor(0x88, 0x88, 0x88)


def _replace_image_slot(slide, caption, style_run, margin_emu=0):
    """Swap the template's sample picture(s) for the dashed image skeleton the
    synthesized image_caption slides draw — the honest marker of «your image
    goes here» — over the box they occupied. The font is copied from
    `style_run` (the slide's title run) so the label is set in the template's
    typeface."""
    import copy

    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    pictures = _image_slot_pictures(slide)
    if not pictures:
        return None
    # Mockups are placed bleeding off the edge (the VK Education laptop runs
    # past the right side); a skeleton cut by the slide edge reads as broken.
    # The frame keeps the margin the template keeps: the title's own left
    # offset, mirrored on the other three sides.
    width, height = int(_slide_width(slide)), int(_slide_height(slide))
    left = max(min(int(p.left) for p in pictures), margin_emu)
    top = max(min(int(p.top) for p in pictures), margin_emu)
    right = min(max(int(p.left + p.width) for p in pictures), width - margin_emu)
    bottom = min(max(int(p.top + p.height) for p in pictures), height - margin_emu)
    if right - left < int(Inches(1)) or bottom - top < int(Inches(1)):
        return None
    for picture in pictures:
        picture._element.getparent().remove(picture._element)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Emu(left), Emu(top), Emu(right - left), Emu(bottom - top))
    frame.name = IMAGE_PLACEHOLDER_NAME
    frame.shadow.inherit = False
    frame.fill.background()
    frame.line.color.rgb = _SKELETON_GREY
    frame.line.width = Pt(1.5)
    ln = frame.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    tf = frame.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = [("ИЗОБРАЖЕНИЕ", 12, True)] + ([(caption, 14, False)] if caption else [])
    for i, (text, size, bold) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        if style_run is not None and style_run._r.rPr is not None:
            run._r.insert(0, copy.deepcopy(style_run._r.rPr))
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _SKELETON_GREY
    return frame


def _fill_image_caption(slide, data, claimed_ids, resolver=None):
    """Native image_caption slide: the template already supplies the artwork, so
    we only write the heading and the caption describing what the image shows.
    (Images still aren't generated — see synthesizer.synthesize_image_caption for
    the from-scratch path, which draws a dashed placeholder frame instead.)

    Slides with no text shape at all never reach here: template_spec.builder
    drops them from their family, because there would be nothing to fill — a
    real T-Zh member is a full-bleed photo with zero text boxes, and reaching
    this filler with it used to raise KeyError before image_caption had one."""
    title_shape = _pick_title_shape(slide, claimed_ids)
    _set_run_text(title_shape, data.get("title", ""), claimed_ids, resolver=resolver)
    caption_shape = _pick_body_shape(slide, claimed_ids)
    _align_left_edges(title_shape, caption_shape)
    description = data.get("image", "") or data.get("caption", "")
    _set_run_text(caption_shape, description, claimed_ids, resolver=resolver)
    # The description goes INTO the skeleton only when no caption box took it —
    # otherwise the slide would say the same sentence twice.
    caption_taken = caption_shape is not None and caption_shape.shape_id in claimed_ids
    style = None
    if title_shape is not None and title_shape.has_text_frame:
        style = next((r for p in title_shape.text_frame.paragraphs for r in p.runs), None)
    margin = int(title_shape.left) if title_shape is not None and title_shape.left else 0
    frame = _replace_image_slot(slide, None if caption_taken else description, style,
                                margin_emu=margin)
    if frame is not None:
        claimed_ids.add(frame.shape_id)


FILLERS = {
    "title": _fill_title,
    # Structurally the same as "title" (a heading ± subtitle) — the distinction is
    # purely semantic (where in the pitch it's used), which is content_parser's job.
    "section_divider": _fill_title,
    "closing": _fill_title,
    "bullet_list": _fill_bullet_list,
    "stats_kpi": _fill_stats_kpi,
    "two_column_comparison": _fill_two_column_comparison,
    "image_caption": _fill_image_caption,
}


# Moved to slide_kit (shared with canvas synthesis); alias kept for callers/tests.
from generator.slide_kit import clone_slide as _clone_slide


def _realign_icons_after_resize(prs, changes, resolver=None):
    """The final QA passes (enforce_text_fits, harmonize_clone_font_sizes) can
    shrink text AFTER the fillers positioned bullet-marker icons for the
    original size — a smaller font means a smaller line pitch, so the stale
    icon positions would drift progressively below their lines. Re-derives
    icon positions from the size the deck actually ships with.

    changes: iterables of (slide_idx, shape_id, ...) as returned by the QA
    passes; shapes without marker icons next to them are skipped unharmed."""
    for slide_idx, shape_id, *_ in changes:
        slide = prs.slides[slide_idx]
        shape = next((s for s in slide.shapes if s.shape_id == shape_id), None)
        if shape is None or not shape.has_text_frame:
            continue
        icons = _find_bullet_icons(slide, shape)
        if not icons:
            continue
        texts = [
            "".join(run.text for run in p.runs)
            for p in shape.text_frame.paragraphs
        ]
        texts = [t for t in texts if t.strip()]
        sizes = [
            run.font.size.pt
            for p in shape.text_frame.paragraphs
            for run in p.runs
            if run.font.size and run.text.strip()
        ]
        if not texts or not sizes:
            continue
        metrics = _metrics_for_run(_reference_run(shape), resolver)
        _reposition_bullet_icons(icons, shape, texts, max(sizes), metrics=metrics)


# Running-header slots that literally ask for the deck's topic. Measured across
# the real templates: "ТЕМА ПРЕЗЕНТАЦИИ" ships on 11 of 12 slides of the T-Zh
# teaching template and 10 of 12 of the universal one, i.e. it is the designer's
# placeholder for the very thing we know — and it went out visible on every
# generated slide. An exact-match allowlist, deliberately: "repeats across
# slides" does NOT separate a prompt from real furniture ("2025" also repeats 10
# times and must stay), so nothing is guessed at. Slots we cannot answer
# ("КОММЕНТАРИЙ", "Название пункта") are left alone rather than blanked — an
# empty box would just leave a hole in the designer's grid.
TOPIC_SLOT_PROMPTS = {
    "тема презентации",
    "название презентации",
    "тема доклада",
}
MAX_RUNNING_HEADER_CHARS = 42


RUNNING_TOPIC_ELLIPSIS = "…"


def _shorten_running_topic(title, shape=None, run=None, resolver=None, cased=None):
    """The deck's topic, cut to the header slot — and marked as cut.

    Trimming at a word boundary alone ships a phrase that LOOKS whole and is
    not: «Единая платформа управленческой отчётности и аналитики для розничной
    сети» became «ЕДИНАЯ ПЛАТФОРМА УПРАВЛЕНЧЕСКОЙ», an adjective governing a
    noun that is no longer there, printed on every slide of the deck. Two other
    real titles cut the same way («…ВНЕДРЕНИЯ В ТРЁХ»).

    So: drop a trailing function word the way iter56 does for bullets — a header
    must not end on a preposition either — and add an ellipsis, which is what
    tells the reader the topic is shortened rather than mangled.

    The budget is the BOX, not a character count, when the box can be measured.
    42 characters is a guess that fits nothing in particular: the T-Zh header
    slot is 1.85x0.12in — one line of 8pt — and holds «ТЕМА ПРЕЗЕНТАЦИИ» at
    84pt of its 133pt budget, while our 42-character topic measures 195pt and
    wrapped to two lines in a one-line box, on every slide of every deck. The
    character cap stays as the fallback for a box we cannot measure.

    `cased` is the transform the caller will apply before writing (these slots
    are set in caps on purpose): the width has to be measured on the string the
    renderer will actually draw, since Cyrillic capitals are the wider ones."""
    fitted = _fit_topic_to_box(title, shape, run, resolver, cased)
    if fitted is not None:
        return fitted
    if len(title) <= MAX_RUNNING_HEADER_CHARS:
        return title
    return _cut_topic(title, MAX_RUNNING_HEADER_CHARS)


def _cut_topic(title, max_chars):
    budget = max_chars - len(RUNNING_TOPIC_ELLIPSIS)
    head = title[:budget + 1]
    if " " not in head:
        return title  # one very long word: _fits_box_width decides its fate
    cut = head.rsplit(" ", 1)[0].rstrip(".,;:—- ")
    cut = drop_dangling_function_words(cut) or cut
    return cut + RUNNING_TOPIC_ELLIPSIS


def _fit_topic_to_box(title, shape, run, resolver, cased=None):
    """`title` shortened to one line of its box, or None when unmeasurable."""
    if shape is None or run is None or resolver is None or not shape.width:
        return None
    metrics = resolver.metrics_for(run.font.name) if run.font.name else None
    size_pt = run.font.size.pt if run.font.size else None
    if metrics is None or not size_pt:
        return None
    as_drawn = cased or (lambda text: text)
    budget_pt = _usable_width_in(
        Emu(shape.width).inches, horizontal_margins_in(shape)) * 72 * SINGLE_LINE_SAFETY
    if budget_pt <= 0:
        return None
    if metrics.text_width_pt(as_drawn(title), size_pt) <= budget_pt:
        return title

    words = title.split()
    for count in range(len(words) - 1, 0, -1):
        cut = drop_dangling_function_words(" ".join(words[:count]).rstrip(".,;:—- "))
        if not cut:
            continue
        candidate = cut + RUNNING_TOPIC_ELLIPSIS
        if metrics.text_width_pt(as_drawn(candidate), size_pt) <= budget_pt:
            return candidate
    return title  # even one word does not fit: _fits_box_width keeps the template's


def _fill_running_topic(prs, deck_title, resolver=None):
    """Replace the template's topic placeholder with the deck's real topic.

    Runs after _reorder_and_prune_slides like the renumbering, and follows the
    same conservative shape: chrome-band shapes only, single run only, exact
    match against the allowlist. Keeps the design's capitalisation — these slots
    are set in caps on purpose."""
    if not deck_title:
        return
    for slide in prs.slides:
        height = prs.slide_height
        for shape in slide.shapes:
            if not shape.has_text_frame or not is_chrome_shape(shape, height):
                continue
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            if len(runs) != 1:
                continue
            original = runs[0].text.strip()
            if original.lower() not in TOPIC_SLOT_PROMPTS:
                continue
            cased = str.upper if original.isupper() else None
            topic = _shorten_running_topic(
                deck_title.strip(), shape, runs[0], resolver, cased)
            topic = cased(topic) if cased else topic
            # Trimming by character count is not enough: a header box is small
            # and one long word can still be wider than it, which the renderer
            # then breaks mid-letter ("ВЫСОКОПРОИЗВОДИТЕЛЬНАЯ" in a 128pt-wide
            # box). If even the first word does not fit, keep the template's own
            # placeholder rather than shipping a broken one.
            if not _fits_box_width(topic, shape, runs[0], resolver):
                continue
            runs[0].text = topic


def _fits_box_width(text, shape, run, resolver):
    """True when every word of `text` fits the shape's usable width at the run's
    own font size — i.e. the renderer will not have to break a word mid-letter."""
    if not shape.width:
        return True
    metrics = resolver.metrics_for(run.font.name) if resolver and run.font.name else None
    if metrics is None:
        return True  # no metrics available: do not block on a guess
    size_pt = run.font.size.pt if run.font.size else 12.0
    budget_pt = max(1.0, (Emu(shape.width).inches - horizontal_margins_in(shape)) * 72)
    return all(metrics.text_width_pt(word, size_pt) <= budget_pt for word in text.split())


def _is_page_number(text, slide_count):
    """Not every all-digit box in the chrome band is a page number. The T-Zh
    universal template carries "2025" in its footer on ten slides, and
    renumbering rewrote the brand year to "0001", "0002", ... on every slide —
    caught by the cross-template smoke test. A page number is small: it cannot
    exceed the deck length, and anything year-sized is furniture."""
    try:
        value = int(text)
    except ValueError:
        return False
    return 0 < value <= max(slide_count, 99)


def _renumber_static_slide_numbers(prs):
    """A cloned canvas brings the template's own page number with it, so a deck
    whose first slide was cut from template slide 5 opens showing "05". The
    number is a plain text box on Google-Slides exports (no SLIDE_NUMBER
    placeholder), so nothing downstream updates it and the chrome filter — quite
    rightly — keeps it as furniture.

    Called only after _reorder_and_prune_slides, when prs.slides is already the
    shipping order, so position i simply means page i+1.

    Deliberately narrow, because "05" can also be real content (a KPI figure):
    only shapes the chrome filter already recognises as furniture, only a single
    run, and only text that is nothing but digits. Zero padding is preserved so
    "05" becomes "01", not "1"."""
    for position, slide in enumerate(prs.slides, start=1):
        height = prs.slide_height
        for shape in slide.shapes:
            if not shape.has_text_frame or not is_chrome_shape(shape, height):
                continue
            if shape.name == SLOT_BADGE_NAME:
                continue  # a list item's badge, numbered by position, not a page
            if shape.is_placeholder:
                try:
                    # A real slide-number placeholder renumbers itself.
                    if shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                        continue
                except (KeyError, ValueError):
                    pass
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            if len(runs) != 1:
                continue
            text = runs[0].text.strip()
            if not text.isdigit() or not _is_page_number(text, len(prs.slides._sldIdLst)):
                continue
            runs[0].text = str(position).zfill(len(text))


def _reorder_and_prune_slides(prs, ordered_slide_indices):
    """Reorders sldIdLst to ordered_slide_indices and permanently drops every
    other slide — clone sources plus any clone left over from stretch/shrink
    slack (see generate()) that a plan didn't end up using.

    Removing a <p:sldId> from sldIdLst alone leaves the slide PART and its
    presentation.xml.rels relationship in the package: still-valid XML, still
    a well-formed zip, but PowerPoint's own package-consistency check flags an
    OPC part reachable from a relationship yet absent from the actual slide
    show as corrupt and offers to "repair" — every generated deck that ever
    exercised the clone path (i.e. any plan reusing a template slide) hit
    this. `part.drop_rel` removes the relationship; `Package.iter_parts()`
    (what save() actually serializes) is a pure reachability walk from there,
    so a slide part with no remaining incoming relationship silently stops
    being written — no separate "delete the part" step needed."""
    sld_id_lst = prs.slides._sldIdLst
    all_ids = list(sld_id_lst)
    keep_ids = [all_ids[i] for i in ordered_slide_indices]
    drop_ids = [el for el in all_ids if el not in keep_ids]
    for el in all_ids:
        sld_id_lst.remove(el)
    for el in keep_ids:
        sld_id_lst.append(el)
    for el in drop_ids:
        prs.part.drop_rel(el.get(qn("r:id")))


def _template_cover_pt(prs):
    """The size the template sets its COVER title at, or None.

    Separate from _template_title_pt on purpose: a cover is deliberately louder
    than a content slide. Measured on slide 0 of the five real templates —
    42 / 68 / 34 / 92 / 92pt against their content titles of 42 / 51 / 26 / 40 /
    32 — so a flat 40pt was wrong on four of them, and less than half the
    designer's size on both survey decks."""
    slides = list(prs.slides)
    if not slides:
        return None
    title = _pick_title_shape(slides[0], set())
    if title is None:
        return None
    pt = _max_font_pt(title)
    return round(pt) if pt else None


def _template_content_gap_in(prs):
    """The gap the template leaves between a title and the content under it, in
    inches — the MEDIAN over its own content slides, or None.

    Synthesized slides centre their block in whatever band is left, which on a
    short block detaches the heading from what it introduces. The five real
    templates measure 0.00 / 0.06 / 0.20 / 0.25 / 1.12in — T-Zh mono really does
    hold its content low, so a flat "hug the title" would be as wrong as a flat
    "always centre". The number is the template's, like cover_pt and the cover's
    left margin.

    Chrome is excluded on both ends, and a slide needs two content boxes before
    it can say anything about the gap between them.
    """
    height = prs.slide_height
    gaps = []
    for slide in prs.slides:
        boxes = []
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.top is None or shape.height is None:
                continue
            if is_chrome_shape(shape, height):
                continue
            boxes.append((int(shape.top), int(shape.top) + int(shape.height)))
        if len(boxes) < 2:
            continue
        boxes.sort()
        gaps.append(Emu(max(0, boxes[1][0] - boxes[0][1])).inches)
    if not gaps:
        return None
    return round(statistics.median(gaps), 2)


def _template_cover_layout(prs):
    """(alignment, left_emu, top_emu, subtitle_top_emu) of the template's cover.

    Synthesized covers and closings were centred unconditionally in a box inset
    10% from each edge. Measured on the five real templates, EVERY title is left
    aligned, at left margins of 0.39 / 0.19 / 0.39 / 1.10 / 1.10in — so a centred
    block at 1.33in matched none of them, and the closing of a deck read as
    coming from a different template than the slides before it.

    Alignment None in the XML means "inherit", which resolves to left; it is
    reported as None here so the caller can tell "the designer said left" from
    "the designer said nothing", and both end up left.

    The TOP is measured for the same reason the size and the left margin are:
    a flat 38% of the slide height matched exactly one of the five real
    templates. Measured — 26 / 37 / 4 / 22 / 26% — mono opens its cover hard
    against the top edge, and a third of a slide of invented air above the
    title is not that template.

    The SUBTITLE is measured too, and it is anchored to the BOTTOM rather than to
    the title: all three T-Zh covers end their subtitle 0.19-0.35in above the
    slide edge, while the gap they leave under the title is 1.21 / 2.55 / 2.28in
    - no gap at all, just whatever is left over. Our synthesis always tucked it
    under the title, which on mono left the bottom two thirds of the slide
    empty. None when the cover carries no subtitle (both survey templates)."""
    slides = list(prs.slides)
    if not slides:
        return None, None, None, None
    cover = slides[0]
    title = _pick_title_shape(cover, set())
    if title is None or title.left is None or title.top is None:
        return None, None, None, None
    paragraphs = title.text_frame.paragraphs
    alignment = paragraphs[0].alignment if paragraphs else None

    below = [s for s in cover.shapes
             if s.has_text_frame and s.text_frame.text.strip()
             and s.shape_id != title.shape_id and s.top is not None
             and not is_chrome_shape(s, prs.slide_height)
             and s.top >= title.top]
    sub_top = int(max(s.top for s in below)) if below else None
    return alignment, int(title.left), int(title.top), sub_top


def _template_title_pt(prs):
    """The size the TEMPLATE sets its own titles at, or None.

    Synthesized slides used a flat 32pt whatever the deck. Measured with the
    same title picker the rest of the pipeline uses, the real templates sit at
    42 / 51 / 26 / 40 / 32pt — on T-Zh mono ours came out a third smaller than
    the designer's, on the study template a fifth larger. Mode, not mean: a
    designed size repeats across the deck (it equals the median on four of the
    five), while a one-off display figure does not."""
    sizes = []
    for slide in prs.slides:
        title = _pick_title_shape(slide, set())
        if title is None:
            continue
        pt = _max_font_pt(title)
        if pt:
            sizes.append(round(pt))
    if len(sizes) < 3:
        return None
    return Counter(sizes).most_common(1)[0][0]


def generate(template_path, plan, out_path, synth_canvas=None, canvas_backgrounds=None):
    """plan: ordered list of (content_block, template_slide_index) — or
    (content_block, SYNTHESIZE) when the matcher found no template slide for that
    archetype but it's one the generator can build from scratch — as produced by
    the matcher. A slide index appearing more than once means the matcher wants
    the same template slide reused for several blocks: every occurrence past the
    first is filled into a clone (see _clone_slide).

    synth_canvas: {plan_position: template_slide_idx} — for SYNTHESIZE items,
    clone that slide as the canvas (plan 9.3) so background art/chrome survive
    instead of drawing on a blank page. Positions absent from the map keep the
    from-scratch path."""
    prs = Presentation(template_path)
    theme = extract_theme(template_path)
    resolver = FontResolver(template_path, theme)
    bounds_in = None
    bounds_computed = False
    synth_theme = theme
    final_order = []

    # Clones must be cut from the still-pristine template BEFORE any filling
    # happens — cloning after the first occurrence is filled would duplicate
    # our own generated content instead of the template's original design.
    repeat_counts = {}
    for _, slide_idx in plan:
        if slide_idx != SYNTHESIZE:
            repeat_counts[slide_idx] = repeat_counts.get(slide_idx, 0) + 1
    spare_clones = {
        idx: [_clone_slide(prs, idx) for _ in range(n - 1)]
        for idx, n in repeat_counts.items() if n > 1
    }
    # Canvas clones likewise: cut from the pristine template up front, before
    # any fill mutates the source slides.
    canvas_clone_by_position = {}
    n_template_slides = len(prs.slides._sldIdLst)
    for position, (block, slide_idx) in enumerate(plan):
        canvas_idx = (synth_canvas or {}).get(position)
        if slide_idx == SYNTHESIZE and canvas_idx is not None and canvas_idx < n_template_slides:
            canvas_clone_by_position[position] = _clone_slide(prs, canvas_idx)

    filled = set()
    fill_groups = {}  # pristine source idx -> every slide filled from it
    for position, (block, slide_idx) in enumerate(plan):
        if slide_idx == SYNTHESIZE:
            if not bounds_computed:
                bounds_in = infer_content_bounds(extract_template(template_path))
                # The theme can lie about the deck's real look (hand-built decks
                # style runs directly and leave theme1.xml at Office defaults) —
                # synthesized slides follow what the slides actually show.
                synth_theme = apply_observed_style(theme, observe_deck_style(prs))
                # Typography scale of THIS template, so a synthesized slide sits
                # at the deck's own title size instead of a flat 32pt.
                synth_theme["title_pt"] = _template_title_pt(prs)
                synth_theme["cover_pt"] = _template_cover_pt(prs)
                synth_theme["content_gap_in"] = _template_content_gap_in(prs)
                (synth_theme["cover_align"], synth_theme["cover_left"],
                 synth_theme["cover_top"],
                 synth_theme["cover_sub_top"]) = _template_cover_layout(prs)
                bounds_computed = True
            # Measured background of the ORIGINAL canvas slide (from the
            # render), so the synthesizer can rescue text that would land
            # invisible on it. Keyed by the template index, not the clone.
            canvas_src = (synth_canvas or {}).get(position)
            new_idx = SYNTHESIZERS[block["type"]](
                prs, synth_theme, bounds_in, block, resolver=resolver,
                canvas_idx=canvas_clone_by_position.get(position),
                canvas_bg=(canvas_backgrounds or {}).get(canvas_src),
            )
            final_order.append(new_idx)
            continue

        source_idx = slide_idx
        if slide_idx in filled:
            slide_idx = spare_clones[slide_idx].pop(0)
        filled.add(slide_idx)
        fill_groups.setdefault(source_idx, []).append(slide_idx)

        slide = prs.slides[slide_idx]
        claimed_ids = set()
        FILLERS[block["type"]](slide, block, claimed_ids, resolver=resolver)
        _clear_unclaimed_text(slide, claimed_ids, prs.slide_height)
        _remove_orphan_marker_columns(slide)
        _remove_oversized_pictures(slide, prs.slide_width, prs.slide_height)
        final_order.append(slide_idx)

    # Final geometric QA over exactly the slides that ship (qa/geometry.py):
    # re-checks every text box against real metrics and shrinks what still
    # overflows — the safety net behind all the per-filler sizing above —
    # then evens out font sizes across clones of the same template slide.
    # Whatever they shrank gets its marker icons re-laid-out for the new size.
    keep_text_inside_slide(prs, resolver, slide_indices=final_order)
    keep_text_clear_of_side_art(prs, resolver, slide_indices=final_order)
    shrink_fixes = enforce_text_fits(prs, resolver, slide_indices=final_order)
    drop_needless_soft_hyphens(prs, resolver, slide_indices=final_order)
    unglue_overwide_pairs(prs, resolver, slide_indices=final_order)
    harmonize_changes = harmonize_clone_font_sizes(prs, fill_groups.values())
    harmonize_changes += _harmonize_slot_sizes(prs, final_order)
    _realign_icons_after_resize(prs, list(shrink_fixes) + harmonize_changes, resolver=resolver)

    _reorder_and_prune_slides(prs, final_order)
    _renumber_static_slide_numbers(prs)
    _fill_running_topic(prs, next((b.get("title") for b, _ in plan if b.get("title")), None),
                        resolver=resolver)
    prs.save(out_path)
    # Package-integrity gate (plan 9.6): three separate "PowerPoint wants to
    # repair this" incidents proved that rendering fine in LibreOffice is no
    # evidence of a valid package. Better to fail the generation loudly than
    # hand the user a file their PowerPoint will call corrupt.
    assert_valid_package(out_path)
    return out_path
