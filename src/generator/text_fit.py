"""Text-to-box fitting. Primary path: real glyph metrics from fonts/metrics.py
(embedded pptx fonts or a metric donor — verified within ~1% of an independent
PIL measurement). Fallback path, when no font could be resolved at all: the old
deliberately-generous character-count estimate.

Every public function keeps its original signature plus an optional `metrics`
(a fonts.metrics.FontMetrics or None), so callers that haven't been wired up
yet keep working unchanged."""

from pptx.oxml.ns import qn
from pptx.util import Emu, Length, Pt

AVG_CHAR_WIDTH_FACTOR = 0.62  # char-count fallback only — deliberately generous
LINE_SPACING = 1.3

# PowerPoint text boxes have default internal insets of 0.1" on each side;
# ignoring them makes every line look wider than it can actually be.
DEFAULT_INSETS_IN = 0.2

# Default top/bottom text-frame inset (0.05" each, in EMU).
DEFAULT_VERTICAL_INSET_EMU = 45720

# Default left/right text-frame inset (0.1" each, in EMU).
DEFAULT_HORIZONTAL_INSET_EMU = 91440


def vertical_insets_emu(shape):
    """Actual top/bottom text-frame insets of a shape (defaults where not
    set). Text must fit the box height MINUS these, not the raw height:
    renderers subtract them before deciding whether to autofit-shrink, so an
    estimate against the raw height accepts sizes that overflow by a hair —
    and LibreOffice then silently shrinks the text at render time, leaving
    every position computed from our font size (bullet icons) misaligned."""
    tf = shape.text_frame
    top = tf.margin_top if tf.margin_top is not None else DEFAULT_VERTICAL_INSET_EMU
    bottom = tf.margin_bottom if tf.margin_bottom is not None else DEFAULT_VERTICAL_INSET_EMU
    return top, bottom


def frame_horizontal_margins_in(shape):
    """Actual left+right text-frame insets of a shape, in inches (PowerPoint
    defaults of 0.1" each where not set)."""
    tf = shape.text_frame
    left = tf.margin_left if tf.margin_left is not None else DEFAULT_HORIZONTAL_INSET_EMU
    right = tf.margin_right if tf.margin_right is not None else DEFAULT_HORIZONTAL_INSET_EMU
    return Emu(left).inches + Emu(right).inches


def paragraph_indent_in(paragraph):
    """Paragraph-level left offset (a:pPr marL/indent), in inches. Icon-bullet
    designs routinely reserve a marL of ~0.25" for the marker column; a width
    budget that ignores it overestimates every line by that much. A negative
    indent is a hanging first line whose text still starts at marL (the bullet
    glyph occupies the overhang), so it doesn't widen the budget."""
    pPr = paragraph._p.find(qn("a:pPr"))
    if pPr is None:
        return 0.0
    marL = int(pPr.get("marL") or 0)
    indent = int(pPr.get("indent") or 0)
    return Emu(max(0, marL + max(indent, 0))).inches


def horizontal_margins_in(shape):
    """Total horizontal space the renderer subtracts from this shape's width
    before wrapping text: real frame insets plus the widest paragraph indent.
    The flat DEFAULT_INSETS_IN guess this replaces was measured to erase most
    of the single-line safety margin on real slides (lIns=0 + rIns=0.1" +
    marL=0.25" ≠ 0.2")."""
    if shape is None or not getattr(shape, "has_text_frame", False):
        return DEFAULT_INSETS_IN
    indents = [
        paragraph_indent_in(p)
        for p in shape.text_frame.paragraphs
        if any(r.text.strip() for r in p.runs)
    ]
    return frame_horizontal_margins_in(shape) + (max(indents) if indents else 0.0)


def _usable_width_in(width_in, margins_in=None):
    margins = DEFAULT_INSETS_IN if margins_in is None else margins_in
    if width_in > margins * 2:
        return width_in - margins
    return width_in


def _wrapped_lines_metric(text, width_in, font_size_pt, metrics, margins_in=None):
    budget_pt = _usable_width_in(width_in, margins_in) * 72
    if budget_pt <= 0:
        return 1
    total = 0
    space_pt = metrics.text_width_pt(" ", font_size_pt)
    for raw_line in text.split("\n"):
        words = raw_line.split()
        if not words:
            total += 1
            continue
        lines = 1
        current = 0.0
        for word in words:
            word_pt = metrics.text_width_pt(word, font_size_pt)
            if word_pt > budget_pt:
                # A single word wider than the box wraps mid-word when rendered.
                extra_full = int(word_pt // budget_pt)
                lines += extra_full
                current = word_pt - extra_full * budget_pt
                continue
            step = word_pt if current == 0 else space_pt + word_pt
            if current + step <= budget_pt:
                current += step
            else:
                lines += 1
                current = word_pt
        total += lines
    return max(1, total)


def _wrapped_lines_charcount(text, width_in, font_size_pt):
    avg_char_width_in = (font_size_pt * AVG_CHAR_WIDTH_FACTOR) / 72
    chars_per_line = max(1, int(width_in / avg_char_width_in))
    lines = 0
    for line in text.split("\n"):
        lines += max(1, -(-max(len(line), 1) // chars_per_line))  # ceil division
    return max(1, lines)


def estimate_wrapped_lines(text, width_in, font_size_pt, metrics=None, margins_in=None):
    if not text or width_in <= 0:
        return 1
    if metrics is not None:
        return _wrapped_lines_metric(text, width_in, font_size_pt, metrics, margins_in=margins_in)
    return _wrapped_lines_charcount(text, width_in, font_size_pt)


def line_height_pt(font_size_pt, metrics=None):
    if metrics is not None:
        return metrics.line_height_pt(font_size_pt)
    return font_size_pt * LINE_SPACING


# Percent line spacing (a:spcPct) is applied by both PowerPoint and LibreOffice
# against the classic "single line = 1.2 em" convention, NOT against the font's
# own metric line height. Measured on real renders (deck 8505a62c, Arial-
# substituted preview): 18pt at 150% renders at exactly 18 × 1.5 × 1.2 = 32.4pt
# per line, while the font-metric model predicted 31.05pt — and on denser slides
# that underestimate let text through that LibreOffice then silently shrank via
# normAutofit, which is what made bullet icons drift off their lines.
PCT_SPACING_BASE = 1.2


def paragraph_pitch_pt(font_size_pt, metrics=None, line_spacing=None):
    """Vertical distance between successive line baselines for a paragraph.
    line_spacing follows python-pptx conventions: a float is a spcPct multiple,
    a Length is an absolute spcPts value, None means "not set" (renderers use
    the font's natural line height)."""
    if line_spacing is not None:
        # Length first: pptx.util.Length subclasses int, so the numeric
        # (percent-multiple) check would silently swallow absolute values.
        if isinstance(line_spacing, Length):
            return line_spacing.pt
        if isinstance(line_spacing, (int, float)):
            return font_size_pt * PCT_SPACING_BASE * line_spacing
    return line_height_pt(font_size_pt, metrics=metrics)


def estimate_block_height_in(lines, width_in, font_size_pt, metrics=None, line_spacing=None, margins_in=None):
    """lines: a single string (its own "\\n"-separated wraps count) or a list of
    separate paragraph strings — either way, returns total estimated height."""
    if isinstance(lines, str):
        lines = [lines]
    total_lines = sum(
        estimate_wrapped_lines(line, width_in, font_size_pt, metrics=metrics, margins_in=margins_in)
        for line in lines
    )
    return total_lines * paragraph_pitch_pt(font_size_pt, metrics=metrics, line_spacing=line_spacing) / 72


# Rendered text comes out a few percent wider than fontTools advance sums
# (hinting/grid-fitting; measured: a 418.6pt-estimated line wrapped inside a
# 435pt box). Single-line guarantees must leave that margin or they're not
# guarantees at all.
SINGLE_LINE_SAFETY = 0.95


def fit_font_size_single_line(texts, width_emu, base_size_pt, min_size_pt=10, metrics=None, margins_in=None):
    """Largest whole-point size <= round(base_size_pt) at which EVERY text fits
    its box width on one line with SINGLE_LINE_SAFETY margin — or None when
    even min_size_pt can't deliver that (caller falls back to wrapped layout).

    Exists for icon-marker lists: an icon column is positioned per line, and
    predicting exactly where a renderer will wrap is fragile at the few-percent
    level — so instead of predicting the wrap, make the wrap impossible."""
    if metrics is None or not width_emu:
        return None
    budget_pt = _usable_width_in(Emu(width_emu).inches, margins_in) * 72 * SINGLE_LINE_SAFETY
    if budget_pt <= 0:
        return None
    size = max(min_size_pt, round(base_size_pt))
    while size >= min_size_pt:
        if all(metrics.text_width_pt(t, size) <= budget_pt for t in texts):
            return Pt(size)
        size -= 1
    return None


def fit_font_size(lines, width_emu, height_emu, base_size_pt, min_size_pt=8, metrics=None, line_spacing=None, margins_in=None):
    """Largest whole-point size <= round(base_size_pt) that's estimated to fit
    `lines` within a box of width_emu x height_emu. Never grows past
    base_size_pt — only shrinks when the box genuinely doesn't have room,
    sized to the box's real dimensions rather than a ratio against old text."""
    base = max(min_size_pt, round(base_size_pt))
    width_in = Emu(width_emu).inches if width_emu else 0
    height_in = Emu(height_emu).inches if height_emu else 0
    if width_in <= 0 or height_in <= 0:
        return Pt(base)

    size = base
    while size > min_size_pt:
        if estimate_block_height_in(lines, width_in, size, metrics=metrics, line_spacing=line_spacing, margins_in=margins_in) <= height_in:
            break
        size -= 1
    return Pt(size)


# Metrics resolve by family name only — a BOLD run renders ~5% wider than the
# regular-weight advances we measured (caught live: a 47pt bold title passed
# the word-fit by 0.4pt and still broke mid-word).
BOLD_WIDTH_FACTOR = 1.05


def cap_size_to_longest_word(lines, width_emu, size_pt, min_size_pt=9, metrics=None,
                             margins_in=None, bold=False):
    """Shrinks size_pt until the LONGEST WORD fits the usable box width on one
    line — height-only fitting happily accepts sizes at which the renderer has
    no choice but to break a word mid-letter («Автоматизаци/я» on a real 51pt
    title, plan 10а). Without metrics an average-width estimate still catches
    the gross cases. Stops at min_size_pt: below that the text itself is the
    problem (title budgets, plan 10б), not the font."""
    if not width_emu:
        return size_pt
    text = " ".join(lines) if isinstance(lines, (list, tuple)) else str(lines)
    words = [w for w in text.split() if w]
    if not words:
        return size_pt
    budget_in = _usable_width_in(Emu(width_emu).inches, margins_in=margins_in)
    if budget_in <= 0:
        return size_pt
    # The renderer wraps 2-4% earlier than fontTools advances predict (see
    # SINGLE_LINE_SAFETY) — a word that "just fits" by metrics still breaks.
    budget_pt = budget_in * 72 * SINGLE_LINE_SAFETY

    weight = BOLD_WIDTH_FACTOR if bold else 1.0

    def longest_word_pt(size):
        if metrics is not None:
            return max(metrics.text_width_pt(w, size) for w in words) * weight
        return max(len(w) for w in words) * size * AVG_CHAR_WIDTH_FACTOR * weight

    size = round(size_pt)
    while size > min_size_pt and longest_word_pt(size) > budget_pt:
        size -= 1
    return size


_VOWELS = set("аеёиоуыэюяАЕЁИОУЫЭЮЯaeiouyAEIOUY")
SOFT_HYPHEN = "­"


def soft_hyphenate_long_words(text, max_word_len=12):
    """Inserts soft hyphens (U+00AD) into words longer than max_word_len at
    vowel boundaries — the escape hatch for words that can't fit the box width
    even at the font floor («Неконсистентность» in a narrow title box).
    Renderers break at SHY with a visible hyphen instead of chopping the word
    mid-letter; short words are left untouched, and the character is invisible
    when no break is needed."""
    out_words = []
    for word in text.split(" "):
        if len(word) <= max_word_len:
            out_words.append(word)
            continue
        pieces = []
        current = ""
        for i, ch in enumerate(word):
            current += ch
            # break after a vowel, at least 4 chars in, with >=3 chars left
            if (len(current) >= 4 and ch in _VOWELS
                    and len(word) - i - 1 >= 3
                    and i + 1 < len(word) and word[i + 1] not in _VOWELS):
                pieces.append(current)
                current = ""
        pieces.append(current)
        out_words.append(SOFT_HYPHEN.join(p for p in pieces if p))
    return " ".join(out_words)
