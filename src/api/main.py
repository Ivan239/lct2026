import hashlib
import json
import os
import shutil
import sys
import tempfile
import uuid

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from dotenv import load_dotenv

load_dotenv()

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from pydantic import BaseModel

from common.model_fallback import TEXT_MODELS
from common.pictures import has_data_object, has_oversized_picture
from common.synthesis import SYNTHESIZE
from content_package import ContentPackageError, extract_numbers, load_package, to_brief_text
from content_parser.parser import RESIZE_FIELD_BY_TYPE, parse_brief, resize_block
from content_parser.two_phase import generate_block, generate_outline, stat_fingerprints
from design_system.extractor import build_archetype_map
from design_system.fingerprint_cache import FingerprintCache
from design_system.extractor import _describe_slide
from design_system.style_card import build_style_card, card_path, card_prompt_preamble, load_card, save_card
from design_system.style_profile import build_measured_profile, rotation_targets
from generator.slide_kit import canvas_content_count, content_text_shapes
from generator.generator import generate, get_capacity
from llm_clients.gigachat import GigaChatClient
from matcher.matcher import match_content_to_slides, plan_from_outline
from evaluation.deterministic import unsourced_numbers
from qa.geometry import find_sparse_slides
from rendering.render import render_pptx_to_pngs
from rendering.export import export_all
from template_parser.parser import extract_template
from template_spec.builder import build_spec

# Chat-capable tiers from list_models() — excludes embeddings and the
# "-preview" duplicates, which aren't meaningfully different choices for a user.
# Must match what api.giga.chat/v1/models actually serves — the old
# gigachat.devices.sberbank.ru line (GigaChat, GigaChat-Plus, GigaChat-Pro,
# GigaChat-Max) 404s on this host and would silently break model selection.
AVAILABLE_MODELS = ["GigaChat-2", "GigaChat-2-Pro", "GigaChat-2-Max", "GigaChat-3-Ultra"]

BASE = os.path.join(os.path.dirname(__file__), "..", "..")
OUTPUT_DIR = os.path.abspath(os.path.join(BASE, "output"))
TEMPLATES_DIR = os.path.join(OUTPUT_DIR, "templates")
RENDERED_DIR = os.path.join(OUTPUT_DIR, "rendered")
PARSED_DIR = os.path.join(OUTPUT_DIR, "parsed")
GENERATED_DIR = os.path.join(OUTPUT_DIR, "generated")
for d in (TEMPLATES_DIR, RENDERED_DIR, PARSED_DIR, GENERATED_DIR):
    os.makedirs(d, exist_ok=True)

# One shared library across every template ever uploaded: fingerprints are
# geometry-only (no client text), so reuse across customers is safe by design.
FINGERPRINT_CACHE_PATH = os.path.join(PARSED_DIR, "fingerprint_cache.json")

PRESET_NAMES = {
    "template_a_corporate": "Корпоративный (навы + оранжевый)",
    "template_b_startup": "Стартап (светлый + фиолетовый)",
}

def _make_client():
    """The service runs on the open-weights endpoint the brief requires
    (LLM_BASE_URL / LLM_MODEL). GigaChat remains only as a development fallback
    when no such endpoint is configured — see docs/MODELS.md."""
    from llm_clients.backends import open_weights_client

    configured = open_weights_client()
    if configured is not None:
        client_, model_ = configured
        print(f"[api] генератор: {model_} через {os.environ.get('LLM_BASE_URL')}", flush=True)
        return client_, model_
    if os.environ.get("GIGACHAT_CLIENT_ID") and os.environ.get("GIGACHAT_CLIENT_SECRET"):
        print("[api] LLM_BASE_URL не задан — используется запасной провайдер разработки",
              flush=True)
        return GigaChatClient(verify_ssl=False), None
    # Без модели сервис всё равно стартует (список шаблонов, превью, скачивания
    # работают), а операции, которым нужна модель, отвечают понятной 503.
    # Раньше импорт падал с KeyError: 'GIGACHAT_CLIENT_ID' — на свежем клоне
    # без .env не поднимался ни сервис, ни тесты API.
    print(f"[api] {MODEL_NOT_CONFIGURED}", flush=True)
    return None, None


MODEL_NOT_CONFIGURED = ("Модель не настроена: задайте LLM_BASE_URL, LLM_MODEL и LLM_API_KEY "
                        "в .env (образец — .env.example)")


def _require_model():
    if client is None:
        raise HTTPException(status_code=503, detail=MODEL_NOT_CONFIGURED)


client, DEFAULT_MODEL = _make_client()
# Every call that does not name a model uses the configured open-weights one;
# without it we fall back to the development chain.
DEFAULT_MODELS = [DEFAULT_MODEL] if DEFAULT_MODEL else TEXT_MODELS
app = FastAPI(title="SlideGen API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/static", StaticFiles(directory=OUTPUT_DIR), name="static")

_template_registry = {}

# Live progress of the generation currently running. One global slot: the demo
# runs one generation at a time, and a stale read merely shows a slightly-off
# stage label — not worth a job-queue architecture at this stage.
_progress = {"active": False, "stage": "", "done": 0, "total": 0}


def _set_progress(stage, done=0, total=0, active=True):
    _progress.update({"active": active, "stage": stage, "done": done, "total": total})


def _slide_urls(template_id):
    pages = sorted(
        f for f in os.listdir(RENDERED_DIR) if f.startswith(f"{template_id}-") and f.endswith(".png")
    )
    return [f"/static/rendered/{f}" for f in pages]


def _archetypes_path(template_id):
    return os.path.join(PARSED_DIR, f"{template_id}_archetypes.json")


def _meta_path(template_id):
    return os.path.join(PARSED_DIR, f"{template_id}_meta.json")


def _load_archetypes(template_id):
    with open(_archetypes_path(template_id), encoding="utf-8") as f:
        return {int(k): v for k, v in json.load(f).items()}


def _build_spec(template_id):
    """The slot spec derives entirely from data we already have (cached
    archetype map + slide geometry) — building it is free, no LLM involved, so
    it's never persisted to disk: a persisted copy would silently go stale
    every time get_capacity's own logic changes (it did, more than once),
    with nothing to invalidate it. Rebuilding fresh on every call costs a
    local pptx traversal, not a request to GigaChat."""
    template_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")
    archetypes = _load_archetypes(template_id)
    # Measured style profile (plan 9.1): palette/rotation from the cached
    # renders — same no-disk-cache policy, milliseconds to rebuild.
    png_paths = [os.path.join(RENDERED_DIR, f) for f in sorted(
        f for f in os.listdir(RENDERED_DIR)
        if f.startswith(f"{template_id}-") and f.endswith(".png")
    )]
    profile = build_measured_profile(png_paths, archetypes) if png_paths else None
    return build_spec(template_path, archetypes, style_profile=profile)


def _current_balance():
    """Best-effort — a balance-check failure must never break the actual
    response it's attached to, so callers get None rather than an exception."""
    try:
        raw = client.get_balance()
        return [{"model": item["usage"], "tokens": item["value"]} for item in raw.get("balance", [])]
    except Exception:
        return None


def _hash_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()[:16]


def _process_template(template_id, pptx_path, model=None):
    """Structure-first: parse OOXML (free), cluster slides by structural
    similarity, and only spend LLM tokens on one medoid slide per cluster —
    unless its geometry fingerprint was already classified on an earlier
    template, in which case that cluster costs nothing (see
    design_system.extractor / clustering / fingerprint_cache for why).
    model, if given, replaces the text-classification model choice (a user pick
    takes priority over the default cost-tiered fallback chain). Vision fallback
    always uses its own chain regardless — a user-picked text model might not
    even support images."""
    _set_progress("Читаем структуру шаблона")
    template_struct = extract_template(pptx_path)
    _set_progress("Рендерим слайды шаблона")
    png_paths = render_pptx_to_pngs(pptx_path, RENDERED_DIR)
    text_models = [model] if model else DEFAULT_MODELS
    archetype_map = build_archetype_map(
        client, template_struct, rendered_png_paths=png_paths, text_models=text_models,
        fingerprint_cache=FingerprintCache(FINGERPRINT_CACHE_PATH),
        progress_cb=lambda done, total: _set_progress(
            "Определяем архетипы слайдов", done=done, total=total
        ),
    )

    # A slide built around a big picture (a real chart/photo from the source deck)
    # can't be reused for our text-only archetypes: keeping the picture shows
    # stale, unrelated old data; removing it leaves a gaping visual hole where it
    # used to be. Either way it's not fixable by touching text alone, so exclude
    # these from matching regardless of what the text classifier called them —
    # free, no LLM needed, and more reliable than hoping the classifier weighs a
    # bare "PICTURE at (x,y) size WxH" line correctly against actual page area.
    prs = Presentation(pptx_path)
    for idx in list(archetype_map.keys()):
        if has_oversized_picture(prs.slides[idx], prs.slide_width, prs.slide_height):
            archetype_map[idx] = "other"

    with open(_archetypes_path(template_id), "w", encoding="utf-8") as f:
        json.dump(archetype_map, f, ensure_ascii=False, indent=2)

    # Style card (plan 9.4): ONE extra LLM call, cached forever for this
    # template. Best-effort — a failed card just means generation prompts run
    # without the design brief, exactly as before the feature existed.
    if load_card(PARSED_DIR, template_id) is None:
        try:
            _set_progress("Составляем дизайн-бриф шаблона")
            descriptions = [_describe_slide(s) for s in template_struct["slides"]]
            card = build_style_card(
                client, descriptions, profile=None, archetype_map=archetype_map,
                models=text_models,
            )
            save_card(PARSED_DIR, template_id, card)
        except Exception as e:
            print(f"[style_card] {template_id}: {e}", flush=True)

    return archetype_map


def _save_meta(template_id, name, is_preset):
    with open(_meta_path(template_id), "w", encoding="utf-8") as f:
        json.dump({"id": template_id, "name": name, "is_preset": is_preset}, f, ensure_ascii=False)


def _load_meta(template_id):
    path = _meta_path(template_id)
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _bootstrap_presets():
    for template_id, name in PRESET_NAMES.items():
        pptx_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")
        if not os.path.exists(pptx_path):
            continue
        if not os.path.exists(_archetypes_path(template_id)) or not _slide_urls(template_id):
            try:
                _process_template(template_id, pptx_path)
            except Exception as e:
                # Bootstrap runs at import time and may reach GigaChat: an
                # unreachable Sber API (routine on this tunnel) must degrade to
                # "this preset is missing for now", never to "the whole server
                # refuses to start".
                print(f"[bootstrap] preset {template_id} skipped: {e}", flush=True)
                continue
        _save_meta(template_id, name, is_preset=True)
        _template_registry[template_id] = {"id": template_id, "name": name, "is_preset": True}


def _bootstrap_customs():
    """Rehydrate previously-uploaded templates after a server restart, so we never
    re-spend tokens re-classifying a file we've already processed."""
    for fname in os.listdir(TEMPLATES_DIR):
        if not fname.startswith("custom_") or not fname.endswith(".pptx"):
            continue
        template_id = fname[: -len(".pptx")]
        if not os.path.exists(_archetypes_path(template_id)):
            continue
        meta = _load_meta(template_id) or {"name": template_id, "is_preset": False}
        _template_registry[template_id] = {"id": template_id, "name": meta["name"], "is_preset": False}


_bootstrap_presets()
_bootstrap_customs()


class GenerateRequest(BaseModel):
    template_id: str
    brief: str
    model: str | None = None
    # Deck size the user asks for; None means the brief's 10-15 slides.
    slides: int | None = None


@app.get("/api/models")
def list_available_models():
    return {"models": [DEFAULT_MODEL] if DEFAULT_MODEL else AVAILABLE_MODELS}


@app.get("/api/balance")
def get_balance():
    balance = _current_balance()
    if balance is None:
        raise HTTPException(status_code=502, detail="Не удалось получить баланс GigaChat")
    return {"balance": balance}


@app.get("/api/templates")
def list_templates():
    result = []
    for template_id, meta in _template_registry.items():
        result.append({
            "id": template_id,
            "name": meta["name"],
            "is_preset": meta.get("is_preset", False),
            "slides": _slide_urls(template_id),
            "archetypes": _load_archetypes(template_id),
        })
    return result


@app.delete("/api/templates/{template_id}")
def delete_template(template_id: str):
    """Bundled presets ship with the app and reappear on next boot regardless
    (see _bootstrap_presets) — deleting their files would just be undone at
    the next restart, so refuse rather than pretend it worked. Fingerprint
    cache is deliberately left untouched: it's geometry-only and shared
    across every template ever uploaded, not this one's to own."""
    meta = _template_registry.get(template_id)
    if meta is None:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    if meta.get("is_preset"):
        raise HTTPException(status_code=403, detail="Встроенный пресет удалить нельзя")

    del _template_registry[template_id]

    paths = [
        os.path.join(TEMPLATES_DIR, f"{template_id}.pptx"),
        _archetypes_path(template_id),
        _meta_path(template_id),
        card_path(PARSED_DIR, template_id),
    ]
    for f in os.listdir(RENDERED_DIR):
        if f.startswith(f"{template_id}-") or f == f"{template_id}.pdf":
            paths.append(os.path.join(RENDERED_DIR, f))
    for path in paths:
        try:
            os.remove(path)
        except FileNotFoundError:
            pass
    return {"deleted": template_id}


@app.post("/api/templates")
def upload_template(file: UploadFile = File(...), model: str | None = Form(None)):
    _require_model()
    # Deliberately sync (runs in FastAPI's threadpool): template processing is
    # minutes of blocking CPU/LLM work, and as an `async def` it froze the
    # whole event loop — /api/progress (and every other request) couldn't get
    # a response until the upload finished, which defeated live progress
    # entirely and would have frozen every other user of the demo too.
    name = os.path.splitext(file.filename)[0]

    tmp_path = os.path.join(TEMPLATES_DIR, f"_upload_{uuid.uuid4().hex[:8]}.pptx")
    with open(tmp_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    template_id = f"custom_{_hash_file(tmp_path)}"
    pptx_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")

    already_processed = os.path.exists(_archetypes_path(template_id))
    if already_processed:
        os.remove(tmp_path)  # identical file was uploaded before — reuse the cached result, spend no tokens
    else:
        os.replace(tmp_path, pptx_path)
        try:
            archetype_map = _process_template(template_id, pptx_path, model=model)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Не удалось разобрать шаблон: {e}")
        finally:
            _set_progress("", active=False)

    _save_meta(template_id, name, is_preset=False)
    _template_registry[template_id] = {"id": template_id, "name": name, "is_preset": False}
    return {
        "id": template_id,
        "name": name,
        "slides": _slide_urls(template_id),
        "archetypes": _load_archetypes(template_id),
        "balance": _current_balance(),
    }


def _measured_backgrounds(template_id):
    """{template_slide_idx: (r,g,b)} measured from the renders — what the slide
    ACTUALLY looks like. Synthesis needs this to keep text readable on a cloned
    canvas; the XML can't answer it (the T-Zh mono template declares the same
    schemeClr on every slide while four of them render solid blue)."""
    try:
        png_paths = [os.path.join(RENDERED_DIR, f) for f in sorted(
            f for f in os.listdir(RENDERED_DIR)
            if f.startswith(f"{template_id}-") and f.endswith(".png")
        )]
        if not png_paths:
            return {}
        return build_measured_profile(png_paths, _load_archetypes(template_id))["backgrounds"]
    except Exception as e:  # noqa: BLE001 — advisory: without it text just keeps the old colour
        print(f"[measured_bg] {template_id}: {e}", flush=True)
        return {}


def _synth_canvas_hints(template_id, plan):
    """{plan_position: template_slide_idx} for SYNTHESIZE positions (plan 9.3):
    the canvas is the template slide with the FEWEST content boxes (covers/
    dividers — maximum free area, minimum stripping), preferring the color the
    deck's discovered rotation wants at that position. Best-effort: any failure
    means synthesized slides just fall back to the from-scratch look."""
    try:
        png_paths = [os.path.join(RENDERED_DIR, f) for f in sorted(
            f for f in os.listdir(RENDERED_DIR)
            if f.startswith(f"{template_id}-") and f.endswith(".png")
        )]
        if not png_paths:
            return {}
        archetypes = _load_archetypes(template_id)
        profile = build_measured_profile(png_paths, archetypes)
        targets = rotation_targets(profile, len(plan))

        template_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")
        prs = Presentation(template_path)
        candidates = []  # (n_content_boxes, slide_idx, bg)
        photo_limit = Inches(1.2)
        for idx in range(len(prs.slides)):
            if idx in profile["breathers"]:
                continue
            slide = prs.slides[idx]
            # A canvas must be furniture-only: slides carrying real photos keep
            # them after content-box stripping, and synthesized text ends up
            # fighting a photo collage for the same area (seen on first 9.3
            # sweep — two_column landed on a 3-photo slide).
            has_photo = any(
                s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width and s.width > photo_limit
                and not (s.width >= prs.slide_width * 0.9)  # full-bleed folder art is background, keep
                for s in slide.shapes
            )
            # Нативная таблица или диаграмма тоже не мебель: снятие текстбоксов
            # её не трогает, и синтезированный слайд выходил поверх таблицы
            # шаблона с «Заголовок / Текст» в ячейках (VK WorkSpace, колоды
            # сдачи). Подбор слайдов исключал их давно (has_data_object в
            # build_spec) — выбор канвы шёл мимо этого правила.
            if has_photo or has_data_object(slide):
                continue
            n_boxes = canvas_content_count(slide)
            candidates.append((n_boxes, idx, profile["backgrounds"].get(idx)))
        if not candidates:
            return {}
        # Sparse slides (covers/dividers/closings, <=3 content boxes) all make
        # good canvases — keeping the whole tier lets the color preference
        # actually choose, instead of a strict "fewest boxes" collapsing the
        # pool to a single slide (which pinned every synthesized slide to the
        # same yellow canvas on the first 9.3 sweep).
        eligible = [c for c in candidates if c[0] <= 3] or sorted(candidates)[:3]

        hints = {}
        for position, (block, slide_idx) in enumerate(plan):
            if slide_idx != SYNTHESIZE:
                continue
            target = targets[position] if targets else None
            best = min(
                eligible,
                key=lambda c: (0 if target is not None and c[2] == target else 1, c[0], c[1]),
            )
            hints[position] = best[1]
        return hints
    except Exception as e:
        print(f"[synth_canvas] hints unavailable: {e}", flush=True)
        return {}


def _plan_two_phase(template_id, brief, model=None, slides=None):
    """Slot-first flow (docs/IMPROVEMENT_PLAN.md item 4): outline against the
    template's actual offering, pick concrete slides, then generate each block's
    text sized to the chosen slide's real capacity. Returns (plan, skipped), or
    (None, None) to signal the caller to fall back to the legacy flow — the new
    path must never make the product less available than the old one was."""
    models = [model] if model else DEFAULT_MODELS
    try:
        spec = _build_spec(template_id)
        # Template design brief (plan 9.4) — cached at upload, advisory only.
        style_preamble = card_prompt_preamble(load_card(PARSED_DIR, template_id))
        _set_progress("Планируем структуру презентации")
        outline = generate_outline(client, brief, spec, models=models, style_preamble=style_preamble,
                                   slides=slides)
        assignments, skipped_items = plan_from_outline(outline, spec)

        plan = []
        # Carry the stat numbers/labels already placed: blocks are separate calls
        # and otherwise repeat the same KPIs on two slides (seen on real decks).
        used_nums, used_labels = set(), set()
        failed = []
        for i, (item, slide_idx, final_count) in enumerate(assignments):
            _set_progress(f"Пишем контент: {item['theme']}", done=i, total=len(assignments))
            try:
                block = generate_block(client, item["role"], item["theme"], brief, count=final_count,
                                       models=models, style_preamble=style_preamble,
                                       used_stats=(used_nums, used_labels))
            except (ValueError, KeyError):
                # One block that stays malformed after every retry costs that
                # slide, not the two-phase deck (the loop lost three runs to one
                # stats block). It is reported in `skipped` for the user to see.
                failed.append({"type": item["role"], "title": item.get("theme")})
                continue
            nums, labels = stat_fingerprints(block)
            used_nums |= nums
            used_labels |= labels
            plan.append((block, slide_idx))

        skipped = [{"type": item["role"], "title": item.get("theme")} for item in skipped_items] + failed
        return plan, skipped
    except Exception:
        # The legacy fallback keeps the product available, but a silent switch
        # made every "почему слайды другие?" report undiagnosable — the reason
        # belongs in the server log even though the user flow continues.
        import traceback
        print(f"[two_phase] falling back to legacy flow for {template_id}:", flush=True)
        traceback.print_exc()
        return None, None


def _plan_legacy(brief, archetype_map, template_path, model=None):
    """Pre-slot-spec flow: single-shot brief parsing, then post-hoc capacity
    resize. Kept as the fallback when any stage of the two-phase path fails."""
    models = [model] if model else DEFAULT_MODELS
    try:
        content_blocks = parse_brief(client, brief, models=models)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Не удалось разобрать бриф через GigaChat: {e}")

    plan, skipped_blocks = match_content_to_slides(content_blocks, archetype_map)

    try:
        prs_for_capacity = Presentation(template_path)
        adjusted_plan = []
        for block, slide_idx in plan:
            if slide_idx != SYNTHESIZE and block["type"] in RESIZE_FIELD_BY_TYPE:
                capacity = get_capacity(prs_for_capacity.slides[slide_idx], block["type"])
                field = RESIZE_FIELD_BY_TYPE[block["type"]]
                if capacity is not None and capacity != len(block.get(field, [])):
                    block = resize_block(client, block, capacity, models=models)
            adjusted_plan.append((block, slide_idx))
        plan = adjusted_plan
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Не удалось подогнать контент под ёмкость слайда: {e}")

    skipped = [{"type": b["type"], "title": b.get("title")} for b in skipped_blocks]
    return plan, skipped


@app.get("/api/progress")
def get_progress():
    return _progress


@app.post("/api/generate")
def generate_presentation(req: GenerateRequest):
    _require_model()
    return _generate_deck(req.template_id, req.brief, req.model, slides=req.slides)


@app.post("/api/generate/package")
def generate_from_package(template_id: str = Form(...), file: UploadFile = File(...),
                          model: str | None = Form(None)):
    """Generation from a content package (docs/CONTENT_PACKAGE.md): a .zip of
    brief.md + optional package.json, facts.md, data/*.csv, images/*.

    The package is decomposed; its parts are flattened into one labelled brief
    for the text pipeline (content_parser.two_phase takes a string), and its
    numbers become the reference for the check «все цифры со слайдов есть в
    исходных материалах» in the response's warnings. A broken package is a 422
    that names the problem, not a 500 three layers down."""
    _require_model()
    if template_id not in _template_registry:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    workdir = tempfile.mkdtemp(prefix="package_upload_")
    try:
        archive = os.path.join(workdir, os.path.basename(file.filename or "package.zip"))
        with open(archive, "wb") as out:
            shutil.copyfileobj(file.file, out)
        try:
            package = load_package(archive, extract_to=workdir)
        except ContentPackageError as e:
            raise HTTPException(status_code=422, detail=f"Контент-пакет: {e}")
        result = _generate_deck(template_id, to_brief_text(package), model,
                                source_numbers=package["numbers"], slides=package["slides"])
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    result["package"] = {
        "title": package["title"],
        "purpose": package["purpose"],
        "facts": len(package["facts"]),
        "numbers": len(package["numbers"]),
        "tables": [t["name"] for t in package["tables"]],
        "images": len(package["images"]),
    }
    return result


def _numbers_warnings(pptx_path, source_numbers):
    """One warning per slide that carries numbers the source does not — the
    Appendix 1 question «все цифры и факты со слайда есть в исходных
    материалах?», asked of every deck the service returns."""
    by_slide = {}
    for number, token in unsourced_numbers(pptx_path, source_numbers):
        by_slide.setdefault(number, []).append(token)
    return [{"slide": number, "kind": "numbers_not_in_source",
             "details": "цифры, которых нет в исходных материалах: " + ", ".join(tokens)}
            for number, tokens in sorted(by_slide.items())]


def _generate_deck(template_id, brief, model=None, source_numbers=None, slides=None):
    """Plan → .pptx → previews → warnings. `source_numbers` — the content
    package's numbers; None means the brief is the whole source."""
    if template_id not in _template_registry:
        raise HTTPException(status_code=404, detail="Шаблон не найден")

    archetype_map = _load_archetypes(template_id)
    template_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")

    try:
        plan, skipped = _plan_two_phase(template_id, brief, model=model, slides=slides)
        if plan is None:
            _set_progress("Разбираем бриф (запасной сценарий)")
            plan, skipped = _plan_legacy(brief, archetype_map, template_path, model=model)

        if not plan:
            raise HTTPException(status_code=422, detail="Ни один блок контента не подошёл ни к одному слайду шаблона")

        return _assemble_deck(template_id, plan, skipped, brief, model, source_numbers)
    finally:
        # Whatever path we exit through — success, quota error, fallback crash —
        # the progress slot must not stay stuck on a stale "active" stage.
        _set_progress("", active=False)


def _session_path(generation_id):
    return os.path.join(GENERATED_DIR, f"{generation_id}.session.json")


def _assemble_deck(template_id, plan, skipped, brief, model, source_numbers):
    """Готовый план → .pptx, превью, .pdf/.html, находки аудита. Общий путь для
    генерации и для «исправить выбранные»: исправленная колода обязана пройти
    те же проверки, что и первая."""
    archetype_map = _load_archetypes(template_id)
    template_path = os.path.join(TEMPLATES_DIR, f"{template_id}.pptx")
    generation_id = uuid.uuid4().hex[:8]
    out_pptx = os.path.join(GENERATED_DIR, f"{generation_id}.pptx")

    try:
        _set_progress("Собираем .pptx в стиле шаблона")
        generate(template_path, plan, out_pptx,
                 synth_canvas=_synth_canvas_hints(template_id, plan),
                 canvas_backgrounds=_measured_backgrounds(template_id))
        _set_progress("Рендерим превью слайдов")
        slide_png_paths = render_pptx_to_pngs(out_pptx, GENERATED_DIR)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Не удалось собрать презентацию: {e}")

    slide_urls = [f"/static/generated/{os.path.basename(p)}" for p in slide_png_paths]

    # .pdf и .html (ТЗ: три формата) — из того же рендера, что и превью:
    # LibreOffice повторно не запускается. Сбой экспорта не отменяет .pptx.
    exports = {}
    try:
        exported = export_all(out_pptx, GENERATED_DIR, png_paths=slide_png_paths)
        exports = {"pdf_url": f"/static/generated/{generation_id}.pdf",
                   "html_url": f"/static/generated/{generation_id}.html",
                   "substituted_fonts": exported["substituted_fonts"],
                   "fallback_font": exported["fallback_font"]}
    except Exception as e:  # noqa: BLE001
        print(f"! экспорт .pdf/.html не собран: {e}", flush=True)

    # Final-deck sanity warnings (free, no LLM): the plan's order IS the final
    # slide order, so each position's role is known — dividers/closings are
    # sparse by design and excluded inside the check.
    try:
        final_prs = Presentation(out_pptx)
        slide_roles = {pos: block["type"] for pos, (block, _) in enumerate(plan)}
        warnings = [
            {"slide": issue["slide"] + 1, "kind": issue["kind"], "details": issue["details"]}
            for issue in find_sparse_slides(final_prs, slide_roles)
        ]
    except Exception:
        warnings = []  # advisory only — must never break a successful generation
    try:
        if source_numbers is None:
            source_numbers = extract_numbers(brief)
        warnings += _numbers_warnings(out_pptx, source_numbers)
    except Exception:  # noqa: BLE001 — advisory, same as above
        pass

    # Сессия — чтобы пользователь мог выбрать находки и исправить их, не
    # перегенерируя всю колоду (ТЗ: «пользователь выбирает, какие исправить»).
    try:
        with open(_session_path(generation_id), "w", encoding="utf-8") as f:
            json.dump({"template_id": template_id, "brief": brief, "model": model,
                       "source_numbers": source_numbers,
                       "plan": [[block, idx] for block, idx in plan], "skipped": skipped},
                      f, ensure_ascii=False)
    except Exception as e:  # noqa: BLE001 — без сессии не будет только исправления
        print(f"! сессия генерации не сохранена: {e}", flush=True)

    return {
        "generation_id": generation_id,
        "download_url": f"/static/generated/{generation_id}.pptx",
        **exports,
        "slides": slide_urls,
        "plan": [
            {
                "type": block["type"],
                "title": block.get("title"),
                "archetype": archetype_map[idx] if idx in archetype_map else f"{block['type']} (synthesized)",
            }
            for block, idx in plan
        ],
        "skipped": skipped,
        "warnings": warnings,
        "balance": _current_balance(),
    }


def _block_count(block):
    """Сколько пунктов/пар у блока — перегенерация держит ту же ёмкость слайда."""
    for key in ("bullets", "stats", "left_points"):
        if isinstance(block.get(key), list) and block[key]:
            return len(block[key])
    return None


# Роли, текст которых пишет generate_block; у остальных (легаси-блоки) нечего
# перезапросить — их находки остаются для ручной правки.
_REGENERABLE = {"title", "section_divider", "bullet_list", "stats_kpi",
                "two_column_comparison", "image_caption", "closing"}


class FixRequest(BaseModel):
    slides: list[int]


@app.post("/api/generate/{generation_id}/fix")
def fix_selected(generation_id: str, req: FixRequest):
    """Исправить выбранные находки аудита: текст выбранных слайдов пишется
    заново — с теми же ограничениями, что при генерации (ёмкость слайда,
    никаких чисел сверх источника, без повторов показателей других слайдов),
    колода пересобирается и проходит аудит заново. Остальные слайды не
    трогаются — их текст дословно тот же."""
    _require_model()
    if not all(c.isalnum() for c in generation_id):
        raise HTTPException(status_code=400, detail="Неверный идентификатор")
    try:
        with open(_session_path(generation_id), encoding="utf-8") as f:
            session = json.load(f)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Генерация не найдена — соберите презентацию заново")

    plan = [(block, idx) for block, idx in session["plan"]]
    brief, model = session["brief"], session.get("model")
    models = [model] if model else DEFAULT_MODELS
    template_id = session["template_id"]
    if template_id not in _template_registry:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    style_preamble = card_prompt_preamble(load_card(PARSED_DIR, template_id))
    targets = sorted({n for n in req.slides if 1 <= n <= len(plan)})

    not_fixed = []
    try:
        for i, number in enumerate(targets):
            block, idx = plan[number - 1]
            role = block.get("type")
            if role not in _REGENERABLE:
                not_fixed.append({"slide": number, "reason": "этот слайд собран без модели"})
                continue
            _set_progress(f"Переписываем слайд {number}", done=i, total=len(targets))
            used_nums, used_labels = set(), set()
            for other_pos, (other, _) in enumerate(plan):
                if other_pos != number - 1:
                    nums, labels = stat_fingerprints(other)
                    used_nums |= nums
                    used_labels |= labels
            try:
                fresh = generate_block(client, role, block.get("title") or "", brief,
                                       count=_block_count(block), models=models,
                                       style_preamble=style_preamble,
                                       used_stats=(used_nums, used_labels))
            except (ValueError, KeyError):
                not_fixed.append({"slide": number, "reason": "модель не дала годного текста"})
                continue
            plan[number - 1] = (fresh, idx)
        result = _assemble_deck(template_id, plan, session.get("skipped", []), brief, model,
                                session.get("source_numbers"))
    finally:
        _set_progress("", active=False)
    result["fixed_slides"] = [n for n in targets if n not in {x["slide"] for x in not_fixed}]
    result["not_fixed"] = not_fixed
    return result
