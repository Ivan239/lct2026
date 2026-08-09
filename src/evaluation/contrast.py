"""Deterministic text/background contrast from the RENDER — the objective
backstop for the one defect that the eye catches instantly but neither the
XML-geometry checks nor the GigaChat judge do: the deck's light ink (orange
accents, light-grey body) landing on a white breather slide, readable on the
navy slides and washed-out on the white ones.

XML can't see it — the text colour is theme-inherited, and the failure only
exists once LibreOffice has painted the actual background under it. So we measure
it on the pixels, with WCAG contrast.

Robustness note (learned the hard way on real renders): a coarse thumbnail
smears text into anti-alias blends, and taking the worst of ALL colour clusters
just measures those dark/light halos, flagging every slide. The fix: sample at a
higher resolution and look only at the DOMINANT ink cluster (>=15% of the ink
pixels) — that's the crisp glyph interior, i.e. the real text colour, never a
halo (halos spread thin across many shades). Its contrast vs the border-modal
background cleanly separates navy slides (~6:1) from the white ones (~2.4:1).
Pure PIL (no numpy in this env).
"""

from collections import Counter

from PIL import Image

_QUANT_SHIFT = 4
_RESIZE = (480, 270)       # enough that glyph interiors survive as real colour, still fast
INK_DISTANCE = 70          # RGB distance from bg to count as ink, not an anti-alias halo
MIN_INK_FRACTION = 0.15    # a real text colour is a big share of ink pixels; halos never are
LOW_CONTRAST_RATIO = 3.0   # WCAG AA large-text; the deck's ink drops under this on a white bg


def _quantize(rgb):
    r, g, b = rgb
    return (r >> _QUANT_SHIFT << _QUANT_SHIFT,
            g >> _QUANT_SHIFT << _QUANT_SHIFT,
            b >> _QUANT_SHIFT << _QUANT_SHIFT)


def _lin(c):
    c /= 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def _rel_luminance(rgb):
    r, g, b = (_lin(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(a, b):
    la, lb = _rel_luminance(a), _rel_luminance(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def _dist(a, b):
    return sum((x - y) ** 2 for x, y in zip(a, b)) ** 0.5


def _pixels(png_path):
    with Image.open(png_path) as img:
        im = img.convert("RGB").resize(_RESIZE)
        raw = im.tobytes()
    w, h = _RESIZE
    px = [(raw[i], raw[i + 1], raw[i + 2]) for i in range(0, len(raw), 3)]
    return px, w, h


def _background(px, w, h):
    """Modal colour of the frame border — the border is always background, even
    on card layouts where the interior is a coloured panel."""
    border = []
    for y in (0, 1, h - 2, h - 1):
        border += [px[y * w + x] for x in range(0, w, 3)]
    for x in (0, 1, w - 2, w - 1):
        border += [px[y * w + x] for y in range(0, h, 3)]
    return Counter(_quantize(p) for p in border).most_common(1)[0][0]


def slide_text_contrast(png_path):
    """WCAG contrast of the slide's DOMINANT ink colour vs its background.
    None when the slide has no substantial ink block (blank / breather / a
    sparse divider) — a safe 'skip', never a false low-contrast flag."""
    px, w, h = _pixels(png_path)
    bg = _background(px, w, h)
    ink = [_quantize(p) for p in px if _dist(p, bg) > INK_DISTANCE]
    if not ink:
        return None
    counts = Counter(ink)
    total = len(ink)
    worst = None
    for color, n in counts.most_common(6):
        if n / total < MIN_INK_FRACTION:
            break  # most_common is descending — nothing further qualifies
        cr = contrast_ratio(color, bg)
        worst = cr if worst is None else min(worst, cr)
    return worst


def evaluate_contrast(png_paths, threshold=LOW_CONTRAST_RATIO):
    """Returns {"per_slide": [ratio|None], "low_contrast_slides": [0-based idx]}."""
    per = [slide_text_contrast(p) for p in png_paths]
    low = [i for i, c in enumerate(per) if c is not None and c < threshold]
    return {"per_slide": per, "low_contrast_slides": low}


# --- Declared-colour check -------------------------------------------------
# The pixel scan above can only see text that differs from the background by at
# least INK_DISTANCE. Text painted almost exactly in the background colour forms
# no distinct cluster at all, so the scan reports "no ink" and the slide passes
# silently — the worst case slips through. Measured on a real deck: subtitle
# #1D1C1D on a (16,16,16) slide is 22 RGB units from the background, invisible
# to the eye AND invisible to the scan.
#
# So read the colours the file DECLARES and compare them with the MEASURED
# background. Only explicit run colours are used; theme/inherited ones are left
# alone rather than guessed at.
#
# Two false-positive guards, both learned from real decks:
#
# 1. A designer may legitimately put light text on a dark button sitting on a
#    light slide, which looks identical to this test from the XML. So a slide is
#    only flagged when the near-invisible runs carry at least
#    INVISIBLE_TEXT_SHARE of its characters — a whole unreadable slide trips it,
#    a small inverted label does not.
# 2. WCAG contrast is a LUMINANCE ratio and ignores hue, so a high-chroma brand
#    pairing scores as badly as invisible text: T-Zh's red-on-blue slide measures
#    1.08, yet the text is perfectly legible because the hues are 339 RGB units
#    apart. "Invisible" must mean the colours are actually CLOSE, so a flag needs
#    low contrast AND small colour distance. The real failures sit far under this
#    (#3D3D3D on (16,16,16) is 78 apart, #1D1C1D on (16,16,16) just 22).
INVISIBLE_TEXT_SHARE = 0.5
INVISIBLE_MAX_DISTANCE = 120


def _declared_run_colours(slide):
    """[(text_length, (r,g,b))] for runs with an EXPLICIT rgb colour."""
    out = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                text = run.text or ""
                if not text.strip():
                    continue
                colour = run.font.color
                try:
                    # .rgb raises for scheme colours — those are inherited, skip.
                    if colour is None or colour.type is None or colour.rgb is None:
                        continue
                    rgb = colour.rgb
                except Exception:  # noqa: BLE001 — scheme colour, see note below
                    # Scheme colours are deliberately NOT resolved through the
                    # theme palette. Tried it and reverted: the theme claimed
                    # #EEEEEE (lt2) for text that renders BLACK, so two legible
                    # card slides came back "invisible". Scheme tokens resolve
                    # through the master's <p:clrMap>, and on top of that the
                    # theme routinely lies about the deck's real look (CLAUDE.md).
                    # Staying blind here is safer than being confidently wrong —
                    # the pixel scan still covers those slides.
                    continue
                out.append((len(text.strip()), (rgb[0], rgb[1], rgb[2])))
    return out


def evaluate_declared_contrast(pptx_path, png_paths, threshold=LOW_CONTRAST_RATIO):
    """Slides whose declared text colour is invisible against the measured
    background. Returns {"per_slide": [worst_ratio|None], "invisible_slides":
    [idx], "coverage": (slides_with_data, total)}.

    `coverage` is reported because this pass reads EXPLICIT rgb only and is
    simply blind on theme-coloured text — on a card template that can be most of
    the deck. Without it in the result, every later reader has to rediscover
    that blindness before trusting a clean verdict."""
    from pptx import Presentation

    from design_system.style_profile import page_base_color, slide_background

    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    page_base = page_base_color(png_paths) if png_paths else None

    per, flagged = [], []
    for i, slide in enumerate(slides):
        bg = slide_background(png_paths[i], page_base) if i < len(png_paths) else None
        runs = _declared_run_colours(slide)
        if bg is None or not runs:
            per.append(None)
            continue
        total = sum(n for n, _ in runs) or 1
        bad_chars, worst = 0, None
        for n, rgb in runs:
            ratio = contrast_ratio(rgb, bg)
            worst = ratio if worst is None else min(worst, ratio)
            if ratio < threshold and _dist(rgb, bg) < INVISIBLE_MAX_DISTANCE:
                bad_chars += n
        per.append(worst)
        if bad_chars / total >= INVISIBLE_TEXT_SHARE:
            flagged.append(i)
    return {
        "per_slide": per,
        "invisible_slides": flagged,
        "coverage": (sum(1 for r in per if r is not None), len(per)),
    }


# --- Per-textbox measurement ------------------------------------------------
# The whole-frame scan cannot work on card layouts and two measured attempts to
# rescue it failed (iter20 moved the background source, iter21 resolved theme
# colours; both made it MORE confidently wrong). The reason is structural: a card
# slide has TWO large colour regions — the page margin and the card — so whichever
# one is called "background", the other becomes a huge "ink" cluster. Measured on
# a deck every slide of which is plainly legible: 5 of 6 slides flagged, ratio
# 1.32 on each, which is white-page-vs-coloured-card and nothing to do with text.
#
# Text, though, sits on whatever is directly behind IT. The .pptx knows where
# every text box is, so sample only inside those rectangles: the box's modal
# colour is the local background (glyphs cover a minority of a text box), and the
# dominant colour far from it is the ink. No page/card ambiguity can arise,
# because a text box is never half margin and half card.
BOX_INSET = 0.04          # trim the box border: rounded corners/edges of art bleed in
MIN_BOX_PIXELS = 200      # below this the crop is too small to cluster reliably
MAX_INK_SHARE = 0.45      # ink above this means the modal colour IS the text
MIN_BG_SHARE = 0.5        # below this the box has no background — it is on a picture


def _box_pixels(img, rect, size):
    """Pixels inside a text box, in image coordinates. rect/size are EMU."""
    w, h = img.size
    x0 = int(w * (rect[0] + rect[2] * BOX_INSET) / size[0])
    x1 = int(w * (rect[0] + rect[2] * (1 - BOX_INSET)) / size[0])
    y0 = int(h * (rect[1] + rect[3] * BOX_INSET) / size[1])
    y1 = int(h * (rect[1] + rect[3] * (1 - BOX_INSET)) / size[1])
    x0, y0 = max(0, x0), max(0, y0)
    x1, y1 = min(w, x1), min(h, y1)
    if x1 - x0 < 2 or y1 - y0 < 2:
        return []
    return list(img.crop((x0, y0, x1, y1)).getdata())


def _box_contrast(pixels):
    """(ratio, ink, bg) for one text box, or None when it can't be measured."""
    if len(pixels) < MIN_BOX_PIXELS:
        return None
    quant = [_quantize(p) for p in pixels]
    counts = Counter(quant)
    bg, bg_n = counts.most_common(1)[0]
    if bg_n / len(quant) < MIN_BG_SHARE:
        # The box has no background — it is sitting on a PICTURE, and the modal
        # colour is just the picture's most common shade. This does not merely
        # blind the check, it makes it confidently wrong: on a synthesized slide
        # whose white text landed across two stock photos, the right column came
        # back at a healthy 7.40 — measured between the photo's black trees and
        # its pink sky, nothing to do with our text. Measured separation is wide:
        # 98 real text boxes across four decks sit at 0.70-0.99 modal share
        # (median 0.92), while those two photo-backed boxes are 0.02 and 0.34.
        return None
    ink = [p for p in quant if _dist(p, bg) > INK_DISTANCE]
    if not ink or len(ink) / len(quant) > MAX_INK_SHARE:
        # No glyphs found, or the box is mostly "ink" — which means the modal
        # colour is the text and the background is what we'd be measuring
        # against. Either way this box cannot answer the question; stay silent
        # rather than report a number that means something else.
        return None
    colour, n = Counter(ink).most_common(1)[0]
    if n / len(ink) < MIN_INK_FRACTION:
        return None
    return contrast_ratio(colour, bg), colour, bg


def evaluate_boxed_contrast(pptx_path, png_paths, threshold=LOW_CONTRAST_RATIO):
    """Contrast measured inside each text box's own rectangle.

    Returns {"per_slide": [worst_ratio|None], "low_contrast_slides": [idx],
    "coverage": (slides_with_data, total)} — same shape as the frame scan, so it
    can stand in for it. A slide's ratio is its WORST measurable box."""
    from pptx import Presentation

    prs = Presentation(pptx_path)
    size = (prs.slide_width, prs.slide_height)
    per, low, blind = [], [], []
    for i, slide in enumerate(prs.slides):
        if i >= len(png_paths):
            per.append(None)
            blind.append(0)
            continue
        with Image.open(png_paths[i]) as raw:
            img = raw.convert("RGB")
            worst, flagged, unmeasured = None, False, 0
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text_frame.text.strip():
                    continue
                if None in (shape.left, shape.top) or not (shape.width and shape.height):
                    continue
                measured = _box_contrast(
                    _box_pixels(img, (shape.left, shape.top, shape.width, shape.height), size))
                if measured is None:
                    # Counted, not just skipped: a box we cannot measure is text
                    # whose readability nobody checked, and silence about it
                    # reads exactly like "checked and fine".
                    unmeasured += 1
                    continue
                ratio, ink, bg = measured
                worst = ratio if worst is None else min(worst, ratio)
                # WCAG contrast is a LUMINANCE ratio and ignores hue, so a
                # high-chroma brand pairing scores like invisible text: this
                # deck's blue headings measure 2.50-2.96 on the mint and beige
                # cards and are perfectly legible. Same guard the declared-colour
                # pass has carried since iter17 — "invisible" has to mean the
                # colours are actually CLOSE. Measured here: those headings sit
                # 144 and 201 RGB units from their card, well outside the
                # threshold, while real invisible text lands far inside it.
                #
                # A box cleared by that guard must still count as MEASURED. The
                # first version returned it as "no data", and a slide whose only
                # text is brand-coloured then fell back to the frame scan — the
                # exact false positive this pass exists to remove. Caught by the
                # synthetic card test, not by the real deck, where such slides
                # happened to carry other text too.
                if ratio < threshold and _dist(ink, bg) < INVISIBLE_MAX_DISTANCE:
                    flagged = True
        per.append(worst)
        blind.append(unmeasured)
        if flagged:
            low.append(i)
    return {
        "per_slide": per,
        "low_contrast_slides": low,
        "unmeasured_boxes": blind,
        "coverage": (sum(1 for r in per if r is not None), len(per)),
    }
