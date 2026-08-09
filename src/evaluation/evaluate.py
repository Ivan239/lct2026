"""Orchestrator: render a deck, run every DETERMINISTIC rubric check (zero
tokens), and leave the LLM-mode criteria for Claude to fill in by eye.

No LLM API is called here to judge anything. GigaChat vision was tried and
measured: it scored a deck with duplicate stat slides and half-empty layouts
90+/100, and it stayed wrong even after several iterations of prompt fixes. The
user looked at the same renders and called it 3-4/10 on the spot. Claude is the
harness's judge now — see evaluation.claude_review.apply_claude_scores, which
merges Claude's own per-criterion scores (assigned after Reading the rendered
PNGs) into the result this module produces and recomputes the weighted total.
That merge step is pure arithmetic; still no network call.
"""

import json
import os
import time

from pptx import Presentation

from evaluation import contrast, deterministic, rubric
from rendering.render import render_pptx_to_pngs

EVAL_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "output", "evaluations")

# Rubric criteria this module can decide are N/A purely from what media the deck
# contains (deck_media()) — no judgment call, just "is there a picture/chart".
_NO_MEDIA_AT_ALL = ("4.1", "4.2", "4.3", "4.4", "4.5", "dop_image_style")
_PLACEHOLDER_ONLY = ("4.2", "dop_image_style")  # skeleton frame: no real image to grade quality/style
_NO_CHARTS = ("7.1", "7.2", "7.3")


def _slide_texts(pptx_path):
    prs = Presentation(pptx_path)
    texts = []
    for slide in prs.slides:
        parts = [s.text_frame.text.strip() for s in slide.shapes
                 if s.has_text_frame and s.text_frame.text.strip()]
        texts.append("\n".join(parts))
    return texts


def evaluate_deck(pptx_path, brief, slide_roles=None, render_dir=None,
                  out_json=None, label=None):
    """Deterministic half of the rubric, always: geometric/typographic checks,
    render-based contrast, and N/A-gating for image/infographic criteria the
    deck doesn't even have the media for. Every llm-mode criterion starts as
    "ожидает оценки Клода" — call claude_review.apply_claude_scores afterward
    with Claude's own judgment to fill those in and get a real total_100.

    Renders the deck once and records the PNG paths in the result so the caller
    (and Claude, reviewing them) doesn't need to render a second time."""
    render_dir = render_dir or os.path.join(EVAL_DIR, "_render")
    os.makedirs(render_dir, exist_ok=True)
    os.makedirs(EVAL_DIR, exist_ok=True)

    scores = {}
    scores.update(deterministic.evaluate(pptx_path, slide_roles=slide_roles))

    png_paths = render_pptx_to_pngs(pptx_path, render_dir)

    # Image/infographics N/A-gating — deterministic (deck_media(), not a guess).
    media = deterministic.deck_media(pptx_path)
    has_pictures = media["substantive_pictures"] > 0
    has_placeholders = media.get("placeholders", 0) > 0
    if not has_pictures and not has_placeholders:
        for cid in _NO_MEDIA_AT_ALL:
            scores[cid] = {"score": None, "detail": "неприменимо: в деке нет такого медиа"}
    elif has_placeholders and not has_pictures:
        for cid in _PLACEHOLDER_ONLY:
            scores[cid] = {"score": None, "detail": "неприменимо: изображение ещё не сгенерировано (скелет)"}
    if media["charts"] == 0:
        for cid in _NO_CHARTS:
            scores[cid] = {"score": None, "detail": "неприменимо: в деке нет диаграмм"}

    # Contrast backstop (zero-token, from the render): downgrade readability (1.1)
    # when the deck's ink washes out on a slide's background — the light-grey-on-
    # white breather defect the eye catches but font/overflow math misses. Only
    # lowers 1.1, never raises it.
    # Measured inside each text box's own rectangle (iter34). The whole-frame
    # scan cannot work on card layouts — two large colour regions mean whichever
    # is called "background", the other becomes a huge "ink" cluster — and two
    # attempts to rescue it were measured and reverted (iter20, iter21).
    #
    # It was kept as a fallback for slides the boxed pass cannot measure. That
    # fallback is now gone, on measurement: across 30 slides generated on all
    # five real templates the boxed pass covered EVERY slide, so the fallback
    # never once fired — while on the card template, asked the same question, it
    # answered "low contrast" for 5 of 6 plainly legible slides. Its expected
    # contribution is no true positives and some false ones. Slides it cannot
    # measure are reported as missing coverage instead of guessed at; a slide
    # with no readable text box has no text-contrast verdict to give.
    cres = contrast.evaluate_boxed_contrast(pptx_path, png_paths)
    # Second, independent pass: the pixel measurement only sees text that
    # differs from its background; text painted almost IN the background colour
    # forms no cluster and slips through. Reading declared run colours against
    # the measured background catches it — found real invisible bullets,
    # (33,63,255) text on a (32,56,248) slide, which the eye reads as "the slide
    # is empty". The two are ADDITIVE: each sees what the other structurally
    # cannot. The old reconcile step existed to suppress the frame scan's card
    # false positives, and suppressing a boxed verdict — measured on the actual
    # rendered glyphs — would now only lose true positives.
    dres = contrast.evaluate_declared_contrast(pptx_path, png_paths)
    cres["invisible_text_slides"] = dres["invisible_slides"]
    cres["declared_per_slide"] = dres["per_slide"]
    cres["declared_coverage"] = dres.get("coverage")

    # Slides the boxed pass could not measure at all — its text sits on a
    # picture, so there is no background to measure against (iter41). Silence
    # there reads exactly like "checked and fine", which is how a harness ends
    # up trusted where it is blind; iter21 added `coverage` to the declared pass
    # for the same reason. Named in the detail so the blindness is visible.
    unmeasured = [i for i, n in enumerate(cres.get("unmeasured_boxes") or []) if n]
    blind_note = (f"; контраст НЕ ПРОВЕРЕН на слайдах {[i + 1 for i in unmeasured]} "
                  "(текст поверх картинки — фона для замера нет)") if unmeasured else ""

    low = sorted(set(cres["low_contrast_slides"]) | set(dres["invisible_slides"]))
    if not low and blind_note:
        cur = scores.get("1.1", {})
        scores["1.1"] = {
            "score": cur.get("score"),
            "detail": (cur.get("detail") or "").rstrip() + blind_note,
        }
    if low:
        frac = len(low) / len(png_paths)
        cscore = max(1, min(5, round(5 - 4 * frac)))
        cur = scores.get("1.1", {}).get("score")
        if cur is None or cscore < cur:
            scores["1.1"] = {
                "score": cscore,
                "detail": (f"низкий контраст текст/фон на слайдах {[i + 1 for i in low]}"
                           + (f"; НЕВИДИМЫЙ текст на {[i + 1 for i in dres['invisible_slides']]}"
                              if dres["invisible_slides"] else "") + blind_note)
                          + (f"; {scores['1.1']['detail']}" if scores.get('1.1', {}).get('detail') else ""),
            }

    # Everything else (composition, tone, hallucinations, ...) waits for Claude.
    for cid in rubric.CRITERIA:
        scores.setdefault(cid, {"score": None, "detail": "ожидает оценки Клода"})

    total = rubric.weighted_total(scores)
    result = {
        "pptx": os.path.abspath(pptx_path),
        "label": label or os.path.basename(pptx_path),
        "brief": brief,
        "timestamp": int(time.time()),
        "total_100": total,
        "buckets": rubric.bucket_breakdown(scores),
        "scores": {
            cid: {"title": rubric.CRITERIA[cid][0], "mode": rubric.CRITERIA[cid][1], **scores[cid]}
            for cid in rubric.CRITERIA
        },
        "llm_evaluated": False,
        "judge": None,
        "contrast": cres,
        "render_dir": os.path.abspath(render_dir),
        "png_paths": png_paths,
    }

    out_json = out_json or os.path.join(EVAL_DIR, f"{result['label']}_{result['timestamp']}.json")
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    result["_json_path"] = os.path.abspath(out_json)
    return result


def format_report(result):
    """Human-readable one-screen summary of an evaluate_deck result."""
    lines = [
        f"# Оценка: {result['label']}",
        f"Итог (харнесс): {result['total_100']}/100"
        + ("" if result["llm_evaluated"] else "  — ТОЛЬКО детерминированные критерии, LLM-часть ждёт оценки Клода"),
        "",
        "## По категориям",
    ]
    for b in result["buckets"].values():
        s = b["score_1_5"]
        lines.append(f"- {b['title']} (вес {b['weight']:.0%}): "
                     + (f"{s}/5" if s is not None else "N/A"))
    lines.append("")
    lines.append("## Слабые места (score ≤ 3)")
    weak = sorted(
        ((cid, v) for cid, v in result["scores"].items()
         if v["score"] is not None and v["score"] <= 3),
        key=lambda kv: kv[1]["score"],
    )
    if weak:
        for cid, v in weak:
            lines.append(f"- [{v['score']}] {cid} {v['title']} — {v['detail']}")
    else:
        lines.append("- нет (все оценённые критерии ≥ 4)")

    # Blind spots are reported even when nothing scored low — the report lists
    # only score<=3, so a healthy deck's 1.1 sits at 5 and the "could not be
    # measured" note it carries would never reach a reader. That is the exact
    # failure this note exists to prevent, one level up.
    blind = [i + 1 for i, n in enumerate(result["contrast"].get("unmeasured_boxes") or []) if n]
    if blind:
        lines.append("")
        lines.append(f"## Не проверено\n- контраст на слайдах {blind}: "
                     "текст лежит на картинке, фона для замера нет — смотри глазами")
    return "\n".join(lines)
