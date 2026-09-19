"""One turn of the improvement loop, headless and per-LLM.

    pick a model -> ensure a template parsed BY that model -> generate a deck
    from a canonical brief -> score it against the rubric.

Deliberately self-contained: it drives the same lower-level pipeline blocks the
API does, but does NOT import api.main (whose import runs preset bootstrap and
reaches GigaChat). Each model gets its own workdir under output/loop/<model>/ so
"a template from LLM X" is a real, cached, separately-owned artifact — the
provenance the loop needs to answer "do we already have one from this LLM?".
"""

import json
from collections import Counter
import os
import re
import time

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from common.pictures import has_oversized_picture
from common.synthesis import SYNTHESIZE
from design_system.extractor import _describe_slide, build_archetype_map
from design_system.fingerprint_cache import FingerprintCache
from design_system.style_card import (
    build_style_card,
    card_path,
    card_prompt_preamble,
    load_card,
    save_card,
)
from design_system.style_profile import build_measured_profile, rotation_targets
from content_parser.two_phase import generate_block, generate_outline, stat_fingerprints
from evaluation.evaluate import evaluate_deck
from generator.generator import generate
from generator.slide_kit import content_text_shapes
from llm_clients.backends import GIGACHAT_MODELS
from matcher.matcher import plan_from_outline
from rendering.render import render_pptx_to_pngs
from template_parser.parser import extract_template
from template_spec.builder import build_spec

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
LOOP_ROOT = os.path.join(BASE, "output", "loop")
# Shared geometry-only cache — same file the API uses; safe to share across every
# template and model (no client text in it), and it saves re-classification.
FINGERPRINT_CACHE_PATH = os.path.join(BASE, "output", "parsed", "fingerprint_cache.json")


def safe_model(model):
    return re.sub(r"[^A-Za-z0-9._-]", "_", model)


def _workdir(model, source_name, out_root=LOOP_ROOT):
    """Per (model, template) — NOT per model alone. The cached archetype map and
    style card describe ONE template; keying them by model only meant that the
    moment the loop rotated templates it would silently reuse template A's
    parse for template B."""
    d = os.path.join(out_root, safe_model(model), safe_model(source_name))
    os.makedirs(os.path.join(d, "rendered"), exist_ok=True)
    os.makedirs(os.path.join(d, "decks"), exist_ok=True)
    return d


def _is_failed_parse(archetype_map):
    """A map that is nothing but "other" is a FAILED classification, not a parse.

    build_archetype_map degrades a slide to "other" on any exception, so a run
    that loses the network mid-way produces a full map of shrugs — and writing
    it to archetypes.json freezes that non-answer for the (model, template) pair
    forever, because later runs reuse the cache. Measured on this machine: 9 of
    20 cached parses were entirely "other", including three of the five real
    templates under GigaChat-2, so those pairs could only ever produce decks
    with zero native slides.

    The project already applies exactly this reasoning one level down —
    extractor's fingerprint cache refuses to store "other" because "a shrug is
    not knowledge". This is the same rule for the per-template cache.

    A template whose every slide genuinely is "other" costs one re-parse and
    then answers the same; a poisoned cache costs every future run."""
    values = set(archetype_map.values())
    return bool(values) and values <= {"other"}


def ensure_template(client, model, source_pptx, source_name, out_root=LOOP_ROOT, force=False):
    """Parse `source_pptx` under `model` if not already cached for it. Returns a
    dict with the archetype map, spec, style preamble, and meta (incl. which LLM
    parsed it and when). Reuses the cached parse on subsequent iterations."""
    workdir = _workdir(model, source_name, out_root)
    arch_path = os.path.join(workdir, "archetypes.json")
    meta_path = os.path.join(workdir, "meta.json")

    render_dir = os.path.join(workdir, "rendered")
    png_paths = render_pptx_to_pngs(source_pptx, render_dir)

    cached = None
    if os.path.exists(arch_path) and not force:
        with open(arch_path, encoding="utf-8") as f:
            cached = {int(k): v for k, v in json.load(f).items()}
        if _is_failed_parse(cached):
            print(f"[loop] кэш разбора {source_name} под {model} — сплошной 'other', "
                  "переразбираю", flush=True)
            cached = None
    if cached is not None:
        archetype_map = cached
        meta = json.load(open(meta_path, encoding="utf-8")) if os.path.exists(meta_path) else {}
    else:
        struct = extract_template(source_pptx)
        archetype_map = build_archetype_map(
            client, struct, rendered_png_paths=png_paths, text_models=[model],
            fingerprint_cache=FingerprintCache(FINGERPRINT_CACHE_PATH),
        )
        # Same exclusion the API applies: a slide built around a big source
        # picture can't host our text-only archetypes.
        prs = Presentation(source_pptx)
        for idx in list(archetype_map):
            if has_oversized_picture(prs.slides[idx], prs.slide_width, prs.slide_height):
                archetype_map[idx] = "other"
        if _is_failed_parse(archetype_map):
            # Do not persist a shrug: the next run would inherit it instead of
            # trying again.
            raise RuntimeError(
                f"классификация {source_name} под {model} вернула только 'other' — "
                "разбор не сохранён")
        with open(arch_path, "w", encoding="utf-8") as f:
            json.dump(archetype_map, f, ensure_ascii=False, indent=2)

        if load_card(workdir, "template") is None:
            try:
                descriptions = [_describe_slide(s) for s in struct["slides"]]
                card = build_style_card(client, descriptions, profile=None,
                                        archetype_map=archetype_map, models=[model])
                save_card(workdir, "template", card)
            except Exception as e:  # noqa: BLE001 — style card is advisory
                print(f"[loop] style_card failed for {model}: {e}", flush=True)

        meta = {
            "model": model,
            "source_pptx": os.path.abspath(source_pptx),
            "source_name": source_name,
            "parsed_at": int(time.time()),
        }
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)

    profile = build_measured_profile(png_paths, archetype_map) if png_paths else None
    spec = build_spec(source_pptx, archetype_map, style_profile=profile)
    style_preamble = card_prompt_preamble(load_card(workdir, "template"))
    return {
        "workdir": workdir,
        "archetype_map": archetype_map,
        "spec": spec,
        "profile": profile,
        "style_preamble": style_preamble,
        "png_paths": png_paths,
        "meta": meta,
        "has_card": load_card(workdir, "template") is not None,
    }


def _synth_canvas_hints(source_pptx, plan, profile):
    """{plan_position: template_slide_idx} for SYNTHESIZE positions — mirrors the
    API's _synth_canvas_hints so the headless loop doesn't draw synthesized
    slides on a blank white page. Without this the synth path leaves light
    template text on the master's white background — unreadable (the defect the
    contrast backstop kept flagging on the loop's own decks, slide 5/6). Canvas =
    a furniture-only template slide (<=3 content boxes), preferring the colour
    the rotation wants at that position."""
    if not profile:
        return {}
    try:
        targets = rotation_targets(profile, len(plan))
        prs = Presentation(source_pptx)
        photo_limit = Inches(1.2)
        candidates = []
        for idx in range(len(prs.slides)):
            if idx in profile["breathers"]:
                continue
            slide = prs.slides[idx]
            has_photo = any(
                s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width and s.width > photo_limit
                and not (s.width >= prs.slide_width * 0.9)
                for s in slide.shapes
            )
            if has_photo:
                continue
            candidates.append((len(content_text_shapes(slide)), idx, profile["backgrounds"].get(idx)))
        if not candidates:
            return {}
        # "Fewest text boxes" alone picks the blandest slide in the file, and the
        # blandest slide is often not a design slide at all: the T-Zh universal
        # template ends with a «Технический слайд» — instructions for whoever
        # uses the template ("we used Roboto and Inter, download them here") —
        # carrying 3 boxes and no furniture. The branded candidates carry 5 and
        # 9, so the <=3 cut left the readme as the ONLY option and all three
        # synthesized slides were cloned from it: bare white pages on a template
        # whose identity is black with yellow artwork.
        #
        # The tell is its background. That page is (248,248,248) and belongs to
        # exactly ONE slide, while every design slide in the deck is (0,0,0), so
        # drop singleton-background candidates before the furniture-only cut.
        # Tried filtering to the deck's MAJORITY background instead and reverted
        # it inside this iteration: that collapses every canvas to one colour and
        # kills the template's own rotation (T-Zh mono went from alternating
        # blue/white to white everywhere). The singleton test is precise — mono
        # is genuinely half white, so its white stays eligible.
        deck_bg = Counter(bg for bg in profile["backgrounds"].values() if bg is not None)
        on_brand = [c for c in candidates if c[2] is None or deck_bg[c[2]] > 1] or candidates
        eligible = [c for c in on_brand if c[0] <= 3] or sorted(on_brand)[:3]
        hints = {}
        for position, (_, slide_idx) in enumerate(plan):
            if slide_idx != SYNTHESIZE:
                continue
            target = targets[position] if targets else None
            best = min(eligible, key=lambda c: (
                0 if target is not None and c[2] == target else 1, c[0], c[1]))
            hints[position] = best[1]
        return hints
    except Exception as e:  # noqa: BLE001 — best-effort, falls back to from-scratch
        print(f"[loop synth_canvas] {e}", flush=True)
        return {}


def _gen_models(model):
    """Primary model + a strong safety-net tier. ANY tier intermittently returns
    a block that fails strict validation; forcing models=[model] left
    call_with_model_fallback no fallback, so one bad block raised ValueError and
    crashed the ENTIRE iteration. The net only engages for the specific blocks the
    primary can't produce validly, so the deck stays essentially `model`'s.

    Every GigaChat tier gets a net now — including Max itself: forcing Max used to
    have NO fallback, and a run of malformed JSON from Max still crashed the run
    (hit for real on `--model GigaChat-2-Max`). Max falls back to Pro (next
    strongest); everyone else falls back to Max. Not for non-GigaChat clients
    (RTX) — they can't serve a GigaChat model name."""
    if model not in GIGACHAT_MODELS:
        return [model]
    net = "GigaChat-2-Pro" if model == "GigaChat-2-Max" else "GigaChat-2-Max"
    return [model, net]


def generate_deck(client, model, source_pptx, spec, brief, style_preamble, out_pptx, profile=None):
    """Two-phase generation under one model (with a validation-failure safety net,
    see _gen_models). Returns (plan, skipped) — plan is [(block, slide_idx)], its
    order is the final slide order."""
    models = _gen_models(model)
    outline = generate_outline(client, brief, spec, models=models, style_preamble=style_preamble)
    assignments, skipped_items = plan_from_outline(outline, spec)
    plan = []
    # Blocks are generated one call at a time and can't see each other — carry the
    # stat numbers/labels already placed so a later stat slide can't repeat them.
    used_nums, used_labels = set(), set()
    # The slide is chosen BEFORE its text is written, so the item budget of that
    # exact slide is known here (SYNTHESIZE positions have no template slide and
    # keep the flat default).
    item_chars = {s["idx"]: s.get("item_chars")
                  for fam in spec["families"] for s in fam["slides"]}
    for item, slide_idx, final_count in assignments:
        block = generate_block(client, item["role"], item["theme"], brief,
                               count=final_count, models=models, style_preamble=style_preamble,
                               used_stats=(used_nums, used_labels),
                               item_chars=item_chars.get(slide_idx))
        nums, labels = stat_fingerprints(block)
        used_nums |= nums
        used_labels |= labels
        plan.append((block, slide_idx))
    synth_canvas = _synth_canvas_hints(source_pptx, plan, profile)
    generate(source_pptx, plan, out_pptx, synth_canvas=synth_canvas,
             canvas_backgrounds=(profile or {}).get("backgrounds"))
    skipped = [{"type": it["role"], "title": it.get("theme")} for it in skipped_items]
    return plan, skipped


def run_iteration(client, model, source_pptx, brief, source_name=None, out_root=LOOP_ROOT,
                  source_numbers=None):
    """Full turn for one model. `client` parses the template and generates the
    deck. Scoring is deterministic-only here (no LLM judge call, ever — Claude
    reviews the renders and calls evaluation.claude_review.apply_claude_scores
    separately). Renders straight into the deck's own "look_*" folder so the
    same PNGs serve both the eval JSON and Claude's visual review — no second
    render. Returns a bundle: template info, deck path, the rubric evaluation,
    and what was skipped. Raises only on a hard failure that leaves no deck to
    score (the caller records that and moves on).

    `source_numbers` — a content package's `numbers`, the reference for
    «все цифры есть в исходных материалах»; None means the brief is the whole
    source."""
    source_name = source_name or os.path.splitext(os.path.basename(source_pptx))[0]
    t = ensure_template(client, model, source_pptx, source_name, out_root=out_root)

    ts = int(time.time())
    out_pptx = os.path.join(t["workdir"], "decks", f"deck_{ts}.pptx")
    # Timed separately from ensure_template: the VK Tech brief caps ONE DECK at
    # 5 minutes, and a first-time template parse (54-55 slides for two of its
    # three templates) is a one-off cost that would otherwise hide inside it.
    started = time.monotonic()
    plan, skipped = generate_deck(client, model, source_pptx, t["spec"], brief,
                                  t["style_preamble"], out_pptx, profile=t["profile"])
    deck_seconds = round(time.monotonic() - started, 1)
    slide_roles = {pos: block["type"] for pos, (block, _) in enumerate(plan)}

    look_dir = os.path.join(t["workdir"], "decks", f"look_deck_{ts}")
    result = evaluate_deck(out_pptx, brief, slide_roles=slide_roles,
                           render_dir=look_dir, label=f"{safe_model(model)}_{ts}",
                           source_numbers=source_numbers)
    # How much of the TEMPLATE the deck actually got to use. Worth surfacing
    # every run: on the two real customer decks the answer is "almost none" and
    # nothing said so. 58% of survey-31 and 62% of survey-69 are forced to
    # "other" because they are built around a chart (a big stale-data picture),
    # so the matcher has barely anything native to offer and nearly every slide
    # is synthesized on a cloned canvas. That is by design — someone else's
    # numbers must not ride along — but a deck that is 90% synthesized is a
    # different product from one that reuses the designer's layouts, and the
    # loop should not have to rediscover which one it is looking at.
    native = sum(1 for _, idx in plan if idx != SYNTHESIZE)
    offered = len({s["idx"] for fam in t["spec"]["families"] for s in fam["slides"]})
    return {
        "model": model,
        "source_name": source_name,
        "deck": out_pptx,
        "n_slides": len(plan),
        "deck_seconds": deck_seconds,
        "native_slides": native,
        "synth_slides": len(plan) - native,
        "template_offered": offered,
        "template_total": len(t["archetype_map"]),
        "skipped": skipped,
        "evaluation": result,
        "template_meta": t["meta"],
    }
